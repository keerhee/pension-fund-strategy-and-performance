# -*- coding: utf-8 -*-
"""W06 M7 LDI 강의본 보강·재구성(2026-10-05) — 사용자 피드백 4건.
 1) 안 A·B·C를 '무엇을 사고/팔고/계약하고, 결과로 무엇이 바뀌나'의 동사 문장으로(도입 9장 · IC 66·68장 · 용어 7장)
 2) Redington 일화 바로 다음 장에 세 조건 PV_A = PV_L · D_A = D_L · C_A > C_L(LaTeX) + 조건마다 막는 것 + 바벨 숫자
 3) 잉여금 최적화: 목적식 · 선택 변수 · 제약(듀레이션은 제약이 아님) → 해(PSP + LHP, 헤지 몫 = DV01 매칭) → 실무적 함의
 4) 교시 대신 소주제 아홉 단원으로 재배열 · 단원 디바이더(이 단원이 푸는 질문) · 목차 표 · 머리 띠 · 상호 참조 · 쪽번호
실행: <PPTX 가능 파이썬> _build/w06_m7_rework/rework_m7.py
원본은 _구판/W06_M7_보강_구판_2026-10-05/에 한 번만 백업하고 항상 백업에서 읽는다(재실행해도 결과가 같다).
수식은 matplotlib mathtext(cm)로 카드 색(#F7F5F0) 위에 잉크(#102027)로 굽는다 — 덱의 기존 수식과 같은 배율(3.12 px/pt-px)."""
import copy, os, shutil, sys, json
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "lecfix"))
from pptx import Presentation
from pptx.util import Emu, Pt
from PIL import Image
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
import lecfix

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", ".."))
NAME = "W06_M7_LDI_부채연계투자_강의본.pptx"
SRC = os.path.join(R, "W06_LDI와GBI", NAME)
BK = os.path.join(R, "_구판", "W06_M7_보강_구판_2026-10-05"); os.makedirs(BK, exist_ok=True)
bk = os.path.join(BK, NAME)
if not os.path.exists(bk):
    shutil.copy2(SRC, bk)
    pdf = SRC[:-5] + ".pdf"
    if os.path.exists(pdf): shutil.copy2(pdf, os.path.join(BK, os.path.basename(pdf)))
EQ = os.path.join(HERE, "eq"); os.makedirs(EQ, exist_ok=True)
EMU = 9525          # 1 px(96dpi) = 9525 EMU
SCALE = 3.12        # 덱 수식 PNG: 이미지 px ÷ 장표 px

P = Presentation(bk)
O = list(P.slides)                    # 원래 번호 n → O[n-1]
def S(n): return O[n - 1]
def sh(slide, sid): return next(x for x in slide.shapes if x.shape_id == sid)
def T(slide, sid, text): lecfix.set_text(sh(slide, sid), text)

# ------------------------------------------------------------------ 도구 ----
def render(tex, name, fontsize=30, color="#102027", bg="#F7F5F0"):
    bad = [c for c in tex if 0xAC00 <= ord(c) <= 0xD7A3]
    assert not bad, f"수식 안 한글: {name}"
    path = os.path.join(EQ, name)
    with plt.rc_context({"mathtext.fontset": "cm"}):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=fontsize, color=color)
        fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.12, facecolor=bg)
        plt.close(fig)
    Image.open(path).convert("RGB").save(path)
    return path

def size_px(path, maxw=None, maxh=None):
    w, h = Image.open(path).size
    w, h = w / SCALE, h / SCALE
    s = min(1.0, (maxw / w) if maxw else 1.0, (maxh / h) if maxh else 1.0)
    return w * s, h * s

def add_pic(slide, path, x, y, w, h, before=None):
    pic = slide.shapes.add_picture(path, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)))
    if before is not None:   # z-순서: 지운 그림 자리에
        before.addprevious(pic._element)
    return pic

def remove(shape): shape._element.getparent().remove(shape._element)

def set_px(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Emu(int(x * EMU))
    if y is not None: shape.top = Emu(int(y * EMU))
    if w is not None: shape.width = Emu(int(w * EMU))
    if h is not None: shape.height = Emu(int(h * EMU))

def container(slide, pic):
    """그림을 감싸는 글자 없는 카드(사각형)를 찾는다."""
    px_, py, pw, ph = lecfix.px(pic); cx, cy = px_ + pw / 2, py + ph / 2
    best = None
    for s in slide.shapes:
        if s is pic or s.shape_type == 13 or (s.has_text_frame and s.text_frame.text.strip()): continue
        x, y, w, h = lecfix.px(s)
        if w > 0 and h > 0 and x <= cx <= x + w and y <= cy <= y + h:
            if best is None or w * h < lecfix.px(best)[2] * lecfix.px(best)[3]: best = s
    return best

def fresh_ids(slide, elements):
    mx = max(int(e.get("id")) for e in slide._element.iter() if e.tag.endswith("}cNvPr"))
    for el in elements:
        for e in el.iter():
            if e.tag.endswith("}cNvPr"):
                mx += 1; e.set("id", str(mx))

def clone(src):
    """같은 덱 안에서 장표를 복제(그림 관계 포함, 배경 포함). 맨 뒤에 붙는다."""
    dst = P.slides.add_slide(src.slide_layout)
    for x in list(dst.shapes): remove(x)
    rmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith("/image"):
            rmap[rId] = dst.part.relate_to(rel._target, rel.reltype)
    cs, cd = src._element.cSld, dst._element.cSld
    if cs.bg is not None:
        cd.insert(0, copy.deepcopy(cs.bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}nvGrpSpPr", "}grpSpPr")): continue
        e2 = copy.deepcopy(el)
        for e in e2.iter():
            for k, v in list(e.attrib.items()):
                if k.endswith("}embed") and v in rmap: e.set(k, rmap[v])
        dst.shapes._spTree.append(e2)
    return dst

def copy_shape(src_shape, dst_slide, dx=0, dy=0):
    e = copy.deepcopy(src_shape._element)
    dst_slide.shapes._spTree.append(e); fresh_ids(dst_slide, [e])
    s = next(x for x in dst_slide.shapes if x._element is e)
    s.left = Emu(s.left + int(dx * EMU)); s.top = Emu(s.top + int(dy * EMU))
    return s

def font(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(pt)

def rep(old, new, n=1):
    k = lecfix.replace_all(P, old, new)
    assert k == n, f"치환 {k}회(기대 {n}): {old}"

# ------------------------------------------------------------------ 수식 ----
E_RED = render(r"PV_A=PV_L,\qquad D_A=D_L,\qquad C_A>C_L", "red_pv.png")
E_OBJ = render(r"\max_{w}\ \mathbb{E}[R_S]-\frac{\gamma}{2}\,\mathrm{Var}(R_S),\qquad R_S=r_A-\frac{L}{A}\,r_L,\qquad r_A=\sum_i w_i\,r_i", "obj.png")
E_SOL = render(r"w_i^{*}=\frac{\mu_i}{\gamma\,\sigma_i^2}+\frac{1}{\mathrm{FR}}\cdot\frac{\mathrm{Cov}(r_i,r_L)}{\sigma_i^2}", "wstar.png")
E_LHP = render(r"w_H^{*}=\frac{1}{\mathrm{FR}}\cdot\frac{D_L}{D_H}\quad\Leftrightarrow\quad w_H^{*}\,A\,D_H=L\,D_L", "lhp.png")

# ============================================================ 원래 장 고치기 ====
# 3장 — 읽는 순서를 단원으로
s = S(3)
T(s, 5, "앞 단원의 결과를 뒤 단원에서 그대로 쓴다 — 막히면 앞 단원으로 돌아간다")
T(s, 19, "단원 ①·②"); T(s, 21, "현재가치 · 적립비율 · 할인율")
T(s, 23, "단원 ③~⑥"); T(s, 25, "듀레이션 · 면역화 · 두 블록 · 스왑")
T(s, 27, "단원 ⑦")
# 7장 — 용어 ②
T(S(7), 5, "먼저 읽는 용어 ② — 단원 ⑥의 수단, ⑦의 2022 영국, ⑨의 IC")
rep("3교시 · 2022년 9월", "단원 ⑦ · 2022년 9월")
rep("안 B는 현물, 안 C는 IRS", "안 B는 국고채를 사고, 안 C는 스왑을 계약")
# 17장 — 비유
T(S(17), 5, "수식이 어렵다면 이 세 문장만 기억한다 — 단원 ③~⑤의 수식은 이 비유의 정밀한 번역이다")
# 24장 — 단원 ①·② 체크포인트
s = S(24)
T(s, 3, "단원 ①·②는 부채를 “할인된 약속”으로 보는 법을 다뤘습니다"); T(s, 5, "단원 ①·② 체크포인트")
T(s, 14, "예상수익률 · 국채 · AA 회사채 — 선택에 따라 부채가 두 배 차이. CalPERS · 공무원연금, 영국 DB의 세 자(130.8% vs 114%).")
# 53장 — 할인율 단원으로 옮기므로 앞선 버퍼 언급을 뺀다
T(S(53), 10, "2022 위기 뒤 흑자 전환\n금리 상승이 부채를 줄였다\n논쟁은 잉여금의 쓰임")
# 32장 — Redington 일화: 다음 장이 세 조건
T(S(32), 16, "답은 다음 장의 세 조건 — 현재가치 · 듀레이션 · 볼록성을 자산과 부채에서 맞춘다")
# 33장 — 조건 2의 근거
T(S(33), 5, "조건 2의 근거 — 잉여금 변화 = 자산 변화 − 부채 변화, 각각에 −D·V·Δy를 넣는다")
T(S(33), 18, "1차만으로는 부족합니다 — 금리가 크게 움직이면 휘어짐(볼록성, 조건 3)이 끼어듭니다")
# 34장 — 단원 ③(재는 자)으로
T(S(34), 5, "볼록성 — 테일러 전개 · 10년 뒤 100을 주는 부채, y = 5%(연속복리): V = 60.65, D = 10, C = 100")

# 35장 — Redington 세 조건(현재가치 형식)으로 다시 쓰고 32장 바로 뒤로 옮긴다
s = S(35)
T(s, 3, "Redington 면역화는 현재가치 · 듀레이션 · 볼록성 세 조건입니다")
T(s, 5, "Redington(1952) — 현재가치가 같으면 조건 2·3은 D_A A = D_L L, C_A A ≥ C_L L과 같다 · 증명은 부록 B-3")
pic = sh(s, 8); box = container(s, pic)
x0, y0, w0, h0 = lecfix.px(pic); w, h = size_px(E_RED, maxw=900)
np_ = add_pic(s, E_RED, 640 - w / 2, y0 + (h0 - h) / 2, w, h, before=pic._element); remove(pic)
if box is not None:
    bx, by, bw, bh = lecfix.px(box); pad = x0 - bx
    set_px(box, x=640 - w / 2 - pad, w=w + 2 * pad)
T(s, 10, "조건 1 — 현재가치 일치"); T(s, 11, "PV_A = PV_L"); T(s, 12, "막는 것: 출발선의 부족")
T(s, 14, "조건 2 — 듀레이션 일치"); T(s, 15, "D_A = D_L"); T(s, 16, "막는 것: 금리 방향의 1차 손실")
T(s, 18, "조건 3 — 볼록성 우위"); T(s, 19, "C_A > C_L"); T(s, 20, "막는 것: 큰 변동의 휘어짐 손실")
T(s, 21, "숫자 — PV 60.65로 금리 −2%p: 바벨(C 125)은 +0.37, 5년물(D 5)은 조건 2를 어겨 −7.05")

# 37장 — 단원 ⑤로 넘기는 말
T(S(37), 17, "실무는 듀레이션 매칭 + 일부 초과수익 — 그것이 단원 ⑤의 두 블록입니다")
# 40장 — 앞 단원 예
T(S(40), 18, "※ 앞 단원의 FR 120% 예(Hedging 83% + Return-Seeking 37%)는 오버레이로 레버리지를 더한 경우다")
# 41장 — 단원 ③~⑥ 체크포인트
s = S(41)
T(s, 3, "단원 ③~⑥은 금리 위험을 재고 막는 수학과 수단을 다뤘습니다"); T(s, 5, "단원 ③~⑥ 체크포인트")
T(s, 17, "실무는 두 블록 + 스왑 오버레이")
T(s, 18, "헤지 몫(LHP)은 부채가, 수익 추구 몫(PSP)은 위험 예산이 정한다. 모자란 헤지를 스왑으로 채우면 증거금 · 버퍼가 따라온다.")
# 47장 — 실무 수단 단원으로
T(S(47), 5, "증거금의 원리 — 단원 ③의 DV01이 그대로 현금 수요가 된다")
T(S(47), 22, "영국 감독당국(TPR · 영란은행 · FCA)은 2022년 위기 뒤 250bp 급등을 견딜 현금 버퍼를 요구했습니다")
# 51장
T(S(51), 3, "단원 ③의 도구로 계산한 SVB 손실은 공시 평가손실과 거의 같습니다")
# 58장 — 단원 ⑦ 체크포인트(같은 연금 다른 숫자는 단원 ②로 옮겼다)
s = S(58)
T(s, 3, "단원 ⑦은 레버리지 헤지가 현금 시험으로 바뀌는 과정을 다뤘습니다"); T(s, 5, "단원 ⑦ 체크포인트")
T(s, 17, "갭의 방향은 달라도 시험은 현금")
T(s, 18, "영국은 짧은 자산을 스왑으로 늘려 증거금에서, SVB는 짧은 예금 앞에 긴 채권을 들고 인출에서 무너졌다.")
# 66장 — 안건 행을 동사로
t = lecfix.find_table(S(66)).table
lecfix.set_tf(t.cell(1, 1).text_frame,
              "LBP를 만들고 헤지 비율을 무엇으로 올리나 — 안 A 새로 사지 않음(3%) · 안 B 보유 채권 131조를 30년 국고채로 교체(5%) · "
              "안 C 30년 고정금리 수취 스왑 1,223조 계약(30%)")
# 67장 — 개념 출처를 단원으로
rep("FR = A/L · 할인율의 두 눈금 (1교시)", "FR = A/L · 할인율의 두 눈금 (단원 ①②)")
rep("DV01 · 듀레이션 갭 (2교시)", "DV01 · 듀레이션 갭 (단원 ③)")
rep("면역화 · 2펀드 · CF 매칭의 한계 (2교시)", "면역화 · 두 블록 · CF 매칭의 한계 (단원 ④⑤)")
rep("2022 교훈 · 증거금 · 버퍼 시뮬레이션", "2022 교훈 · 증거금 · 버퍼 시뮬레이션 (단원 ⑥~⑧)")
# 68장 — 기본 답: 안마다 하는 일로
s = S(68)
T(s, 8, "안 A · 현행 유지 (헤지 3%)"); T(s, 10, "새로 사지 않아 갭 −34년을 그대로 진다\n부채 벤치마크가 없어 갭을 재지도 못한다")
T(s, 12, "안 B · 30년 국고채로 교체 (헤지 5%)")
T(s, 14, "부채 벤치마크를 만들어 공시한다\n시장이 받아 줄 131조까지만 30년물로 바꾼다\n현물이라 증거금이 없다")
T(s, 16, "안 C · 금리 스왑 계약 (헤지 30%)")
T(s, 18, "스왑 1,223조는 시장 상한(299조)의 4배\n사흘 +200bp 증거금 401조 > 유동자산 123조\n지침에 파생 조항이 없다")
for sid in (8, 12, 16): set_px(sh(s, sid), w=320)
for sid in (10, 14, 18): set_px(sh(s, sid), w=318); font(sh(s, sid), 15)

# ============================================================ 새 장 만들기 ====
# (가) 새 2장 — 아홉 단원 목차(4장 표 장표 복제)
UNITS = [  # (번호, 머리 띠 이름, 디바이더 제목, 이 단원이 푸는 질문, 다루는 것)
    ("①", "부채와 적립비율", "부채와 적립비율 — 회계가 아니라 관점",
     "미래의 약속은 오늘 얼마이고, 기금의 성적표는 무엇인가?", "현재가치 · 적립비율 FR = A/L · 잉여금 수익률", "현재가치 · FR = A/L · 잉여금 수익률"),
    ("②", "할인율 — 같은 약속, 다른 숫자", "할인율 — 같은 약속, 다른 숫자",
     "같은 약속이 왜 할인율(자)에 따라 다른 숫자가 되는가?", "할인율의 정치학 · CalPERS · 공무원연금 · 영국 DB의 세 자", "할인율의 정치학 · 영국 DB의 세 자"),
    ("③", "금리 위험을 재는 자", "금리 위험을 재는 자 — 듀레이션 · DV01 · 볼록성",
     "금리가 움직이면 자산과 부채는 얼마나, 어떻게 움직이는가?", "듀레이션 갭 · DV01 · 볼록성 · 헤지비율", "듀레이션 갭 · DV01 · 볼록성 · 헤지비율"),
    ("④", "면역화와 Redington", "면역화와 Redington — 수학적 방패",
     "금리가 어느 쪽으로 움직여도 잉여금을 지키려면 무엇을 맞추는가?", "Redington 세 조건 · 바벨 검증 · TDF · 면역화의 한계", "Redington 세 조건 · 바벨 · 한계"),
    ("⑤", "잉여금 최적화와 두 블록", "잉여금 최적화 — 헤지 블록과 수익 추구 블록",
     "헤지 몫과 수익 추구 몫은 각각 무엇이 정하는가?", "목적식 · 해(PSP + LHP) · 실무적 함의 · 적립비율과 헤지 비중", "목적식 · 해(PSP + LHP) · 함의"),
    ("⑥", "실무 헤지 수단", "실무 헤지 수단 — 현물 · 스왑 · 버퍼",
     "부채 DV01을 무엇으로 채우고, 그 대가는 무엇인가?", "부채 벤치마크(LBP) · IRS 오버레이 · LDI 산업 · 변동증거금", "LBP · IRS 오버레이 · 증거금"),
    ("⑦", "실패 사례 — 2022 영국 · 2023 SVB", "2022 영국 · 2023 SVB — 이론의 실전 시험",
     "평상시 작동하던 헤지가 왜 위기에는 현금 시험으로 바뀌는가?", "레버리지 LDI · 타임라인 · Doom Loop · BOE 개입 · SVB", "Doom Loop · BOE 개입 · SVB"),
    ("⑧", "실습 · LDI 미니 시뮬레이터", "손으로 하지 말고 Claude Code에게 시켜라",
     "IC 발언의 숫자 근거를 어떻게 20분 만에 만드는가?", "메타프롬프트 · LDI 미니 시뮬레이터 · 버퍼 100 vs 250bp", "LDI 미니 시뮬레이터 · 버퍼 비교"),
    ("⑨", "모의 IC", "모의 투자위원회 — 국민연금 LDI 도입 심의",
     "헤지 비율을 몇 %로, 무엇으로, 얼마의 버퍼로 올리는가?", "안 A · B · C 심의 · 판정 조건 ①~⑤ · 과제", "안 A · B · C · 판정 조건 ①~⑤"),
]
N2 = clone(S(4))
T(N2, 3, "아홉 단원은 오후 IC의 헤지 비율 한 숫자로 모입니다")
T(N2, 5, "단원마다 푸는 질문이 하나씩 있다 — 1교시 ①②, 2교시 ③~⑥, 3교시 ⑦, 4교시 ⑧⑨")
tb = lecfix.find_table(N2)
lecfix.set_table(tb, [["단원", "이 단원이 푸는 질문", "다루는 것"]] +
                 [[f"{u[0]} {u[1]}", u[3], u[5]] for u in UNITS])
for ci, wpx in enumerate((322, 490, 320)): tb.table.columns[ci].width = Emu(wpx * EMU)
for r in tb.table.rows:
    for c in r.cells:
        for p in c.text_frame.paragraphs:
            for rr in p.runs: rr.font.size = Pt(13)
T(N2, 8, "답은 “LDI를 하느냐”가 아니라 “몇 %를, 무엇으로, 얼마의 버퍼로”입니다"); set_px(sh(N2, 8), y=598)

# (나) 새 9장 — 안 A·B·C를 동사로(68장 카드 장표 복제, 가운데 카드는 흰 카드로)
N9 = clone(S(68))
by = {x.shape_id: x for x in N9.shapes}
for sid in (11, 12, 13, 14): remove(by[sid])
for sid in (7, 8, 9, 10): copy_shape(by[sid], N9, dx=390)
T(N9, 3, "세 시간 뒤 여러분은 국민연금 LDI 안건에 표결합니다")
T(N9, 5, "제1호 — 부채 벤치마크(LBP)를 만들고, 헤지 비율을 무엇으로 얼마나 올리는가")
cards = [s_ for s_ in N9.shapes if s_.has_text_frame and s_.text_frame.text.strip()]
def card_text(x_lo, x_hi, y):
    return next(s_ for s_ in N9.shapes if s_.has_text_frame and x_lo <= lecfix.px(s_)[0] < x_hi and abs(lecfix.px(s_)[1] - y) < 6)
SPEC = [(60, 440, "안 A · 현행 유지 · 헤지 3%", "새로 사지 않는다",
         "보유 국내채권을 그대로 둔다\n부채 벤치마크도 만들지 않는다\n→ 헤지 3% · 갭 −34년이 그대로"),
        (440, 840, "안 B · 현물 연장 · 헤지 5%", "30년 국고채로 갈아탄다",
         "보유 채권 131조를 30년 국고채로 바꿔 산다 (5년간)\n부채 벤치마크를 만들어 공시한다\n→ 헤지 3 → 5% · 증거금 없음"),
        (840, 1200, "안 C · IRS 오버레이 · 헤지 30%", "금리 스왑을 계약한다",
         "30년 고정금리를 받고 변동금리를 내는 스왑을 명목 1,223조 맺는다\n부채 벤치마크를 만들어 공시한다\n→ 헤지 3 → 30% · 사흘 +200bp에 증거금 401조")]
for lo, hi, lab, val, body in SPEC:
    a, b, c = card_text(lo, hi, 250), card_text(lo, hi, 290), card_text(lo, hi, 344)
    lecfix.set_text(a, lab); set_px(a, w=320)
    lecfix.set_text(b, val); set_px(b, w=320)
    lecfix.set_text(c, body); set_px(c, w=318, h=174); font(c, 14)
foot = sh(N9, 19); set_px(foot, y=576)
cond = copy_shape(sh(S(9), 34), N9); set_px(cond, y=533)
lecfix.set_text(cond, "판정 조건 ① 갭 실재 · ② 시장 수용 · ③ 버퍼 · ④ 허들 5.5% · ⑤ 지침")
lecfix.set_text(foot, "발언은 수익률 언어가 아니라 부채 언어(FR · 갭 · 잉여금)로 합니다")
for x in N9.shapes:           # 카드 높이를 글 높이에 맞춘다
    if not (x.has_text_frame and x.text_frame.text.strip()) and abs(lecfix.px(x)[1] - 223) <= 2 and lecfix.px(x)[2] > 300: set_px(x, h=298)

# (다) 단원 ⑤ — 목적식(35장 카드 장표 복제)
N38a = clone(S(35))
T(N38a, 3, "잉여금 최적화의 목적은 잉여금이고, 듀레이션은 제약이 아닙니다")
T(N38a, 5, "Sharpe–Tint(1990) 잉여금 평균-분산 — 단원 ④의 Redington은 ‘맞출 조건’, 여기서는 ‘최적화의 답’으로 같은 헤지에 도달한다")
pic = next(x for x in N38a.shapes if x.shape_type == 13); box = container(N38a, pic)
x0, y0, w0, h0 = lecfix.px(pic); w, h = size_px(E_OBJ, maxw=1060, maxh=96)
ytop = 236; add_pic(N38a, E_OBJ, 640 - w / 2, ytop, w, h, before=pic._element); remove(pic)
pad = 14
set_px(box, x=640 - w / 2 - pad, y=ytop - 8, w=w + 2 * pad, h=h + 16)
T(N38a, 10, "목적 — 무엇을 키우나"); T(N38a, 11, "잉여금 수익률 R_S"); T(N38a, 12, "평균은 크게, 분산에는 벌점 γ/2")
T(N38a, 14, "선택 변수 — 무엇을 고르나"); T(N38a, 15, "자산 비중 w"); T(N38a, 16, "주식 · 장기채 · 현금")
T(N38a, 18, "제약 — 무엇이 묶이나"); T(N38a, 19, "비중의 합 = 1뿐"); T(N38a, 20, "듀레이션 매칭은 제약이 아니라 해")
T(N38a, 21, "부채 L이 R_S 안에 들어 있어 최적화가 스스로 부채를 따라가는 자산을 찾습니다 — 해는 다음 장")
for sid in (10, 14, 18): set_px(sh(N38a, sid), w=300)
for sid in (12, 16, 20): set_px(sh(N38a, sid), w=300)

# (라) 단원 ⑤ — 해(30장 수식+표 장표 복제)
N38b = clone(S(30))
T(N38b, 3, "최적해는 수익 추구(PSP)와 헤지(LHP) 두 블록으로 갈라집니다")
T(N38b, 5, "한 요인 모형 — 주식 기대초과수익 4% · 변동성 15% · 위험회피 γ = 3.6, 30년 국채 D_H = 20(초과수익 0), A = 100 · L = 80")
pic = next(x for x in N38b.shapes if x.shape_type == 13); box = container(N38b, pic)
w1, h1 = size_px(E_SOL, maxh=78); w2, h2 = size_px(E_LHP, maxh=60)
gap = 46; tot = w1 + gap + w2; xl = 640 - tot / 2; ytop = 268
add_pic(N38b, E_SOL, xl, ytop, w1, h1, before=pic._element)
add_pic(N38b, E_LHP, xl + w1 + gap, ytop + (h1 - h2) / 2, w2, h2, before=pic._element); remove(pic)
lab = sh(N38b, 8); lecfix.set_text(lab, "해 — 최적 비중 = 투기적 수요(PSP) + 헤지 수요(LHP)  ·  한 요인 모형이면 헤지 몫은 DV01 매칭")
set_px(lab, x=xl, w=tot)
set_px(box, x=xl - 14, y=224, w=tot + 28, h=ytop + h1 + 10 - 224)
tb = lecfix.find_table(N38b)
lecfix.set_table(tb, [["몫", "계산", "자산 대비"],
                      ["주식 (PSP · 수익 추구)", "4% ÷ (3.6 × 0.15²)", "49% ≈ 50%"],
                      ["30년 국채 (LHP · 헤지)", "(1 ÷ 1.25) × (20 ÷ 20)", "80%"],
                      ["합계", "모자라는 30%는 고정금리 수취 스왑", "130%"]])
set_px(tb, y=ytop + h1 + 30)
for ci, wpx in enumerate((330, 520, 282)): tb.table.columns[ci].width = Emu(wpx * EMU)
c0, c1 = tb.table.cell(0, 0), tb.table.cell(0, 1)       # 원래 30장 머리 첫 칸은 비어 있어 글자색이 없다 → 옆 칸 서식으로
r0 = c0.text_frame.paragraphs[0].runs[0]; r1 = c1.text_frame.paragraphs[0].runs[0]
r0._r.remove(r0._r.get_or_add_rPr()); r0._r.insert(0, copy.deepcopy(r1._r.get_or_add_rPr()))
tcPr0 = c0._tc.get_or_add_tcPr(); tcPr1 = c1._tc.get_or_add_tcPr()
c0._tc.remove(tcPr0); c0._tc.append(copy.deepcopy(tcPr1))
T(N38b, 11, "헤지 몫 80% × D_H 20 = 부채 80 × D_L 20 — 듀레이션 매칭은 제약이 아니라 이 해의 결과입니다")

# (마) 단원 ⑤ — 실무적 함의(41장 체크포인트 장표 복제)
N38c = clone(S(41))
T(N38c, 3, "헤지 몫은 부채가 정하고, 남는 위험 예산만 수익 추구에 씁니다")
T(N38c, 5, "해의 실무적 함의 — 헤지 비율을 감이 아니라 규칙으로 정한다")
T(N38c, 9, "헤지 몫은 위험회피도와 무관하게 부채가 정한다")
T(N38c, 10, "LHP = (1/FR) × (D_L/D_H)에는 μ도 γ도 없다. 공격적인 기금도 보수적인 기금도 헤지 몫은 같고, 성향을 이유로 헤지를 깎으면 그것은 금리 베팅이다.")
T(N38c, 13, "남는 위험 예산만 수익 추구에 쓴다")
T(N38c, 14, "보상 없는 금리 위험을 지우면 그 예산을 주식에 쓸 수 있다. 주식 50%일 때 잉여금 변동성은 헤지 19%에서 12.8%, 헤지 100%에서 7.5%.")
T(N38c, 17, "합이 100%를 넘으면 스왑 — 증거금과 버퍼도 해의 일부")
T(N38c, 18, "PSP 50 + LHP 80 = 130%. 모자란 30%를 고정금리 수취 스왑으로 채우는 순간 증거금이 생긴다 — 단원 ⑥의 수단, 단원 ⑦의 2022 영국.")
T(N38c, 19, "운용팀은 두 블록을 만들고, 이사회는 헤지 비율 하나를 정합니다")

# (바) 디바이더 — 원래 다섯 장 + 복제 네 장, 질문 한 줄과 다루는 것 한 줄
DIV = {}
for k, n in (("①", 10), ("③", 25), ("⑦", 42), ("⑧", 59), ("⑨", 65)): DIV[k] = S(n)
for k in ("②", "④", "⑤", "⑥"): DIV[k] = clone(S(10))
for i, u in enumerate(UNITS, 1):
    d = DIV[u[0]]
    T(d, 2, f"{i:02d}"); T(d, 3, u[2])
    q = sh(d, 5); lecfix.set_text(q, f"이 단원이 푸는 질문 — {u[3]}")
    for p in q.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb = __import__("pptx").dml.color.RGBColor(0xFF, 0xFF, 0xFF)
    tp = copy_shape(q, d, dy=46); lecfix.set_text(tp, u[4])
    for p in tp.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb = __import__("pptx").dml.color.RGBColor(0x9F, 0xB0, 0xB2); r.font.size = Pt(16)

# ============================================================ 새 순서 ====
NEW = "new"
ORDER = [  # (단원, 장표)
    ("도입", S(1)), ("도입", N2), ("도입", S(3)), ("도입", S(4)), ("도입", S(5)), ("도입", S(6)), ("도입", S(7)), ("도입", S(8)), ("도입", N9),
    ("div", DIV["①"]), *[("①", S(n)) for n in (11, 12, 13, 14, 21, 22, 23)],
    ("div", DIV["②"]), *[("②", S(n)) for n in (15, 16, 18, 19, 20, 53, 54, 55, 17, 24)],
    ("div", DIV["③"]), *[("③", S(n)) for n in (26, 27, 28, 29, 30, 34, 31)],
    ("div", DIV["④"]), *[("④", S(n)) for n in (32, 35, 33, 36, 56, 57, 37)],
    ("div", DIV["⑤"]), ("⑤", N38a), ("⑤", N38b), ("⑤", N38c), ("⑤", S(39)),
    ("div", DIV["⑥"]), *[("⑥", S(n)) for n in (40, 52, 47, 41)],
    ("div", DIV["⑦"]), *[("⑦", S(n)) for n in (43, 44, 45, 46, 48, 49, 50, 51, 58)],
    ("div", DIV["⑧"]), *[("⑧", S(n)) for n in (60, 61, 62, 63, 64)],
    ("div", DIV["⑨"]), *[("⑨", S(n)) for n in (66, 67, 68, 69)],
    ("div", S(70)), *[("부록", S(n)) for n in (71, 72, 73, 74, 75, 76)], ("끝", S(77)),
]
DROP = [S(2), S(9), S(38)]
ids = P.slides._sldIdLst
by_el = {}
for el in list(ids):
    by_el[id(P.part.related_part(el.rId).slide._element)] = el
for s_ in DROP:
    el = by_el[id(s_._element)]; P.part.drop_rel(el.rId); ids.remove(el)
for el in list(ids): ids.remove(el)
for _, s_ in ORDER: ids.append(by_el[id(s_._element)])
assert len(ids) == len(ORDER) == len({id(s_._element) for _, s_ in ORDER})

# 머리 띠 · 쪽번호
UN = {u[0]: u[1] for u in UNITS}
for i, (u, s_) in enumerate(ORDER, 1):
    for x in s_.shapes:
        if not x.has_text_frame: continue
        t = x.text_frame.text.strip()
        if t.startswith("BAF.60080 · W06 LDI ·"):
            tag = "도입" if u == "도입" else "부록" if u == "부록" else f"{u} {UN[u]}"
            lecfix.set_text(x, f"BAF.60080 · W06 LDI · {tag}")
        xx, yy, ww, hh = lecfix.px(x)
        if xx >= 1000 and yy >= 640 and t.isdigit(): lecfix.set_text(x, f"{i:02d}")

left = [t for _, t in lecfix.grep(P, "교시") if not t.startswith("단원마다")]
assert not left, left
P.save(SRC)
# 장 순서 표(원래 번호)
orig_no = {id(s_._element): n for n, s_ in enumerate(O, 1)}
rows = []
for i, (u, s_) in enumerate(ORDER, 1):
    title = ""
    for x in s_.shapes:
        if x.has_text_frame and lecfix.px(x)[1] in range(60, 75) and lecfix.px(x)[2] > 1000: title = x.text_frame.text.strip(); break
    if u == "div" or not title:
        title = " / ".join(x.text_frame.text.strip() for x in s_.shapes if x.has_text_frame and x.text_frame.text.strip())[:60]
    rows.append({"new": i, "unit": u, "title": title, "orig": orig_no.get(id(s_._element), "신설")})
json.dump(rows, open(os.path.join(HERE, "order.json"), "w"), ensure_ascii=False, indent=1)
print("saved", SRC, len(ORDER), "장")
