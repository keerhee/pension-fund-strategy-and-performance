# -*- coding: utf-8 -*-
"""W06 M8 GBI 강의본 보강·재구성(2026-10-05) — 사용자 피드백 7건.
 1) H대 케이스 장마다 '결정할 것 = 세 숫자 x · y · z를 정한다' + GBI 대응표(Floor · GHP · 위험 예산 · Flexicure식 재설정)
 2) 심리계정: Thaler · Shefrin–Statman BPT 옆에 Das–Markowitz–Scheid–Statman(2010) — 문제와 해 · 합산 효율 · 비교표 · 실무 장점
 3) 옛 23장: 문제를 말로 · 결정 변수 · 목적식과 제약(LaTeX) · 기호표 · '자금을 최대로 남기나'에 대한 답 → 두 장으로
 4) 옛 33장: 상황 → 질문 → 답의 문장으로
 5) 표 안의 텍스트 수식(옛 45 · 47 · 74장)을 수식 PNG로(셀을 비우고 행 색을 구운 그림을 얹는다)
 6) '모방 3장벽' 정의와 숫자 조건 · 앞쪽 '먼저 읽는 용어' · H대 ↔ GBI 기법 한 장
 8) [2차 2026-10-05] Das–Markowitz를 단원 ②는 소개만, 단원 ③의 방법 2(손계산 · 방법 1과 비교)로 옮김
 9) [2차] H대: 강의 대응표 · Flexicure를 비유동성까지 옮기는 다섯 단계(과제 4팀 B안 숫자) · 기본 답에 숫자 조건
 10) [2차 추가] 부록 B-6 — Das–Markowitz 원 논문 예(보충교재 05 4-2절 재계산) 세 장, 단원 ③에서 상호 참조
 11) [3차] 단원 ③ 두 방법의 손계산을 한 장씩 더(방법 1 목표 하나 다섯 단계 · 방법 2 효율선 상수와 γ 방정식)
 12) [4차] Market 필요 자본을 공식값(w* 67% 혼합, g 2.90%, K 2.244억)으로 통일 — 합 9.13억 · 남는 0.87억 · 49/22/28%, 균형형 2.22억은 후보 비교·각주로만
 13) [5차] 단원 ⑤에 VaR로 floor · m 정하기 + 단원 ③ 두 방법과 같은 “평균 − z × 변동성” 원리(보충교재 3-2·3-5·3-6·9-3·9-5)
 14) [6차] VaR 두 장을 레버 ① m · 레버 ② 위험 자본(floor) · 세 방법은 같은 문장 · 같은 숫자 74 네 장으로(표기: 기호 = 값, 결과 옆에 식)
 15) [8차] 방법 1 문헌 계보(Kataoka · Telser · Roy · Chhabra · Merton) · 단원 ⑦에 한국 퇴직연금 · 디폴트옵션 · TDF 세 장 · 부록 C-2 한국 제도 용어 · 부록 D 참고문헌
 16) [9차] 58장을 VaR 연동 floor(위험 자본 방식 CPPI — 위험 ↑ → floor ↑)로 바로잡음 · 디지털 TDF 정의 · IRP 장 신설 · 한국 제도 장을 로보 장 앞으로
 17) [10차] Das–Markowitz를 손계산(한 포트폴리오 검사 → 두 자산 축소) → 행렬 일반화 순서로 · 롱온리 영향(은퇴 · 교육 0, 합산 연 12bp)
 18) [11차] 원문(DMSS 2010, JFQA 45(2) 311–334) 대조 감사 반영 — 합산 효율은 공매도 허용 시 · 12bp는 계정별 롱온리 · A·B·C·D 표기 · 19.89% 오기 · 표 3 · Telser 1956 · 인용 범위
 19) [12차] 롱온리를 엑셀 해 찾기 · scipy로 푸는 장(solver_fig.py) · 바닥은 CPPI, 그 위 위험 블록은 Das–Markowitz로 — 두 층 구조 장
 20) [13차] 방법 3 몬테카를로 네 장(원리 · 절차 · 결과 · Claude Code) — 숫자는 W06_LDI와GBI/W06_M8_실습_몬테카를로/gbi_mc.py 실행 결과(gbi_mc_results.json) · 비교표 세 열 · 실무 장마다 방법 1 · 2 · 3 연결
 21) [14차] 방법 3을 초보자용으로 — 규칙을 경로에 적용(3년 표) · 10경로 장난감 K 찾기 · 오해 바로잡기 · 여러 목표 격자(gbi_mc.py --demo)
 22) [15차] 목표가 여럿일 때 조합 폭발과 줄이는 법 두 장(손으로 세기 · 다섯 방법 표 · 우선순위 순서 실제 풀이 · 공통 난수 효과)
 23) [16차] 47장 재현 실습 장 — 프롬프트 요약 · priority_mc.py 실제 출력(코드 그림) · 검증 넷
 24) [17차] 변수 묶기 장 — 고를 숫자의 개수를 줄인다(Das–Markowitz γ · CPPI m)
 25) [18차] 몬테카를로로 방법 2 문제를 푼다 두 장(dm_mc.py — 격자 · γ 한 축 · 롱온리 · t5 꼬리)
 28) [21차] 티일 설명 블록 — 한 줄 한 생각 · 머리 표시(① • ▶ ※) · 굵은 머리말 · spcPts 줄 간격(38 · 45 · 47 · 49 · 53 · 55 · 56 · 61 · 81장), 격자 읽는 법 한 장 신설
 27) [20차] 가독성 — 붙여 쓴 문장을 카드 · 단계로(46 · 48 · 51 · 52 · 55 · 80장 등), Colab 실습 폴더 안내
 26) [19차] 수강생 혼란 정리 — 한 계좌 격자 읽는 법 · 변수 묶기 두 장(g · v 계산) · MC 결과표 열 이름 · 실습 A · B · 도구 구분 · 단원 ③ 지도 · 머리 띠 소주제
 7) 교시 대신 소주제 열 단원 · 단원 디바이더(이 단원이 푸는 질문) · 목차 표 · 머리 띠 · 상호 참조 · 쪽번호
실행: <PPTX 가능 파이썬> _build/w06_m8_rework/rework_m8.py
원본은 _구판/W06_M8_보강_구판_2026-10-05/에 한 번만 백업하고 항상 백업에서 읽는다(재실행해도 결과가 같다).
수식은 matplotlib mathtext(cm)로 카드 색 위에 잉크(#102027)로 굽는다 — 덱의 기존 수식과 같은 배율(3.12 px/pt-px)."""
import copy, os, shutil, sys, json, re
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "lecfix"))
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from PIL import Image
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
from matplotlib import font_manager
import lecfix
from solver_fig import solver_svg, layers_svg, png as svg_png

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", ".."))
NAME = "W06_M8_GBI_목표기반투자_강의본.pptx"
SRC = os.path.join(R, "W06_LDI와GBI", NAME)
BK = os.path.join(R, "_구판", "W06_M8_보강_구판_2026-10-05"); os.makedirs(BK, exist_ok=True)
bk = os.path.join(BK, NAME)
if not os.path.exists(bk):
    shutil.copy2(SRC, bk)
    pdf = SRC[:-5] + ".pdf"
    if os.path.exists(pdf): shutil.copy2(pdf, os.path.join(BK, os.path.basename(pdf)))
EQ = os.path.join(HERE, "eq"); os.makedirs(EQ, exist_ok=True)
EMU = 9525          # 1 px(96dpi) = 9525 EMU
SCALE = 3.12        # 덱 수식 PNG: 이미지 px ÷ 장표 px
FONTDIR = os.path.expanduser("~/Library/Fonts")
for f in ("Pretendard-Regular.otf", "Pretendard-SemiBold.otf", "Pretendard-Bold.otf"):
    p = os.path.join(FONTDIR, f)
    if os.path.exists(p): font_manager.fontManager.addfont(p)

P = Presentation(bk)
O = list(P.slides)                    # 원래 번호 n → O[n-1]
def S(n): return O[n - 1]
def sh(slide, sid): return next(x for x in slide.shapes if x.shape_id == sid)
def T(slide, sid, text): lecfix.set_text(sh(slide, sid), text)

# ------------------------------------------------------------------ 도구 ----
def _korean(s): return any(0xAC00 <= ord(c) <= 0xD7A3 for c in s)

def render(tex, name, fontsize=30, color="#102027", bg="#F7F5F0"):
    assert not _korean(tex), f"수식 안 한글: {name}"
    assert r"\;" not in tex, f"\\; 금지: {name}"
    path = os.path.join(EQ, name)
    with plt.rc_context({"mathtext.fontset": "cm"}):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=fontsize, color=color)
        fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.12, facecolor=bg)
        plt.close(fig)
    Image.open(path).convert("RGB").save(path)
    return path

def render_mixed(s, name, fontsize=21, color="#102027", bg="#FFFFFF", weight="regular"):
    """표 칸용 — 글(Pretendard)과 $수식$(cm)을 한 줄에. 한글은 $ 밖에만."""
    for i, part in enumerate(s.split("$")):
        if i % 2: assert not _korean(part), f"수식 안 한글: {name}"
    assert r"\;" not in s
    path = os.path.join(EQ, name)
    with plt.rc_context({"mathtext.fontset": "cm", "font.family": "Pretendard"}):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, s, fontsize=fontsize, color=color, fontweight=weight)
        fig.savefig(path, dpi=220, bbox_inches="tight", pad_inches=0.05, facecolor=bg)
        plt.close(fig)
    Image.open(path).convert("RGB").save(path)
    return path

def code_png(text, name, fontsize=13, bg="#071A1D", fg="#FFFFFF", hi="#B7F34A"):
    """실행 출력 그림 — 고정폭 글꼴, 어두운 배경(수식이 아니라 출력이라 mathtext를 끈다)."""
    path = os.path.join(EQ, name)
    lines = text.split("\n")
    with plt.rc_context({"font.family": ["Menlo", "Pretendard"], "text.parse_math": False}):
        fig = plt.figure(figsize=(9.6, 0.25 * len(lines) + 0.3)); fig.patch.set_facecolor(bg)
        for i, ln in enumerate(lines):
            fig.text(0.02, 1 - (i + 0.9) / (len(lines) + 0.6), ln, fontsize=fontsize, color=hi if ln.startswith("$") or ln.startswith("합계") else fg,
                     family=["Menlo", "Pretendard"])
        fig.savefig(path, dpi=220, facecolor=bg); plt.close(fig)
    return path

def size_px(path, maxw=None, maxh=None):
    w, h = Image.open(path).size
    w, h = w / SCALE, h / SCALE
    s = min(1.0, (maxw / w) if maxw else 1.0, (maxh / h) if maxh else 1.0)
    return w * s, h * s

def add_pic(slide, path, x, y, w, h, before=None):
    pic = slide.shapes.add_picture(path, Emu(int(x * EMU)), Emu(int(y * EMU)), Emu(int(w * EMU)), Emu(int(h * EMU)))
    if before is not None:
        before.addprevious(pic._element)
    return pic

def remove(shape): shape._element.getparent().remove(shape._element)

def set_px(shape, x=None, y=None, w=None, h=None):
    if x is not None: shape.left = Emu(int(x * EMU))
    if y is not None: shape.top = Emu(int(y * EMU))
    if w is not None: shape.width = Emu(int(w * EMU))
    if h is not None: shape.height = Emu(int(h * EMU))

def container(slide, pic):
    """그림을 감싸는 글자 없는 카드(사각형)를 찾는다."""
    px_, py, pw, ph = lecfix.px(pic); cx, cy = px_ + pw / 2, py + ph / 2
    best = None
    for s in slide.shapes:
        if s is pic or s.shape_type == 13 or getattr(s, "has_table", False) and s.has_table: continue
        if s.has_text_frame and s.text_frame.text.strip(): continue
        x, y, w, h = lecfix.px(s)
        if w > 0 and h > 0 and x <= cx <= x + w and y <= cy <= y + h:
            if best is None or w * h < lecfix.px(best)[2] * lecfix.px(best)[3]: best = s
    return best

def swap_eq(slide, sid, path, maxw=1000, maxh=None, cx=None, pad=None):
    """수식 그림을 바꾼다 — 같은 세로 중심, 가로는 cx(기본: 옛 그림 중심)에 맞추고 카드 폭을 그림에 맞춘다."""
    pic = sh(slide, sid); box = container(slide, pic)
    x0, y0, w0, h0 = lecfix.px(pic)
    w, h = size_px(path, maxw=maxw, maxh=maxh)
    cx = cx if cx is not None else x0 + w0 / 2
    add_pic(slide, path, cx - w / 2, y0 + (h0 - h) / 2, w, h, before=pic._element); remove(pic)
    if box is not None:
        bx, by, bw, bh = lecfix.px(box)
        pad = pad if pad is not None else max(20, x0 - bx)
        nw = max(w + 2 * pad, 0)
        set_px(box, x=cx - nw / 2, w=nw)
        for s in slide.shapes:   # 카드 안의 머리 글(라벨)도 카드 왼쪽에 맞춘다
            if s.has_text_frame and s.text_frame.text.strip():
                sx, sy, sw, shh = lecfix.px(s)
                if by <= sy <= by + 40 and bx - 5 <= sx <= bx + 60:
                    set_px(s, x=cx - nw / 2 + (sx - bx), w=nw - 2 * (sx - bx))
    return box

def fresh_ids(slide, elements):
    mx = max(int(e.get("id")) for e in slide._element.iter() if e.tag.endswith("}cNvPr"))
    for el in elements:
        for e in el.iter():
            if e.tag.endswith("}cNvPr"):
                mx += 1; e.set("id", str(mx))

def clone(src):
    """같은 덱 안에서 장표를 복제(그림 관계 포함, 배경 포함). 맨 뒤에 붙는다."""
    dst = P.slides.add_slide(src.slide_layout)
    for x in list(dst.shapes): remove(x)
    rmap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype.endswith("/image"):
            rmap[rId] = dst.part.relate_to(rel._target, rel.reltype)
    cs, cd = src._element.cSld, dst._element.cSld
    if cs.bg is not None:
        cd.insert(0, copy.deepcopy(cs.bg))
    for el in src.shapes._spTree.iterchildren():
        if el.tag.endswith(("}nvGrpSpPr", "}grpSpPr")): continue
        e2 = copy.deepcopy(el)
        for e in e2.iter():
            for k, v in list(e.attrib.items()):
                if k.endswith("}embed") and v in rmap: e.set(k, rmap[v])
        dst.shapes._spTree.append(e2)
    return dst

def copy_shape(src_shape, dst_slide, dx=0, dy=0):
    e = copy.deepcopy(src_shape._element)
    dst_slide.shapes._spTree.append(e); fresh_ids(dst_slide, [e])
    s = next(x for x in dst_slide.shapes if x._element is e)
    s.left = Emu(s.left + int(dx * EMU)); s.top = Emu(s.top + int(dy * EMU))
    return s

def font(shape, pt):
    for p in shape.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(pt)

def table_font(tb, pt, head_pt=None):
    for ri, r in enumerate(tb.table.rows):
        for c in r.cells:
            for p in c.text_frame.paragraphs:
                for rr in p.runs: rr.font.size = Pt(head_pt if (ri == 0 and head_pt) else pt)

def col_widths(tb, widths):
    for ci, wpx in enumerate(widths): tb.table.columns[ci].width = Emu(int(wpx * EMU))
    tb.width = Emu(int(sum(widths) * EMU))

def row_heights(tb, hs):
    for r, hpx in zip(tb.table.rows, hs): r.height = Emu(int(hpx * EMU))
    tb.height = Emu(int(sum(hs) * EMU))

def drop_col(tb, ci):
    tbl = tb.table._tbl
    grid = tbl.tblGrid; grid.remove(grid.gridCol_lst[ci])
    for tr in tbl.tr_lst: tr.remove(tr.tc_lst[ci])

def stripe(tb):
    """본문 행을 흰색 · 미색 번갈아, 글자는 잉크로(복제 원본의 강조 행을 지운다)."""
    t = tb.table
    tpl = {k: [copy.deepcopy(t.cell(k, ci)._tc.tcPr) for ci in range(len(t.columns))] for k in (1, 2)}
    for ri in range(1, len(t.rows)):
        for ci in range(len(t.columns)):
            c = t.cell(ri, ci); c._tc.remove(c._tc.tcPr)
            c._tc.append(copy.deepcopy(tpl[1 if ri % 2 else 2][ci]))
            for p in c.text_frame.paragraphs:
                for r in p.runs:
                    r.font.color.rgb = RGBColor(0x10, 0x20, 0x27); r.font.bold = (ci == 0)

def rep(old, new, n=1):
    k = lecfix.replace_all(P, old, new)
    assert k == n, f"치환 {k}회(기대 {n}): {old}"

def white(shape):
    for p in shape.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def cell_box(tb, ri, ci):
    x, y, _, _ = lecfix.px(tb)
    cx = x + sum(tb.table.columns[i].width for i in range(ci)) / EMU
    cy = y + sum(tb.table.rows[i].height for i in range(ri)) / EMU
    return cx, cy, tb.table.columns[ci].width / EMU, tb.table.rows[ri].height / EMU

def cell_fill(cell):
    f = cell._tc.tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
    return "#" + f[0].get("val")

def eq_in_cell(slide, tb, ri, ci, s, name, fontsize=21, maxh=None):
    """표 칸을 비우고 행 색을 구운 수식 그림을 칸 왼쪽(여백 10px)에 세로 가운데로 얹는다."""
    cell = tb.table.cell(ri, ci)
    color = "#0B6B68" if cell_fill(cell) == "#EAF7D6" else "#102027"
    path = render_mixed(s, name, fontsize=fontsize, bg=cell_fill(cell), color=color)
    lecfix.set_tf(cell.text_frame, " ")
    x, y, w, h = cell_box(tb, ri, ci)
    pw, ph = size_px(path, maxw=w - 16, maxh=maxh or h - 6)
    return add_pic(slide, path, x + 8, y + (h - ph) / 2, pw, ph)

def line(slide, y, text, pt=15, src=None):
    l = copy_shape(src or sh(S(19), 14), slide); set_px(l, x=79, y=y, w=1132, h=30); lecfix.set_text(l, text); font(l, pt); return l
def _nobullet(tf):
    from pptx.oxml.ns import qn
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        for tag in ("a:buChar", "a:buAutoNum", "a:buNone", "a:buFont", "a:buSzPct"):
            for e in pPr.findall(qn(tag)): pPr.remove(e)
        pPr.set("marL", "0"); pPr.set("indent", "0")
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
def rich(shape):
    """문단 안의 **굵게** 표시를 굵은 run으로 나눈다."""
    from pptx.oxml.ns import qn
    for p in shape.text_frame.paragraphs:
        full = "".join(r.text for r in p.runs)
        if "**" not in full: continue
        r0 = p.runs[0]
        for r in p.runs[1:]: r._r.getparent().remove(r._r)
        parts = full.split("**"); r0.text = parts[0]; prev = r0._r
        for k, part in enumerate(parts[1:], 1):
            if not part: continue
            nr = copy.deepcopy(r0._r); prev.addnext(nr); prev = nr
            nr.find(qn("a:t")).text = part
            rp = nr.find(qn("a:rPr"))
            if rp is None: rp = nr.makeelement(qn("a:rPr"), {}); nr.insert(0, rp)
            rp.set("b", "1" if k % 2 else "0")
    return shape
def card(slide, x, y, w, h, head, lines, dark=False, pt=16, bullets=False, head_pt=15):
    """20차 — 카드(흰 · 잉크) 한 장: 머리 한 줄 + 본문 줄들(한 줄에 하나, **굵게** 지원). 원본은 74장(옛 36장)의 카드."""
    ids = (15, 16, 18) if dark else (7, 8, 10)
    r = copy_shape(sh(S(36), ids[0]), slide); set_px(r, x=x, y=y, w=w, h=h)
    hd = copy_shape(sh(S(36), ids[1]), slide); set_px(hd, x=x + 26, y=y + 12, w=w - 52, h=30)
    lecfix.set_text(hd, head); font(hd, head_pt)
    for p in hd.text_frame.paragraphs:
        for rr in p.runs: rr.font.bold = True
    b = copy_shape(sh(S(36), ids[2]), slide); set_px(b, x=x + 26, y=y + 48, w=w - 52, h=h - 56)
    lecfix.set_text(b, "\n".join(lines)); font(b, pt)
    for k, p in enumerate(b.text_frame.paragraphs):
        p.line_spacing = Pt(round(pt * 1.3)); p.space_before = Pt(0 if k == 0 else 4)
    if not bullets: _nobullet(b.text_frame)
    rich(b); return r
INK, TEAL = RGBColor(0x1F, 0x2A, 0x2E), RGBColor(0x0B, 0x6B, 0x68)
def blk(slide, x, y, w, items, pt=16, h=None, before=5):
    """21차 — 설명 블록: 한 줄 = 한 생각, '**머리 표시 머리말:**'은 굵은 티일, 본문은 잉크. 줄 간격은 spcPts로 고정."""
    box = copy_shape(sh(S(19), 14), slide); set_px(box, x=x, y=y, w=w, h=h or 30 * len(items))
    lecfix.set_text(box, "\n".join(items)); font(box, pt)
    for p in box.text_frame.paragraphs:
        for rr in p.runs: rr.font.bold = False
    for k, p in enumerate(box.text_frame.paragraphs):
        p.line_spacing = Pt(round(pt * 1.3)); p.space_before = Pt(0 if k == 0 else before)
    rich(box)
    from pptx.oxml.ns import qn
    for p in box.text_frame.paragraphs:
        for rr in p.runs:
            rp = rr._r.find(qn("a:rPr")); b_ = rp is not None and rp.get("b") == "1"
            rr.font.bold = b_; rr.font.color.rgb = TEAL if b_ else INK
    return box

def note(slide, y, text):
    n = copy_shape(sh(S(21), 8), slide); set_px(n, x=79, y=y, w=1132, h=26); lecfix.set_text(n, text); font(n, 12); return n

# ------------------------------------------------------------------ 수식 ----
E_OBJ = render(r"\max_{K_j,\ w_j}\ \ W-\sum_{j}K_j\quad \mathrm{s.t.}\quad "
               r"\mathrm{Pr}\left(K_j\,e^{m_jT+\sigma_j\sqrt{T}\,Z}\geq G_j\right)\geq p_j,\qquad j=1,2,3", "obj23.png")
E_K = render(r"K_j=G_j\,e^{-T\,g_p(w_j)}\quad\Rightarrow\quad \min\ K_j\ \ \Leftrightarrow\ \ \max_{w_j}\ g_p(w_j)", "k23.png")
E_DM1 = render(r"\max_{w}\ w^{\top}\mu\quad \mathrm{s.t.}\quad \mathrm{Pr}(R_p<H)\leq\alpha,\quad \mathbf{1}^{\top}w=1", "dm_obj.png")
E_DM2 = render(r"\mathrm{Pr}(R_p<H)\leq\alpha\quad\Leftrightarrow\quad \mu_p-z_{\alpha}\,\sigma_p\geq H\qquad (z_{0.05}=1.645)", "dm_cons.png")
E_B2 = render(r"\mathrm{Pr}(R_p<H)=\Phi\left(\frac{H-m_p}{\sigma_p}\right)\leq\alpha\quad\Leftrightarrow\quad m_p+\Phi^{-1}(\alpha)\,\sigma_p\geq H", "b6_cons.png")
E_B3 = render(r"A=\mathbf{1}^{\top}\Sigma^{-1}\mu,\quad C=\mathbf{1}^{\top}\Sigma^{-1}\mathbf{1},\quad K=\mu^{\top}\Sigma^{-1}\mu-\frac{A^2}{C},\qquad m(\gamma)=\frac{A}{C}+\frac{K}{\gamma},\quad \sigma^2(\gamma)=\frac{1}{C}+\frac{K}{\gamma^2}", "b6_front.png")
E_B4 = render(r"\frac{A}{C}+\frac{K}{\gamma}+\Phi^{-1}(\alpha)\sqrt{\frac{1}{C}+\frac{K}{\gamma^2}}=H\qquad (A=22.917,\ \ C=426.04,\ \ K=0.18399)", "b6_gamma.png")
E_V1 = render(r"E\,x\leq C\ \ \Leftrightarrow\ \ m\,C\,x\leq C\ \ \Leftrightarrow\ \ m\leq\frac{1}{x}", "v1_m.png")
E_RC = render(r"E\,x=L,\quad E=m\,C\ \ \Rightarrow\ \ C=\frac{L}{m\,x},\quad F=A-\frac{L}{m\,x}\ \ (F\geq F_{\mathrm{GHP}}),\quad x=1.645\,\sigma-\mu", "rc.png")
E_DM3 = render(r"\sum_k a_k\,w(\gamma_k)=g+v\sum_k\frac{a_k}{\gamma_k},\qquad \frac{1}{\gamma_{\mathrm{total}}}=\sum_k\frac{a_k}{\gamma_k}", "dm_sum.png")

# ============================================================ 원래 장 고치기 ====
P23 = clone(S(23))                 # 새 23a는 옛 23장을 고치기 전에 복제한다
# 4장 — 메시지 ②에 Das–Markowitz
T(S(4), 13, "3계층 버킷 — 계좌마다 달성 확률을 건다")
T(S(4), 14, "심리계정(Thaler)을 공학으로 — 계좌마다 확률 한도(Das–Markowitz): Safety(≥95%)가 확보되어야 Market(≥70%), 그다음 Aspirational(≥30%).")
# 6장 — 두 안건: 결정할 것을 동사로
s = S(6)
T(s, 10, "디폴트옵션 Flexicure형 도입")
T(s, 11, "결정: Floor · m 상한 · 총보수를 정한다\nCash Trap — 쿠션 0으로 회복에 못 탐\nLTCM · 2022 TDF 반면교사")
T(s, 13, "H대학기금 — 세 숫자를 정한다")
T(s, 14, "결정: 비유동 상한 x · 유동 하한 y · 지출률 z\n모방 3장벽 — 접근권 · 시간 지평 · 인력\nGBI로 읽기 — Floor · GHP · 위험 예산")
font(sh(s, 11), 14); font(sh(s, 14), 14); set_px(sh(s, 10), h=50); set_px(sh(s, 11), y=362)
# 9장 — 카드의 텍스트 수식을 말로
T(S(9), 20, "Safety 미달 < 5%\nMarket 미달 < 30%\nAspirational 미달 < 70%"); font(sh(S(9), 20), 14)
# 10장
T(S(10), 5, "수식이 어렵다면 이 세 문장만 기억한다 — 단원 ②·③의 수식은 이 비유의 정밀한 번역이다")
# 11장 — BPT 형식
T(S(11), 11, "말로 — 계좌별 미달 확률 ≤ α")
T(S(11), 14, "α = 1 − 달성 확률")
T(S(11), 19, "합치면 “비효율”이지만 행동을 정확히 담습니다 — 효율을 되살린 것이 Das–Markowitz")
# 17장 — 지배 검정 정의
T(S(17), 14, "GHP보다 나을 때만 위험 계층")
T(S(17), 15, "지배 검정 — 25년 인출형은 GHP가 더 쌀 수 있다")
# 18장 · 19장 · 21장 · 26장 · 72장 — e^ 텍스트 수식을 exp( )로
T(S(18), 9, "exp(mT)"); T(S(18), 10, "연 m씩 T년 복리\nexp(0.2) ≈ 1.22")
t = lecfix.find_table(S(19)).table; lecfix.set_tf(t.cell(2, 0).text_frame, "p · z")
T(S(19), 14, "지수의 −mT는 성장할 몫만큼의 할인, +zσ√T는 확률 조건의 안전마진 — p < 50%면 z가 음수라 오히려 덜 든다")
T(S(21), 8, "※ 예: Safety·GHP 6 × exp(−0.2) = 4.91억 · Safety·균형형 6 × exp(0.02) = 6.11억 — 아홉 칸은 부록 B-2")
T(S(72), 11, "② 확률 조건 — 하위 (1 − p)의 나쁜 운(Z = −z)에서도 G")
T(S(72), 19, "검산 — Safety를 균형형에 넣으면 6 × exp(−0.45 + 1.645 × 0.285) = 6 × exp(0.02) = 6.11억 (확률 정확히 95%)")
T(S(73), 5, "부록 B-2 — 지수 = −mT + zσ√T, K = G × exp(지수) · √10 = 3.162")
lecfix.set_tf(lecfix.find_table(S(73)).table.cell(0, 3).text_frame, "exp(지수)")

# 23장(옛) → 새 23b: 풀이
s = S(23)
T(s, 3, "목표마다 따로 풀면 확률 조정 성장률이 가장 큰 쪽이 가장 쌉니다")
T(s, 5, "방법 1의 풀이 — 목표마다 쪼개진다: 목표마다 K를 가장 작게 하면 합도 가장 작고, 남는 자금은 가장 크다")
T(s, 8, "목표 하나의 필요 자본 — 단원 ③의 K 공식을 주식 비중 w의 함수로")
swap_eq(s, 9, E_K, maxw=760, maxh=70, cx=640)
T(s, 11, "로그를 취하면 — K 최소 = 확률 조정 성장률 g 최대")
lecfix.set_tf(lecfix.find_table(s).table.cell(0, 1).text_frame, "z값")

# 25장 · 26장 — 참조
T(S(25), 8, "단원 ③의 “가장 싼 것”과 같은 답")
s = S(26)
T(s, 3, "단원 ①~③은 위험을 목표 미달 확률로 다시 정의했습니다"); T(s, 5, "단원 ①~③ 체크포인트")
T(s, 13, "필요 자본 K = G × exp(−mT + zσ√T)")
T(s, 18, "σ가 아니라 P(W < G). 같은 문제를 방법 1(가장 싼 K)과 방법 2(Das–Markowitz, (H, α) → γ)로 푼다 — 두 IC 안건 모두 이 언어로 심의한다.")

# 33장 — 상황 → 질문 → 답
s = S(33)
T(s, 3, "필수 소득을 지키는 길은 ‘이상 → 복제 → 가격 → 운용’ 네 걸음입니다")
T(s, 5, "상황 — 이순자 씨는 국민연금 외에 매달 150만 원이 평생 더 필요하다 · 질문 — 무엇을, 얼마에 사고, 남는 자금은 어떻게 굴리나?")
T(s, 8, "① 이상적인 답\n평생 소득 채권(RB)을 산다")
T(s, 11, "② 현실의 답\n없으니 국채로 복제한다(GHP)")
T(s, 14, "③ 가격표\n복제 비용 3.50억이 Floor")
T(s, 17, "④ 운용\n남는 1.50억만 위험을 진다")
T(s, 19, "상황"); T(s, 20, "목표는 잔고가 아니라 평생 소득"); T(s, 21, "그 소득을 주는 RB는 한국에 없다")
T(s, 23, "답"); T(s, 24, "국채로 복제하고 나머지만 굴린다"); T(s, 25, "쿠션 1.50억을 Flexicure 규칙으로 굴린다")
T(s, 26, "Floor는 임의의 비율이 아니라 ‘필수 소득을 오늘 사는 값’ — 그 위의 쿠션만 위험을 집니다")

# 35장 · 46장 — G_S를 말로
T(S(35), 12, "최악에도 Safety 목표"); T(S(35), 16, "“최악에도 Safety 목표”를 수학으로 보장한 뒤 그 위에서 자유")
T(S(46), 11, "하락 시 PSP 자동 감소\n“최악에도 Safety 목표” 심리적 안정\n감정 편향 차단")

# 45장 — 표 안의 식을 수식 그림으로
s = S(45); tb = lecfix.find_table(s)
row_heights(tb, [40] + [52] * 5)
lecfix.set_tf(tb.table.cell(4, 2).text_frame, "자산이 C에 처음 닿는 날 전액 GHP로")
for ri, tex, nm in ((1, r"$F_t=\max(F_{t-1},\ k\,A_t)$", "t45_1.png"),
                    (2, r"$F_t=(1-\delta)\,\max_{s\leq t}\ A_s$", "t45_2.png"),
                    (3, r"$E_t=m\,\min(A_t-F_t,\ C-A_t)$", "t45_3.png"),
                    (5, r"$m_t=1\,/\,\mathrm{VaR}_t$", "t45_5.png")):
    eq_in_cell(s, tb, ri, 2, tex, nm, fontsize=25)
T(s, 5, "GBI 일반 틀의 장치 — F는 Floor, A는 자산, E는 위험자산 금액, C는 열망 목표의 현재가치")

# 47장 — 표 안의 식
s = S(47); tb = lecfix.find_table(s)
row_heights(tb, [40] + [52] * 5)
eq_in_cell(s, tb, 2, 1, r"$m=\min(m_{\mathrm{util}},\ 1/L)$ + 몬테카를로", "t47_2.png", fontsize=24)
eq_in_cell(s, tb, 3, 1, r"$m\leq 1/L$ · 변동성 연동 $m_t$", "t47_3.png", fontsize=24)

# 49장 — 단원 ④~⑥ 체크포인트
s = S(49)
T(s, 3, "단원 ④~⑥은 Floor를 만들고 쿠션만큼 위험을 지는 법을 다뤘습니다"); T(s, 5, "단원 ④~⑥ 체크포인트")
T(s, 18, "m = min(1/L, (μ − r)/γσ²) + 몬테카를로. 실무는 VaR로 m ≤ 1/x 또는 손실 예산 L을 고정해 F = A − L/(m x) — 위험 ↑면 floor ↑ — 단원 ③과 같은 “나쁜 경우에도 floor”.")

# 58장 — 안건 2로 넘기는 말
T(S(58), 16, "안건 2는 H대 기금의 세 숫자 — 비유동 상한 · 안전 유동 하한 · 지출률을 이 틀로 정합니다")
# 60장 — 단원 ⑦ 체크포인트
T(S(55), 9, "적립금의 75.4%")
T(S(55), 10, "2025 말 원리금보장 비중(금감원)\n실질 수익률 ~1% · 부족액 누적")
t59 = sh(S(59), 10); lecfix.set_text(t59, "2024.12 일임 혁신금융 지정\n2025.3 첫 IRP 일임\n원리금보장 75.4%(2025 말)")
s = S(60)
T(s, 17, "한국 — 제도는 열렸고 관행은 원리금보장")
T(s, 18, "DB 45.7 · DC 28.2 · IRP 26.1%, 원리금보장 75.4% · 디폴트옵션도 85.4%가 안정형(2025 말). 2025.3 IRP 로보 일임 개시.")
T(s, 3, "단원 ⑦은 이상과 대중화 사이의 거리를 다뤘습니다"); T(s, 5, "단원 ⑦ 체크포인트")
T(s, 19, "다음 — 같은 틀을 기관에: H대 기금은 무엇을 어떤 숫자로 정하나")
T(S(78), 7, "출처 — 부록 D 참고문헌 (EDHEC Flexicure · Roy · Telser · Kataoka · Chhabra · Merton · Black-Jones · Shefrin-Statman · Das 외 · 금감원 · 고용노동부)")

# 67장 — IC 개요표: 결정할 것을 동사로
t = lecfix.find_table(S(67)).table
lecfix.set_tf(t.cell(1, 2).text_frame, "H대 기금 5,000억의 비유동 상한 x · 안전 유동 하한 y · 지출률 z를 정한다")
lecfix.set_tf(t.cell(2, 2).text_frame, "Flexicure를 비유동까지 — 충격 뒤 m ≤ 2 · 분모 효과 · 모방 3장벽")
lecfix.set_tf(t.cell(4, 1).text_frame, "Floor · m 상한 · 총보수를 정한다")
lecfix.set_tf(t.cell(4, 2).text_frame, "세 숫자 (x, y, z)와 부수 조건을 정한다")
# 68장 — 개념 출처를 단원으로
tb = lecfix.find_table(S(68))
lecfix.set_table(tb, [["이번 강의의 개념", "IC에서 꽂히는 쟁점"],
                      ["위험 = P(W < G) · 필요 자본 K · 방법 1 · 2 · 3 (단원 ①~③)", "공통 — 수익률이 아니라 달성 확률로 판 가르기"],
                      ["RB → GHP → Floor 연쇄 (단원 ④)", "안건 1 — Floor를 한국에서 무엇으로 만드는가"],
                      ["쿠션 공식 · m의 세 기준 · Cash Trap (단원 ⑤⑥)", "안건 1 — 승수 m 상한, 급락 락인 방어"],
                      ["4대 도전 · 비용 0.8~1.5% (단원 ⑦)", "안건 1 — 수수료가 디폴트옵션 자격이 있는가"],
                      ["Flexicure 다섯 단계 · 모방 3장벽 (단원 ⑧)", "안건 2 — Floor 1,536억 위의 쿠션으로 x · y · z를 정한다"]])
table_font(tb, 15, 14); row_heights(tb, [38] + [46] * 5)
# 69장 — 제2호의 숫자 라벨
s = S(69)
T(s, 17, "제2호 · H대 기금 — 세 숫자 (x, y, z)")
T(s, 19, "비유동 x"); T(s, 20, "≤ 20% (재간접 · 세컨더리)")
T(s, 21, "유동 y"); T(s, 22, "≥ 30%")
T(s, 23, "지출 z"); T(s, 24, "4.5% (고정 3.5%가 Floor)")
fx = copy_shape(sh(s, 24), s); set_px(fx, x=682, y=486, w=486, h=28)
lecfix.set_text(fx, "Flexicure: 충격 뒤 m ≤ 2 · 사모 22% 넘으면 약정 중단"); font(fx, 14)

# 74장 — 표 안의 식
s = S(74); tb = lecfix.find_table(s)
row_heights(tb, [40, 56, 56, 56, 56])
lecfix.set_tf(tb.table.cell(1, 1).text_frame, "4.91 + 2.24 + 1.97")
lecfix.set_tf(tb.table.cell(1, 2).text_frame, "9.13억 ≤ 10억 → 0.87억 남음")
lecfix.set_tf(tb.table.cell(2, 2).text_frame, "목표 7.22억으로 상향")
lecfix.set_tf(tb.table.cell(3, 2).text_frame, "달성 확률 약 52%")
lecfix.set_tf(tb.table.cell(4, 2).text_frame, "상속 목표 약 2.1억 (30%)")
eq_in_cell(s, tb, 2, 1, r"$G=2.843\times e^{0.932}=7.22$", "t74_2.png", fontsize=25)
eq_in_cell(s, tb, 3, 1, r"$z=\left(\ln(2.843/5)+0.60\right)/\,0.632=0.06$", "t74_3.png", fontsize=25)
eq_in_cell(s, tb, 4, 1, r"$8-4.91-2.24=0.84$억 → $G=0.84\times e^{0.932}$", "t74_4.png", fontsize=24)
# 76장 — 한계 ①의 참조
T(S(76), 13, "※ 한계 ① 계좌 간 분산효과를 무시한다 — Das–Markowitz(단원 ③ 방법 2 · 부록 B-6): 합친 포트폴리오도 효율선 위, 손실은 작다")
T(S(71), 5, "B-1~B-5 유도와 손계산 · B-6 Das–Markowitz 원 논문 예 · C 용어집 · C-2 한국 제도 용어 · D 참고문헌")
# 77장 — 용어집
tb = lecfix.find_table(S(77))
rows = [[c.text for c in r.cells] for r in tb.table.rows]
rows = [r if r[0] != "필요 자본 K" else ["필요 자본 K", "G를 T년 뒤 확률 p로 이룰 오늘 자금 = G × exp(−mT + zσ√T)"] for r in rows]
rows.insert(3, ["심리계정 포트폴리오", "Das–Markowitz — 계좌마다 (H, α)를 묻고 평균-분산으로 푼다 · 공매도 허용 시 합쳐도 효율선 위 (단원 ③ · 부록 B-6)"])
rows.append(["모방 3장벽", "작은 기금이 Yale을 따라 할 때 막히는 접근권 · 시간 지평 · 인력"])
rows.append(["VaR (밸류앳리스크)", "정한 확률(예: 95%)에서의 최악 손실 = zσ − μ · CPPI 승수 상한 m ≤ 1/VaR"])
rows.append(["위험 자본", "위험이 커질수록 묶어 두는 자본 — 여기서는 VaR 손실 예산 L에 맞춰 올리고 내리는 floor, F = A − L/(m x)"])
lecfix.set_table(tb, rows); table_font(tb, 13, 13); row_heights(tb, [34] + [32] * (len(rows) - 1))


# 4차 — Market = w* 67% 혼합(2.24억)으로 통일한 숫자 전파
t = lecfix.find_table(S(17)).table          # 13장 표: Market 담기는 자산
lecfix.set_tf(t.cell(2, 4).text_frame, "GHP + 주식 (주식 약 67%)")
lecfix.set_tf(lecfix.find_table(S(20)).table.cell(5, 1).text_frame, "전체 배분 (예: 49 / 22 / 28)")
T(S(20), 5, "49 / 22 / 28은 입력이 아니라 결과다")
T(S(21), 8, "※ 예: Safety·GHP 6 × exp(−0.2) = 4.91억 · 아홉 칸은 부록 B-2 · Market의 균형형 2.22억은 후보 비교, 본문 계산은 주식 67% 혼합 2.24억(손계산 장)")
s = S(22)
T(s, 3, "목표마다 공식 비중으로 사면 9.13억이 들고 남는 0.87억은 꿈의 계좌로 갑니다")
T(s, 5, "합계 4.91 + 2.24 + 1.97 = 9.13억 ≤ 10억 · Market은 GHP + 주식 67%")
T(s, 13, "Market → 주식 67% 혼합"); T(s, 14, "2.24억"); T(s, 15, "22%")
T(s, 19, "1.97 + 0.87 = 2.84억"); T(s, 20, "28%")
T(s, 23, "목표 상향 5억 → 7.22억"); T(s, 27, "확률 상향 30% → 약 52%")
T(s, 33, "※ 합쳐 보면(look-through) 주식 4.35억(43%) · 채권·GHP 5.65억(57%)")
t = lecfix.find_table(S(23)).table; lecfix.set_tf(t.cell(2, 3).text_frame, "주식 약 67% 혼합")
t = lecfix.find_table(S(24)).table; lecfix.set_tf(t.cell(2, 2).text_frame, "67% (GHP + 주식)")
T(S(26), 14, "Safety 4.91억(GHP) · Market 2.24억(주식 67% 혼합) · Aspirational 1.97억(주식) — 49 / 22 / 28은 결과다.")
T(S(73), 9, "아홉 칸에서는 70%가 균형형(2.22억) — 본문 Market은 주식 67% 혼합 2.24억")

# 13차 — 실무 장에서 방법 1 · 2 · 3 연결
_t = lecfix.find_table(S(52)).table
lecfix.set_tf(_t.cell(2, 1).text_frame, "개인별 MC 1,000회 × 동적 최적화 × 10만 명 × 매월 — 단원 ③ 방법 3(몬테카를로)의 규모 문제")
T(S(53), 16, "버킷 배분은 방법 1 · 2로, 확률 대시보드는 방법 3으로 — 남은 병목은 목표 수집과 신뢰입니다")
T(S(63), 5, "Step 3 필요 자본은 방법 1(가장 싼 K), Step 4~6은 방법 3(몬테카를로)과 Flexicure — 약 20분 · 데이터는 W6 실습데이터 덱")

# ============================================================ 새 장 만들기 ====
UNITS = [  # (번호, 머리 띠 이름, 디바이더 제목, 이 단원이 푸는 질문, 다루는 것(디바이더), 다루는 것(목차))
    ("①", "목표와 위험", "목표와 위험 — σ가 아니라 확률",
     "같은 나이 · 같은 자산인데 왜 답이 달라야 하고, 위험은 무엇으로 재는가?",
     "GBI 3원칙 · 세 비유 · 미달 확률과 기대 부족분 · 95 · 70 · 30", "3원칙 · 미달 확률 · 95 · 70 · 30"),
    ("②", "심리계정", "심리계정 — Thaler · BPT · Das–Markowitz",
     "자금을 계좌로 나눠 생각하는 습관을 어떻게 최적화 문제로 바꾸는가?",
     "Thaler 심리계정 · Shefrin–Statman BPT · Das–Markowitz 소개 · 세 방식 비교 · 3계층 버킷", "BPT · Das–Markowitz 소개 · 버킷"),
    ("③", "필요 자본", "필요 자본과 가장 싼 포트폴리오",
     "목표 G를 확률 p로 이루려면 오늘 얼마를, 무엇에 떼어 두는가?",
     "같은 문제, 세 풀이 — 방법 1 가장 싼 필요 자본 K · 방법 2 Das–Markowitz (H, α) → γ · 방법 3 몬테카를로 · 세 방법 비교", "방법 1 가장 싼 K · 방법 2 Das–Markowitz · 방법 3 몬테카를로"),
    ("④", "Floor 만들기", "Floor 만들기 — RB에서 GHP로",
     "필수 목표를 무엇으로, 얼마에 지키는가?",
     "Retirement Bond · GHP 복제 · Floor = 필수 목표의 시장가 · 한국의 현실", "RB · GHP 복제 · Floor의 값"),
    ("⑤", "쿠션과 승수", "쿠션과 승수 — CPPI에서 Flexicure로",
     "Floor 위의 쿠션으로 위험을 얼마나, 어떤 규칙으로 지는가?",
     "Flexicure · 쿠션 공식 · CPPI · m의 세 기준 · VaR로 m · 위험 자본으로 floor · 세 방법은 같은 문장 · Cash Trap", "쿠션 공식 · VaR · 위험 자본 · Cash Trap"),
    ("⑥", "Floor 운용 규칙", "Floor 운용 규칙과 한계",
     "Floor를 언제 다시 정하고, 규칙 기반 보호의 한계는 무엇인가?",
     "매년 재설정 · Ratchet · Cap · Stop-gain · 한계와 쏠림 · VaR 운용", "재설정 · Ratchet · 한계 · VaR"),
    ("⑦", "GBI와 현실 세계", "GBI와 현실 세계 — 이상과 대중화 사이",
     "가장 개인화된 이 틀이 왜 대중화가 어렵고, 한국은 어디까지 왔나?",
     "MVO · TDF · LDI · GBI 비교 · 4대 도전 · 로보 2세대 · 한국 퇴직연금 · 디폴트옵션 · TDF · 한국 DC의 덫", "4대 도전 · 한국 퇴직연금 · 디폴트옵션 · TDF"),
    ("⑧", "기관의 GBI · H대", "기관의 GBI — H대 기금의 세 숫자",
     "대학기금의 지출 약속을 GBI로 읽으면 무엇을 어떤 숫자로 정하는가?",
     "세 숫자 x · y · z · 강의 대응표 · 모방 3장벽 · Flexicure를 비유동성까지 옮기는 다섯 단계", "x · y · z · 대응표 · Flexicure 다섯 단계"),
    ("⑨", "실습 · GBI 설계기", "손으로 하지 말고 Claude Code에게 시켜라",
     "이순자 씨의 GBI를 20분 만에, IC에서 쓸 숫자로 어떻게 만드는가?",
     "메타프롬프트 · 6단계 설계기 · 승수 민감도 · 목표 언어 발언", "설계기 · 승수 민감도"),
    ("⑩", "모의 IC", "모의 투자위원회 — 이번 주는 2건",
     "디폴트옵션의 Floor · m · 보수와 H대의 x · y · z를 몇으로 정하는가?",
     "제1호 디폴트옵션 Flexicure형 · 제2호 H대 기금 · 기본 답 · 과제", "안건 2건 · 기본 답 · 과제"),
]

# (가) 목차 — 3장 표 장표 복제
N2 = clone(S(3))
T(N2, 3, "열 개 단원은 오후 IC 두 안건의 숫자로 모입니다")
T(N2, 5, "단원마다 푸는 질문이 하나씩 있다 — 1교시 ①~③, 2교시 ④~⑥, 3교시 ⑦⑧, 4교시 ⑨⑩")
tb = lecfix.find_table(N2)
lecfix.set_table(tb, [["단원", "이 단원이 푸는 질문", "다루는 것"]] + [[f"{u[0]} {u[1]}", u[3], u[5]] for u in UNITS])
col_widths(tb, (230, 590, 312)); table_font(tb, 12, 12); row_heights(tb, [30] + [31] * 10)
stripe(tb)
T(N2, 8, "발언은 수익률이 아니라 목표 언어 — Floor · 쿠션 · 달성 확률로 합니다"); set_px(sh(N2, 8), y=600)

# (나) 먼저 읽는 용어 — 77장 용어집 장표 복제
GL = clone(S(77))
T(GL, 3, "이번 주의 말은 처음 보는 사람 기준으로 여기서 먼저 정의합니다")
T(GL, 5, "먼저 읽는 용어 — 본문에서 설명 없이 쓰는 말 · 식과 숫자는 부록 C 용어집")
tb = lecfix.find_table(GL)
lecfix.set_table(tb, [["용어", "뜻"],
    ["목표 언어", "수익률 대신 Floor · 쿠션 · 달성 확률로 말하는 법 — 이번 주 IC의 공용어"],
    ["3계층 버킷", "자금을 Safety(필수) · Market(여유) · Aspirational(꿈) 계좌로 나누고 달성 확률 95 · 70 · 30%를 건다"],
    ["GHP · PSP", "목표를 확실히 맞추는 안전자산 묶음(Goal-Hedging) · 수익을 추구하는 위험자산 묶음(Performance-Seeking)"],
    ["Floor · 쿠션", "필수 목표를 지금 시장에서 사는 값 · 자산 − Floor (잃어도 되는 금액)"],
    ["승수 m · Cash Trap", "쿠션의 몇 배를 위험자산에 두나 · 쿠션이 0이 되어 회복장에 참여하지 못하는 상태"],
    ["Flexicure", "Floor 위 쿠션에 비례해 위험자산을 두고 매년 Floor를 다시 정하는 은퇴 설계 (EDHEC, Flexibility + Security)"],
    ["지배 검정", "같은 자금으로 GHP가 목표를 더 확실히 이루면 위험 계좌를 쓰지 않는다는 점검"],
    ["모방 3장벽", "작은 기금이 Yale 모델을 따라 할 때 막히는 접근권 · 시간 지평 · 인력 — 단원 ⑧에서 숫자로"]])
table_font(tb, 14, 13); col_widths(tb, (220, 912)); row_heights(tb, [36] + [42] * 8)

# (다) 단원 ② — Das–Markowitz 문제와 해(24장 복제: 수식 둘 + 표)
DM1 = clone(S(24))
T(DM1, 3, "방법 2 — Das–Markowitz는 (H, α)를 평균-분산 조건으로 바꿔 풉니다")
T(DM1, 5, "H 최소 수익률 · α 미달 확률 한도 · 원 논문(2010)의 자산 셋으로 — 먼저 손으로(다음 두 장), 그다음 행렬로 일반화")
T(DM1, 8, "계좌 하나의 문제 — 미달 확률 한도 안에서 기대수익률을 최대로")
swap_eq(DM1, 9, E_DM1, maxw=700, maxh=52, cx=640)
T(DM1, 11, "정규분포면 확률 한도 = 평균 · 표준편차 조건 → 효율선 위 한 점, γ 하나")
note(DM1, 630, "※ 강의 식의 z(α가 5%면 1.645, 양수)는 원문의 −Φ⁻¹(α)")
swap_eq(DM1, 12, E_DM2, maxw=720, maxh=46, cx=640)
tb = lecfix.find_table(DM1)
lecfix.set_table(tb, [["자산 (원 논문의 가상 자산)", "기대수익률 · 표준편차", "상관"],
                      ["채권형", "5% · 5%", "두 주식형과 0"],
                      ["저위험 주식형", "10% · 20%", "고위험 주식형과 0.20 (공분산 0.02)"],
                      ["고위험 주식형", "25% · 50%", "저위험 주식형과 0.20"]])
col_widths(tb, (330, 360, 442))

# (라) 단원 ② — 합쳐도 효율(16장 복제: 수식 하나 + 카드 둘)
DM2 = clone(S(16))
T(DM2, 3, "공매도를 허용하면 계좌별 해를 합쳐도 효율선 위에 있습니다")
T(DM2, 5, "계좌의 해 = 최소분산 포트폴리오 g + (1/γ) × 위험을 더 지는 방향 v — 같은 두 재료를 비율만 달리 섞는다")
T(DM2, 8, "계좌 배분 a로 합치면 — 전체도 같은 꼴, γ는 역수로 평균한다")
swap_eq(DM2, 9, E_DM3, maxw=760, maxh=72, cx=640)
T(DM2, 11, "숫자 — 은퇴 60 · 교육 20 · 상속 20%"); T(DM2, 12, "전체 γ = 2.174 (평균 2.99 아님)")
T(DM2, 14, "합친 비중 — 여전히 효율선 위"); T(DM2, 15, "채권 24 · 저위험 42 · 고위험 34%")
T(DM2, 16, "BPT의 “합치면 비효율”이 정규분포 + 공매도 허용에서는 사라집니다(계정별 롱온리면 12bp, {p:DMB}장)")
note(DM2, 630, "※ 원문은 전체 γ가 계정 γ의 가중평균과 다르다고만 쓴다 — 역수 가중평균은 식 (3)에서 바로 나오는 강의 도출 · 원문 p.315 ii)b · pp. 320–321")

# (마) 단원 ② — 세 방식 비교(51장 표 복제, 4열)
CMP = clone(S(51))
T(CMP, 3, "심리계정 세 방식 — 관찰 · 행동의 최적화 · 운용의 최적화")
T(CMP, 5, "같은 심리계정, 세 가지 쓰임 — 무엇을 받고, 무엇을 풀고, 해는 어떤 꼴이며, 실무에 무엇이 좋고 나쁜가")
tb = lecfix.find_table(CMP); drop_col(tb, 4)
lecfix.set_table(tb, [["구분", "Thaler (1985) 심리계정", "Shefrin–Statman (2000) BPT", "Das–Markowitz 외 (2010)"],
    ["입력", "없음 — 사람의 습관을 묘사", "계좌별 열망 수준 · 미달 확률 한도 α", "계좌별 (H, α) + 자본시장 가정 μ · Σ"],
    ["푸는 문제", "풀지 않는다 — 자금에 꼬리표가 붙는다는 관찰", "미달 확률 ≤ α 아래 기대 부를 최대로", "미달 확률 ≤ α 아래 기대수익률 최대 → γ 하나로"],
    ["해의 성질", "—", "계좌마다 안전자산 + 복권형 자산, 합치면 평균-분산 비효율", "계좌마다 효율선 위 한 점, 공매도 허용 시 합쳐도 효율선 위"],
    ["실무 장점", "패닉 매도 · 꼬리표 행동을 설명한다", "행동을 그대로 담는다", "묻기 쉽고 MVO 도구 그대로, 전체 위험을 한 번에 보고"],
    ["실무 한계", "설계 지침이 없다", "해가 복잡해 계산이 어렵고 비효율", "정규분포 가정 · 미달의 크기는 못 막는다"]])
col_widths(tb, (150, 300, 330, 352)); table_font(tb, 15, 14)
T(CMP, 8, "Das–Markowitz의 풀이는 단원 ③ — 가장 싼 필요 자본(방법 1)과 나란히 방법 2로 풉니다")
set_px(sh(CMP, 8), y=590)
note(CMP, 548, "※ Das–Markowitz는 BPT의 위험추구(복권형) 특징은 모형에 넣지 않았다(원문 p.313)")

# (바) 단원 ② — 왜 Das–Markowitz인가(4장 메시지 장표 복제)
WHY = clone(S(4))
T(WHY, 3, "실무에서 방법 2(Das–Markowitz)가 편한 이유는 세 가지입니다")
T(WHY, 5, "BPT는 왜 그렇게 행동하는지를 설명하고, Das–Markowitz는 그 행동을 어떻게 운용할지를 답한다")
T(WHY, 9, "묻기 쉬운 두 숫자 (H, α)")
T(WHY, 10, "“위험회피도가 얼마인가” 대신 “이 계좌는 5% 확률로도 −10% 아래면 안 된다”를 묻는다 — 운용사가 그것을 γ = 3.795로 번역한다.")
T(WHY, 13, "쓰던 평균-분산 도구 그대로, 합쳐도 효율(공매도 허용 시)")
T(WHY, 14, "계좌마다 같은 자본시장 가정으로 풀고, 공매도를 허용하면 더해도 효율선 위다. 위험회피도를 10 · 20 · 30% 잘못 말하면 연 2.5~44bp를 잃으니(원문 표 3) (H, α)로 묻는 편이 낫다.")
T(WHY, 17, "안 되는 목표를 먼저 거른다")
T(WHY, 18, "이 세 자산으로 “원금 손실 확률 5% 이하”는 불가능하다 — 하위 5% 경계의 최댓값이 −2.31%다(강의 계산). 그러면 목표 · 기간 · 저축을 다시 묻는다.")
T(WHY, 19, "한계 — 미달의 크기는 막지 못하므로 필수 목표의 바닥은 Floor(CPPI)로 따로 지킵니다")

# (사) 새 23a — 문제 · 결정 변수 · 목적식(23장 복제)
T(P23, 3, "방법 1 — 필요 자본은 남는 자금을 최대로 하는 최적화입니다")
T(P23, 5, "같은 문제를 두 방법으로 푼다 — 방법 1: G와 p는 고정, 목표마다 K와 주식 비중 w를 고른다 · 방법 2는 뒤에서")
T(P23, 8, "목적식 · 결정 변수 · 제약 — 고르는 것은 K와 w, 지키는 것은 확률 p")
by = {x.shape_id: x for x in P23.shapes}
box2 = container(P23, by[12])
for x in (by[11], by[12], box2): remove(x)
box1 = swap_eq(P23, 9, E_OBJ, maxw=1060, maxh=70, cx=640, pad=24)
pic1 = next(x for x in P23.shapes if x.shape_type == 13)
bx, byy, bw, bh = lecfix.px(box1); set_px(box1, y=200, h=128)
set_px(sh(P23, 8), y=210); set_px(pic1, y=246)
tb = lecfix.find_table(P23)
lecfix.set_table(tb, [["기호", "뜻", "기호", "뜻"],
                      ["K", "목표마다 오늘 떼어 둘 자금 — 고른다", "w", "그 계좌의 주식 비중(나머지는 GHP) — 고른다"],
                      ["G · T", "목표 금액 · 남은 연수 — 주어진다", "p · z", "달성 확률 · 표준정규 값(95% → 1.645) — 주어진다"],
                      ["m · σ", "계좌의 연 성장률(로그) · 변동성 — w가 정한다", "Z", "운 — 평균 0 · 표준편차 1인 표준정규"],
                      ["W", "총자산 — W − ΣK가 남는 자금", "g", "확률 조정 성장률 = m − zσ/√T — 다음 장"]])
col_widths(tb, (100, 466, 100, 466)); table_font(tb, 14, 13); row_heights(tb, [34] + [38] * 4)
set_px(tb, y=346)
l1 = copy_shape(sh(S(19), 14), P23); set_px(l1, y=540, h=34)
lecfix.set_text(l1, "읽는 법 — 목표끼리는 서로 묶이지 않아, 목표마다 K를 가장 작게 하면 합도 가장 작고 남는 자금은 가장 크다")
l2 = copy_shape(sh(S(45), 8), P23); set_px(l2, y=580, h=44)
lecfix.set_text(l2, "“자금을 최대로 남긴다”는 맞습니다 — 단, 목표와의 차이를 줄이는 것이 아니라 확률 p를 딱 채우는 가장 싼 K를 찾습니다")
font(l1, 16); font(l2, 17)
n28 = copy_shape(sh(S(21), 8), P23); set_px(n28, x=79, y=630, w=1132, h=24); font(n28, 12)
lecfix.set_text(n28, "※ 근거: Kataoka(1963) 안전우선 기준(미달 확률 고정 · 하위 경계 최대)을 로그정규에 적용 · 버킷은 Chhabra(2005) · 무조건 항은 Merton(1969/1971) — 부록 D")

# (아) 단원 ⑧ — H대가 정할 세 숫자(51장 표 복제, 5열)
H1 = clone(S(51))
T(H1, 3, "H대 IC는 철학의 이름을 고르지 않고 세 숫자 x · y · z를 정합니다")
T(H1, 5, "문제 — H대(가상) 5,000억 기금이 장학 · 연구 고정 지출(필수 목표)을 지키면서, 팔 수 없는 사모투자를 어디까지 담고(x) 현금 · 국채를 얼마 두고(y) 얼마나 꺼내 쓸지(z)를 정한다")
tb = lecfix.find_table(H1)
lecfix.set_table(tb, [["세 숫자", "IC가 정하는 것", "기본 답", "GBI로 읽으면", "지키는 판정 조건"],
    ["x 비유동 대체 상한", "PE · VC · 사모부동산을 기금의 몇 %까지 담을지 상한을 정한다", "20%\n재간접 · 세컨더리", "PSP 중 팔 수 없는 몫의 위험 예산 — 충격 뒤에도 (주식 + 사모) ÷ 쿠션 ≤ 2", "① 충격 후 버티는 연수 13년 ≥ 락업 10년"],
    ["y 안전 유동자산 하한", "현금 · 국채를 최소 몇 % 둘지 하한을 정한다", "30%", "GHP(Safety 버킷) — 충격 뒤 3년 지출 + 공사비를 매도 없이 댄다", "④ 3년 지출 유지 확률 100% ≥ 95%"],
    ["z 지출률", "해마다 기금의 몇 %를 꺼내 쓸지 정한다 (전년 80% + 목표 20% 평활)", "4.5%", "고정 지출 3.5%(175억) = 필수 목표 · Floor, 4.5%까지는 희망 목표", "② 30년 실질가치 유지 확률 59% ≥ 50%"],
    ["부수 조건", "분모 효과 때 신규 약정을 멈추고 매년 다시 계산한다", "의결문에 명시", "Flexicure식 재설정 — 쿠션이 줄면 위험을 늘리지 않고 매년 다시 잰다", "⑤ 거버넌스"]])
col_widths(tb, (160, 300, 130, 330, 212)); table_font(tb, 14, 13)
T(H1, 8, "기본 답 (20, 30, 4.5) = 과제 B안 (주식 50 · 국채 25 · 현금 5 · 사모 20)")
set_px(sh(H1, 8), y=596)

# (자) 단원 ⑧ — 모방 3장벽(55장 카드 장표 복제)
H2 = clone(S(55))
T(H2, 3, "Yale을 흉내 낼 때 막히는 세 장벽은 접근권 · 시간 지평 · 인력입니다")
T(H2, 5, "모방 3장벽 = 작은 기금이 Yale의 대체투자 60%를 따라 할 때 숫자로 막히는 세 조건 (케이스 판정 조건 ①·③)")
T(H2, 8, "장벽 ① 접근권"); T(H2, 9, "좋은 매니저에 들어가나")
T(H2, 10, "일류 펀드는 기존 출자자 우선\n직접: 15곳 × 100억 = 30%\n안 되면 재간접 −2%p")
T(H2, 12, "장벽 ② 시간 지평"); T(H2, 13, "못 팔아도 버티나")
T(H2, 14, "락업 10년 — 팔 수 없다\n충격 후 버틴 연수 ≥ 10년\nx 20: 13년 · x 60: 2년")
T(H2, 16, "장벽 ③ 인력"); T(H2, 17, "골라낼 사람이 있나")
T(H2, 18, "Yale: 전문 인력 수십 명\nH대: 대체 전담 1명\n직접 투자: 3명 이상")
for sid in (10, 14, 18): font(sh(H2, sid), 14)
T(H2, 19, "하나라도 못 넘으면 x를 낮추거나 간접(재간접 · 세컨더리)으로 갑니다")

# (차) 단원 ⑧ — H대 ↔ GBI 기법(51장 표 복제, 4열)
H3 = clone(S(51))
T(H3, 3, "정리 — H대 문제는 기관판 GBI입니다: 같은 기법, 다른 주인")
T(H3, 5, "Floor · GHP · 쿠션 · 확률 제약 · Flexicure 재설정 — 이순자 씨의 노후 소득과 H대의 지출 약속을 나란히 놓는다")
tb = lecfix.find_table(H3); drop_col(tb, 4)
lecfix.set_table(tb, [["GBI 기법", "이순자 씨 (개인)", "H대 기금 (기관)", "정하는 숫자"],
    ["필수 목표 → Floor", "월 부족분 150만 원 → Floor 3.50억", "앞 3년 고정 지출 536 + 공사비 600 + 남은 출자 요청 400 = 1,536억 (B안)", "z의 바닥"],
    ["GHP (Safety 버킷)", "물가연동국채 사다리", "현금 · 국채 30% — 충격 뒤 3년 지출 + 공사비", "y"],
    ["쿠션 · 위험 예산", "쿠션 1.50억의 m배만 위험자산", "쿠션 2,464억 — (주식 + 사모) ≤ 쿠션의 2배, 충격 뒤에도", "x"],
    ["확률 제약", "Safety ≥ 95% · Market ≥ 70%", "3년 지출 유지 ≥ 95% · 30년 실질가치 유지 ≥ 50%", "y · z"],
    ["Flexicure 재설정", "매년 Floor = 자산의 80%로 다시 정한다", "쿠션이 기준 아래면 희망 지출 동결 · 80/20 평활 · 연 1회 재계산", "z · ⑤"]])
col_widths(tb, (220, 330, 432, 150)); table_font(tb, 15, 14)
T(H3, 8, "차이 하나 — 비유동 자산은 쿠션이 줄어도 팔 수 없어, 줄이는 수단은 신규 약정 중단뿐입니다")
set_px(sh(H3, 8), y=596)


# ======================================================= 2차 보수 — 새 장 ====
# (타) 단원 ② — Das–Markowitz 소개(16장 복제)
DMI = clone(S(16))
T(DMI, 3, "Das–Markowitz는 계좌마다 (H, α) 두 숫자만 묻습니다")
T(DMI, 5, "Das–Markowitz–Scheid–Statman(2010) — BPT의 계좌 나누기는 그대로 두고, 평균-분산으로 풀 수 있는 꼴로 바꿨다")
T(DMI, 8, "계좌 하나의 문제 — 미달 확률 한도 안에서 기대수익률을 최대로")
swap_eq(DMI, 9, E_DM1, maxw=700, maxh=52, cx=640)
T(DMI, 11, "묻는 것 — 계좌마다 두 숫자"); T(DMI, 12, "H 최소 수익률 · α 미달 확률 한도")
T(DMI, 14, "바뀌는 꼴 — 효율선 위 한 점"); T(DMI, 15, "풀이는 단원 ③의 방법 2")
T(DMI, 16, "위험회피도 γ를 묻는 대신 (H, α)를 묻는다 — 둘은 같은 포트폴리오를 가리키는 다른 말입니다")

# (파) 10차 — 세 계정 결과 + 롱온리(3장 표 장표 복제)
DMB = clone(S(3))
T(DMB, 3, "세 계정에 같은 순서를 적용하면 — γ 3.795 · 2.706 · 0.877")
T(DMB, 5, "계정마다 (H, α)만 다르다 — 원 논문 표 재계산(부록 B-6) · 미달 기준이 느슨할수록 γ가 작아 주식이 는다")
tb = lecfix.find_table(DMB)
lecfix.set_table(tb, [["계정 (H, α)", "γ", "비중 채권 · 저위험 · 고위험 → 기대수익 · 표준편차"],
    ["은퇴 (−10%, 5%)", "3.795", "53.9 · 26.6 · 19.5% → 10.23% · 12.30%"],
    ["교육 (−5%, 15%)", "2.706", "37.9 · 35.0 · 27.1% → 12.18% · 16.57%"],
    ["상속 (−15%, 20%)", "0.877", "−78.9 · 96.2 · 82.7% (채권 공매도) → 26.35% · 49.13%"]])
col_widths(tb, (300, 160, 672)); table_font(tb, 15, 14); row_heights(tb, [36] + [52] * 3); stripe(tb)
row_heights(tb, [34] + [40] * 3)
blk(DMB, 79, 398, 1132, [
    "**• 계정마다 롱온리:** 은퇴 · 교육은 원래 양수라 그대로, 상속만 0 · 8.9 · 91.1% (26.35% → 23.67%) — 푸는 법은 다음 장",
    "**• 합치면:** σ 19.38% · 기대 13.30% → 같은 σ의 제약 효율선 13.43% (무제약과 같다)보다 연 12bp 낮다",
    "**• 합산에서만 막으면:** 손실 0",
    "**▶ 비교:** 위험회피도를 10 · 20 · 30% 잘못 말하면 연 2.5~44bp를 잃는다 (원문 표 3, 낮은 γ일수록 큼) — 롱온리 12bp보다 클 수 있다"], pt=16)
T(DMB, 8, "“5% 확률로도 −10% 아래는 안 된다”는 목표 언어가 γ = 3.795와 같은 말입니다"); set_px(sh(DMB, 8), y=574)
note(DMB, 618, "※ 계정별 롱온리 해는 원문에 없는 강의 재계산(직접 최적화) · 12bp는 계정마다 막을 때, 합산 수준에서만 막으면 이 예에서 손실 0(합산이 이미 롱, 원문 p.329)")
note(DMB, 642, "※ 원문 p.328의 합산 σ 19.89%는 19.38%의 오기로 보인다 — 19.89%면 효율선 13.65%, 손실 약 34bp로 원문의 12bp와 모순")

# (파2) 10차 — 손계산 ① 한 포트폴리오 검사, ② 두 자산 축소(3장 표 장표 복제)
HC1 = clone(S(3))
T(HC1, 3, "손계산 ① — 포트폴리오 하나가 조건을 지키는지 검사한다")
T(HC1, 5, "은퇴 계정(H = −10%, α = 5%)에 채권 60 · 저위험 30 · 고위험 10%를 넣어 본다 — 보충교재 4-2 단계 1의 확인")
tb = lecfix.find_table(HC1)
lecfix.set_table(tb, [["항목", "계산", "값"],
    ["기대수익 m", " ", "8.50%"],
    ["분산 σ²", " ", "0.0082 → σ = 9.06%"],
    ["하위 5% 수익률", " ", "−6.40% ≥ H = −10% → 통과"],
    ["미달 확률", " ", "2.05% ≤ α = 5%"]])
col_widths(tb, (220, 640, 272)); table_font(tb, 15, 14); row_heights(tb, [36] + [58] * 4); stripe(tb)
for ri, tex, nm in ((1, r"$0.6\times 5\%+0.3\times 10\%+0.1\times 25\%$", "hc1_1.png"),
                    (2, r"$0.6^2\times 0.05^2+0.3^2\times 0.20^2+0.1^2\times 0.50^2+2\times 0.3\times 0.1\times 0.02$", "hc1_2.png"),
                    (3, r"$8.50\%-1.645\times 9.06\%$", "hc1_3.png"),
                    (4, r"$\Phi\left((-10\%-8.50\%)/9.06\%\right)=\Phi(-2.04)$", "hc1_4.png")):
    eq_in_cell(HC1, tb, ri, 1, tex, nm, fontsize=23)
line(HC1, 500, "분산의 네 항: 0.0009 + 0.0036 + 0.0025 + 0.0012 = 0.0082 — 마지막 항이 두 주식의 공분산(0.02) 몫", 15)
T(HC1, 8, "통과하지만 아직 여유가 있다 — 기대수익을 더 올릴 수 있다(다음 장)"); set_px(sh(HC1, 8), y=560)

HC2 = clone(S(3))
T(HC2, 3, "손계산 ② — 두 자산이면 비중 하나를 조건에 딱 맞추면 됩니다")
T(HC2, 5, "교육용 축소 — 원문 자산 중 채권형(5% · 5%)과 저위험 주식형(10% · 20%), 상관 0 · 은퇴 계정 H = −10%, α = 5%")
tb = lecfix.find_table(HC2)
lecfix.set_table(tb, [["단계", "계산", "값"],
    ["평균과 변동성을 w 하나로", " ", "주식 비중 w"],
    ["시험 w = 0.5", " ", "통과 — 여유 0.54%p"],
    ["시험 w = 0.6", " ", "위반"],
    ["등호로 놓고 제곱", " ", "2차식"],
    ["근의 공식", " ", "w* = 52.2% (m 7.61% · σ 10.70%)"]])
col_widths(tb, (230, 600, 302)); table_font(tb, 15, 14); row_heights(tb, [36] + [54] * 5); stripe(tb)
for ri, tex, nm in ((1, r"$m=5\%+5\%\,w,\quad \sigma=\sqrt{(1-w)^2\times 0.05^2+w^2\times 0.20^2}$", "hc2_1.png"),
                    (2, r"$7.5\%-1.645\times 10.31\%=-9.46\%$", "hc2_2.png"),
                    (3, r"$8.0\%-1.645\times 12.17\%=-12.01\%$", "hc2_3.png"),
                    (4, r"$(15\%+5\%\,w)^2=1.645^2\,[(1-w)^2\times 0.05^2+w^2\times 0.20^2]$", "hc2_4.png"),
                    (5, r"$-0.1125\,w^2+0.0285\,w+0.0157=0\ \Rightarrow\ w=0.522$", "hc2_5.png")):
    eq_in_cell(HC2, tb, ri, 1, tex, nm, fontsize=23)
T(HC2, 8, "기대수익을 최대로 하려면 제약이 등호로 걸리는 데까지 w를 올린다 — 자산이 셋 이상이면 이 일을 행렬로(다음 장)")
set_px(sh(HC2, 8), y=586); font(sh(HC2, 8), 16)


# (파3) 11차 — 롱온리 풀이 예(상속 계정, 3장 표 장표 복제)
LO = clone(S(3))
T(LO, 3, "롱온리면 이렇게 푼다 — 음수 자산을 0에 묶고 다시 푼다")
T(LO, 5, "상속 계정 H = −15%, α = 20% → z = Φ⁻¹(0.80) = 0.8416 · 무제약 해의 채권 −78.9%는 롱온리 위반")
tb = lecfix.find_table(LO)
lecfix.set_table(tb, [["단계", "계산", "결과"],
    ["① 무제약 해", "채권 −78.9 · 저위험 96.2 · 고위험 82.7%", "음수 → 롱온리 위반"],
    ["② 걸린 제약을 등호로", "채권 = 0으로 묶는다(KKT의 활성 제약)", "남는 자산 둘"],
    ["③ 두 자산으로 다시({p:HC2}장 방식)", " ", "w = 고위험 비중"],
    ["시험 두 개", " ", "0.8 통과 · 1.0 위반 → 답은 사이"],
    ["등호로 놓고 제곱", " ", "−0.1546w² + 0.1033w + 0.0342 = 0 → w* = 91.1% (다른 근 −0.243은 음수라 버림)"],
    ["④ 확인", "0 · 8.9 · 91.1% — 모두 0 이상이면 끝(음수면 또 0에 묶고 반복)", "기대 23.67% · σ 45.94%"]])
col_widths(tb, (220, 560, 352)); table_font(tb, 14, 13); row_heights(tb, [34] + [48] * 6); stripe(tb); set_px(tb, y=204)
for ri, tex, nm in ((3, r"$m=10\%+15\%\,w,\quad \sigma^2=0.04-0.04\,w+0.25\,w^2$", "lo_3.png"),
                    (4, r"$w=0.8:\ 22.0\%-0.8416\times 40.99\%=-12.50\%,\quad w=1.0:\ -17.08\%$", "lo_4.png"),
                    (5, r"$(25\%+15\%\,w)^2=0.8416^2\,(0.04-0.04\,w+0.25\,w^2)$", "lo_5.png")):
    eq_in_cell(LO, tb, ri, 1, tex, nm, fontsize=22)
T(LO, 8, "⑤ 자산이 많으면 이 일을 반복하거나 수치 최적화(엑셀 해 찾기 · SLSQP) — 닫힌 해 g + v/γ는 공매도 허용일 때만")
set_px(sh(LO, 8), y=572); font(sh(LO, 8), 16)
note(LO, 626, "※ w* = 0.911은 1을 넘지 않아 경계해(w = 1)가 아니다 · 은퇴 · 교육 계정은 해가 원래 양수라 롱온리 풀이가 필요 없다 · 강의 재계산(원문 §V는 결과만)")


# (파4) 12차 — 롱온리를 엑셀 해 찾기로(30장 그림 장표 복제)
_sv, _w, _h = solver_svg(); SOLVER_PNG = svg_png(_sv, os.path.join(EQ, "solver.png"), _w, _h)
_ly, _w2, _h2 = layers_svg(); LAYER_PNG = svg_png(_ly, os.path.join(EQ, "layers.png"), _w2, _h2)
XL = clone(S(30))
T(XL, 3, "롱온리는 엑셀 ‘해 찾기’ 대화창 하나로 풉니다")
T(XL, 5, "앞 장의 상속 계정을 그대로 — 체크 하나(음이 아닌 수)가 공매도 금지 · 같은 문제를 파이썬 SLSQP로도")
pic = next(x for x in XL.shapes if x.shape_type == 13)
h_ = 418; w_ = h_ * _w / _h; add_pic(XL, SOLVER_PNG, 640 - w_ / 2, 196, w_, h_, before=pic._element); remove(pic)
T(XL, 8, "※ 한국어판 엑셀 기준 · B8은 365에서 그대로 Enter(구버전은 Ctrl + Shift + Enter) · 해법은 GRG 비선형 · scipy 검산: 0 · 8.893 · 91.107%, 기대 23.67%, 하위 20% 경계 −15.00%")
set_px(sh(XL, 8), y=626); font(sh(XL, 8), 12)

# (거2) 12차 — 바닥은 CPPI, 위험 블록 안은 Das–Markowitz(3장 표 장표 복제)
FD = clone(S(3))
T(FD, 3, "바닥은 CPPI가, 위험 블록의 구성은 Das–Markowitz가 정합니다")
T(FD, 5, "Das–Markowitz는 미달 확률만 묶고 크기는 못 막는다 → 필수 목표는 floor로 · 숫자는 {p:V1} · {p:RC}장과 같다(A = 100, F = 80, m = 3)")
add_pic(FD, LAYER_PNG, 60, 196, 560, 560 * 500 / 1100)
tb = lecfix.find_table(FD)
lecfix.set_table(tb, [["", "Das–Markowitz 단독", "CPPI + Das–Markowitz"],
    ["무엇을 막나", "미달의 빈도(확률 ≤ α)", "빈도 + 크기(floor 아래로 안 감)"],
    ["비용", "—", "쿠션이 줄면 위험 블록 축소 · Cash Trap"],
    ["재조정", "정적 — 한 번 정한다", "동적 — 쿠션에 따라 E를 다시 계산"]])
col_widths(tb, (230, 420, 482)); table_font(tb, 14, 13); row_heights(tb, [32, 34, 34, 34]); set_px(tb, y=448); stripe(tb)
blk(FD, 630, 192, 590, [
    "**① CPPI:** C = A − F = 100 − 80 = 20, E = m × C = 3 × 20 = 60",
    "**② 위험 블록:** E 60 = Das–Markowitz 합산 (39.9 · 24.7 · 35.4%), 40 = GHP",
    "**• −20%:** E = 60 → 48, A = 40 + 48 = 88, C = 88 − 80 = 8",
    "     → E = 3 × 8 = 24로 줄인다 (24 매도)",
    "**• −40%:** E = 60 → 36, A = 76 < F = 80 → 뚫림",
    "**▶ 갭 위험:** 한 번에 1/m = 33%를 넘으면 → {p:V1} · {p:RC}장의 VaR로",
    "**※ 합산:** x = 1.645 × 19.38% − 13.30% = 18.6% → 1/x = 5.4 ≥ 3"], pt=14, before=2)
T(FD, 8, "얼마나 위험을 지나는 CPPI가, 무엇으로 지나는 Das–Markowitz가 — 바닥의 크기와 빈도를 함께 막습니다")
set_px(sh(FD, 8), y=592); font(sh(FD, 8), 16)
note(FD, 630, "※ 위험 블록 = 계정별 롱온리 해의 합산 · 갭 위험의 실제 크기는 방법 3으로 센다 — m = 3 · 매년 재조정이면 floor 위반 1.9%, 두꺼운 꼬리(t5)면 3.3%(gbi_mc.py)")


# (파5) 13차 — 방법 3 몬테카를로 네 장 · 숫자는 gbi_mc_results.json
MC = json.load(open(os.path.join(R, "W06_LDI와GBI", "W06_M8_실습_몬테카를로", "gbi_mc_results.json")))
G_ = {g["goal"]: g for g in MC["goals"]}
E_MC1 = render(r"\hat{p}=\frac{1}{N}\sum_{i=1}^{N}\mathbf{1}\left\{W_T^{(i)}\geq G\right\},\qquad \mathrm{SE}=\sqrt{\frac{\hat{p}\,(1-\hat{p})}{N}}", "mc1.png", fontsize=42)
MC1 = clone(S(16))
T(MC1, 3, "방법 3 — 몬테카를로는 경로를 수만 개 만들어 달성 확률을 직접 셉니다")
T(MC1, 5, "방법 1은 로그정규 · 고정 비중, 방법 2는 한 기간 · 정규 — 실무의 두꺼운 꼬리 · 납입과 인출 · CPPI 같은 동적 규칙 · 세금 · 수수료에는 닫힌 해가 없다")
T(MC1, 8, "원리 — 목표에 닿은 경로의 비율 = 달성 확률")
swap_eq(MC1, 9, E_MC1, maxw=820, maxh=88, cx=640, pad=24)
T(MC1, 11, "정확도 — N = 10,000, p = 70%"); T(MC1, 12, "√(0.7 × 0.3 ÷ 10,000) = ±0.46%p")
T(MC1, 14, "강의 실행 — N = 100,000"); T(MC1, 15, "±0.14%p · 1초 안에 끝난다")
T(MC1, 16, "‘로그정규는 손계산용 단순화, 실무는 몬테카를로’(부록 B-5) · ‘개인별 MC 1,000회’(단원 ⑦)가 바로 이 방법 3입니다")

MC2 = clone(S(3))
T(MC2, 3, "방법 3의 절차 — 가정 · 경로 · 규칙 · 집계 · K 찾기 · 검증")
T(MC2, 5, "같은 문제(55세 · 10억 · 세 목표 · 10년)를 경로로 푼다 — 검증의 기준은 방법 1의 닫힌 해")
tb = lecfix.find_table(MC2)
lecfix.set_table(tb, [["단계", "하는 일", "이 예에서"],
    ["① 가정", "수익률 분포 · 상관 · 현금흐름 · 운용 규칙을 적는다", "GHP 2% · 주식 λ 6% · σ 20%, 정규와 t분포(자유도 5)"],
    ["② 경로 생성", "난수 시드를 고정하고 연 수익률 경로를 만든다", "seed 2026 · 100,000 경로 × 10년"],
    ["③ 규칙 적용", "경로 하나를 한 해씩 굴린다 — 해마다 비중을 정하고 그 해 수익을 반영(다음 장 표)", "고정 비중(매년 되돌림) · CPPI(floor = 목표의 현가, m = 3)"],
    ["④ 집계", "달성 확률 · 미달했을 때의 평균 부족액을 센다", "P(W ≥ G) · 부족액"],
    ["⑤ K 찾기", "후보 K마다 같은 경로 전체에서 성공 비율을 세고, 비율 ≥ p인 가장 작은 K", "10경로 장난감 예 → 100,000경로"],
    ["⑥ 검증", "로그정규 · 고정 비중이면 방법 1 닫힌 해를 재현해야 한다 → 그다음 가정을 바꾼다", "닫힌 해 재현 → t분포로 교체"]])
col_widths(tb, (170, 540, 422)); table_font(tb, 14, 13); row_heights(tb, [34] + [46] * 6); stripe(tb)
T(MC2, 8, "⑥을 먼저 통과해야 ⑤의 숫자를 믿을 수 있습니다 — 재현하지 못하는 시뮬레이션은 쓰지 않는다"); set_px(sh(MC2, 8), y=590)

MC3 = clone(S(51))
T(MC3, 3, "MC는 닫힌 해를 재현하고, 꼬리의 비용은 갭 위험에서 드러납니다")
T(MC3, 5, f"seed {MC['seed']} · {MC['N']:,} 경로 · 10년 연 단위 · t분포는 자유도 5를 분산 1로 맞춰 같은 평균 · 변동성 — gbi_mc.py 실행 결과")
tb = lecfix.find_table(MC3)
def _r(g, name):
    n, t = g["normal"], g["t5"]
    return [name, f'{g["K_closed"]:.2f}억', f'{n["K_mc"]:.2f}억', f'{n["p_at_K_closed"]*100:.1f}% ± {n["se"]*100:.2f}%p', f'{t["K_mc"]:.2f}억']
lecfix.set_table(tb, [["목표 금액 · 기준 확률", "공식 K\n방법 1 식의 필요 자본", "MC K 정규\n몬테카를로로 찾은 K (방법 3)", "공식 K로 센 확률 (검증)\np에 붙으면 공식이 맞다", "MC K 두꺼운 꼬리\nt5로 바꿨을 때"],
                      _r(G_["Safety"], "Safety 6억 · 95%"), _r(G_["Market"], "Market 3억 · 70%"), _r(G_["Aspirational"], "Aspirational 5억 · 30%")])
col_widths(tb, (230, 210, 240, 250, 202)); table_font(tb, 15, 13); row_heights(tb, [48, 34, 34, 34]); set_px(tb, y=190)
cn, ct = MC["cppi"]["normal"], MC["cppi"]["t5"]
card(MC3, 69, 352, 556, 248, "필요 자본 K — 정규와 t5",
     ["• **정규:** K가 닫힌 해와 같다",
      "• **Market 2.26 vs 2.24:** 닫힌 해는 연속 재조정, MC는 매년 되돌림이라 생기는 차이",
      "• **Safety:** GHP뿐 — 확정이라 확률 100% · 표준오차 0",
      "• **t5:** K는 ±1% 안팎 — 10년 합산이 꼬리를 희석한다 (Market은 조금 줄고 Aspirational은 조금 는다)"], pt=15)
card(MC3, 641, 352, 560, 248, "CPPI — 자산 10억 · m = 3 · 매년 재조정",
     ["• **floor:** Safety 6억의 현가",
      f"• **floor 위반:** {cn['P_floor_breach']*100:.1f}% → t5 {ct['P_floor_breach']*100:.1f}%",
      f"• **6억 달성:** {cn['P_W_ge_6']*100:.1f}% → t5 {ct['P_W_ge_6']*100:.1f}%",
      f"• **9억 달성:** {cn['P_W_ge_9']*100:.1f}% · **14억 (세 목표 합):** {cn['P_W_ge_14']*100:.1f}%",
      f"• **중위:** {cn['median_W']:.1f}억",
      "▶ 동적 규칙의 확률은 방법 3으로만 센다"], dark=True, pt=15)
T(MC3, 8, "두꺼운 꼬리는 필요 자본보다 한 해 급락에 걸리는 갭 위험을 두 배로 만듭니다 — {p:V1} · {p:RC}장의 VaR 규칙이 필요한 이유")
set_px(sh(MC3, 8), y=606); font(sh(MC3, 8), 16)
note(MC3, 648, "※ 가정: 무위험 실질 2% · 주식 초과 6% · σ 20% · 혼합 m = r + wλ − ½w²σ² · Market 비중 w 67.08% · 연 단위 경로 — 월 단위 · 비정규 상관이면 차이가 더 클 수 있다")

MC4 = clone(S(63))
T(MC4, 3, "[도구] 방법 3 전체 — gbi_mc.py")
T(MC4, 5, "방법 1 재현 · 한 계좌 격자 · CPPI · --demo — 실습 A · B는 한 문제씩, 이 도구는 방법 3의 숫자를 모두 만든다")
T(MC4, 8, "검증 체크리스트"); T(MC4, 9, "넷을 통과해야 쓴다")
T(MC4, 10, "시드 고정 — 재현되나\nN↑ → 표준오차↓\n방법 1 닫힌 해 재현\n확률이 p에 붙나"); font(sh(MC4, 10), 14)
T(MC4, 14, "너는 GBI 몬테카를로 분석가다. numpy만 써서 gbi_mc.py를 만들어줘.\n"
           "1) 가정: 55세 · 10억 · 10년, GHP 실질 2%, 주식 초과 6% · 변동성 20%, seed 2026 · 100,000 경로(정규 · t5)\n"
           "2) 규칙 함수 두 개: 고정 비중(매년 되돌림), CPPI(floor = 목표의 현가, 쿠션 × m, 0과 자산 사이로 자름)\n"
           "3) 공통 난수: 모든 후보에 같은 경로를 쓴다\n"
           "4) 목표별 K: 후보 K마다 성공 비율을 세고 비율 ≥ p인 최소 K(이분법) — 경로별 최적 K를 고르지 말 것\n"
           "5) 한 계좌 격자: (총자본 K, 주식 비중 w) 후보마다 세 목표 확률을 동시에 세고, 조건을 다 만족하는 최소 K\n"
           "   목표가 많으면 후보는 우선순위 순서로 1차원씩(Safety 확정 → Market → Aspirational) — 실습 {p:CC}장\n"
           "6) --demo: 경로 하나 3년 표 · 경로 10개 장난감 예 · 격자 표 출력\n"
           "7) 검증: 정규면 방법 1 닫힌 해와 비교표 · 표준오차, 한국어 해석 세 줄")
font(sh(MC4, 14), 11)
p14m = sh(MC4, 14); set_px(p14m, h=196)
x14m, y14m, _, _ = lecfix.px(p14m)
for x in list(MC4.shapes):
    xx, yy, ww, hh = lecfix.px(x)
    if x.shape_type != 13 and not (x.has_text_frame and x.text_frame.text.strip()) and xx >= 400 and hh > 200: set_px(x, h=(y14m + 204) - yy)
rel_el = copy.deepcopy(lecfix.find_table(S(51))._element); MC4.shapes._spTree.append(rel_el); fresh_ids(MC4, [rel_el])
rel = next(x for x in MC4.shapes if x._element is rel_el); drop_col(rel, 4)
lecfix.set_table(rel, [["파일", "푸는 문제 · 방법", "강의 장", "프롬프트"],
    ["priority_mc.py", "방법 1 문제를 MC로 — 우선순위", "{p:CB2} · {p:CC}장", "실습_우선순위MC"],
    ["dm_mc.py", "방법 2 문제를 MC로 — Das–Markowitz", "{p:DMC1} · {p:CCB}장", "실습_DasMarkowitzMC"],
    ["gbi_mc.py", "방법 3 전체 — 방법 1 재현 · 격자 · CPPI", "{p:MC1}~{p:MC3}장", "이 장"],
    ["Colab 노트북", "같은 예제를 Colab 노트북으로", "{p:CC} · {p:CCB} · {p:MC4}장", "W06_M8_실습_Colab/"]])
col_widths(rel, (130, 310, 130, 207)); table_font(rel, 11, 11); row_heights(rel, [24, 24, 24, 24, 24]); set_px(rel, x=x14m - 2, y=y14m + 212)

# (파6) 14차 — 방법 3 초보자용 세 장 · 숫자는 gbi_mc_demo.json
DM = json.load(open(os.path.join(R, "W06_LDI와GBI", "W06_M8_실습_몬테카를로", "gbi_mc_demo.json")))
RA = clone(S(3))
T(RA, 3, "‘규칙을 경로에 적용한다’ = 경로 하나를 한 해씩 굴린다는 뜻입니다")
T(RA, 5, "seed 2026의 경로 1, 앞 3년(본 계산과 같은 난수) · 주식 연 수익률 R은 표대로, GHP는 해마다 2.02%")
tb = lecfix.find_table(RA)
fx = DM["A_fixed"]
lecfix.set_table(tb, [["고정 비중 — Market 몫, 주식 w = 67%", "시작 자산 → 주식 · GHP", "끝 자산 = 주식 × (1 + R) + GHP × 1.0202"]] +
    [[f'{a["year"]}년차 · R = {a["R"]*100:+.1f}%', f'{a["W0"]:.3f} → {a["stock"]:.3f} · {a["ghp"]:.3f}',
      f'{a["stock"]:.3f} × {1+a["R"]:.3f} + {a["ghp"]:.3f} × 1.020 = {a["W1"]:.3f}'] for a in fx])
col_widths(tb, (300, 330, 502)); table_font(tb, 14, 13); row_heights(tb, [32, 32, 32, 32]); set_px(tb, y=196); stripe(tb)
tb2_el = copy.deepcopy(tb._element); tb._element.addnext(tb2_el); fresh_ids(RA, [tb2_el])
tb2 = next(x for x in RA.shapes if x._element is tb2_el)
cp_ = DM["A_cppi"]
lecfix.set_table(tb2, [["CPPI — 10억, floor 9억의 현가, m = 3", "F = 9 × exp(−0.02 × (10 − t)) → C = W − F → E = 3 × C", "끝 자산 = E × (1 + R) + (W − E) × 1.0202"]] +
    [[f'{a["year"]}년차 · R = {a["R"]*100:+.1f}%', f'F {a["F"]:.2f} → C = {a["W0"]:.2f} − {a["F"]:.2f} = {a["C"]:.2f} → E = {a["E"]:.2f}',
      f'{a["E"]:.2f} × {1+a["R"]:.3f} + {a["ghp"]:.2f} × 1.020 = {a["W1"]:.2f}'] for a in cp_])
set_px(tb2, y=340); col_widths(tb2, (270, 470, 392))
for r_ in tb2.table.rows:
    for c in r_.cells:
        for p in c.text_frame.paragraphs:
            for rr in p.runs: rr.font.size = Pt(13)
T(RA, 8, "고정 비중은 해마다 같은 비율로 되돌리고, CPPI는 해마다 쿠션을 다시 재서 위험 노출을 바꿉니다")
set_px(sh(RA, 8), y=540); font(sh(RA, 8), 16)
note(RA, 586, "※ 이 일을 10년 × 100,000 경로 반복한 것이 방법 3 · Flexicure는 floor를 매년 자산의 일정 비율로 다시 잡는 CPPI · CPPI의 floor 9억(Safety + Market)은 설명용")
note(RA, 612, "※ 표의 exp(−0.02 × (10 − t))는 목표 금액을 남은 해만큼 2%로 할인한 값 = 목표의 현재가치")

KT = clone(S(51))
T(KT, 3, "K 찾기 — 후보 K마다 같은 경로 전체에서 성공 비율을 셉니다")
T(KT, 5, "장난감 예 — Market(G 3억 · p 70%) · 경로 10개 · 후보 K 2.0 · 2.2 · 2.4억 · 10년 뒤 자산 = K × 성장 배수")
tb = lecfix.find_table(KT); drop_col(tb, 4); drop_col(tb, 3); drop_col(tb, 2)
g10 = DM["B_paths"]
rows = [["경로 i (10년 성장 배수)", "K = 2.0 · 2.2 · 2.4억일 때 10년 뒤 자산(억) — 3억 이상이면 성공"]]
for i_, g in enumerate(g10):
    vals = " · ".join(f"{K * g:.2f}{' ✓' if K * g >= 3 else ''}" for K in (2.0, 2.2, 2.4))
    rows.append([f"경로 {i_ + 1}  ({g:.3f})", vals])
h_ = DM["B_hits"]
rows.append(["성공 수 → 센 달성 확률 (기준 70%)", " · ".join(f'{h_[k]}/10 = {h_[k]*10}% ' + ("✓" if h_[k] * 10 >= 70 else "미달") for k in ("2.0", "2.2", "2.4"))])
lecfix.set_table(tb, rows)
col_widths(tb, (330, 802)); table_font(tb, 11, 11); row_heights(tb, [24] + [23] * 11); set_px(tb, y=194)
blk(KT, 79, 504, 1132, [
    "**▶ 답:** 비율 ≥ 70%인 가장 작은 후보는 K = 2.2억 (80%)",
    f'**• 100,000 경로로 하면:** K = {DM["B_full"]["K_mc"]:.2f}억 (닫힌 해 {DM["B_full"]["K_closed"]:.2f}억)',
    "**• 이분법:** 후보 구간을 반씩 좁힌다 — 2.0 (실패)과 2.4 (성공) 사이 2.2를 보고",
    "     성공이면 [2.0, 2.2], 실패면 [2.2, 2.4]로 · 60번이면 충분"], pt=16, before=2)
T(KT, 8, "중요 — 모든 후보에 같은 경로를 쓴다(공통 난수). 그래야 후보끼리의 차이가 운이 아니라 K의 차이입니다")
set_px(sh(KT, 8), y=626); font(sh(KT, 8), 16)

MX = clone(S(51))
T(MX, 3, "다수결은 틀립니다 — 후보를 먼저 정하고 같은 경로로 셉니다")
T(MX, 5, "수강생 질문에 대한 답 · 목표가 여럿이면 (i) 계좌를 나눠 목표별로 K를 더하거나 (ii) 한 계좌의 후보 격자를 함께 센다")
remove(lecfix.find_table(MX))
card(MX, 69, 200, 490, 256, "✕ 틀린 방법 — 경로마다 다수결",
     ["경로마다 가장 좋은 K · 비중을 고르고 가장 많이 뽑힌 것을 택한다",
      "**사후 최적** — 미래 경로를 안 뒤 고른다",
      "**실행 불가** — 그 경로를 미리 알 수 없다",
      "**극단 해** — 오른 경로는 주식 100%, 내린 경로는 GHP 100%로 갈린다"], bullets=True)
card(MX, 575, 200, 626, 256, "✓ 올바른 방법 — 네 단계",
     ["① 후보(K · 비중의 조합)를 **미리** 정한다",
      "② 각 후보를 **같은 경로 전체**에 적용한다",
      "③ 후보마다 목표별 성공 비율과 부족액을 센다",
      "④ 확률 조건을 모두 만족하는 후보 중 **K가 가장 작은**(남는 자금이 가장 큰) 것을 고른다"], dark=True)
card(MX, 69, 470, 1132, 122, "목표가 여럿이면 — 두 길",
     ["(i) 계좌를 나누면 목표별로 따로 → 4.91 + 2.26 + 1.97 = **9.14억**",
      "(ii) 한 계좌로 함께 운용하면 후보 = (총자본 K, 주식 비중 w) 격자 → **{p:MXR}장**"])
T(MX, 8, "미래를 보고 고르지 말고, 미리 정한 후보를 같은 경로로 세어 고릅니다")
set_px(sh(MX, 8), y=608); font(sh(MX, 8), 18)

MXR = clone(S(51))
T(MXR, 3, "한 계좌 격자 읽는 법 — 후보 · 기준선 · 두 가지 %")
T(MXR, 5, "다음 장의 표는 후보 (K, w)마다 세 기준선에 닿을 확률을 센다 — 표를 보기 전에 넷을 정해 둔다")
remove(lecfix.find_table(MXR))
blk(MXR, 79, 192, 1132, ["**• 후보 K:** 이 계좌에 오늘 넣는 총자본",
                         "**• 후보 w:** 주식 비중 0~100% (나머지 GHP, 레버리지 없음)",
                         "**• 남는 자금:** 10 − K"], pt=16, before=3)
card(MXR, 69, 300, 344, 270, "기준선 — 쌓은 목표 금액",
     ["목표 금액 6 · 3 · 5억을 우선순위대로 쌓은 누적", "**Safety** 6억", "**Market** 6 + 3 = 9억", "**Aspirational** 9 + 5 = 14억"])
card(MXR, 427, 300, 384, 270, "두 가지 %",
     ["**열 이름의 %** = 지켜야 할 기준 (목표 확률 p)", "**표 안의 %** = 경로 100,000개 중 그 기준선을 넘은 비율 (센 달성 확률)", "**빨간 칸** = 기준 미달"])
card(MXR, 825, 300, 376, 270, "세 칸은 포함 관계",
     ["≥ 6억 ⊃ ≥ 9억 ⊃ ≥ 14억", "한 경로가 여러 칸에서 성공으로 세어진다", "▶ 합은 의미 없고 100%를 넘어도 정상", "▶ 늘 Safety ≥ Market ≥ Aspirational", "구간으로 나눈 표는 {p:MXP}장"], dark=True, pt=15)
T(MXR, 8, "후보마다 세 기준선과 각각 비교해, 셋 다 넘는 가장 싼 후보를 찾습니다"); set_px(sh(MXR, 8), y=592); font(sh(MXR, 8), 18)

MXG = clone(S(51))
T(MXG, 3, "한 계좌 격자 — 후보 (K, w)마다 세 기준선에 닿을 확률을 셉니다")
T(MXG, 5, "읽는 법은 {p:MXR}장 — 열 이름의 % = 기준, 표 안의 % = 센 달성 확률, 빨간 칸 = 기준 미달")
tb = lecfix.find_table(MXG)
_NM = ("Safety", "Market", "Aspirational"); _TH = (0.95, 0.70, 0.30)
def _gr(g, tag=""):
    ps = (g["pS"], g["pM"], g["pA"])
    miss = [nm for nm, p_, th in zip(_NM, ps, _TH) if p_ < th]
    return [f'{tag}K = {g["K"]:.1f}억 · w = {g["w"]*100:.0f}%', *[f'{p_*100:.1f}%' for p_ in ps],
            f'{g["left"]:.1f}억 → ' + ("통과" if not miss else "탈락 — " + " · ".join(miss))]
_rows = list(DM["D_grid"]) + [DM["D_edge"][0], DM["D_edge"][1]]
for g in _rows: assert g["pass"] == all(p_ >= th for p_, th in zip((g["pS"], g["pM"], g["pA"]), _TH))
lecfix.set_table(tb, [["후보(예시) — 총자본 K · 주식 w", "Safety 달성 확률\n자산 ≥ 6억 (기준 95%)", "Market 달성 확률\n자산 ≥ 9억 (기준 70%)",
                       "Aspirational 달성 확률\n자산 ≥ 14억 (기준 30%)", "남는 자금 → 통과? (미달 목표)"]] +
    [_gr(g) for g in DM["D_grid"]] + [_gr(DM["D_edge"][0], "최소 "), _gr(DM["D_edge"][1], "경계 ")])
col_widths(tb, (250, 190, 190, 210, 292)); table_font(tb, 15, 13); row_heights(tb, [50] + [33] * 8); set_px(tb, y=196)
_tt = tb.table
for ri, g in enumerate(_rows, 1):
    for ci, (p_, th) in enumerate(zip((g["pS"], g["pM"], g["pA"]), _TH), 1):
        if p_ < th:
            for p in _tt.cell(ri, ci).text_frame.paragraphs:
                for rr in p.runs: rr.font.color.rgb = RGBColor(0xC6, 0x28, 0x28); rr.font.bold = True
    if not g["pass"]:
        for p in _tt.cell(ri, 4).text_frame.paragraphs:
            for rr in p.runs: rr.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
T(MXG, 8, f'최소 통과 K = {DM["D_best"]["K"]:.1f}억 · w = {DM["D_best"]["w"]*100:.0f}% (7.0억은 Aspirational {DM["D_edge"][1]["pA"]*100:.1f}%로 탈락) — 나눈 계좌 9.14억보다 싸지만 Safety는 확률 95%')
set_px(sh(MXG, 8), y=532, h=60); font(sh(MXG, 8), 16)
note(MXG, 612, "※ 표의 앞 여섯 행은 예시 후보 · 최소 · 경계 두 행은 주식 비중 0~100%(5%p) × 총자본 6~10억(0.1억) 전 격자(21 × 41 = 861개)에서 찾은 값 · gbi_mc.py --demo")

_g = next(g for g in DM["D_grid"] if abs(g["K"] - 8.5) < 1e-9 and abs(g["w"] - 0.6) < 1e-9)
_a, _b, _c = (round(g * 100, 1) for g in (_g["pS"], _g["pM"], _g["pA"]))
_iv = [round(100 - _a, 1), round(_a - _b, 1), round(_b - _c, 1), _c]
assert abs(sum(_iv) - 100) < 1e-6, _iv
MXP = clone(S(51))
T(MXP, 3, "격자 표의 %를 더하면 100%를 넘습니다 — 그래도 정상입니다")
T(MXP, 5, f"같은 행(K = 8.5억 · w = 60%) {_a} · {_b} · {_c}%를 더하면 {_a + _b + _c:.1f}% — 세 칸이 서로 포함되는 사건이기 때문")
card(MXP, 69, 200, 600, 196, "격자 표 = 누적(≥ 기준선) 확률",
     ["≥ 14억인 경로는 ≥ 9억에도, ≥ 6억에도 들어간다",
      "→ 한 경로가 **세 칸 모두**에서 성공으로 세어진다",
      "→ 합은 의미가 없고, 늘 Safety ≥ Market ≥ Aspirational", "(오른쪽으로 갈수록 작다)"])
card(MXP, 685, 200, 516, 196, "아래 표 = 구간 확률",
     ["겹치지 않는 구간으로 나누면", "한 경로는 **한 칸에만** 들어간다", "→ 그때 합이 **100%**가 된다"], dark=True)
tb = lecfix.find_table(MXP)
lecfix.set_table(tb, [["10년 뒤 자산 구간", "6억 미만", "6 ~ 9억", "9 ~ 14억", "14억 이상"],
    ["뺄셈", f"100 − {_a}", f"{_a} − {_b}", f"{_b} − {_c}", "격자 표 그대로"],
    ["확률", *[f"{v:.1f}%" for v in _iv]]])
col_widths(tb, (260, 218, 218, 218, 218)); table_font(tb, 15, 15); row_heights(tb, [34, 40, 40]); set_px(tb, y=412); stripe(tb)
line(MXP, 536, f"합 = {_iv[0]:.1f} + {_iv[1]:.1f} + {_iv[2]:.1f} + {_iv[3]:.1f} = 100% · 격자 표가 쓰는 것은 누적(≥ 기준선) 확률, 이 표는 구간 확률", 15)
T(MXP, 8, "격자 표의 세 칸은 각각 따로 기준(95 · 70 · 30%)과 비교합니다 — 더하지 않습니다")
set_px(sh(MXP, 8), y=574); font(sh(MXP, 8), 17)
note(MXP, 620, "※ 숫자는 gbi_mc.py --demo의 같은 경로(K = 8.5억 · w = 60%)에서 센 값")


# (파7) 15차 — 조합 폭발과 줄이는 법 두 장 · 숫자는 gbi_mc_demo.json
CB1 = clone(S(51))
T(CB1, 3, "목표가 여럿이면 조합이 폭발합니다 — 1차원 탐색 몇 번으로 줄인다")
T(CB1, 5, "예 ① 손으로 세기 — 목표마다 K 후보 10개 × 주식 비중 w 후보 10개 = 100개, 한 계좌로 묶으면 목표 수만큼 곱해진다")
tb = lecfix.find_table(CB1); drop_col(tb, 4)
cnt_el = copy.deepcopy(tb._element); tb._element.addprevious(cnt_el); fresh_ids(CB1, [cnt_el])
cnt = next(x for x in CB1.shapes if x._element is cnt_el); drop_col(cnt, 3)
lecfix.set_table(cnt, [["목표 수 n", "한 계좌 격자 후보 수 = 100ⁿ", "계좌 분리 + 이분법 = 20 × n회"],
                       ["1", "100", "20"], ["2", "100 × 100 = 10,000", "40"],
                       ["3", "100 × 100 × 100 = 1,000,000 — 후보당 0.01초면 약 2.8시간", "60"],
                       ["5", "10의 10제곱 = 100억 — 경로 100,000을 곱하면 불가능", "100"]])
col_widths(cnt, (150, 640, 342)); table_font(cnt, 13, 12); row_heights(cnt, [28] + [27] * 4); set_px(cnt, y=190); stripe(cnt)
lecfix.set_table(tb, [["줄이는 법", "하는 일", "차원 · 후보 수", "강의 어디"],
    ["① 계좌 분리", "목표마다 따로 K를 구해 더한다", "목표당 1차원 · 약 20회", "방법 1 · 2와 같은 구조"],
    ["② 우선순위 순서", "Safety를 GHP로 확정 → 남은 자금에서 Market → Aspirational", "1차원 탐색 몇 번", "다음 장 예 ②"],
    ["③ 공통 난수", "경로는 한 번만 만들고 모든 후보에 재사용(행렬 연산 한 번)", "경로 생성 1회", "44장 · 다음 장 예 ③"],
    ["④ 격자 대신 최적화", "이분법 · Nelder–Mead 같은 미분 없는 최적화", "수십 회", "44장 이분법"],
    ["⑤ 변수 묶기", "(H, α) → γ 하나 · CPPI는 floor + 위험 블록 하나", "변수 1~2개", "예시 {p:VB} · {p:VB2}장"]])
col_widths(tb, (200, 470, 210, 252)); table_font(tb, 13, 12); row_heights(tb, [28] + [30] * 5); set_px(tb, y=346)
T(CB1, 8, "조합 폭발은 ‘한 계좌 + 모든 변수를 격자’일 때만 — 분리 · 순서 · 공통 난수 · 최적화로 줄입니다")
set_px(sh(CB1, 8), y=560); font(sh(CB1, 8), 16)
note(CB1, 600, "※ {p:MXG}장 격자는 변수를 둘(총자본 K, 주식 w)로 묶어 21 × 41 = 861개 · 해마다 비중을 바꾸는 동적 문제는 동적계획법(자산 수준 × 시점 격자를 거꾸로 푼다)")
note(CB1, 626, "※ Das · Ostrov · Radhakrishnan · Srivastav (2018) A New Approach to Goals-Based Wealth Management, JOIM 16(3) · 동적계획법: 같은 저자 (2020) CMS 17")

DE = json.load(open(os.path.join(R, "W06_LDI와GBI", "W06_M8_실습_몬테카를로", "gbi_mc_demo.json")))
EP, CR = DE["E_priority"], DE["F_crn"]
CB2 = clone(S(51))
T(CB2, 3, "우선순위 순서로 풀면 탐색 10번 — 대신 Safety는 확정입니다")
T(CB2, 5, "예 ② gbi_mc.py --demo 실행 결과 — 1단계 Safety → 2단계 Market(이분법 시도를 그대로) → 3단계 Aspirational")
tb = lecfix.find_table(CB2); drop_col(tb, 4); drop_col(tb, 3); drop_col(tb, 2)
tr = EP["trials"]
show = [0, 1, 2, 3, 7, len(tr) - 1]
rows = [["단계 · 시도", "후보 K → 3억에 닿은 비율(10만 경로)"],
        ["1단계 Safety 6억 · 95%", f'GHP로 확정 — K = 6 × exp(−0.2) = {EP["K_S"]:.2f}억, 탐색 0회 → 남은 자금 10 − {EP["K_S"]:.2f} = {EP["left_after_S"]:.2f}억']]
for i_ in show:
    t_ = tr[i_]
    rows.append([f"2단계 Market — 시도 {i_ + 1}", f'K = {t_["K"]:.3f}억 → {t_["p"]*100:.2f}% ' + ("≥ 70% → 위를 줄인다" if t_["p"] >= 0.70 else "< 70% → 아래를 올린다")])
rows.append(["3단계 Aspirational 5억 · 30%", f'남은 {EP["left_after_M"]:.2f}억이면 달성 {EP["pA_with_left"]*100:.1f}% ≥ 30% · 필요 K는 {EP["K_A"]:.2f}억 → 여유 {EP["spare"]:.2f}억'])
lecfix.set_table(tb, rows)
col_widths(tb, (330, 802)); table_font(tb, 13, 12); table_font(tb, 12, 12); row_heights(tb, [24] + [24] * (len(rows) - 1)); set_px(tb, y=186); stripe(tb)
for _r in tb.table.rows:
    for _c in _r.cells: _c.margin_top = _c.margin_bottom = Emu(18000)
card(CB2, 69, 410, 430, 204, "합계 — 우선순위 순서",
     [f'• **합계:** {EP["K_S"]:.2f} + {EP["K_M"]:.2f} + {EP["K_A"]:.2f} = {EP["total"]:.2f}억',
      f'• **탐색:** Market 이분법 {EP["searches"]}회 · {EP["seconds"]:.3f}초',
      "• **한 계좌 격자 ({p:MXG}장):** 7.1억 · w 60% — 약 2억 싸지만 Safety가 확률 95%"], pt=15)
card(CB2, 515, 410, 686, 204, "예 ③ 공통 난수",
     [f'• **묶음을 바꾸면:** 1만 경로 묶음에서 K 2.26억이 {CR["K226_setA"]*100:.2f} · {CR["K226_setB"]*100:.2f}%로 흔들려 K 2.24억 ({CR["K224_setA"]*100:.2f}%)보다 낮게도 나온다',
      "• **같은 경로를 쓰면:** 같은 후보는 언제나 같은 비율, 큰 K는 언제나 같거나 높은 비율",
      "▶ 후보 비교가 흔들리지 않는다"], dark=True, pt=15)
T(CB2, 8, "목표를 우선순위대로 하나씩 풀면 조합 대신 1차원 탐색이 목표 수만큼 — 확정할 것은 확정하고 남은 것만 확률로")
set_px(sh(CB2, 8), y=620); font(sh(CB2, 8), 16)
note(CB2, 658, "※ 이분법은 0과 남은 자금 사이를 반씩 좁혀 폭이 0.005억 아래가 될 때까지 · 표는 10번 중 여섯 번 · seed 2026 · 매년 비중 되돌림")


# (파8) 16차 — 47장을 직접: Claude Code 실습 장(63장 프롬프트 장표 복제)
import subprocess as _sp
_out = _sp.run([sys.executable, os.path.join(R, "W06_LDI와GBI", "W06_M8_실습_몬테카를로", "priority_mc.py")], capture_output=True, text=True, check=True).stdout.strip()
_out = re.sub(r"걸린 시간 [0-9.]+초", "걸린 시간 0.03초", _out)          # 실행마다 다른 시간은 대표값으로
OUT_PNG = code_png("$ python3 priority_mc.py\n" + _out, "prio_out.png", fontsize=12)
CC = clone(S(63))
T(CC, 3, "[실습 A] 방법 3으로 방법 1 문제 — 목표 여럿 · 우선순위({p:CB2}장)")
T(CC, 5, "W06_M8_실습_몬테카를로/ — priority_mc.py · 전문 W06_M8_실습_우선순위MC_ClaudeCode_프롬프트.md")
T(CC, 8, "검증 체크"); T(CC, 9, "넷이 맞으면 성공")
T(CC, 10, "Safety 4.91억\nMarket 2.26억 ± 0.01\n합계 9.14억\n시드를 바꿔도 둘째 자리만"); font(sh(CC, 10), 14)
T(CC, 12, "[입력] 프롬프트 요약 — 빈 폴더의 Claude Code에 전문을 붙여 넣는다")
T(CC, 14, "numpy만 써서 priority_mc.py를 만들고 실행해줘. 가정: 10억 · 10년 · 매년 되돌림, GHP 2%, 주식 초과 6% · σ 20%, seed 2026 · 100,000 경로(한 번만 생성)\n"
           "① Safety 6억은 식으로 확정 ② 남은 자금 안에서 Market 3억 · 70%의 K를 이분법 10번(시도마다 K · 비율 출력)\n"
           "③ Aspirational 5억 · 30%의 확률과 필요 몫 ④ 합계 · 시간 ⑤ 공통 난수 비교 — 4.91 · 2.26 · 9.14가 안 나오면 고쳐줘")
font(sh(CC, 14), 12)
p14 = sh(CC, 14); set_px(p14, h=96)
x14, y14, w14, _ = lecfix.px(p14)
for x in list(CC.shapes):          # 프롬프트 카드(배경 상자)를 줄인다
    xx, yy, ww, hh = lecfix.px(x)
    if x.shape_type != 13 and not (x.has_text_frame and x.text_frame.text.strip()) and xx >= 400 and hh > 200: set_px(x, h=(y14 + 104) - yy)
ow = 777; oh = ow * Image.open(OUT_PNG).size[1] / Image.open(OUT_PNG).size[0]
add_pic(CC, OUT_PNG, x14 - 2, y14 + 112, ow, min(oh, 640 - (y14 + 112)))
note(CC, 640, "※ 출력은 실제 실행 결과(걸린 시간만 대표값) · [실습 B]는 {p:CCB}장 · [도구] gbi_mc.py는 {p:MC4}장")


# (파9) 17차 — 변수 묶기(51장 표 장표 복제)
E_GV = render(r"g=\frac{\Sigma^{-1}\mathbf{1}}{C},\qquad v=\Sigma^{-1}\left(\mu-\frac{A}{C}\mathbf{1}\right),\qquad w(\gamma)=g+\frac{v}{\gamma}", "gv.png")
VB = clone(S(51))
T(VB, 3, "변수 묶기 ① — Das–Markowitz는 비중 6개 대신 γ 3개를 고릅니다")
T(VB, 5, "묶기 전에는 계좌마다 비중 둘을 직접 고른다(고위험 = 1 − 두 비중) · 묶은 뒤에는 효율선 위 한 점을 정하는 γ 하나만")
tb = lecfix.find_table(VB); drop_col(tb, 4)
lecfix.set_table(tb, [["계좌", "묶기 전 고르던 것", "후보(10칸씩)", "묶은 뒤 고를 것"],
    ["은퇴", "채권 비중 · 저위험 비중", "10 × 10", "γ 하나 (3.795)"],
    ["교육", "채권 비중 · 저위험 비중", "10 × 10", "γ 하나 (2.706)"],
    ["상속", "채권 비중 · 저위험 비중", "10 × 10", "γ 하나 (0.877)"],
    ["합", "6개", "10의 6제곱 = 1,000,000", "3개, 각각 1차원 방정식"]])
col_widths(tb, (150, 380, 300, 302)); table_font(tb, 14, 13); row_heights(tb, [30] + [30] * 4); set_px(tb, y=188); stripe(tb)
w_, h_ = size_px(E_GV, maxw=760, maxh=58); add_pic(VB, E_GV, 640 - w_ / 2, 352, w_, h_)
blk(VB, 79, 410, 1132, [
    "**※ 출처:** g, v는 ‘비중 합 = 1’만 둔 평균-분산 해 (공매도 허용 · 닫힌 해, {p:M2D}장) — 제약 하의 최적화를 한 번 풀어 둔 결과",
    "**① g = 최소분산 포트폴리오:** Σ⁻¹1 = (400, 23.96, 2.08), C = 426.04 → g = Σ⁻¹1 ÷ C = (0.939, 0.056, 0.005)",
    "**② v = 위험을 더 지는 방향:** A/C = 0.0538 → μ − (A/C)1 = (−0.0038, 0.0462, 0.1962)",
    "     → v = Σ⁻¹(μ − (A/C)1) = (−1.516, 0.795, 0.721)",
    "**③ 은퇴 γ = 3.795:** w = (0.939, 0.056, 0.005) + (−1.516, 0.795, 0.721) ÷ 3.795",
    "     = **53.9 · 26.6 · 19.5%**"], pt=15, before=3)
T(VB, 8, "비중 셋을 고르지 않고 γ 하나를 고르면, 그 γ가 비중을 모두 정합니다"); set_px(sh(VB, 8), y=612); font(sh(VB, 8), 17)

VB2 = clone(S(51))
T(VB2, 3, "변수 묶기 ② — CPPI는 해마다의 비중 대신 승수 m 하나를 고릅니다")
T(VB2, 5, "묶기 전에는 매년 비중을 직접 고른다(연도 10 × 비중 2 = 20개) · 묶은 뒤에는 floor는 식, 위험 노출은 규칙이 계산")
card(VB2, 69, 200, 364, 200, "floor F — 식이 정한다",
     ["필수 목표의 현재가치", "식으로 정해짐 — 고를 것 아님", "위험 블록의 구성은 Das–Markowitz 합산으로 **한 번**"])
card(VB2, 449, 200, 364, 200, "위험 노출 E — 규칙이 계산",
     ["해마다 E = m × (W − F)", "고를 것은 **승수 m 하나**", "({p:V1}장 VaR로 상한)"])
card(VB2, 829, 200, 372, 200, "예 — A = 100 · F = 80 · m = 3",
     ["1년차 E = 3 × (100 − 80) = **60**", "위험 블록 −20% → A = 88", "2년차 E = 3 × (88 − 80) = **24**"], dark=True)
tb = lecfix.find_table(VB2)
lecfix.set_table(tb, [["무엇을 묶나", "묶기 전 고를 것", "묶은 뒤 고를 것", "무엇이 대신 정하나", "강의 어디"],
    ["계좌별 비중 (Das–Markowitz)", "6개 · 10의 6제곱 후보", "γ 3개 · 각각 1차원", "효율선 식 g + v/γ", "{p:VB} · {p:M2D}장"],
    ["해마다의 비중 (CPPI)", "20개 이상", "m 1개 + 위험 블록 구성 1번", "floor 식 + 규칙 E = m × C", "{p:FD} · {p:V1}~{p:RC}장"]])
col_widths(tb, (300, 200, 240, 240, 152)); table_font(tb, 14, 13); row_heights(tb, [32, 40, 40]); set_px(tb, y=424)
T(VB2, 8, "매년 비중을 ‘고르지’ 않고 규칙이 ‘계산’합니다 — 그래서 목표가 늘어도 고를 숫자는 거의 늘지 않습니다")
set_px(sh(VB2, 8), y=560); font(sh(VB2, 8), 17)


# (파10) 18차 — 방법 3으로 방법 2 문제를 푼다(51장 표 장표 복제) · 숫자는 dm_mc_results.json
DQ = json.load(open(os.path.join(R, "W06_LDI와GBI", "W06_M8_실습_몬테카를로", "dm_mc_results.json")))
A_ = DQ["accounts"]
def _w(v): return " · ".join(f"{x:.1f}" for x in v)
DMC1 = clone(S(51))
T(DMC1, 3, "몬테카를로로 방법 2 문제를 푼다 — 비중 후보마다 미달 비율을 셉니다")
T(DMC1, 5, "방법 1 문제(목표별 K)는 {p:KT} · {p:MC3}장에서 풀었다(닫힌 해 2.24억 → MC 2.26억) — 같은 일을 Das–Markowitz 계좌에")
card(DMC1, 69, 194, 1132, 172, "절차 — 네 단계",
     [f"① 1년 수익률 표본 **{DQ['N']:,}개**를 한 번만(원문 자산 셋, seed {DQ['seed']})",
      "② 후보 비중을 정한다 — 격자 또는 효율선 위 γ 한 축",
      "③ 후보마다 기대수익(μ로 바로)과 미달 비율(표본 중 R < H인 비율 — **정규식이 아니라 센다**)",
      "④ 미달 비율 ≤ α인 후보 중 **기대수익 최대**"], pt=15)
for p in DMC1.shapes[-1].text_frame.paragraphs: p.line_spacing = Pt(19); p.space_before = Pt(2)
tb = lecfix.find_table(DMC1)
rows = [["계좌 (H, α)", "닫힌 해 (방법 2)", "MC 격자 (17,565 후보)", "MC γ 한 축 (400 후보)", "MC 롱온리 격자"]]
for nm in ("은퇴", "교육", "상속"):
    r_ = A_[nm]; n_ = r_["normal"]
    rows.append([f'{nm} ({r_["H"]*100:.0f}%, {r_["alpha"]*100:.0f}%)', _w(r_["closed"]), _w(n_["grid"]["w"]),
                 _w(n_["gamma_line"]["w"]) + f' (γ {n_["gamma_line"]["gamma"]:.2f})', _w(n_["long_only"]["w"])])
lecfix.set_table(tb, rows)
col_widths(tb, (200, 220, 230, 270, 212)); table_font(tb, 13, 12); row_heights(tb, [30, 32, 32, 32]); set_px(tb, y=374)
se5 = A_["은퇴"]["normal"]["se_miss"]
blk(DMC1, 79, 508, 1132, [
    f"**• 차이:** 목적 (기대수익)이 제약 근처에서 평평해, 미달 비율의 표본오차 (α 5%에서 ±{se5*100:.2f}%p)만으로 격자 해가 2~3%p 움직인다",
    "**• 변수 묶기 ({p:VB}장):** 효율선 위 γ 한 축으로 찾으면 후보 400개로 닫힌 해와 0.1~2%p 안",
    f"**• 롱온리 상속:** {_w(A_['상속']['normal']['long_only']['w'])} (닫힌 식 풀이 0 · 8.9 · 91.1)"], pt=15, before=3)
T(DMC1, 8, "같은 문제를 식 대신 세기로 — 가정이 바뀌어도 그대로 쓸 수 있습니다(다음 장)")
set_px(sh(DMC1, 8), y=608); font(sh(DMC1, 8), 16)
note(DMC1, 646, f"※ dm_mc.py — numpy만 · {DQ['seconds']:.0f}초 · 격자는 −100~200% 2%p 간격 뒤 근처 0.25%p로 다시 · 롱온리는 0~100% · γ는 0.3~30 사이 400개")

DMC2 = clone(S(51))
T(DMC2, 3, "두꺼운 꼬리로 바꾸면 — 닫힌 해가 못 하는 일을 세기로 합니다")
T(DMC2, 5, "같은 평균 · 공분산의 다변량 t(자유도 5)로 표본만 바꾸고 같은 (H, α)를 지키는 비중을 다시 센다(γ 한 축)")
tb = lecfix.find_table(DMC2); drop_col(tb, 4)
TC = DQ["tail_check"]
rows = [["계좌 (H, α)", "정규 — 채권 · 저위험 · 고위험", "t5 — 채권 · 저위험 · 고위험", "읽는 법"]]
for nm in ("은퇴", "교육", "상속"):
    r_ = A_[nm]
    rows.append([f'{nm} ({r_["H"]*100:.0f}%, {r_["alpha"]*100:.0f}%)', _w(r_["normal"]["gamma_line"]["w"]), _w(r_["t5"]["gamma_line"]["w"]), "채권이 준다 — 덜 보수적"])
rows.append([f'꼬리 점검 ({TC["H"]*100:.0f}%, {TC["alpha"]*100:.0f}%)', _w(TC["normal"]["w"]), _w(TC["t5"]["w"]), "채권이 는다 — 더 보수적"])
lecfix.set_table(tb, rows)
col_widths(tb, (230, 320, 320, 262)); table_font(tb, 14, 13); row_heights(tb, [32] + [38] * 4); set_px(tb, y=196)
blk(DMC2, 79, 396, 1132, [
    "**• 왜:** 분산을 맞춘 t는 가운데가 뾰족하고 어깨 (5~20%)가 얇다",
    "**• 5% 분위수:** −1.56σ (정규 −1.645σ)",
    "**• 1% 분위수:** −2.61σ (정규 −2.33σ)",
    "**▶ 그래서:** α가 5~20%면 같은 (H, α)가 오히려 위험을 더 허용하고, α가 1%처럼 작으면 더 보수적으로 바뀐다",
    "**▶ 결론:** 어떤 α를 쓰는지가 결론을 바꾼다"], pt=16, before=4)
T(DMC2, 8, "닫힌 해는 정규를 가정해야 나오지만, 세기는 분포만 바꾸면 그대로 — 방법 3이 방법 1 · 2를 검증하고 넓히는 방법입니다")
set_px(sh(DMC2, 8), y=596); font(sh(DMC2, 8), 16)
note(DMC2, 640, "※ t 표본 = 정규 × √(3/5) ÷ √(카이제곱(5)/5) — 분산 1로 맞춘 자유도 5 · 공분산은 원문과 같다 · seed 2026 · 200,000 표본")


# (파11) 19차 — [실습 B] dm_mc.py(63장 프롬프트 장표 복제) · 단원 ③ 지도(51장 표 복제)
_ao = DQ["accounts"]
_txt = "$ python3 dm_mc.py\n" + "\n".join(
    f'{nm} (H {a_["H"]*100:.0f}%, α {a_["alpha"]*100:.0f}%) 닫힌 해 {_w(a_["closed"])} | γ 한 축 {_w(a_["normal"]["gamma_line"]["w"])} (γ {a_["normal"]["gamma_line"]["gamma"]:.2f})'
    for nm, a_ in _ao.items()) + "\n롱온리  " + " | ".join(f'{nm} {_w(a_["normal"]["long_only"]["w"])}' for nm, a_ in _ao.items()) + \
    "\nt5 γ 한 축  " + " | ".join(f'{nm} {_w(a_["t5"]["gamma_line"]["w"])}' for nm, a_ in _ao.items()) + \
    f'\n꼬리 점검 (−20%, 1%): 정규 {_w(DQ["tail_check"]["normal"]["w"])} · t5 {_w(DQ["tail_check"]["t5"]["w"])}\n초 (약 {DQ["seconds"]:.0f})'
OUTB = code_png(_txt, "dm_out.png", fontsize=11)
CCB = clone(S(63))
T(CCB, 3, "[실습 B] 방법 3으로 방법 2 문제 — Das–Markowitz({p:DMC1}장)")
T(CCB, 5, "W06_M8_실습_몬테카를로/ — dm_mc.py · 전문 W06_M8_실습_DasMarkowitzMC_ClaudeCode_프롬프트.md")
T(CCB, 8, "검증 체크"); T(CCB, 9, "넷이 맞으면 성공")
T(CCB, 10, "γ 한 축이 닫힌 해 ±1%p\n(상속 ±2%p)\n롱온리 상속 0 · 8.9 · 91.1 (±1%p)\n확률은 세기(정규식 금지)"); font(sh(CCB, 10), 14)
T(CCB, 12, "[입력] 프롬프트 요약 — 빈 폴더의 Claude Code에 전문을 붙여 넣는다")
T(CCB, 14, "numpy만 써서 dm_mc.py를 만들고 실행해줘. 자산 셋(μ 5 · 10 · 25%, Σ)의 1년 수익률 200,000개(seed 2026)\n"
            "① 후보: 격자(공매도 허용 · 롱온리) + 효율선 γ 한 축 400개 ② 기대수익은 μ로, 미달 비율은 표본에서 센다\n"
            "③ 미달 ≤ α 중 기대수익 최대 ④ 다변량 t(자유도 5)로 다시 ⑤ 닫힌 해와 비교표")
font(sh(CCB, 14), 12)
p14b = sh(CCB, 14); set_px(p14b, h=96)
x14b, y14b, _, _ = lecfix.px(p14b)
for x in list(CCB.shapes):
    xx, yy, ww, hh = lecfix.px(x)
    if x.shape_type != 13 and not (x.has_text_frame and x.text_frame.text.strip()) and xx >= 400 and hh > 200: set_px(x, h=(y14b + 104) - yy)
owb = 777; ohb = owb * Image.open(OUTB).size[1] / Image.open(OUTB).size[0]
add_pic(CCB, OUTB, x14b - 2, y14b + 112, owb, min(ohb, 640 - (y14b + 112)))
note(CCB, 640, "※ 출력은 dm_mc_results.json의 실제 값 · [실습 A]는 {p:CC}장 · [도구] gbi_mc.py는 {p:MC4}장")

U3 = clone(S(51))
T(U3, 3, "단원 ③ 지도 — 같은 문제를 세 방법으로, 그다음 조합 · 실습 · 비교")
T(U3, 5, "질문은 하나: 목표 G를 확률 p로 이루려면 오늘 얼마를, 무엇에 떼어 두는가 · 소주제마다 머리 띠에 이름을 적었다")
tb = lecfix.find_table(U3); drop_col(tb, 4); drop_col(tb, 3)
lecfix.set_table(tb, [["소주제 (머리 띠 이름)", "푸는 질문", "장"],
    ["필요 자본 기초", "GHP · PSP, 운의 크기, 필요 자본 K 공식", "{p:o15}~{p:o22}장"],
    ["방법 1 가장 싼 K", "목표마다 확률 조건을 지키는 가장 싼 K — 목적식 · 손계산 · 최적 비중", "{p:P23}~{p:M1}장"],
    ["방법 2 Das–Markowitz", "계좌마다 (H, α) → 효율선 위 한 점 — 손계산 · 행렬 · 세 계정", "{p:DM1}~{p:DMB}장"],
    ["롱온리 · 합산", "음수 자산을 0에 묶어 다시 푼다 · 해 찾기 · 합쳐도 효율", "{p:LO}~{p:DM2}장"],
    ["방법 3 몬테카를로", "경로를 세어 확률을 직접 — 원리 · 절차 · K 찾기 · 한 계좌 격자 · 결과", "{p:MC1}~{p:MC3}장"],
    ["조합 · 묶기", "목표가 여럿일 때 조합 폭발을 줄이는 법 · 변수 묶기 · 우선순위 순서", "{p:CB1}~{p:CB2}장"],
    ["방법 3으로 방법 2", "Das–Markowitz 문제를 세기로 · 두꺼운 꼬리", "{p:DMC1}~{p:DMC2}장"],
    ["실습 · 도구", "[실습 A] priority_mc.py · [실습 B] dm_mc.py · [도구] gbi_mc.py", "{p:CC} · {p:CCB} · {p:MC4}장"],
    ["세 방법 비교 · 정리", "비교표 · 방법 2가 편한 이유 · 바닥 + 계정 · 체크포인트", "{p:CMP2}~{p:o26}장"]])
col_widths(tb, (260, 662, 210)); table_font(tb, 14, 13); row_heights(tb, [30] + [34] * 9); set_px(tb, y=186); stripe(tb)
rich(line(U3, 534, "**파일** — 스크립트 · 프롬프트는 W06_M8_실습_몬테카를로/, 같은 예제의 Colab 노트북은 W06_M8_실습_Colab/", 16))
T(U3, 8, "막히면 이 지도로 — 어느 방법의 어느 단계인지 머리 띠에서 확인합니다"); set_px(sh(U3, 8), y=584); font(sh(U3, 8), 17)

# (하) 단원 ③ — 방법 1 vs 방법 2(3장 표 장표 복제)
CMP2 = clone(S(51))
T(CMP2, 3, "같은 문제, 세 풀이 — K를, 비중을, 확률을 셉니다")
T(CMP2, 5, "숫자는 각 방법의 강의 예제(방법 1 · 3: 55세 · 10억 · 세 목표, 방법 2: 원문 가상 자산 · 세 계좌)")
tb = lecfix.find_table(CMP2); drop_col(tb, 4)
lecfix.set_table(tb, [["구분", "방법 1 — 가장 싼 필요 자본", "방법 2 — Das–Markowitz", "방법 3 — 몬테카를로"],
    ["묻는 것", "목표 G · 시점 T · 확률 p", "최소 수익률 H · 확률 한도 α", "목표 · 확률 + 현금흐름 · 운용 규칙"],
    ["가정", "로그정규 · 고정 비중", "한 기간 · 정규 · 고정 비중", "아무 분포 · 동적 규칙 가능"],
    ["답의 형태", "목표별 K → 49 / 22 / 28%", "계좌별 비중 · γ 3.795 · 2.706 · 0.877", "경로를 세어 낸 확률 · K · 부족액"],
    ["답", " ", " ", "정규면 방법 1과 같은 K · t5면 갭 위험 2배"],
    ["장점", "손으로 풀린다 · 직관", "묻기 쉽다 · MVO 도구 그대로", "현실 그대로 · 검증 가능"],
    ["한계", "꼬리 · 현금흐름 못 담음", "미달 크기 못 막음 · 한 기간", "계산량 · 가정 의존 · 재현 관리"],
    ["언제 쓰나", "상담 첫 장 · 손계산 점검", "계좌별 배분 설계", "상품 운용 · 공시 · IC 검증"],
    ["이론 계보", "Kataoka(1963)형", "Telser(1956)형", "Roy(1952) 확률을 직접 계산"]])
col_widths(tb, (150, 320, 330, 332)); table_font(tb, 13, 12); row_heights(tb, [30] + [35] * 8); stripe(tb)
T(CMP2, 8, "방법 1 · 2는 확률을 평균 · 변동성 조건으로 번역하고, 방법 3은 그 번역이 맞는지 경로로 확인합니다"); set_px(sh(CMP2, 8), y=596); font(sh(CMP2, 8), 16)

# (거) 단원 ⑧ — H대 판정 기준 ↔ 강의 단원(51장 표 복제, 4열)
MAP = clone(S(51))
T(MAP, 3, "H대의 판정 기준마다 쓰는 강의 단원이 정해져 있습니다")
T(MAP, 5, "케이스 · 과제(4팀)의 기준을 강의 도구로 — 중심은 Flexicure: Floor · GHP · 쿠션과 승수를 비유동성까지 옮긴다")
tb = lecfix.find_table(MAP); drop_col(tb, 4)
lecfix.set_table(tb, [["판정 기준", "묻는 것 · 통과 조건", "쓰는 강의 단원", "그 단원의 도구"],
    ["① 버틸 수 있는 기간", "충격 직후 사모를 팔지 않고 유동자산만으로 지출 · 출자 요청 · 공사비 — 10년(락업) 이상", "⑤⑥ 쿠션 · Cash Trap", "쿠션과 승수 — 사모는 팔 수 없다"],
    ["② 기금 가치 유지", "30년 실질가치를 유지할 확률 50% 이상", "① 위험 = P(W < G) · ③ 확률 제약", "미달 확률 · 필요 자본 · 지출 상한"],
    ["③ 운용 역량", "직접 투자는 인력 3명 · 펀드 15개(1,500억) 이상", "⑧ 모방 3장벽", "접근권 · 시간 지평 · 인력"],
    ["④ 위기 직후 현금", "충격 직후 현금 · 국채만으로 3년 지출 + 공사비 600억", "④ Floor · GHP", "Floor = 거절할 수 없는 유출, GHP로 덮는다"],
    ["지출률 z", "고정 3.5%는 지키고 그 위를 얼마까지 쓰나", "③ 필요 자본 · ⑤ Flexicure 재설정", "필수 · 희망 목표, 매년 재설정(80/20 평활)"]])
col_widths(tb, (190, 430, 250, 262)); table_font(tb, 14, 13)
T(MAP, 8, "다음 두 장 — Flexicure를 H대에 옮기는 다섯 단계 (과제 데이터 기준 교육용 숫자)"); set_px(sh(MAP, 8), y=596)

# (너) 단원 ⑧ — Flexicure 다섯 단계 ①~③(51장 표 복제, 5열)
FX1 = clone(S(51))
T(FX1, 3, "Flexicure를 H대에 옮기는 다섯 단계 ① — Floor · GHP · 쿠션과 승수")
T(FX1, 5, "단위 억 원 · 배분 주식/국채/현금/사모 — A 65/30/5/0 · z 4.0%, B 50/25/5/20 · z 4.5%(기본 답), C 55/10/5/30 · z 5.25%")
tb = lecfix.find_table(FX1)
lecfix.set_table(tb, [["단계", "하는 일", "A안 보수형", "B안 균형형", "C안 Yale 추종형"],
    ["① Floor", "거절할 수 없는 유출 = 앞 3년 고정 지출(175 + 178.5 + 182.1 = 536) + 공사비 600 + 남은 출자 요청(사모 목표 × 40%)", "536 + 600 + 0\n= 1,136", "536 + 600 + 400\n= 1,536", "536 + 600 + 600\n= 1,736"],
    ["② GHP로 덮기", "2022형 충격 뒤 현금 · 국채가 3년 고정 지출 + 공사비 1,136을 판매 없이 댄다 (기준 ④)", "1,530 — 통과", "1,318 — 통과", "680 — 미달"],
    ["③ 쿠션과 승수", "쿠션 = 유동자산(주식 + 국채 + 현금) − Floor, 승수 m = (주식 + 사모) ÷ 쿠션, 상한 m ≤ 2 (디폴트옵션 기본 답과 같다)", "쿠션 3,864\nm 0.84", "쿠션 2,464\nm 1.42", "쿠션 1,764\nm 2.41 — 출발부터 초과"]])
col_widths(tb, (150, 470, 160, 160, 192)); table_font(tb, 15, 14); row_heights(tb, [36, 104, 70, 104])
note(FX1, 630, "※ 이 표의 확률(기준 ② 30년 실질가치 유지 · ④ 3년 지출)은 방법 3(몬테카를로)으로 센다 — 과제 도구의 4,000 경로가 그것")
T(FX1, 8, "①~③은 디폴트옵션과 같습니다 — 다른 것은 사모를 팔 수 없다는 점 (다음 장 ④)"); set_px(sh(FX1, 8), y=596)

# (더) 단원 ⑧ — 다섯 단계 ④⑤(51장 표 복제, 3열)
FX2 = clone(S(51))
T(FX2, 3, "다섯 단계 ② — 사모는 팔 수 없으니 충격 뒤 쿠션으로 다시 잽니다")
T(FX2, 5, "④ 비유동성 — 사모는 못 팔고 줄이는 수단은 신규 약정 중단뿐 · 2008형: 주식 −40 · 국채 +8 · 현금 +3 · 사모 −10%")
tb = lecfix.find_table(FX2); drop_col(tb, 4); drop_col(tb, 3)
lecfix.set_table(tb, [["2008형 충격 뒤 (억 원)", "B안 (기본 답)", "C안"],
    ["유동자산 (주식 + 국채 + 현금)", "1,500 + 1,350 + 258 = 3,108", "1,650 + 540 + 258 = 2,448"],
    ["쿠션 = 유동자산 − Floor", "3,108 − 1,536 = 1,572", "2,448 − 1,736 = 712"],
    ["승수 m = (주식 + 사모) ÷ 쿠션", "(1,500 + 900) ÷ 1,572 = 1.53 — 상한 2 안", "(1,650 + 1,350) ÷ 712 = 4.2 — 초과"],
    ["주식을 다 팔아도 (사모 ÷ 쿠션)", "900 ÷ 1,572 = 0.57", "1,350 ÷ 712 = 1.9 — 사실상 갇힘"],
    ["사모 비중 (분모 효과)", "22.5% — 목표 20% + 2%p 초과", "35.5% — 목표 30% 초과"]])
col_widths(tb, (360, 386, 386)); table_font(tb, 16, 14); row_heights(tb, [40] + [52] * 5)
r1 = copy_shape(sh(S(19), 14), FX2); set_px(r1, y=546, h=34); font(r1, 15)
lecfix.set_text(r1, "규칙 ④ — 비유동 상한 x = 최악 충격 뒤에도 m ≤ 2를 지키는 최대치 · 사모가 목표 + 2%p를 넘으면 신규 약정 중단")
T(FX2, 8, "⑤ 지출 — 고정 3.5%는 Floor, z까지는 희망 몫: 매년 쿠션을 다시 재고 기준 아래면 희망 몫 동결(80/20 평활이 완충)")
set_px(sh(FX2, 8), y=586); font(sh(FX2, 8), 16)


# (러) 부록 B-6 — Das–Markowitz 원 논문 예(보충교재 05 4-2절) 세 장
SRC_NOTE = "출처: Das · Markowitz · Scheid · Statman (2010), JFQA 45(2), pp. 311–334 · 식 (1)–(11) · 표 1–2 · 가능성 점검은 식 (10)–(19)의 닫힌 해 · 롱온리는 §V(식 (20)–(29), pp. 328–329)"
def src_note(slide):
    n = copy_shape(sh(S(21), 8), slide); set_px(n, x=79, y=648, w=1000, h=24); lecfix.set_text(n, SRC_NOTE); font(n, 11)
B6A = clone(S(24))
T(B6A, 3, "원 논문 예는 가상 자산 셋과 계정별 (H, α)로 시작합니다")
T(B6A, 5, "부록 B-6 ① — Das · Markowitz · Scheid · Statman(2010) 식 (1)–(11) · 표 1–2 재계산 · 무위험 자산 없음 · 공매도 허용")
T(B6A, 8, "계정 하나의 문제 — 미달 확률 한도 안에서 기대수익률을 최대로")
swap_eq(B6A, 9, E_DM1, maxw=700, maxh=52, cx=640)
T(B6A, 11, "정규분포면 확률 제약 = 평균 · 표준편차 조건 (α 5%이면 역함수 값 −1.645)")
swap_eq(B6A, 12, E_B2, maxw=900, maxh=66, cx=640)
tb = lecfix.find_table(B6A)
lecfix.set_table(tb, [["자산 (원문의 가상 자산)", "기대수익률 · 표준편차", "상관"],
                      ["채권형", "5% · 5%", "두 주식형과 0"],
                      ["저위험 주식형", "10% · 20%", "고위험 주식형과 0.20 (공분산 0.02)"],
                      ["고위험 주식형", "25% · 50%", "저위험 주식형과 0.20"]])
col_widths(tb, (330, 360, 442)); src_note(B6A)

B6H = clone(S(3))
T(B6H, 3, "먼저 손으로 — 한 포트폴리오를 검사하고, 두 자산으로 줄여 직접 풉니다")
T(B6H, 5, "부록 B-6 ② — 은퇴 계정(H = −10%, α = 5%) · 보충교재 4-2 단계 1의 확인과 단계 1-2")
tb = lecfix.find_table(B6H)
lecfix.set_table(tb, [["손계산", "계산", "결과"],
    ["검사 — 채권 60 · 저위험 30 · 고위험 10", " ", "하위 5% −6.40% ≥ −10% 통과 (미달 2.05%)"],
    ["두 자산 축소 — 채권형 + 저위험 주식형", " ", "w* = 52.2% · m 7.61% · σ 10.70%"],
    ["n개 자산 — 행렬로 일반화", "같은 일을 비중 셋으로 → 다음 장", "γ = 3.795 · 53.9 · 26.6 · 19.5%"]])
col_widths(tb, (330, 470, 332)); table_font(tb, 14, 13); row_heights(tb, [36] + [62] * 3); stripe(tb)
for ri, tex, nm in ((1, r"$m=8.50\%,\ \ \sigma=\sqrt{0.0082}=9.06\%$", "b6h_1.png"),
                    (2, r"$(15\%+5\%\,w)^2=1.645^2\,[(1-w)^2\,0.05^2+w^2\,0.20^2]$", "b6h_2.png")):
    eq_in_cell(B6H, tb, ri, 1, tex, nm, fontsize=22)
T(B6H, 8, "손으로 한 일과 행렬로 한 일은 같다 — 비중이 늘어 식을 묶을 뿐"); set_px(sh(B6H, 8), y=520)
src_note(B6H)

B6B = clone(S(24))
T(B6B, 3, "효율선을 γ로 쓰고 제약을 등호로 거는 γ를 찾으면 원문 표가 재현됩니다")
T(B6B, 5, "부록 B-6 ③ — 원문 각주 6: A · B(= μ′Σ⁻¹μ = 1.4167) · C · D(= BC − A² = 78.39), 여기 K = B − A²/C = D/C는 강의 표기 · 원문 p.320의 3.9750은 오기(p.317 · 표 1 · 2는 3.7950)")
T(B6B, 8, "효율선의 상수와 γ의 식 — 평균 m과 분산이 γ 하나로 정리된다")
swap_eq(B6B, 9, E_B3, maxw=1080, maxh=56, cx=640)
T(B6B, 11, "확률 제약을 등호로 — γ 하나의 방정식")
swap_eq(B6B, 12, E_B4, maxw=900, maxh=66, cx=640)
tb = lecfix.find_table(B6B)
lecfix.set_table(tb, [["계정 (H, α)", "γ", "비중 채권 · 저위험 · 고위험 → 기대수익률 · 표준편차"],
                      ["은퇴 (−10%, 5%)", "3.795", "53.9 · 26.6 · 19.5% → 10.23% · 12.30%"],
                      ["교육 (−5%, 15%)", "2.706", "37.9 · 35.0 · 27.1% → 12.18% · 16.57%"],
                      ["상속 (−15%, 20%)", "0.877", "−78.9 · 96.2 · 82.7% (채권 공매도 = 차입) → 26.35% · 49.13%"]])
col_widths(tb, (300, 160, 672)); src_note(B6B)

B6C = clone(S(4))
T(B6C, 3, "공매도를 허용하면 합쳐도 효율선 위 — 가능성 점검과 세 한계")
T(B6C, 5, "부록 B-6 ④ — 계정 배분 은퇴 60 · 교육 20 · 상속 20%")
T(B6C, 9, "합쳐도 효율 — 전체 γ는 역수의 가중평균")
T(B6C, 10, "1/γ = 0.6/3.795 + 0.2/2.706 + 0.2/0.877 → γ = 2.174 (가중평균 2.99가 아니다). 전체 비중 채권 24.2 · 저위험 42.2 · 고위험 33.7%(원문 표 1), 기대수익률 13.84% · 표준편차 20.32%. 역수 가중평균은 강의 도출.")
T(B6C, 13, "가능성 점검 — H가 도달 가능한 하위 경계 이하여야 한다")
T(B6C, 14, "이 자산들로 만들 수 있는 하위 5% 수익률의 최댓값은 −2.31%다(식 (10)–(19)를 닫힌 해로 — 강의 계산). “원금 손실 확률 5% 이하”(H 0, α 5%)는 불가능 — H · α · 자산 · 목표를 다시 묻는다.")
T(B6C, 17, "한계 — 한 기간 · 정규분포 · 공매도 허용")
T(B6C, 18, "꼬리가 두꺼우면 미달이 더 크고 크기는 못 막는다. 롱온리면 은퇴 · 교육은 그대로, 상속만 0 · 8.9 · 91.1%(26.35 → 23.67%) — 합친 포트폴리오는 같은 위험의 효율선보다 연 12bp 낮을 뿐이다. 푸는 법은 본문 ‘롱온리면 이렇게 푼다’ 장.")
T(B6C, 19, "그래서 바닥(필수 목표)은 Floor(CPPI)로, 그 위의 계정은 Das–Markowitz로 지킵니다")
for x in B6C.shapes:
    if 7 <= x.shape_id <= 18:
        xx, yy, ww, hh = lecfix.px(x); set_px(x, y=yy - 26)
src_note(B6C)


# (머) 3차 — 방법 1 손계산(3장 표 장표 복제): 덱의 예 55세 · 10억 · 세 목표, T 10년
M1 = clone(S(51))
T(M1, 3, "방법 1의 손계산 — 목표마다 네 단계, 합은 다섯째 단계")
T(M1, 5, "가정(부록 B-2): GHP r 2%, 주식 λ 6% · σ 20% · 혼합 m = r + wλ − ½w²σ², 변동성 wσ · σ√T = 0.632")
tb = lecfix.find_table(M1); drop_col(tb, 4)
lecfix.set_table(tb, [["단계", "Safety 6억 · 95% (끝까지)", "Market 3억 · 70%", "Aspirational 5억 · 30%"],
    ["① 목표 읽기 — G · T · p → z", "G 6 · T 10 · p 95% → z 1.645", "G 3 · p 70% → z 0.524", "G 5 · p 30% → z −0.524"],
    ["② 주식 비중 w* = λ/σ² − z/(σ√T), 0~1로 자른다", " ", " ", " "],
    ["③ 확률 조정 성장률 g = m − zσ/√T", " ", " ", " "],
    ["④ 필요 자본 K = G × exp(−T g)", " ", " ", " "]])
col_widths(tb, (300, 300, 270, 262)); table_font(tb, 14, 13); row_heights(tb, [36, 50, 66, 66, 56]); stripe(tb)
for ri, ci, tex, nm in ((2, 1, r"$1.5-1.645/0.632=-1.10\ \Rightarrow\ w=0$", "m1_21.png"),
                        (2, 2, r"$1.5-0.524/0.632=0.67$", "m1_22.png"),
                        (2, 3, r"$1.5+0.524/0.632=2.33\ \Rightarrow\ w=1$", "m1_23.png"),
                        (3, 1, r"$g=r=2.0\%$ (전액 GHP)", "m1_31.png"),
                        (3, 2, r"$5.13\%-0.524\times 13.4\%/3.162=2.90\%$", "m1_32.png"),
                        (3, 3, r"$6\%+0.524\times 20\%/3.162=9.31\%$", "m1_33.png"),
                        (4, 1, r"$6\times e^{-10\times 0.020}=4.91$", "m1_41.png"),
                        (4, 2, r"$3\times e^{-10\times 0.0290}=2.24$", "m1_42.png"),
                        (4, 3, r"$5\times e^{-10\times 0.0931}=1.97$", "m1_43.png")):
    eq_in_cell(M1, tb, ri, ci, tex, nm, fontsize=22)
l1 = copy_shape(sh(S(19), 14), M1); set_px(l1, y=520, h=34); font(l1, 15)
lecfix.set_text(l1, "⑤ 합 4.91 + 2.24 + 1.97 = 9.13억 → 남는 0.87억은 Aspirational로 → 4.91 · 2.24 · 2.84억 = 49 / 22 / 28%")
remove(sh(M1, 8)); fn = copy_shape(sh(S(21), 8), M1); set_px(fn, x=79, y=566, w=1132, h=28)
lecfix.set_text(fn, "※ 같은 성장률 4.5%에 변동성 9%인 균형형을 쓰면 2.22억 — 변동성이 낮은 상품이 같은 목표를 더 싸게 산다."); font(fn, 13)

# (버) 3차 — 방법 2의 ②③④를 풀어 쓰기(3장 표 장표 복제): 보충교재 4-2절 · numpy 검산
M2D = clone(S(3))
T(M2D, 3, "일반화 — 자산이 n개면 같은 일을 행렬로 합니다")
T(M2D, 5, "②와 같은 일을 비중 셋으로 — μ 기대수익 벡터 · Σ 공분산 행렬 · 1 모두 1인 벡터 · Σ⁻¹ 역행렬(채권 400, 주식 2×2)")
tb = lecfix.find_table(M2D)
lecfix.set_table(tb, [["단계", "식", "숫자"],
    ["② 상수 c와 a", " ", "c = 400 + 23.96 + 2.08 = 426.04 · a = 20 + 2.08 + 0.83 = 22.917"],
    ["② 상수 d", " ", "1.4167 − 22.917² ÷ 426.04 = 0.18399"],
    ["③ 제약에 넣는다 (x = 1/γ)", " ", "a/c = 0.0538 · 1/c = 0.002347"],
    ["③ 양변을 제곱 → x의 2차식", " ", "근의 공식 → 양수 근 x = 0.2635 → γ = 1/x = 3.795"],
    ["④ 비중 = 최소분산 + x × 위험 방향", " ", "(0.939, 0.056, 0.005) + 0.2635 × (−1.516, 0.795, 0.721) = 53.9 · 26.6 · 19.5%"]])
col_widths(tb, (250, 470, 412)); table_font(tb, 14, 13); row_heights(tb, [36, 58, 58, 64, 58, 64]); stripe(tb)
for ri, tex, nm in ((1, r"$c=\mathbf{1}^{\top}\Sigma^{-1}\mathbf{1},\quad a=\mathbf{1}^{\top}\Sigma^{-1}\mu$", "m2_1.png"),
                    (2, r"$d=\mu^{\top}\Sigma^{-1}\mu-a^2/c$", "m2_2.png"),
                    (3, r"$0.0538+0.18399\,x-1.645\sqrt{0.002347+0.18399\,x^2}=-0.10$", "m2_3.png"),
                    (4, r"$-0.4639\,x^2+0.0566\,x+0.0173=0$", "m2_4.png"),
                    (5, r"$w=\Sigma^{-1}\mathbf{1}/c+x\,\Sigma^{-1}(\mu-(a/c)\mathbf{1})$", "m2_5.png")):
    eq_in_cell(M2D, tb, ri, 1, tex, nm, fontsize=22)
T(M2D, 8, "제곱하면 근이 둘 — 원래 식을 만족하는 양수 근 하나만 씁니다(음수 근 −0.142는 버린다)"); set_px(sh(M2D, 8), y=580)
note(M2D, 626, "※ 원문 각주 6(p.317): A = 1′Σ⁻¹μ = 22.917, B = μ′Σ⁻¹μ = 1.4167, C = 1′Σ⁻¹1 = 426.04, D = BC − A² = 78.39 · 강의의 a = A, c = C, d = B − A²/C = D/C(원문 D와 다른 양)")


# (서) 6차 — 레버 ① 승수 m을 VaR로(42장 복제: 수식 + 표, 카드는 걷고 숫자 흐름 세 줄)
def strip_cards(slide, y_lo, y_hi):
    for x in list(slide.shapes):
        xx, yy, ww, hh = lecfix.px(x)
        if y_lo <= yy <= y_hi and not (getattr(x, "has_table", False) and x.has_table) and x.shape_type != 13:
            t = x.text_frame.text.strip() if x.has_text_frame else ""
            if not t.isdigit() and not t.startswith("BAF"): remove(x)
def line(slide, y, text, pt=15, src=None):
    l = copy_shape(src or sh(S(19), 14), slide); set_px(l, x=79, y=y, w=1132, h=30); lecfix.set_text(l, text); font(l, pt); return l
def note(slide, y, text):
    n = copy_shape(sh(S(21), 8), slide); set_px(n, x=79, y=y, w=1132, h=26); lecfix.set_text(n, text); font(n, 12); return n

V1 = clone(S(42))
T(V1, 3, "레버 ① 승수 m을 VaR로 정한다 — 손실이 쿠션을 넘지 않게")
T(V1, 5, "쿠션 C = A − F, 위험 노출 E = m × C · 최악 하락률 x(리밸런싱 사이 VaR)에서 손실 E × x ≤ C이면 floor가 지켜진다")
T(V1, 8, "조건 — 최악의 손실이 쿠션보다 작게")
swap_eq(V1, 9, E_V1, maxw=500, maxh=72, cx=371, pad=22)
tb = lecfix.find_table(V1)
lecfix.set_table(tb, [["최악 하락률 x", "m 상한 = 1/x"],
                      ["x = 10% (월간 99%, 보충교재 3-5)", "1 ÷ 10% = 10"],
                      ["x = 20% (하루 −20%, 1987년형)", "1 ÷ 20% = 5"],
                      ["x = 33%", "1 ÷ 33% = 3"],
                      ["x = 50% (기본 답 m ≤ 2의 역산)", "1 ÷ 50% = 2"]])
table_font(tb, 14, 13)
strip_cards(V1, 430, 640)
card(V1, 69, 404, 364, 156, "숫자 흐름",
     ["A = 100 · F = 80 · m = 3", "쿠션 C = 100 − 80 = **20**", "위험 노출 E = 3 × 20 = **60**"])
card(V1, 449, 404, 364, 156, "x = 20% → floor 지킴",
     ["손실 = E × x = 60 × 20%", "= **12** ≤ C = 20", "m = 3 ≤ 1/x = 5"])
card(V1, 829, 404, 372, 156, "x = 40% → floor 뚫림",
     ["손실 = E × x = 60 × 40%", "= **24** > C = 20", "m = 3 > 1/x = 2.5"], dark=True)
r = line(V1, 584, "Flexicure 재설정 — 매년 x를 다시 재면 m = 1/x도 다시 정해진다: 변동성이 오르면 m이 내려간다", 17, src=sh(S(45), 8))

# (어) 9차 — 레버 ② floor를 위험 자본으로: VaR 연동 floor(위험 자본 방식 CPPI, 51장 표 복제)
RC = clone(S(51))
T(RC, 3, "레버 ② floor를 위험 자본으로 — 위험이 오르면 floor를 올린다")
T(RC, 5, "손실 예산 L(95% VaR 손실)을 고정하고 floor를 다시 정한다 — 바젤처럼 위험 ↑ → 묶어 두는 자본(floor) ↑")
box = copy_shape(sh(S(16), 7), RC); set_px(box, x=140, y=194, w=1000, h=92)
lab = copy_shape(sh(S(16), 8), RC); set_px(lab, x=160, y=200, w=960); lecfix.set_text(lab, "규칙 — VaR 손실 = 손실 예산 L, floor F는 필수 목표 Floor(GHP 가치) 아래로 내리지 않는다")
w_, h_ = size_px(E_RC, maxw=960, maxh=50); add_pic(RC, E_RC, 640 - w_ / 2, 232, w_, h_)
tb = lecfix.find_table(RC)
lecfix.set_table(tb, [["σ (위험)", "VaR 비율 x = 1.645σ − μ", "쿠션 C = L ÷ (m × x)", "floor F = A − C", "위험 노출 E = m × C (VaR 손실 E × x)"],
    ["σ = 15% (하락)", "1.645 × 15% − 6% = 18.7%", "16.1 ÷ (3 × 18.7%) = 28.8", "100 − 28.8 = 71.2 ↓", "3 × 28.8 = 86.4 ↑ (× 18.7% = 16.1)"],
    ["σ = 20% (기준)", "1.645 × 20% − 6% = 26.9%", "16.1 ÷ (3 × 26.9%) = 20.0", "100 − 20.0 = 80.0", "3 × 20.0 = 60.0 (× 26.9% = 16.1)"],
    ["σ = 25% (상승)", "1.645 × 25% − 6% = 35.1%", "16.1 ÷ (3 × 35.1%) = 15.3", "100 − 15.3 = 84.7 ↑", "3 × 15.3 ≈ 46.0 ↓ (× 35.1% = 16.1)"]])
col_widths(tb, (150, 240, 240, 190, 312)); table_font(tb, 13, 13); row_heights(tb, [36, 42, 42, 42]); set_px(tb, y=300)
blk(RC, 79, 470, 1132, [
    "**• 기준:** 자산 A = 100, 승수 m = 3, μ = 6%, σ = 20% → C = 20이 되게 L = m × x × C = 3 × 26.9% × 20 = 16.1",
    "**• 두 레버:** {p:V1}장은 floor를 고정하고 m을, {p:RC}장은 m을 고정하고 floor를 움직인다 — 둘 다 E × x를 묶는다",
    "**• 같은 원리:** {p:o48}장 ‘위험자본 방식의 CPPI’"], pt=15, before=3)
T(RC, 8, "위험 ↑ → floor ↑ · 쿠션 ↓ · 노출 ↓, VaR 손실은 16.1 그대로 — floor가 바젤의 ‘필요 자본’ 역할을 한다")
set_px(sh(RC, 8), y=580); font(sh(RC, 8), 17)
note(RC, 626, "※ 위험이 오를 때 floor를 낮추면 위험 자본과 반대 방향 · 변동성 급등 때는 모두 함께 판다({p:o48}장) — 필수 목표는 고정 Floor로 · 무위험 0 · 1년 · 정규({p:V2B}장과 같음)")
note(RC, 650, "※ 바젤 시장위험 자본도 VaR 기반 — 위험이 크면 필요 자본이 커진다(일반 지식, 배수 생략)")

# (저) 6차 — 세 방법은 같은 문장(51장 표 복제, 4열)
V2 = clone(S(51))
T(V2, 3, "세 방법 모두 “나쁜 경우에도 floor를 지킨다”는 같은 문장입니다")
T(V2, 5, "나쁜 경우 = 평균 − z × 변동성 · 그 경우에도 floor 위에 있게 한다 — 다른 것은 floor를 금액 · 수익률 · 목표 중 무엇으로 쓰느냐")
tb = lecfix.find_table(V2); drop_col(tb, 4)
lecfix.set_table(tb, [["방법", "나쁜 경우의 조건", "floor는 어디에", "정적 · 동적"],
    ["CPPI + VaR (단원 ⑤)", " ", "F 그대로 — 나쁜 경우 자산 ≥ F", "동적 — 매일 · 매년"],
    ["방법 2 Das–Markowitz (단원 ③)", " ", "H가 수익률로 쓴 floor: F = A × (1 + H), A = 100 · H = −20% ⇔ F = 80", "정적 — 한 번"],
    ["방법 1 필요 자본 (단원 ③)", " ", "목표 G = T년 뒤의 floor — 나쁜 경우 성장률로도 G에 닿게 오늘 떼어 둘 K", "정적 — 한 번"]])
col_widths(tb, (260, 330, 372, 170)); table_font(tb, 14, 13); row_heights(tb, [36, 66, 66, 66]); stripe(tb)
for ri, tex, nm in ((1, r"$E\,(z\sigma-\mu)\leq C=A-F$", "v2_1.png"),
                    (2, r"$\mu_p-z\,\sigma_p\geq H$", "v2_2.png"),
                    (3, r"$K=G\,e^{-T\,(m-z\sigma/\sqrt{T})}$", "v2_3.png")):
    eq_in_cell(V2, tb, ri, 1, tex, nm, fontsize=24)
line(V2, 486, "읽는 법 — 셋 다 변동성 σ에 z를 곱해 벌점으로 빼고, 그 ‘나쁜 경우’가 floor 위에 있으면 통과")
note(V2, 548, "※ 셋 다 정규 · 닫힌 식 — 두꺼운 꼬리 · 동적 규칙이면 단원 ③ 방법 3(몬테카를로)으로 같은 확률을 센다")
note(V2, 524, "※ 정규분포면 위험회피도 · 변동성 · VaR · (H, α)가 같은 포트폴리오를 가리킨다(보충교재 9-3: γ 8.0 ⇔ 변동성 5.35% ⇔ VaR 7.9%) · 꼬리가 두꺼우면 깨진다")
T(V2, 8, "다음 장 — 같은 숫자로 풀면 세 방법이 같은 답에 닿는다"); set_px(sh(V2, 8), y=578)

# (처) 6차 — 같은 숫자 예로 세 방법이 같은 답(3장 표 장표 복제)
V2B = clone(S(3))
T(V2B, 3, "같은 숫자로 풀면 세 방법이 같은 답 74에 닿습니다")
T(V2B, 5, "공통 가정({p:RC}장과 같은 x = 26.9%) — μ = 6%, σ = 20%, z = 1.645, 무위험 0 · 이번엔 F = 80을 고정")
tb = lecfix.find_table(V2B)
lecfix.set_table(tb, [["방법", "계산", "답"],
    ["나쁜 경우 하락률", " ", "x = 26.9%"],
    ["CPPI + VaR", " ", "위험 노출 E = 74 (m ≤ 3.7)"],
    ["방법 2 Das–Markowitz", " ", "위험자산 비중 74%"],
    ["방법 1 필요 자본", " ", "떼어 둘 금액 K = 100 = A"]])
col_widths(tb, (250, 560, 322)); table_font(tb, 15, 14); row_heights(tb, [36, 60, 60, 60, 60]); stripe(tb)
for ri, tex, nm in ((1, r"$x=z\sigma-\mu=1.645\times 20\%-6\%=32.9\%-6\%=26.9\%$", "v2b_1.png"),
                    (2, r"$E\leq C/x=(100-80)/0.269=74,\quad m\leq 1/x=3.7$", "v2b_2.png"),
                    (3, r"$w\,(\mu-z\sigma)\geq -20\%\ \Rightarrow\ w\leq 20\%/26.9\%=74\%$", "v2b_3.png"),
                    (4, r"$K=G/(1+\mu_p-z\sigma_p)=80/(1+4.5\%-24.5\%)=80/0.80=100$", "v2b_4.png")):
    eq_in_cell(V2B, tb, ri, 1, tex, nm, fontsize=23)
line(V2B, 552, "방법 1: μp = 74% × 6% = 4.5%, σp = 74% × 20% = 14.9% → 나쁜 경우 −20% — 자산 A 전부가 필요 자본")
T(V2B, 8, "차이 — 단원 ③은 한 번 정하고, CPPI · Flexicure는 쿠션이 변하면 E를 매일 · 매년 다시 계산한다(첫날만 같다, 9-5)")
set_px(sh(V2B, 8), y=590); font(sh(V2B, 8), 16)
note(V2B, 636, "※ 가정: 무위험 수익 0 · 1년 · 정규분포 · 방법 1은 1년 단리 근사(단원 ③ 본문은 연속복리) — 교육용 예시")


T(S(53), 5, "디지털 TDF = 설문 몇 개로 표준 포트폴리오 3~5개 중 하나를 배정 — 나이 · 성향 하나로 비중을 정하는 TDF의 온라인판")
# (커) 8차 — 단원 ⑦ 한국 제도 세 장 · 부록 C-2 · 부록 D
def note8(slide, y, text):
    n = copy_shape(sh(S(21), 8), slide); set_px(n, x=79, y=y, w=1132, h=24); lecfix.set_text(n, text); font(n, 11); return n
K1 = clone(S(51))
T(K1, 3, "한국 퇴직연금 501조 — DB는 회사, DC · IRP는 개인이 위험을 진다")
T(K1, 5, "누가 운용 결과를 떠안느냐가 제도를 가른다 — DC · IRP는 개인이 은퇴 목표(부채)를 지는 구조라 GBI의 무대")
tb = lecfix.find_table(K1)
lecfix.set_table(tb, [["제도", "누가 위험을 지나", "적립금 · 비중 (2025 말)", "2025 수익률", "GBI로 읽으면"],
    ["DB 확정급여", "회사 — 퇴직급여가 미리 정해져 있다", "228.9조 · 45.7%", "3.53%", "기관의 LDI(M7) 무대"],
    ["DC 확정기여", "근로자 — 쌓인 적립금과 운용 결과가 급여", "141.6조 · 28.2%", "8.47%", "개인이 목표를 지는 GBI 무대"],
    ["IRP 개인형", "근로자 — 퇴직금 이전 · 개인 추가 납입", "130.9조 · 26.1%", "9.44%", "GBI 무대 · 로보 일임 대상"],
    ["합계", "—", "501.4조", "6.47%", "원리금보장 3.09% · 실적배당 16.80%"]])
col_widths(tb, (170, 330, 230, 130, 272)); table_font(tb, 14, 13)
line(K1, 470, "관행 — 적립금의 75.4%가 원리금보장(2025 말) · 그 수익률 3.09%는 실적배당 16.80%의 5분의 1도 안 된다", 15)
T(K1, 8, "DC · IRP 가입자는 스스로 목표를 정하지 않으면 원리금보장에 머문다 — IRP(다음 장)와 디폴트옵션(그다음 장)")
set_px(sh(K1, 8), y=576); font(sh(K1, 8), 17)
note8(K1, 636, "출처: 금융감독원 2025년 퇴직연금 운용현황(2026.5.20 발표) · 금융감독원 2025 퇴직연금 투자백서 — 제도별 금액 · 비중 · 수익률")

IRPS = clone(S(52))
T(IRPS, 3, "IRP는 개인이 운용하는 퇴직연금 계좌 — 이 강의에서 세 역할을 합니다")
T(IRPS, 5, "개인형 퇴직연금(IRP) — 퇴직 · 이직 때 퇴직급여를 받아 두는 계좌이자, 본인이 더 넣어 세액공제를 받는 계좌")
tb = lecfix.find_table(IRPS)
lecfix.set_table(tb, [["항목", "내용"],
    ["누가 운용하나", "가입자 본인 — DC처럼 운용 결과(위험)를 개인이 진다 · 2025 말 130.9조(26.1%), 수익률 9.44%"],
    ["세제 · 인출", "연금저축과 합쳐 연 900만 원까지 세액공제(총급여 5,500만 원 이하 16.5%, 초과 13.2%) · 중도 인출은 법정 사유만(주택 구입 · 6개월 이상 요양 · 파산 등)"],
    ["역할 ① 디폴트옵션", "DC와 함께 디폴트옵션 대상 계좌 — 지시가 없으면 사전 지정 상품으로(다음 장)"],
    ["역할 ② 로보 일임", "퇴직연금 중 로보 일임이 허용된 계좌(2024.12 지정 · 2025.3 출시) — Floor 재설정 같은 GBI 규칙을 매번 승인 없이 실행할 수 있다"],
    ["역할 ③ 새 상품의 무대", "이 단원 끝 새 상품 ‘목표 계좌형 IRP’ — 은퇴 · 교육 · 주거 계정별 (H, α)"]])
col_widths(tb, (230, 902)); table_font(tb, 14, 13); row_heights(tb, [36] + [52] * 5)
T(IRPS, 8, "IRP는 개인이 목표(부채)를 지면서, 규칙 기반 운용을 맡길 수 있는 자리 — GBI가 들어갈 첫 계좌입니다")
set_px(sh(IRPS, 8), y=568); font(sh(IRPS, 8), 17)
note8(IRPS, 636, "출처: 금감원 2025 퇴직연금 운용현황(2026.5) · 금감원 금융꿀팁 125(연금계좌 중도인출) · 소득세법 연금계좌 세액공제(뉴시스 2025.12.12 보도) · 덱 로보 일임 장")

K2 = clone(S(52))
T(K2, 3, "디폴트옵션은 목표를 묻지 않고 대신 정해 주는 장치입니다")
T(K2, 5, "사전지정운용제도 — DC · IRP 가입자가 운용 지시를 하지 않으면 미리 고른 방법으로 자동 운용")
tb = lecfix.find_table(K2)
lecfix.set_table(tb, [["항목", "내용"],
    ["대상 · 시행", "DC · IRP 가입자의 미지시 적립금 · 2022.7.12 시행, 1년 유예 뒤 2023.7 본격 적용"],
    ["작동", "만기 후 지시 없이 4주 → 사업자 통지 → 2주 더 지나면 사전 지정 방법으로 운용 · 언제든 다른 지시 가능"],
    ["상품", "원리금보장 · TDF · 밸런스드 · 스테이블밸류 · 사회간접자본(인프라) 펀드 또는 혼합 — 위험도 4단계(초저 · 저 · 중 · 고)"],
    ["승인", "고용노동부 심의위원회가 상품을 심의 · 승인"],
    ["실제 선택 (2025 말)", "적립금 53.3조 중 안정형(원리금보장) 85.4% · 지정 가입자 734만 명 중 583만 명(79.4%)"]])
col_widths(tb, (220, 912)); table_font(tb, 14, 13); row_heights(tb, [36] + [48] * 5)
T(K2, 8, "열에 여덟이 원리금보장을 고르니 목표에 못 미친다 — 케이스 1부는 Floor를 지키는 투자형(Flexicure)을 심의한다")
set_px(sh(K2, 8), y=560); font(sh(K2, 8), 17)
note8(K2, 636, "출처: 고용노동부 보도자료(사전지정운용제도 도입, 2022.7) · 한국금융신문 2026.9.10(고용노동부 · 금감원 공시 집계) · 케이스 1부 덱")

K3 = clone(S(52))
T(K3, 3, "한국 TDF는 나이 하나로 비중을 줄이는 글라이드패스입니다")
T(K3, 5, "디폴트옵션 투자형의 주력 · 목표 금액과 확률은 묻지 않고, 바닥(Floor)도 없다")
tb = lecfix.find_table(K3)
lecfix.set_table(tb, [["항목", "내용"],
    ["구조", "목표 시점(빈티지) 하나로 주식 비중을 나이에 따라 줄인다 — 운용사 22곳(2024 말)"],
    ["규모 · 보수", "순자산 16.6조 · 평균 총보수 0.59%(2024 말) · 설정액 20.2조(2026.7) · 퇴직연금 공모펀드 중 30.3%"],
    ["규정", "DC · IRP는 위험자산 70% 한도 — 적격 TDF(안전자산 20% 이상, 목표 시점 이후 60% 이상)는 100%까지 가능"],
    ["2022년", "Floor가 없어 국내 TDF2025 −14.0% · TDF2045 −16.8% — 은퇴 직전 손실이 그대로"],
    ["GBI · Flexicure와", "TDF는 나이만 묻는다 · GBI는 목표와 확률을 묻고, Flexicure는 Floor를 GHP로 지키며 쿠션만큼만 위험"]])
col_widths(tb, (220, 912)); table_font(tb, 14, 13); row_heights(tb, [36] + [48] * 5)
T(K3, 8, "TDF는 ‘나이’로, GBI는 ‘목표와 확률’로 — 케이스 1부의 비교 대상이 바로 이 TDF입니다")
set_px(sh(K3, 8), y=560); font(sh(K3, 8), 17)
note8(K3, 636, "출처: 자본시장연구원 이슈보고서 26-08(제로인) · 금융투자협회 · KG제로인(파이낸셜뉴스 2026.8.17) · 퇴직연금감독규정 시행세칙 · 한국금융신문 2026.5.30")

C2 = clone(S(77))
T(C2, 3, "한국 퇴직연금의 말도 시험과 과제에서 그대로 씁니다")
T(C2, 5, "부록 C-2 — 한국 제도 용어 (단원 ⑦)")
tb = lecfix.find_table(C2)
lecfix.set_table(tb, [["용어", "뜻"],
    ["DB (확정급여형)", "퇴직급여가 미리 정해지고 회사가 운용 위험을 지는 제도"],
    ["DC (확정기여형)", "회사가 정해진 부담금을 넣고 근로자가 운용 결과를 지는 제도"],
    ["IRP (개인형 퇴직연금)", "퇴직급여를 받아 두거나 본인이 추가 납입해 세액공제를 받는 계좌 — 운용은 본인, 로보 일임 허용"],
    ["디지털 TDF", "로보 1세대 — 설문 몇 개로 표준 포트폴리오 3~5개 중 하나를 배정(나이 · 성향 하나로 정하는 TDF의 온라인판)"],
    ["디폴트옵션 (사전지정운용제도)", "DC · IRP 가입자가 운용 지시를 하지 않으면 미리 고른 방법으로 자동 운용(2022.7 시행)"],
    ["원리금보장상품", "예금 · 보험 등 원금과 약정 이자를 지급하는 상품 — 2025 말 적립금의 75.4%"],
    ["TDF (타깃데이트펀드)", "목표 시점에 맞춰 주식 비중을 나이에 따라 줄이는 펀드"],
    ["적격 TDF", "안전자산 비중 요건을 갖춰 위험자산 70% 한도의 예외로 100%까지 담을 수 있는 TDF"]])
table_font(tb, 13, 13); col_widths(tb, (300, 832)); row_heights(tb, [34] + [40] * 8)

RD = clone(S(77))
T(RD, 3, "이번 주의 근거 문헌은 안전우선 계열과 Flexicure 계열로 나뉩니다")
T(RD, 5, "부록 D — 참고문헌")
tb = lecfix.find_table(RD)
lecfix.set_table(tb, [["문헌", "이 덱에서"],
    ["Roy, A.D. (1952) Safety First and the Holding of Assets, Econometrica 20(3)", "안전우선 — 미달 확률을 가장 작게"],
    ["Telser, L.G. (1956) Safety First and Hedging, Review of Economic Studies 23(1), 1–16 (권호 1955–56)", "방법 2의 계보 — 미달 한도 아래 기대수익 최대"],
    ["Kataoka, S. (1963) A Stochastic Programming Model, Econometrica 31(1–2)", "방법 1의 계보 — 하위 경계 최대"],
    ["Chhabra, A. (2005) Beyond Markowitz, Journal of Wealth Management 7(4)", "3계층 버킷"],
    ["Merton, R.C. (1969 · 1971) 연속시간 최적 소비 · 포트폴리오", "무조건 항(Merton 비중)"],
    ["Das, Markowitz, Scheid, Statman (2010) Portfolio Optimization with Mental Accounts, JFQA 45(2)", "방법 2 · 부록 B-6"],
    ["Thaler (1985) · Shefrin-Statman (2000)", "심리계정 · BPT"],
    ["Black-Jones (1987) · Estep-Kritzman (1988) · Grossman-Zhou (1993) · EDHEC Flexicure", "CPPI · Ratchet · 최대낙폭 Floor · Flexicure"]])
table_font(tb, 13, 13); col_widths(tb, (720, 412)); row_heights(tb, [34] + [44] * 8)


# (터) 8차 추가 — 개선점 · 발전 방향 · 새 상품(제안 · 교육용)
I1 = clone(S(51))
T(I1, 3, "한국 제도의 개선점 — 문제는 목표를 묻지 않는 데서 시작합니다")
T(I1, 5, "왼쪽 두 열은 공시 사실, 오른쪽 열은 제안(교육용)")
tb = lecfix.find_table(I1); drop_col(tb, 4); drop_col(tb, 3)
lecfix.set_table(tb, [["문제 (공시 사실)", "원인", "개선 방향 (제안 · 교육용)"],
    ["원리금보장 쏠림 — 적립금 75.4%, 수익률 3.09% vs 실적배당 16.80%(2025)", "목표를 묻지 않아 손실 회피가 기본값", "목표 기반 디폴트 — Safety는 GHP로, 쿠션만큼만 위험"],
    ["디폴트옵션 안정형 85.4% — 위험 등급으로 고른다", "상품이 초저 · 저 · 중 · 고위험으로만 진열", "은퇴 소득 목표 · 달성 확률로 진열"],
    ["성과를 수익률로만 공시", "가입자의 목표 정보를 모으지 않는다", "목표 달성 확률 · Floor 확보율 공시"],
    ["일시금 관행 — 2024 수령 개시 계좌의 87%(금액 기준 연금 57%)", "적립 단계만 설계, 인출 상품 부족", "인출 단계까지 — Flexicure 소득 Floor"],
    ["수수료 — TDF 평균 0.59%, GBI 예상 0.8~1.5%(교육용 가정)", "개인화 비용", "총보수 상한(케이스 1부 ≤ 0.5%) · 로보 일임"]])
col_widths(tb, (470, 310, 352)); table_font(tb, 14, 13)
T(I1, 8, "해외 비교(미국 QDIA · 호주 MySuper)는 케이스 1부 글로벌자문 팀의 과제입니다"); set_px(sh(I1, 8), y=590); font(sh(I1, 8), 16)
note8(I1, 636, "출처: 금감원 2025 퇴직연금 운용현황(2026.5) · 한국금융신문 2026.9.10 · 금감원 · 고용노동부 2024년 퇴직연금 통계(2025) · 자본시장연구원 26-08")

I2 = clone(S(4))
T(I2, 3, "발전 방향은 세 줄 — 목표를 묻고, 목표로 운용하고, 소득으로 돌려준다")
T(I2, 5, "제안 · 교육용 — 줄마다 이번 강의의 개념 하나와 연결한다")
T(I2, 9, "수익률 공시 → 목표 달성 확률 공시")
T(I2, 10, "가입자마다 은퇴 소득 목표를 받고, 달성 확률과 Floor 확보율을 보여 준다 — 단원 ① 위험 = P(W < G), 단원 ③ 필요 자본 K.")
T(I2, 13, "위험 등급 디폴트 → 목표 기반 디폴트")
T(I2, 14, "필수 소득은 GHP로 지키고 쿠션 × m만 위험자산, m은 VaR · 위험 자본으로 상한 — 단원 ④⑤, 계좌 나누기는 Das–Markowitz(단원 ③ 방법 2).")
T(I2, 17, "적립 단계만 → 인출 단계까지")
T(I2, 18, "은퇴 뒤에도 소득 Floor를 매년 다시 재고 그 위에서만 위험을 진다 — Flexicure 재설정(단원 ⑤⑥), 일시금 대신 소득으로.")
T(I2, 19, "도입에는 공시 기준 · 디폴트옵션 상품 승인 · 일임 규정의 변경이 필요합니다(제안)")

I3 = clone(S(51))
T(I3, 3, "강의로 만든 새 상품 세 가지 — 교육용 설계안")
T(I3, 5, "숫자는 덱 기존 값만 — 케이스 1부 기본 답(Floor 0.90 · m ≤ 2 · 총보수 ≤ 0.5%), {p:V1}~{p:V2B}장 VaR 예, 부록 B-6")
tb = lecfix.find_table(I3)
lecfix.set_table(tb, [["상품 (교육용)", "작동 규칙", "강의의 장", "가입자에게 보여 줄 지표", "제도 변경 · 위험"],
    ["(a) Flexicure형 목표 소득 디폴트옵션", "은퇴 목표의 90%를 국채(GHP)로 복제, 쿠션 × m(m ≤ 2)만 위험자산, 매년 재설정, 총보수 ≤ 0.5%", "방법 1 + CPPI · 단원 ④⑤ · 케이스 1부", "Safety 미달 ≤ 5% · Cash Trap ≤ 10%", "상품 승인 · ‘보장’ 표현 금지 / Cash Trap · 수수료"],
    ["(b) 목표 계좌형 IRP", "은퇴 · 교육 · 주거 계정마다 (H, α) → 계정별 평균-분산 배분, 합쳐도 효율", "방법 2 · 단원 ②③ · 부록 B-6", "계정별 하위 5% 수익률 · 가능성 점검", "한 계좌 안 하위 계정 · 일임 / 정규 가정 · 미달 크기"],
    ["(c) 인출 단계 소득 상품", "필수 소득 Floor(RB · GHP) + 희망 소득(쿠션 × m) 두 층, 매년 Floor 재측정", "방법 3으로 공시 · 단원 ④⑤⑥", "필수 소득 확보율 · 희망 소득 달성 확률", "연금 수령 유인 · 장기 국채 공급 / 금리 · 장수"]])
col_widths(tb, (190, 330, 180, 200, 232)); table_font(tb, 13, 13)
note(I3, 630, "※ 가입자에게 보여 줄 지표(달성 확률 · Cash Trap)는 방법 3(몬테카를로)으로 계산해 공시한다 — 설계 근거는 (a) 방법 1 + CPPI · (b) 방법 2")
T(I3, 8, "과제(팀 — 한국 GBI 플랫폼 설계)의 출발점: 셋 중 하나를 골라 시장 진단 → UX → 금융 모델 → 규제까지"); set_px(sh(I3, 8), y=590); font(sh(I3, 8), 16)

# (카) 디바이더 — 원래 다섯 장 + 복제 다섯 장
DIV = {"①": S(7), "④": S(27), "⑦": S(50), "⑨": S(61), "⑩": S(66)}
for k in ("②", "③", "⑤", "⑥", "⑧"): DIV[k] = clone(S(7))
for i, u in enumerate(UNITS, 1):
    d = DIV[u[0]]
    T(d, 2, f"{i:02d}"); T(d, 3, u[2])
    q = sh(d, 5); lecfix.set_text(q, f"이 단원이 푸는 질문 — {u[3]}"); white(q)
    tp = copy_shape(q, d, dy=46); lecfix.set_text(tp, u[4])
    for p in tp.text_frame.paragraphs:
        for r in p.runs: r.font.color.rgb = RGBColor(0x9F, 0xB0, 0xB2); r.font.size = Pt(16)

# ============================================================ 새 순서 ====
ORDER = [
    ("도입", S(1)), ("도입", N2), ("도입", S(3)), ("도입", S(4)), ("도입", S(5)), ("도입", GL), ("도입", S(6)),
    ("div", DIV["①"]), *[("①", S(n)) for n in (8, 9, 10, 16, 17)],
    ("div", DIV["②"]), ("②", S(12)), ("②", S(11)), ("②", S(13)), ("②", DMI), ("②", CMP), ("②", S(14)),
    ("div", DIV["③"]), ("③", U3), *[("③", S(n)) for n in (15, 18, 19, 20, 21, 22)], ("③", P23), *[("③", S(n)) for n in (23, 24, 25)], ("③", M1),
    ("③", DM1), ("③", HC1), ("③", HC2), ("③", M2D), ("③", DMB), ("③", LO), ("③", XL), ("③", DM2), ("③", MC1), ("③", MC2), ("③", RA), ("③", KT), ("③", MX), ("③", MXR), ("③", MXG), ("③", MXP), ("③", MC3), ("③", CB1), ("③", VB), ("③", VB2), ("③", CB2), ("③", CC), ("③", DMC1), ("③", DMC2), ("③", CCB), ("③", MC4), ("③", CMP2), ("③", WHY), ("③", FD), ("③", S(26)),
    ("div", DIV["④"]), *[("④", S(n)) for n in (28, 29, 30, 31, 34, 41, 32, 33)],
    ("div", DIV["⑤"]), *[("⑤", S(n)) for n in (35, 36, 37, 38, 39, 40, 42)], ("⑤", V1), ("⑤", RC), ("⑤", V2), ("⑤", V2B), ("⑤", S(43)),
    ("div", DIV["⑥"]), *[("⑥", S(n)) for n in (44, 45, 46, 47, 48, 49)],
    ("div", DIV["⑦"]), *[("⑦", S(n)) for n in (51, 52)], ("⑦", K1), ("⑦", IRPS), ("⑦", K2), ("⑦", K3), *[("⑦", S(n)) for n in (53, 54, 55, 56, 57, 59)], ("⑦", I1), ("⑦", I2), ("⑦", I3), ("⑦", S(60)),
    ("div", DIV["⑧"]), ("⑧", S(58)), ("⑧", H1), ("⑧", MAP), ("⑧", H2), ("⑧", FX1), ("⑧", FX2), ("⑧", H3),
    ("div", DIV["⑨"]), *[("⑨", S(n)) for n in (62, 63, 64, 65)],
    ("div", DIV["⑩"]), *[("⑩", S(n)) for n in (67, 68, 69, 70)],
    ("div", S(71)), *[("부록", S(n)) for n in (72, 73, 74, 75, 76)], ("부록", B6A), ("부록", B6H), ("부록", B6B), ("부록", B6C), ("부록", S(77)), ("부록", C2), ("부록", RD), ("끝", S(78)),
]
DROP = [S(2)]
ids = P.slides._sldIdLst
by_el = {}
for el in list(ids):
    by_el[id(P.part.related_part(el.rId).slide._element)] = el
for s_ in DROP:
    el = by_el[id(s_._element)]; P.part.drop_rel(el.rId); ids.remove(el)
for el in list(ids): ids.remove(el)
for _, s_ in ORDER: ids.append(by_el[id(s_._element)])
assert len(ids) == len(ORDER) == len({id(s_._element) for _, s_ in ORDER})

# 비교표 답 행 — 손계산 장 번호(새 순서에서 계산)
pos = {id(s_._element): i for i, (_, s_) in enumerate(ORDER, 1)}
nM1, nDMB, nM2D, nHC1, nHC2 = pos[id(M1._element)], pos[id(DMB._element)], pos[id(M2D._element)], pos[id(HC1._element)], pos[id(HC2._element)]
t = lecfix.find_table(CMP2).table
lecfix.set_tf(t.cell(4, 1).text_frame, f"K 4.91 · 2.24 · 1.97억 (손계산 {nM1}장)")
lecfix.set_tf(t.cell(4, 2).text_frame, f"합치면 γ 2.174 (손계산 {nHC1} · {nHC2}장, 행렬 {nM2D}장)")
nMC3 = pos[id(MC3._element)]
nDMC1 = pos[id(DMC1._element)]
lecfix.set_tf(t.cell(4, 3).text_frame, f"방법 1 · 2의 문제를 그대로 푼다 — K 2.26억({nMC3}장) · 은퇴 54 · 27 · 20%({nDMC1}장)")
T(S(76), 14, f"※ 한계 ② 로그정규는 손계산을 위한 단순화 — 실무는 방법 3(몬테카를로, {nMC3}장)으로 같은 계산을 한다")
nFD = pos[id(FD._element)]
T(S(26), 18, f"σ가 아니라 P(W < G). 같은 문제를 세 방법으로 — 방법 1 가장 싼 K · 방법 2 Das–Markowitz · 방법 3 몬테카를로 · 바닥은 CPPI({nFD}장).")
T(WHY, 19, f"한계 — 미달의 크기는 막지 못하므로 필수 목표의 바닥은 Floor(CPPI)로 따로 지킵니다({nFD}장)")
T(B6C, 19, f"그래서 바닥(필수 목표)은 Floor(CPPI)로, 그 위의 계정은 Das–Markowitz로 지킵니다 — 방법은 {nFD}장")
# 장 번호 참조 {p:KEY} → 실제 번호(KEY = 신설 장 이름 또는 o옛번호)
_KEYS = {id(v._element): k for k, v in dict(TOC=N2, GLOSS=GL, DMI=DMI, DM1=DM1, DMB=DMB, DM2=DM2, CMP=CMP, CMP2=CMP2, WHY=WHY, P23=P23,
         H1=H1, MAP=MAP, H2=H2, FX1=FX1, FX2=FX2, H3=H3, M1=M1, M2D=M2D, HC1=HC1, HC2=HC2, LO=LO, XL=XL, FD=FD,
         MC1=MC1, MC2=MC2, MC3=MC3, MC4=MC4, RA=RA, KT=KT, MX=MX, CB1=CB1, CB2=CB2, CC=CC, VB=VB, DMC1=DMC1, DMC2=DMC2, MXG=MXG, MXR=MXR, MXP=MXP, VB2=VB2, CCB=CCB, U3=U3, V1=V1, RC=RC, V2=V2, V2B=V2B, K1=K1, IRPS=IRPS, K2=K2, K3=K3, I1=I1, I2=I2, I3=I3,
         B6A=B6A, B6H=B6H, B6B=B6B, B6C=B6C, C2=C2, RD=RD).items()}
PAGE = {}
for _i, (_u, _s) in enumerate(ORDER, 1):
    PAGE[_KEYS.get(id(_s._element))] = _i
    for _n, _o in enumerate(O, 1):
        if _o is _s: PAGE[f"o{_n}"] = _i
import re as _re
def _sub(m):
    assert m.group(1) in PAGE, m.group(1); return str(PAGE[m.group(1)])
for _s in P.slides:
    for _x in _s.shapes:
        for _tf in lecfix._tf_iter(_x):
            for _p in _tf.paragraphs:
                for _r in _p.runs:
                    if "{p:" in _r.text: _r.text = _re.sub(r"\{p:(\w+)\}", _sub, _r.text)
assert not lecfix.grep(P, r"\{p:"), lecfix.grep(P, r"\{p:")
json.dump({k: v for k, v in PAGE.items() if k}, open(os.path.join(HERE, "pages.json"), "w"), ensure_ascii=False, indent=1)

# 머리 띠 소주제(19차)
SUBT = {}
def _st(objs, name):
    for o in objs: SUBT[id(o._element)] = name
_st([U3], "단원 지도"); _st([S(n) for n in (15, 18, 19, 20, 21, 22)], "필요 자본 기초")
_st([P23, S(23), S(24), S(25), M1], "방법 1 가장 싼 K"); _st([DM1, HC1, HC2, M2D, DMB], "방법 2 Das–Markowitz")
_st([LO, XL, DM2], "롱온리 · 합산"); _st([MC1, MC2, RA, KT, MX, MXR, MXG, MXP, MC3], "방법 3 몬테카를로")
_st([CB1, VB, VB2, CB2], "조합 · 묶기"); _st([CC], "실습 A"); _st([DMC1, DMC2], "방법 3으로 방법 2"); _st([CCB], "실습 B"); _st([MC4], "도구")
_st([CMP2, WHY, FD], "세 방법 비교 · 정리"); _st([S(26)], "체크포인트")
_st([S(n) for n in (35, 36, 37, 38, 39, 40, 42)], "CPPI · Flexicure"); _st([V1, RC], "VaR · 위험 자본"); _st([V2, V2B], "세 방법과 같은 문장"); _st([S(43)], "Cash Trap")
_st([S(51), S(52)], "4대 도전"); _st([K1, IRPS, K2, K3], "한국 제도"); _st([S(n) for n in (53, 54, 55, 56, 57, 59)], "로보 · 한국 DC")
_st([I1, I2, I3], "개선 · 새 상품"); _st([S(60)], "체크포인트")
_st([S(58), H1, MAP], "세 숫자 · 대응표"); _st([H2], "모방 3장벽"); _st([FX1, FX2, H3], "Flexicure 다섯 단계")
for _n, _nm in zip((72, 73, 74, 75, 76), ("B-1 K 유도", "B-2 아홉 칸", "B-3 잉여와 부족", "B-4 변동성", "B-5 최적 비중")): _st([S(_n)], _nm)
_st([B6A, B6H, B6B, B6C], "B-6 Das–Markowitz 원 논문"); _st([S(77)], "C 용어집"); _st([C2], "C-2 한국 제도 용어"); _st([RD], "D 참고문헌")

# 머리 띠 · 쪽번호
UN = {u[0]: u[1] for u in UNITS}
for i, (u, s_) in enumerate(ORDER, 1):
    for x in s_.shapes:
        if not x.has_text_frame: continue
        t = x.text_frame.text.strip()
        if t.startswith("BAF.60080 · W06 GBI ·"):
            tag = "도입" if u == "도입" else "부록" if u == "부록" else f"{u} {UN[u]}"
            sub = SUBT.get(id(s_._element))
            if sub: tag = (f"부록 · {sub}" if u == "부록" else f"{tag} · {sub}")
            lecfix.set_text(x, f"BAF.60080 · W06 GBI · {tag}")
        xx, yy, ww, hh = lecfix.px(x)
        if xx >= 1000 and yy >= 640 and t.isdigit(): lecfix.set_text(x, f"{i:02d}")

left = [t for _, t in lecfix.grep(P, "교시") if not t.startswith("단원마다")]
assert not left, left
under = [(i, t) for i, t in lecfix.grep(P, r"[A-Za-zα-ωΑ-Ω]_[A-Za-z0-9{]") if not any(k in t for k in ("fml_w8", "gbi_mc", "W06_LDI", "priority_mc", "dm_mc", "DasMarkowitzMC", "ClaudeCode"))]
assert not under, under
P.save(SRC)
orig_no = {id(s_._element): n for n, s_ in enumerate(O, 1)}
KEY = {id(v._element): k for k, v in dict(TOC=N2, GLOSS=GL, DMI=DMI, DM1=DM1, DMB=DMB, DM2=DM2, CMP=CMP, CMP2=CMP2, WHY=WHY, P23=P23,
       H1=H1, MAP=MAP, MC1=MC1, MC2=MC2, MC3=MC3, MC4=MC4, RA=RA, KT=KT, MX=MX, CB1=CB1, CB2=CB2, CC=CC, VB=VB, DMC1=DMC1, DMC2=DMC2, MXG=MXG, MXR=MXR, MXP=MXP, VB2=VB2, CCB=CCB, U3=U3, XL=XL, FD=FD, LO=LO, HC1=HC1, HC2=HC2, B6H=B6H, K1=K1, IRPS=IRPS, K2=K2, K3=K3, I1=I1, I2=I2, I3=I3, C2=C2, RD=RD, V1=V1, RC=RC, V2=V2, V2B=V2B, M1=M1, M2D=M2D, B6A=B6A, B6B=B6B, B6C=B6C, H2=H2, FX1=FX1, FX2=FX2, H3=H3, **{"DIV" + k: v for k, v in DIV.items()}).items()}
rows = []
for i, (u, s_) in enumerate(ORDER, 1):
    title = ""
    for x in s_.shapes:
        if x.has_text_frame and lecfix.px(x)[1] in range(60, 75) and lecfix.px(x)[2] > 1000: title = x.text_frame.text.strip(); break
    if u == "div" or not title:
        title = " / ".join(x.text_frame.text.strip() for x in s_.shapes if x.has_text_frame and x.text_frame.text.strip())[:60]
    rows.append({"new": i, "unit": u, "title": title, "orig": orig_no.get(id(s_._element), "신설"), "key": KEY.get(id(s_._element), "")})
json.dump(rows, open(os.path.join(HERE, "order.json"), "w"), ensure_ascii=False, indent=1)
print("saved", SRC, len(ORDER), "장")
