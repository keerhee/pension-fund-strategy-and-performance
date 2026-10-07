# -*- coding: utf-8 -*-
"""7주차 용어 통일: '재균형' -> '리밸런싱' (2026-10-07, 사용자 지시).
범위: W07 폴더 PPTX(본문·표·그룹·노트)·md, W07 팀과제 docx·py, 파일명(프라이머·SS2 강의본 pptx/pdf/mp4),
README·사이트·커리큘럼 배치표·보충교재 README의 7주차 표기.
원본은 _구판/재균형to리밸런싱_구판_2026-10-07/ 에 같은 상대경로로 보관(있으면 거기서 되돌려 재실행 결정적).
실행: .venv/bin/python _build/rebal_term/apply.py [--apply]"""
import glob, os, shutil, sys, urllib.parse
from pptx import Presentation
import docx
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))); os.chdir(ROOT)
APPLY = "--apply" in sys.argv
A, B = "재균형", "리밸런싱"
BK = "_구판/재균형to리밸런싱_구판_2026-10-07"
W = "W07_동적포트폴리오와장기투자"

def backup_restore(f):
    b = os.path.join(BK, f)
    if os.path.exists(b): shutil.copy2(b, f)
    else:
        os.makedirs(os.path.dirname(b), exist_ok=True); shutil.copy2(f, b)

def para_replace(p):
    t = p.text
    if A not in t: return 0
    runs = p.runs
    if any(A in r.text for r in runs) and "".join(r.text.replace(A, B) for r in runs) == t.replace(A, B):
        for r in runs: r.text = r.text.replace(A, B)
    elif runs:                       # 낱말이 run 경계에 걸침: 첫 run에 모은다
        runs[0].text = t.replace(A, B)
        for r in runs[1:]: r.text = ""
    return t.count(A)

def tfs(sh):
    if sh.has_text_frame: yield sh.text_frame
    if getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells: yield c.text_frame
    if sh.shape_type == 6:
        for s in sh.shapes: yield from tfs(s)

log = []
for f in sorted(glob.glob(f"{W}/*.pptx")):
    if APPLY: backup_restore(f)
    prs = Presentation(f); n = 0
    for s in prs.slides:
        src = [tf for sh in s.shapes for tf in tfs(sh)]
        if s.has_notes_slide: src.append(s.notes_slide.notes_text_frame)
        for tf in src:
            for p in tf.paragraphs: n += para_replace(p)
    if n:
        log.append((n, f))
        if APPLY: prs.save(f)

def docx_paras(d):
    yield from d.paragraphs
    for t in d.tables:
        for r in t.rows:
            for c in r.cells: yield from c.paragraphs
    for s in d.sections:
        for part in (s.header, s.footer):
            yield from part.paragraphs
for f in sorted(glob.glob("Assignment/W07_Team_Assignment/**/*.docx", recursive=True)):
    d = docx.Document(f); n = sum(p.text.count(A) for p in docx_paras(d))
    if n:
        log.append((n, f))
        if APPLY:
            backup_restore(f); d = docx.Document(f)
            for p in docx_paras(d): para_replace(p)
            d.save(f)

TEXT = [f"{W}/W07_특별세션_SS2_케이스_모범답안_교수용.md",
        "Assignment/W07_Team_Assignment/1_Student/Data/team3_4_Rebalancing/rebal_tool.py",
        "README.md", "site/index.html", "_과정운영/16주_커리큘럼_배치표.md", "보충교재/README.md"]
qa, qb = urllib.parse.quote(A), urllib.parse.quote(B)
for f in TEXT:
    s = open(f, encoding="utf-8").read()
    n = s.count(A) + s.count(qa)
    if n:
        log.append((n, f))
        if APPLY:
            backup_restore(f); s = open(f, encoding="utf-8").read()
            open(f, "w", encoding="utf-8").write(s.replace(A, B).replace(qa, qb))

RENAME = [f for f in sorted(glob.glob(f"{W}/*{A}*"))]
for f in RENAME:
    log.append((0, f"rename {f} -> {os.path.basename(f).replace(A, B)}"))
    if APPLY:
        b = os.path.join(BK, f)
        if not os.path.exists(b):
            os.makedirs(os.path.dirname(b), exist_ok=True); shutil.copy2(f, b)
        os.rename(f, f.replace(A, B))
for n, f in log: print(f"{n:3d}  {f}")
print("(APPLIED)" if APPLY else "(dry run)")
