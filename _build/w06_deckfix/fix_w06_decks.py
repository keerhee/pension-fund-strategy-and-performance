"""W06 덱 7개 보수(2026-10-04) — 나레이션 집필 중 발견한 오류와 '돈' 표기.
실행: .venv/bin/python _build/w06_deckfix/fix_w06_decks.py   (원본은 _구판/W06_덱보수_구판_2026-10-04/에 한 번만 백업)"""
import os, re, shutil, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "lecfix"))
from pptx import Presentation
from pptx.util import Inches
import lecfix
R = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
W = os.path.join(R, "W06_LDI와GBI"); BK = os.path.join(R, "_구판", "W06_덱보수_구판_2026-10-04")
DECKS = {"m7": "W06_M7_LDI_부채연계투자_강의본", "m7-data": "W06_M7_실습데이터_LDI부채연계투자", "m7-case": "W06_M7_케이스_국민연금LDI도입_IC",
         "m8": "W06_M8_GBI_목표기반투자_강의본", "m8-data": "W06_M8_실습데이터_GBI목표기반투자",
         "m8-case1": "W06_M8_케이스1_퇴직연금디폴트옵션도입_IC", "m8-case2": "W06_M8_케이스2_대학발전기금Yale모델도입_IC"}
FIX = {"m7": [("3-4의 DV01", "2교시의 DV01"), ("3장의 도구로", "2교시의 도구로")],
       "m7-case": [("30분 안에", "15분 안에")],
       "m8": [("다음 — 위험관리와 성과평가", "다음 주 W07 — 동적 포트폴리오"), ("만든 포트폴리오를 지키고 측정하는 법", "시간이 흐르면 최적 비중도 변하는가")],
       "m8-case1": [("m ≥ 3은 ①②④에서 탈락", "m ≥ 3은 ①에서(m ≥ 5는 ②④도) 탈락")],
       "m8-case2": [("국채도 1년차 −12%", "안전 유동자산(국채 + 현금)도 1년차 −12%")]}
AMOUNT = ("줄 ", "받는 ", "물어줄 ", "번 ", "늘어난 ", "할 ")

def money(p):
    """문단 안 '돈'을 문맥에 맞는 말로(받침이 같아 조사가 바뀌지 않는다). 바꾼 수를 돌려준다."""
    full = "".join(r.text for r in p.runs); edits = []   # (전체 위치, 지울 길이, 새 말)
    for m in re.finditer("돈", full):
        i = m.start(); after = full[i:i + 5]; before = full[:i]
        if after.startswith("돈다"): continue
        if before.endswith("목"): edits.append((i - 1, 2, "일시금")); continue
        if after.startswith("돈의 시간"): new = "화폐"
        elif before.endswith("가진 "): new = "자산"
        elif before.endswith(AMOUNT): new = "금액"
        else: new = "자금"
        edits.append((i, 1, new))
    for pos, ln, new in sorted(edits, reverse=True):   # 뒤에서부터 run 단위로 고친다
        acc = 0
        for r in p.runs:
            t = r.text
            if acc <= pos < acc + len(t):
                k = pos - acc
                if k + ln <= len(t): r.text = t[:k] + new + t[k + ln:]
                else:  # '목|돈'처럼 run에 걸치면 앞 run 끝을 지우고 다음 run의 '돈'을 바꾼다
                    r.text = t[:k] + new; rest = ln - (len(t) - k)
                    nxt = p.runs[p.runs.index(r) + 1] if False else None
                    for r2 in list(p.runs)[list(p.runs).index(r) + 1:]:
                        if rest <= 0: break
                        cut = min(rest, len(r2.text)); r2.text = r2.text[cut:]; rest -= cut
                break
            acc += len(t)
    return len(edits)

os.makedirs(BK, exist_ok=True)
for k, base in DECKS.items():
    src = os.path.join(W, base + ".pptx"); bk = os.path.join(BK, base + ".pptx")
    if not os.path.exists(bk): shutil.copy2(src, bk)
    P = Presentation(bk)
    log = []
    for old, new in FIX.get(k, []):
        n = lecfix.replace_all(P, old, new); log.append(f"{old!r}→{n}")
        if n == 0: sys.exit(f"[{k}] 찾지 못함: {old}")
    nm = 0
    for s in P.slides:
        for sh in s.shapes:
            for tf in lecfix._tf_iter(sh):
                for p in tf.paragraphs: nm += money(p)
    if k == "m8":   # 41장 다크 패널 불릿이 패널 밖으로 넘친다 → 블록을 위로 올리고 불릿 상자를 키운다
        s = P.slides[40]; by = {sh.shape_id: sh for sh in s.shapes}
        by[11].top, by[12].top, by[13].top = Inches(4.10), Inches(4.46), Inches(4.98); by[13].height = Inches(0.86)
        log.append("s41 패널 재배치")
    left = sum(1 for s in P.slides for sh in s.shapes for tf in lecfix._tf_iter(sh) for p in tf.paragraphs
               if re.search(r"돈(?!다)", "".join(r.text for r in p.runs)))
    P.save(src)
    print(f"{k}: {' · '.join(log) or '-'} · 돈 {nm}곳 · 남은 돈 {left}")
