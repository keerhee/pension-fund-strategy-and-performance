"""보충교재 17 「LDI · GBI 수식 도출」 받은 덱 02·04 보수(2026-10-05) — 영상 대본 집필 전 점검에서 나온 것.

02 LDI 수식 도출 II
  · 8장 '같은 글자를 두 뜻으로 쓰지 않습니다' 정정 + 겹치는 글자 목록 각주, 18·40·42·53장에 그 장의 뜻 ※ 각주
    (κ: 40·42장 가격 충격 ↔ 53장 헤지 반영 비율 · S: 40장 강제 매도량 ↔ 잉여금 · D₁: 18장 배당 ↔ 듀레이션)
  · 44장 SVB 손계산 '117 × 6 × 0.025 ≈ 175억' → '1,170억 × 6 × 0.025 ≈ 176억'(175.5, 01 덱과 같게)
  · 52장 같은 h = 100%의 두 구현(3-3 국채 50 + 스왑 30 ↔ 3-4 사다리 20 + 국채 30 + 스왑 45) ※ 각주
  · 화면의 '돈' → 자금·금액(글자 3·7·22·26·27·31·34·36장, 그림 9장 표·70장 행 머리)
04 Martellini GBI 수식 도출 III
  · 7장 같은 문구 정정 + 목록 각주, 38·39·40·43·68·69·70장 ※ 각주 (A·C·K 상수 · 평균 mₚ · 인출액 Cₜ)
  · 36장 표 그림을 읽히게(그래프를 줄이고 표를 오른쪽 아래에 키워 둔다)
  · 11장 강조 문구가 둘째 수식 상자와 겹침 → 상자를 올리고 강조 두 줄을 내린다
  · 9·63장 마지막 글머리 줄이 상자 아래 경계에 걸림 → 상자를 늘리고 내용을 올린다
  · 화면의 '돈'(글자 6·18·19·48·51·53·54·55·69장, 그림 7·8·18·69장 표)
실행: <python-pptx 파이썬> _build/sup17_fix/fix_sup17.py
원본은 _구판/보충17_덱보수_구판_2026-10-05/에 한 번만 백업하고 늘 백업에서 읽으므로 재실행해도 결과가 같다.
그림 낱말 교체는 glyphswap.py(같은 폴더, Pillow·numpy, 글꼴 ~/Library/Fonts/NotoSansCJK.ttc).
"""
import copy, io, os, re, shutil, sys
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "lecfix")); sys.path.insert(0, HERE)
from pptx import Presentation
from pptx.util import Pt, Emu
from PIL import Image
import lecfix
from glyphswap import swap_word

R = os.path.abspath(os.path.join(HERE, "..", ".."))
D = os.path.join(R, "보충교재", "17_LDI_GBI 수식도출")
BK = os.path.join(R, "_구판", "보충17_덱보수_구판_2026-10-05"); os.makedirs(BK, exist_ok=True)
DECKS = {"02": os.path.join(D, "LDI", "02_LDI_Derivations_II_zeroone.pptx"),
         "04": os.path.join(D, "GBI", "04_Martellini_GBI_Derivations_III_zeroone.pptx")}
PX = 9525  # EMU per px (96 dpi)

# ---------------------------------------------------------------- 공용 ----
AMOUNT = ("줄 ", "받는 ", "물어줄 ", "번 ", "늘어난 ", "할 ", "의 ")
def money(p):
    """문단 안 '돈'을 문맥에 맞는 말로(받침이 같아 조사가 안 바뀐다). 바꾼 수를 돌려준다. (W06 fix_w06_decks.money 확장)"""
    full = "".join(r.text for r in p.runs); edits = []
    for m in re.finditer("돈", full):
        i = m.start(); before = full[:i]
        if full[i:i + 2] == "돈다": continue
        if before.endswith(("가진 ", "지금 ")): new = "자산"
        elif before.endswith(AMOUNT): new = "금액"
        else: new = "자금"
        edits.append((i, new))
    for pos, new in sorted(edits, reverse=True):
        acc = 0
        for r in p.runs:
            t = r.text
            if acc <= pos < acc + len(t):
                k = pos - acc; r.text = t[:k] + new + t[k + 1:]; break
            acc += len(t)
    return len(edits)

def money_all(P):
    n = sum(money(p) for s in P.slides for sh in s.shapes for tf in lecfix._tf_iter(sh) for p in tf.paragraphs)
    left = [i for i, _ in lecfix.grep(P, r"돈(?!다)")]
    return n, left

def rep(P, old, new, want=1):
    n = lecfix.replace_all(P, old, new)
    if n != want: sys.exit(f"치환 수 {n} ≠ {want}: {old}")

def shape(P, n, sid):
    return next(sh for sh in P.slides[n - 1].shapes if sh.shape_id == sid)

def move(P, n, sid, x=None, y=None, w=None, h=None):
    sh = shape(P, n, sid)
    if x is not None: sh.left = Emu(int(x * PX))
    if y is not None: sh.top = Emu(int(y * PX))
    if w is not None: sh.width = Emu(int(w * PX))
    if h is not None: sh.height = Emu(int(h * PX))

def note(P, n, text, tmpl, y=630, h=44):
    """※ 각주: 하단 강조 문구(584~624) 아래, 쪽번호(678) 위. 소제목 상자(회색 16pt)를 본떠 12pt로."""
    s = P.slides[n - 1]
    el = copy.deepcopy(tmpl._element)
    el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set(
        "id", str(max(sh.shape_id for sh in s.shapes) + 1))
    el.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr").set("name", "Note sup17")
    s.shapes._spTree.append(el)
    sh = s.shapes[-1]
    sh.left, sh.top, sh.width, sh.height = Emu(79 * PX), Emu(y * PX), Emu(1016 * PX), Emu(h * PX)
    sh.text_frame._txBody.bodyPr.set("anchor", "t")
    lecfix.set_text(sh, text)
    for r in sh.text_frame.paragraphs[0].runs: r.font.size = Pt(12)

def swap_pic(P, n, sid, jobs):
    """그림 표 안의 '돈' 교체. jobs: [(x, y, 새 낱말, 굵게, 칸 오른쪽 끝)] — 좌표는 그림 원본 픽셀."""
    sh = shape(P, n, sid)
    im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
    for x, y, new, bold, rl in jobs:
        im, info = swap_word(im, x, y, new, bold, rl)
    b = io.BytesIO(); im.save(b, "PNG")
    sh.part.related_part(sh._element.blipFill.blip.rEmbed)._blob = b.getvalue()

FOOT_OLD, FOOT_NEW = "같은 글자를 두 뜻으로 쓰지 않습니다.", "같은 글자를 두 뜻으로 쓰는 곳은 그 장에서 밝힙니다."

# ---------------------------------------------------------------- 02 ----
def fix02(P):
    tm = shape(P, 8, 5)   # 소제목(회색 16pt) — 각주 서식 원형
    rep(P, FOOT_OLD, FOOT_NEW)
    note(P, 8, "※ 두 뜻으로 쓰는 글자 — κ: 40·42장은 강제 매도의 가격 충격, 53장은 헤지 반영 비율 · "
               "S: 40장은 강제 매도량, 그 밖은 잉여금 · D₁: 18장은 내년 배당, 그 밖의 D는 듀레이션", tm)
    note(P, 18, "※ 이 장의 D₁은 내년 배당(dividend)이고 D(eq)가 주식 듀레이션입니다 — 다른 장의 D는 모두 듀레이션입니다.", tm)
    note(P, 40, "※ 이 장의 Sₜ는 t일의 강제 매도량, κ는 매도가 금리를 밀어 올리는 정도(가격 충격)입니다 — "
                "다른 장의 S(잉여금), 53장의 κ(헤지 반영 비율)와 다른 뜻입니다.", tm)
    note(P, 42, "※ κ는 강제 매도가 금리를 밀어 올리는 정도(40장)입니다 — 53장의 κ(헤지 반영 비율)와 다른 뜻입니다.", tm)
    note(P, 53, "※ 이 장의 κ는 부채 중 헤지에 반영하는 비율입니다 — 40·42장의 κ(강제 매도의 가격 충격)와 다른 뜻입니다.", tm)
    rep(P, "117 × 6 × 0.025 ≈ 175억 달러 손실", "1,170억 × 6 × 0.025 ≈ 176억 달러 손실")
    note(P, 52, "※ 같은 h = 100%를 이루는 조합은 여럿입니다 — 여기(3-3)는 사다리 없이 30년 국채 50 + 스왑 명목 30, "
                "35장(3-4)은 사다리 20 + 30년 국채 30 + 스왑 명목 45로 채웁니다.", tm)
    swap_pic(P, 9, 10, [(1192, 141, "금액", False, 1700)])     # 회사가 망하면 물어줄 돈은? → 금액은?
    swap_pic(P, 70, 7, [(30, 343, "자금", True, 440)])         # 행 머리 '돈' → '자금'

# ---------------------------------------------------------------- 04 ----
def fix04(P):
    tm = shape(P, 9, 5)
    rep(P, FOOT_OLD, FOOT_NEW)
    note(P, 7, "※ 두 뜻으로 쓰는 글자 — A·C·K: 39~43장은 평균-분산 해의 상수, 그 밖은 자산·쿠션·행사가(56장) · "
               "m: 38~40장은 평균(mₚ, m(γ)), 그 밖은 CPPI 승수 · Cₜ: 68~70장은 그해 인출액, 그 밖의 C는 쿠션", tm)
    note(P, 38, "※ 이 장의 mₚ는 포트폴리오의 평균 수익률입니다 — CPPI 승수 m(3절)과 다른 뜻입니다.", tm)
    note(P, 39, "※ 39~43장의 A·C·K는 평균-분산 해를 정리하는 상수, m(γ)는 평균입니다 — "
                "자산 A·쿠션 C·행사가 K·CPPI 승수 m과 다른 뜻입니다.", tm)
    for n in (40, 43):
        note(P, n, "※ A·C·K는 39장의 평균-분산 상수입니다 — 자산 A·쿠션 C·행사가 K와 다른 뜻입니다.", tm)
    for n in (68, 69, 70):
        note(P, n, "※ 이 장의 Cₜ는 t해의 인출액입니다 — CPPI의 쿠션 C(3절)와 다른 뜻입니다.", tm)
    # 36장: 그래프를 줄여 오른쪽 위에, 표를 그 아래에 키워 둔다(원래 표는 109×26px)
    move(P, 36, 14, x=750, y=220, w=361, h=232)
    move(P, 36, 13, x=680, y=458, w=500, h=round(500 * 308 / 1320))
    # 11장: 둘째 상자를 올리고 하단 강조 두 줄을 내린다
    move(P, 11, 10, h=136); move(P, 11, 11, y=305); move(P, 11, 12, y=341)
    move(P, 11, 13, y=440, h=130); move(P, 11, 14, y=452); move(P, 11, 15, y=486, h=72)
    move(P, 11, 16, y=578); move(P, 11, 17, y=616)
    # 9·63장: 마지막 상자를 위로 늘리고 내용을 올린다
    move(P, 9, 13, y=423, h=155); move(P, 9, 14, y=440); move(P, 9, 15, y=480, h=88)
    move(P, 63, 13, y=428, h=150); move(P, 63, 14, y=445); move(P, 63, 15, y=485, h=83)
    rep(P, "인출 뒤 남는 돈은", "인출 뒤 남는 자산은")
    swap_pic(P, 7, 6, [(805, 128, "금액", False, 1620), (905, 768, "자산", False, 1620)])
    swap_pic(P, 8, 6, [(935, 323, "금액", False, 1600)])
    swap_pic(P, 18, 7, [(650, 280, "자금", True, 1400), (1525, 278, "자금", True, 2330)])
    swap_pic(P, 69, 13, [(910, 47, "자산", True, 1040)])

for k, src in DECKS.items():
    bk = os.path.join(BK, os.path.basename(src))
    if not os.path.exists(bk): shutil.copy2(src, bk)
    P = Presentation(bk); n0 = len(P.slides)
    (fix02 if k == "02" else fix04)(P)
    nm, left = money_all(P)
    assert len(P.slides) == n0
    P.save(src)
    print(f"{k}: saved · 글자 '돈' {nm}곳 · 남은 '돈' 장 {left}")
