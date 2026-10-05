# -*- coding: utf-8 -*-
"""W06 팀 과제 — 팀 번호 교체(3팀 = 디폴트옵션, 4팀 = H대) + 3·4팀 「강의와 연결」 + H대 Flexicure 가이드(2026-10-05).

사용자 피드백: "H대 케이스는 강의의 어떤 부분을 이용할지까지 구체적으로 명시해야 한다. Flexicure의 개념을 이용하되
(디폴트 옵션처럼) 비유동성을 고려하는 가이드를 주어야 한다. 디폴트 옵션을 팀3, H대를 팀4로."
원본은 _구판/W06_Assignment_팀번호교체전_2026-10-05/에서 읽는다(재실행해도 결과가 같다).
숫자는 Data/team4_H_University(h_univ_tool.py와 같은 가정)에서 손계산 — 아래 CHECK가 다시 검산한다.
실행: <python-docx 있는 파이썬> _build/w06_assignment/swap_teams_guide.py → 그다음 make_pdfs.sh"""
import copy, os, re
import docx

HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", ".."))
SRC = os.path.join(R, "_구판", "W06_Assignment_팀번호교체전_2026-10-05")
DST = os.path.join(R, "Assignment", "W06_Team_Assignment")

# ------------------------------------------------------------------ 검산
W = 5000; SP3 = 175.0 + 178.5 + 182.1
def flex(e, b, c, x):
    calls = 0.4 * x / 100 * W; floor = SP3 + 600 + calls
    liq = W * (e + b + c) / 100; cush = liq - floor
    l08 = W * (e / 100 * 0.6 + b / 100 * 1.08 + c / 100 * 1.03); c08 = l08 - floor
    return dict(floor=floor, cush=cush, m=(W * e / 100 + W * x / 100) / cush, l08=l08, c08=c08,
                m08=(W * e / 100 * 0.6 + W * x / 100 * 0.9) / c08, alt08=W * x / 100 * 0.9 / c08,
                share08=W * x / 100 * 0.9 / (l08 + W * x / 100 * 0.9))
FA, FB, FC = flex(65, 30, 5, 0), flex(50, 25, 5, 20), flex(55, 10, 5, 30)
CHECK = [(round(SP3), 536), (round(FB["floor"]), 1536), (round(FB["cush"]), 2464), (round(FB["m"], 2), 1.42),
         (round(FB["l08"]), 3108), (round(FB["c08"]), 1572), (round(FB["m08"], 2), 1.53),
         (round(FA["m"], 2), 0.84), (round(FA["m08"], 2), 0.72), (round(FC["floor"]), 1736), (round(FC["cush"]), 1764),
         (round(FC["m"], 2), 2.41), (round(FC["c08"]), 712), (round(FC["m08"], 1), 4.2), (round(FC["alt08"], 2), 1.9),
         (round(FB["alt08"], 2), 0.57), (round(FB["share08"] * 100, 1), 22.5), (round(FC["share08"] * 100, 1), 35.5)]
for got, want in CHECK:
    assert got == want, (got, want)

# ------------------------------------------------------------------ 도구
def all_paras(d):
    yield from d.paragraphs
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                yield from c.paragraphs
    for s in d.sections:
        yield from s.header.paragraphs; yield from s.footer.paragraphs

SWAP = [("W06_Team3_H_University", "@F4H@"), ("W06_Team4_Default_Option", "@F3D@"),
        ("team3_H_University", "@D4H@"), ("team4_Default_Option", "@D3D@"),
        ("3팀", "@T4@"), ("4팀", "@T3@")]
BACK = [("@F4H@", "W06_Team4_H_University"), ("@F3D@", "W06_Team3_Default_Option"),
        ("@D4H@", "team4_H_University"), ("@D3D@", "team3_Default_Option"),
        ("@T4@", "4팀"), ("@T3@", "3팀"), ("4팀 ↔ 3팀", "3팀 ↔ 4팀")]
def swap(d):
    for p in all_paras(d):
        for r in p.runs:
            t = r.text
            for a, b in SWAP + BACK: t = t.replace(a, b)
            if t != r.text: r.text = t

def para_text(p, text):
    """첫 run 서식으로 문단 전체를 text로."""
    rs = p.runs
    rs[0].text = text
    for r in rs[1:]: r._r.getparent().remove(r._r)

def find(d, prefix):
    return next(p for p in d.paragraphs if p.text.startswith(prefix))

def clone_after(anchor_el, proto, texts):
    """proto 문단을 복제해 anchor 뒤에 넣는다. texts = run별 글(남는 run은 지움)."""
    el = copy.deepcopy(proto._p); anchor_el.addnext(el)
    p = docx.text.paragraph.Paragraph(el, proto._parent)
    rs = p.runs
    for i, r in enumerate(rs):
        if i < len(texts): r.text = texts[i]
        else: r._r.getparent().remove(r._r)
    return p

def clone_table_after(anchor_el, proto_tbl, rows):
    el = copy.deepcopy(proto_tbl._tbl); anchor_el.addnext(el)
    t = docx.table.Table(el, proto_tbl._parent)
    trs = el.tr_lst; body_proto = copy.deepcopy(trs[1])
    for tr in trs[1:]: el.remove(tr)
    for _ in rows[1:]: el.append(copy.deepcopy(body_proto))
    for r, vals in zip(t.rows, rows):
        for c, v in zip(r.cells, vals):
            para_text(c.paragraphs[0], v)
            for extra in c.paragraphs[1:]: extra._p.getparent().remove(extra._p)
    return t

def swap_rows(tbl, i, j):
    trs = tbl._tbl.tr_lst; a, b = trs[i], trs[j]
    a.addprevious(b) if i < j else b.addprevious(a)

# ------------------------------------------------------------------ 4팀 H대
d = docx.Document(os.path.join(SRC, "1_Student", "W06_Team3_H_University.docx")); swap(d)
H2 = find(d, "상황"); BODY = find(d, "미국 Yale"); FIRST = find(d, "H대학교(가상)")
crit_head = find(d, "4팀 판정 기준")
anchor = crit_head._p.getprevious()

h = clone_after(anchor, H2, ["강의와 연결 — 이 과제는 M8 GBI 강의본의 무엇을 쓰나"])
b1 = clone_after(h._p, FIRST, ["이 과제는 대학기금의 지출 약속을 GBI의 목표 언어로 읽는다. 장학 · 연구 고정 지출(기금의 3.5%, 175억)이 반드시 지킬 필수 목표이고, 그 위의 지출(목표 지출률 z까지)과 사모투자의 높은 수익은 희망 목표다. 문제마다 쓰는 강의 단원은 아래와 같다."])
t = clone_table_after(b1._p, d.tables[2], [
    ["이 과제에서", "M8 GBI 강의본 단원", "쓰는 개념"],
    ["문제 1 · 기준 ④ 위기 직후 현금", "④ Floor 만들기 — RB에서 GHP로", "Floor(거절할 수 없는 지출) · GHP(현금 · 국채)"],
    ["문제 2 · 지출률 상한", "③ 필요 자본과 가장 싼 포트폴리오", "목표를 지키는 데 드는 자본 · 지출 상한"],
    ["문제 3 · 기준 ① 버틸 수 있는 기간", "⑤ 쿠션과 승수 · ⑥ Floor 운용 규칙과 한계", "쿠션(여유분) · Cash Trap · 팔 수 없는 자산"],
    ["문제 4 · 기준 ② 30년 실질가치", "① 목표와 위험", "위험 = 목표 미달 확률 P(W < G)"],
    ["문제 5 · 기준 ③ 운용 역량", "⑧ 기관의 GBI — 모방 3장벽", "접근권 · 시간 지평 · 인력"],
    ["문제 6 · Flexicure로 다시 읽기", "⑤ CPPI에서 Flexicure로 (3팀 디폴트옵션과 같은 도구)", "Floor · 쿠션 · 승수 m · 매년 재설정"],
])
h2 = clone_after(t._tbl, H2, ["Flexicure를 대학기금에 옮기는 법 — 비유동성까지"])
b2 = clone_after(h2._p, FIRST, ["3팀 디폴트옵션의 CPPI · Flexicure는 “지킬 바닥(Floor)을 먼저 떼어 두고, 남은 여유분(쿠션)의 m배만 위험자산에 넣고, 해마다 다시 잰다”는 규칙이다. 대학기금에 옮길 때 다른 점은 하나다 — 사모투자는 여유분이 줄어도 팔 수 없다. 아래 다섯 단계로 옮긴다(숫자는 B안 예시, 교육용 가정)."])
STEPS = [
    ("1단계 Floor.", " 거절할 수 없는 현금 유출을 더한다 — 앞 3년 고정 지출(175 + 178.5 + 182.1 = 536억) + 공사비 600억 + 남은 출자 요청(사모투자 금액 × 40%). B: 536 + 600 + 400 = 1,536억."),
    ("2단계 GHP.", " 현금 · 국채가 충격 뒤에도 3년 지출 + 공사비를 팔지 않고 댈 수 있어야 한다(기준 ④). 출자 요청은 주식까지 포함한 유동자산에서 나간다(기준 ①)."),
    ("3단계 쿠션과 승수.", " 쿠션 = 유동자산(주식 + 국채 + 현금) − Floor, 위험 노출 = 주식 + 사모투자, 승수 m = 위험 노출 ÷ 쿠션. B: 쿠션 4,000 − 1,536 = 2,464억, m = (2,500 + 1,000) ÷ 2,464 = 1.42. 이 과제에서는 3팀 B 설계와 같은 m ≤ 2를 상한으로 본다(교육용 가정)."),
    ("4단계 비유동성 — 충격 뒤에 다시 잰다.", " 디폴트옵션은 쿠션이 줄면 주식을 팔아 m을 되돌린다. 사모투자는 팔 수 없으므로 최악 충격 직후의 쿠션으로 m을 다시 잰다. B, 2008형(주식 −40% · 국채 +8% · 현금 +3% · 사모투자 장부 −10%): 유동자산 3,108억, 쿠션 1,572억, m = (1,500 + 900) ÷ 1,572 = 1.53. 비유동 상한 x는 최악 충격 뒤에도 m ≤ 2를 지키는 최대치로 정한다 — 쿠션이 줄었을 때 디폴트옵션처럼 주식을 팔아 m을 되돌린다는 보장이 사모투자에는 없기 때문이다. 주식을 다 팔았을 때 남는 ‘사모투자 ÷ 쿠션’(B 900 ÷ 1,572 = 0.57)은 마지막 안전판으로 함께 본다. 줄일 수단은 신규 약정 중단뿐이며, 충격 뒤 사모투자 비중이 목표를 넘는 분모 효과(B 20% → 22.5%)가 그 신호다."),
    ("5단계 지출은 두 층.", " 고정 3.5%가 Floor(필수 목표), 목표 지출률 z까지는 희망 목표다. Flexicure처럼 해마다 다시 잰다 — 쿠션이 기준 아래로 내려가면 희망 몫을 동결한다. 평활 규칙(전년 80% + 목표 20%)은 그 사이의 완충이다."),
]
STEP_PROTO = find(d, "할 일: (a) B안")
prev = b2._p
for a, b in STEPS:
    prev = clone_after(prev, STEP_PROTO, [a, b])._p

# 문제 제목에 강의 단원을 붙이고, 문제 6(Flexicure)을 새로 넣는다
for old, new in [("문제 1. 세 안의 현금 흐름", "문제 1. 세 안의 현금 흐름 — 강의 ④ Floor"),
                 ("문제 2. 기대수익과 지출률 상한", "문제 2. 기대수익과 지출률 상한 — 강의 ③ 필요 자본"),
                 ("문제 3. 위기 직후 버틸 수 있는 기간 (기준 ①)", "문제 3. 위기 직후 버틸 수 있는 기간 (기준 ①) — 강의 ⑤⑥ 쿠션 · Cash Trap"),
                 ("문제 4. 30년 기금 가치 (기준 ②)", "문제 4. 30년 기금 가치 (기준 ②) — 강의 ① 목표 미달 확률"),
                 ("문제 5. 운용 역량과 위기 직후 현금 (기준 ③ · ④)", "문제 5. 운용 역량과 위기 직후 현금 (기준 ③ · ④) — 강의 ⑧ 모방 3장벽 · ④ GHP"),
                 ("문제 6. 최종 결정", "문제 7. 최종 결정")]:
    para_text(find(d, old), new)
q5_ans = find(d, "답할 것: 인력 3명 증원으로")
H3 = find(d, "문제 1. 세 안의 현금 흐름")
q6 = clone_after(q5_ans._p, H3, ["문제 6. Flexicure로 다시 읽기 — 비유동성 (강의 ⑤ Flexicure)"])
DO = find(d, "할 일: 실행 결과 [3]"); ANS = find(d, "답할 것: 세 안의 실질가치")
p = clone_after(q6._p, DO, ["할 일:", " ", "위 「Flexicure를 대학기금에 옮기는 법」 1~4단계를 A와 C에 대해 손으로 계산하고, B는 예시와 맞는지 확인한다. C는 공사비를 8년차로 미뤘지만 Floor에는 넣는다(기준 ④와 같은 이유)."])
clone_after(p._p, ANS, ["답할 것:", " ", "세 안의 Floor · 쿠션 · 평시 m, 2008형 직후 m, 주식을 다 팔았을 때의 ‘사모투자 ÷ 쿠션’, 2008형 직후 사모투자 비중을 한 표로 정리하라. m ≤ 2를 지키는 안은 어느 것인가? 지키지 못하는 안은 무엇으로 m을 낮출 수 있으며, 그 수단은 위기에 쓸 수 있는가?"])
dec = find(d, "A · B · C 가운데 하나를 고르고")
dec.runs[-1].text = dec.runs[-1].text + " 승인 조건에는 Flexicure 규칙(충격 뒤 m 상한 · 신규 약정 중단 트리거 · 희망 지출 동결 조건) 가운데 하나 이상을 숫자로 넣는다."
d.save(os.path.join(DST, "1_Student", "W06_Team4_H_University.docx")); print("4팀 H대 ok")

# ------------------------------------------------------------------ 3팀 디폴트옵션
d = docx.Document(os.path.join(SRC, "1_Student", "W06_Team4_Default_Option.docx")); swap(d)
H2 = find(d, "상황"); FIRST = find(d, "퇴직연금 디폴트옵션은")
crit_head = find(d, "3팀 판정 기준"); anchor = crit_head._p.getprevious()
h = clone_after(anchor, H2, ["강의와 연결 — 이 과제는 M8 GBI 강의본의 무엇을 쓰나"])
b1 = clone_after(h._p, FIRST, ["이 과제의 세 설계는 M8 GBI 강의본의 Floor · 쿠션 · 승수를 그대로 쓴다. 4팀(H대)은 같은 도구를 팔 수 없는 사모투자가 있는 대학기금에 옮긴다."])
clone_table_after(b1._p, d.tables[3], [
    ["이 과제에서", "M8 GBI 강의본 단원", "쓰는 개념"],
    ["문제 1 · CPPI 출발점", "④ Floor 만들기 · ⑤ 쿠션과 승수", "보장선 비용(금리) · 여유분 · m × 여유분"],
    ["문제 2 · 목표 미달 · 하위 10%", "① 목표와 위험", "위험 = 목표 미달 확률 P(W < G)"],
    ["문제 3 · 4 · 안전자산 고착", "⑥ Floor 운용 규칙과 한계", "Cash Trap — 여유분 0이면 회복에 못 탄다"],
    ["문제 6 · 등록 조건", "⑤ CPPI에서 Flexicure로 · ⑦ GBI와 현실 세계", "매년 재설정 · 수수료"],
])
d.save(os.path.join(DST, "1_Student", "W06_Team3_Default_Option.docx")); print("3팀 디폴트 ok")

# ------------------------------------------------------------------ 전체 가이드
d = docx.Document(os.path.join(SRC, "1_Student", "W06_Assignment_Guide.docx")); swap(d)
for ti in (1, 2, 6): n = len(d.tables[ti].rows); swap_rows(d.tables[ti], n - 2, n - 1)
para_text(d.tables[0].rows[2].cells[3].paragraphs[0], "3팀 퇴직연금 · 4팀 대학기금")
para_text(find(d, "1팀과 2팀은 같은 국민연금"), "1팀과 2팀은 같은 국민연금을 반대 방향의 금리 상황에서 다룬다. 3팀은 개인(퇴직연금 가입자)의 목표에서 출발해 Floor · 쿠션 · 승수로 상품을 설계하고, 4팀은 같은 도구를 팔 수 없는 사모투자가 있는 대학기금에 옮긴다.")
p = find(d, "이 과제는 네 기관의")
para_text(p, "이 과제는 네 기관의 실제 결정 문제에 이 원리를 적용해 본다. 개념 설명은 각 팀 과제에 모두 들어 있으므로 따로 공부할 자료는 필요 없다. 3 · 4팀 문서에는 「강의와 연결」 표로 M8 GBI 강의본의 해당 단원을 적어 두었다.")
p = find(d, "단계별 문제(문제 1~6)")
for r in p.runs: r.text = r.text.replace("단계별 문제(문제 1~6)", "단계별 문제(문제 1~6, 4팀은 1~7)")
p = find(d, "각 팀 문서는 같은 구조")
p.runs[-1].text = "상황 → 기본 지식 → (3 · 4팀: 강의와 연결) → 세 가지 안 → 판정 기준 → 계산 도구 사용법 → 단계별 문제 → 최종 결정 → 상대 팀이 물을 질문."
d.save(os.path.join(DST, "1_Student", "W06_Assignment_Guide.docx")); print("가이드 ok")

# ------------------------------------------------------------------ 교수용 해설(로컬 전용, 저장소 제외)
d = docx.Document(os.path.join(SRC, "2_Instructor", "W06_Answer_Key.docx")); swap(d)
swap_rows(d.tables[0], 3, 4)
body = d.element.body
hH = find(d, "4팀 — H대학교")._p; hD = find(d, "3팀 — 퇴직연금")._p; hQ = find(d, "질의 짝 강평")._p
block_D = []; el = hD
while el is not hQ: block_D.append(el); el = el.getnext()
for e in block_D: hH.addprevious(e)          # 3팀 디폴트 묶음을 4팀 H대 앞으로
tH = next(t for t in d.tables if t.rows[1].cells[1].text.startswith("A: 사모 0억"))
last = tH._tbl.tr_lst[-1]; new = copy.deepcopy(last); last.addprevious(new)
row6 = docx.table._Row(new, tH)
for c, v in zip(row6.cells, ["6",
        f"A: Floor 1,136 · 쿠션 3,864 · m 0.84 (2008형 뒤 0.72) / B: 1,536 · 2,464 · m 1.42 (2008형 뒤 1.53, 사모÷쿠션 0.57, 사모 비중 22.5%) / C: 1,736 · 1,764 · m 2.41 (2008형 뒤 4.2, 사모÷쿠션 1.90, 사모 비중 35.5%)",
        "A는 사모가 없어 주식만 팔면 되고, B는 충격 뒤에도 m ≤ 2. C는 출발부터 m 상한(2)을 넘고 2008형 뒤 4.2로 탈락한다. 주식을 다 팔아도 사모만으로 1.9 — 상한에 거의 닿아 사실상 갇히고, 위기에 쓸 수단은 신규 약정 중단뿐이다."]):
    para_text(c.paragraphs[0], v)
para_text(tH._tbl.tr_lst and docx.table._Row(tH._tbl.tr_lst[-1], tH).cells[0].paragraphs[0], "7")
d.save(os.path.join(DST, "2_Instructor", "W06_Answer_Key.docx")); print("해설 ok")
