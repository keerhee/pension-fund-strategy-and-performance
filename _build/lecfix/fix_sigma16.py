"""60/40 예시의 입력을 주식 σ 16% · 채권 σ 5% · ρ 0.2 로 통일한다 (2026-09-27).

W03 2교시(22·25쪽)가 이 값을 데이터에서 끌어내 '3교시 입력값'이라고 못 박는다. 그런데 결과 숫자가
어긋나 있었다 — W03·W05 강의본은 93/7 · ERC 22/78, 프라이머는 σ 18%로 그려 94/6 · 22%.
σ 16% · 5% · ρ 0.2 를 그대로 계산하면
    σ_p = 10.2% ,  주식 RC 비율 92.4% (≈ 92/8) ,  두 자산 ERC 주식 23.8% (≈ 24/76)
두 자산 ERC 는 상관과 무관하게 w ∝ 1/σ 이다. 그래서 결과만 고친다.

  W03 강의본  35 · 37(글 + 수식 그림) · 38(그림) · 43 · 49 · 57 · 59 · 61 · 62
  W05 강의본  2(+ w_i(Σw)_i 표기) · 5 · 7 · 9 · 11 · 12(글 + 그림) · 16
  W04 M4 강의본 11 ('W3에서는 위험의 93%')
  W05 프라이머 4 · 5 · 9(그림은 w05_art.py 를 σ 16% 로 다시 그림) · 13 · 15
원본은 _구판/σ16통일_구판_2026-09-27/. 거기서 다시 돌리면 같은 결과가 나온다.
그림 속 숫자(강의본 38·12쪽)는 원본 그림 파일이 없어 같은 폰트로 덧그린다(relabel.py).
"""
import io
import shutil
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
BAK = ROOT / '_구판' / 'σ16통일_구판_2026-09-27'
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(ROOT / '_build'))
from relabel import ink_box, repaint          # noqa: E402

EMU_PER_PX = 4315
W03 = 'W03_연기금모델과CMA/W03_연기금모델과CMA_강의본.pptx'
W05 = 'W05_리스크패리티와HRP/W05_리스크패리티와HRP_강의본.pptx'
PRI = 'W05_리스크패리티와HRP/W05_0_프라이머_리스크패리티는왜주식을줄이나.pptx'
W04 = 'W04_MVO와블랙리터맨/W04_M4_MVO와공분산추정_강의본.pptx'

# deck -> {slide: [(old, new), ...]}  — 문단 안의 부분 문자열 치환
TEXT = {
    W03: {
        35: [('~93%', '~92%')],
        37: [('위험은 93/7', '위험은 92/8'), ('주식 ~22 / 채권 ~78', '주식 ~24 / 채권 ~76')],
        43: [('~93%', '~92%')],
        49: [('ERC(주식 22 / 채권 78)', 'ERC(주식 24 / 채권 76)')],
        57: [('93/7 재현', '92/8 재현'), ('주식 RC ~93%', '주식 RC ~92%'), ('주식 ~22/채권 ~78', '주식 ~24/채권 ~76')],
        59: [('~93%', '~92%')],
        61: [('주식 RC ~93% (비중 60%) — ERC면 22/78', '주식 RC ~92% (비중 60%) — ERC면 24/76')],
        62: [('60/40의 위험 93%가 주식', '60/40의 위험 92%가 주식')],
    },
    W05: {
        2: [('w_i(Σw)_i / σ_p', 'w_i·Cov(r_i, r_p) / σ_p'), ('60/40의 주식 RC 93%', '60/40의 주식 RC 92%')],
        5: [('(주식 93%)', '(주식 92%)')],
        7: [('자본 22%인 채권이 위험 50%를 담당', '자본 24%인 주식이 위험 50%를 담당')],
        9: [('주식 RC ≈ 93% / 채권 RC ≈ 7%', '주식 RC ≈ 92% / 채권 RC ≈ 8%')],
        11: [('대략 24 : 76 (W3의 ERC 22/78)', '대략 24 : 76 (W3의 ERC와 같다)')],
        12: [('자본 22%인 채권 중심 배분이', '주식 24%인 채권 중심 배분이')],
        16: [('60/40의 주식 RC 93%', '60/40의 주식 RC 92%')],
    },
    W04: {
        11: [('W3에서는 위험의 93%가 주식', 'W3에서는 위험의 92%가 주식')],
    },
    PRI: {
        4: [('위험은 94 대 6', '위험은 92 대 8')],
        9: [('주식은 22%만', '주식은 24%만')],
        13: [('위험은 94 대 6', '위험은 92 대 8')],
        15: [('94%', '92%'), ('주식 22%만', '주식 24%만')],
    },
}

# 그림 속 숫자 덧그리기: deck -> slide -> [(search region, new label, size, face)]
RELABEL = {
    (W03, 38): [((640, 295, 880, 345), '주식 92%', 30, 'Bold'), ((690, 520, 830, 558), '채권 8%', 30, 'Bold'),
                ((1110, 170, 1350, 212), '주식 24%', 30, 'Bold'), ((1110, 355, 1350, 400), '채권 76%', 30, 'Bold')],
    (W05, 12): [((560, 285, 690, 322), '주식 92%', 23, 'Bold'),
                ((900, 160, 1090, 198), '주식 24%', 23, 'Bold'), ((900, 340, 1090, 380), '채권 76%', 23, 'Bold')],
}


def paragraphs(slide):
    for sh in slide.shapes:
        if sh.has_text_frame:
            yield from sh.text_frame.paragraphs
        elif sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    yield from cell.text_frame.paragraphs


def replace_in_para(p, old, new):
    """run 하나 안에 있으면 그 run 만, 여러 run 에 걸치면 걸친 run 들을 첫 run 으로 합친다."""
    runs = p.runs
    for r in runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            return True
    text = ''.join(r.text for r in runs)
    i = text.find(old)
    if i < 0:
        return False
    j = i + len(old)
    pos, first = 0, None
    for r in runs:
        a, b = pos, pos + len(r.text)
        pos = b
        if b <= i or a >= j:
            continue
        if first is None:
            first = r
            r.text = r.text[:i - a] + new + (r.text[j - a:] if b > j else '')
        else:
            rest = r.text[j - a:] if b > j else ''
            r.text = rest
    return True


def swap_blob(slide, pic, png_bytes, keep_box=True):
    part, rid = slide.part.get_or_add_image_part(io.BytesIO(png_bytes))
    pic._element.blipFill.find(qn('a:blip')).set(qn('r:embed'), rid)


def pic_image(pic):
    return Image.open(io.BytesIO(pic.image.blob)).convert('RGB')


def main():
    BAK.mkdir(parents=True, exist_ok=True)
    for rel in (W03, W05, PRI, W04):
        src = ROOT / rel
        if not (BAK / src.name).exists():
            shutil.copy2(src, BAK / src.name)
            if src.with_suffix('.pdf').exists():
                shutil.copy2(src.with_suffix('.pdf'), BAK / src.with_suffix('.pdf').name)

    # 프라이머 그림 — σ 16% 로 다시 그린다
    import w05_art
    import numpy as np
    assert abs(w05_art.SD[0] - 0.16) < 1e-9, 'w05_art.py 의 SD 를 먼저 0.16 으로'
    tmp = HERE / '_sig_png'
    tmp.mkdir(exist_ok=True)
    w05_art._OUT = str(tmp)
    import primer_lib
    primer_lib._OUT = str(tmp)
    figs = {4: w05_art.capital_vs_risk(), 5: w05_art.seesaw(), 9: w05_art.equalize()}

    for rel in (W03, W05, PRI, W04):
        prs = Presentation(BAK / Path(rel).name)
        slides = list(prs.slides)
        for n, pairs in TEXT[rel].items():
            paras = list(paragraphs(slides[n - 1]))
            for old, new in pairs:
                hit = [p for p in paras if old in p.text]
                assert len(hit) >= 1, (rel, n, old)
                for p in hit:
                    assert replace_in_para(p, old, new), (rel, n, old)
        # 그림 덧그리기
        for (deck, n), jobs in RELABEL.items():
            if deck != rel:
                continue
            pic = next(sh for sh in slides[n - 1].shapes if sh.shape_type == 13)
            im = pic_image(pic)
            for region, label, size, face in jobs:
                box, bg, col = ink_box(im, region)
                repaint(im, box, bg, col, label, size, face)
            buf = io.BytesIO(); im.save(buf, 'PNG')
            swap_blob(slides[n - 1], pic, buf.getvalue())
        # W03 37쪽 수식 그림
        if rel == W03:
            import matplotlib
            matplotlib.use('Agg')
            import matplotlib.pyplot as plt
            plt.rcParams.update({'mathtext.fontset': 'cm'})
            tex = (r'$\sigma_p\;=\;\sqrt{w^{\top}\Sigma w}\;\approx\;10.2\%\qquad\Rightarrow\qquad'
                   r' \mathrm{RC}_{\mathrm{eq}}\;\approx\;92\%,\quad \mathrm{RC}_{\mathrm{bond}}\;\approx\;8\%$')
            f = plt.figure(figsize=(16, 2.5), facecolor='#F4F5F9')
            f.text(.5, .5, tex, ha='center', va='center', fontsize=30, color='#1B2C5E')
            out = tmp / 'w03p37.png'
            f.savefig(out, bbox_inches='tight', facecolor='#F4F5F9', pad_inches=0.26, dpi=200)
            plt.close(f)
            pic = next(sh for sh in slides[36].shapes if sh.shape_type == 13)
            w_px, h_px = Image.open(out).size
            cx, cy = pic.left + pic.width // 2, pic.top + pic.height // 2
            w, h = w_px * EMU_PER_PX, h_px * EMU_PER_PX
            k = min(1.0, pic.width * 1.05 / w)
            swap_blob(slides[36], pic, out.read_bytes())
            pic.width, pic.height = int(w * k), int(h * k)
            pic.left, pic.top = cx - pic.width // 2, cy - pic.height // 2
        # 프라이머 그림 교체 — 같은 함수로 그렸으니 비율이 같다
        if rel == PRI:
            for n, png in figs.items():
                pic = next(sh for sh in slides[n - 1].shapes if sh.shape_type == 13)
                ow, oh = pic_image(pic).size
                nw, nh = Image.open(png).size
                assert abs(ow / oh - nw / nh) < 0.02, (n, ow / oh, nw / nh)
                swap_blob(slides[n - 1], pic, Path(png).read_bytes())
        prs.save(ROOT / rel)
        print('saved', Path(rel).name)
    shutil.rmtree(tmp)


if __name__ == '__main__':
    main()
