# -*- coding: utf-8 -*-
"""본문 ↔ 부록 이동 버튼 — 본문에서 부록으로 가는 버튼과, 부록에서 그 본문으로 돌아오는 버튼을 단다.

· 버튼은 꼬리 줄(워드마크와 쪽번호 사이)에 놓는 알약 도형이고, 도형 클릭 = 그 장으로 이동.
· 본문의 '부록 B-3' 같은 글자도 같은 곳으로 가는 링크로 만든다.
· LibreOffice PDF 변환이 슬라이드 이동을 PDF 내부 링크(GoTo)로 옮기므로 PDF에서도 눌린다.
· 다시 돌려도 같은 결과 — 전에 단 NAV_ 도형은 지우고 새로 단다. fix_* 스크립트를 다시 돌린 뒤에는 이것을 마지막에 돌린다.
"""
import copy, os, re, subprocess, sys, zipfile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# 덱 → [(본문 쪽, 부록 쪽들, 본문에 적힌 부록 이름)]  — 쪽은 1부터
DECKS = {
    "W15_TPA/W15_TPA_제1부_사일로해체와설계_강의본": [
        (21, [71, 72], "부록 B-3"),
        (23, [73], "부록 B-4"),
        (23, [74], "B-5"),
        (47, [70], "부록 B-2"),
    ],
}

NAVY, TINT, FONT = RGBColor(0x1B, 0x2C, 0x5E), RGBColor(0xE8, 0xEE, 0xF8), "Noto Sans CJK KR"
Y, H, RIGHT, GAP = 7.10, 0.30, 11.55, 0.12      # 꼬리 줄: 워드마크(0.62~8.62)와 쪽번호(11.71~) 사이


def pill(slide, x, w, text, target, tag):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(Y), Inches(w), Inches(H))
    sh.name = "NAV_" + tag
    sh.adjustments[0] = 0.5
    # PDF에서는 링크 글자가 테마의 링크 색으로 칠해진다 — 그래서 바탕은 밝게, 링크 색은 남색으로 맞춘다(theme_link_color)
    sh.fill.solid(); sh.fill.fore_color.rgb = TINT
    sh.line.color.rgb = NAVY; sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(12); r.font.bold = True; r.font.name = FONT; r.font.color.rgb = NAVY
    sh.click_action.target_slide = target                         # PowerPoint·Keynote: 도형 어디를 눌러도
    rPr = r._r.get_or_add_rPr()                                   # PDF: LibreOffice는 도형 클릭을 안 옮기고 글자 링크만 옮긴다
    rId = slide.part.relate_to(target.part, RT.SLIDE)
    rPr.append(rPr.makeelement(qn("a:hlinkClick"), {qn("r:id"): rId, "action": "ppaction://hlinksldjump"}))
    rPr.set("u", "none")
    return sh


def theme_link_color(prs, rgb="1B2C5E"):
    """테마의 링크 색(hlink · folHlink)을 덱 남색으로 — 본문 링크도 파랑 대신 남색 밑줄이 된다."""
    from lxml import etree
    for rel in prs.slide_master.part.rels.values():
        if rel.reltype == RT.THEME:
            part = rel.target_part
            root = etree.fromstring(part.blob)
            for tag in ("a:hlink", "a:folHlink"):
                el = root.find(".//" + qn(tag))
                for c in list(el): el.remove(c)
                etree.SubElement(el, qn("a:srgbClr"), val=rgb)
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def pills(slide, items):
    """items = [(글자, 목표 슬라이드, 폭)] — 오른쪽 끝(쪽번호 앞)에서부터 왼쪽으로 늘어놓는다."""
    x = RIGHT
    for text, target, w in reversed(items):
        x -= w
        pill(slide, x, w, text, target, re.sub(r"\W", "", text))
        x -= GAP


def link_text(slide, needle, target):
    """본문 글자 needle을 target 장으로 가는 링크로 만든다 (런을 셋으로 나눠 가운데에만 링크)."""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for p in sh.text_frame.paragraphs:
            for r in p.runs:
                if needle not in r.text:
                    continue
                if r._r.find(qn("a:rPr")) is not None and r._r.find(qn("a:rPr")).find(qn("a:hlinkClick")) is not None:
                    return True                                  # 이미 링크
                pre, post = r.text.split(needle, 1)
                runs = []
                for t in (pre, needle, post):
                    if t:
                        nr = copy.deepcopy(r._r); nr.find(qn("a:t")).text = t; runs.append((t, nr))
                for t, nr in runs:
                    r._r.addprevious(nr)
                    if t == needle:
                        rPr = nr.find(qn("a:rPr"))
                        if rPr is None:
                            rPr = nr.makeelement(qn("a:rPr"), {}); nr.insert(0, rPr)
                        rId = slide.part.relate_to(target.part, RT.SLIDE)
                        h = rPr.makeelement(qn("a:hlinkClick"), {qn("r:id"): rId, "action": "ppaction://hlinksldjump"})
                        # hlinkClick은 rPr 안에서 채우기·글꼴 뒤에 온다
                        rPr.append(h)
                        rPr.set("u", "sng")
                p._p.remove(r._r)
                return True
    return False


def to_pdf(pptx):
    names = zipfile.ZipFile(pptx).namelist()
    assert len(names) == len(set(names)), "zip 안에 중복 파일"
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(pptx), pptx],
                       check=True, capture_output=True, text=True)
    assert "Error" not in (r.stdout + r.stderr), r.stdout + r.stderr


def build(deck, refs):
    path = os.path.join(ROOT, deck + ".pptx")
    prs = Presentation(path)
    S = prs.slides
    for s in S:                                                   # 전에 단 버튼 지우기
        for sh in [x for x in s.shapes if x.name.startswith("NAV_")]:
            sh._element.getparent().remove(sh._element)

    theme_link_color(prs)
    fwd, back = {}, {}
    for src, dsts, label in refs:
        name = label if label.startswith("부록") else "부록 " + label
        fwd.setdefault(src, []).append((f"▶ {name} · {dsts[0]}쪽", S[dsts[0] - 1], 1.78 if len(name) <= 7 else 2.0))
        assert link_text(S[src - 1], label, S[dsts[0] - 1]), (src, label)
        for d in dsts:
            back.setdefault(d, [])
            if src not in [b for b, _ in back[d]]:
                back[d].append((src, S[src - 1]))

    for src, items in fwd.items():
        pills(S[src - 1], items)
    for d, srcs in back.items():
        pills(S[d - 1], [(f"◀ 본문 {src}쪽으로", t, 1.62) for src, t in srcs])

    prs.save(path); to_pdf(path)
    print(deck, "→", {k: [i[0] for i in v] for k, v in fwd.items()}, "| 돌아오기", {k: [s for s, _ in v] for k, v in back.items()})


if __name__ == "__main__":
    for deck, refs in DECKS.items():
        if len(sys.argv) < 2 or sys.argv[1] in deck:
            build(deck, refs)
