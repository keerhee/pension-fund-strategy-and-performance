"""W05 실습데이터 6쪽 '데이터 출처' 표 정렬 (2026-09-27).

글상자 + 줄무늬 사각형으로 만든 표라 칸 폭이 고정이다. 데이터 칸의 'NPS 2027 목표 · CMA'가 두 줄로 넘어가
아래 행과 겹치고, 받는 법 칸의 'auto_adjust=True'가 'Tru / e'로 끊겼다. 출처 칸이 넉넉하므로 그 폭을
데이터 · 받는 법 칸에 나눠 주어 모든 셀을 한 줄로 만들고, 행 간격을 고르게(620k EMU) 다시 놓는다.
원본: _구판/실습데이터6쪽_구판_2026-09-27/
"""
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / '_구판' / '실습데이터6쪽_구판_2026-09-27' / 'W05_실습데이터_리스크패리티와HRP.pptx'
OUT = ROOT / 'W05_리스크패리티와HRP' / 'W05_실습데이터_리스크패리티와HRP.pptx'

COLS = [(694944, 2350000), (3200000, 3450000), (6820000, 2200000), (9200000, 2266576)]   # (left, width)
HEAD = ['TextBox 9', 'TextBox 10', 'TextBox 11', 'TextBox 12']
ROWS = [['TextBox 13', 'TextBox 14', 'TextBox 15', 'TextBox 16'],
        ['TextBox 18', 'TextBox 19', 'TextBox 20', 'TextBox 21'],
        ['TextBox 22', 'TextBox 23', 'TextBox 24', 'TextBox 25'],
        ['TextBox 27', 'TextBox 28', 'TextBox 29', 'TextBox 30'],
        ['TextBox 31', 'TextBox 32', 'TextBox 33', 'TextBox 34'],
        ['TextBox 36', 'TextBox 37', 'TextBox 38', 'TextBox 39']]
BANDS = {1: 'Rectangle 17', 3: 'Rectangle 26', 5: 'Rectangle 35'}   # 줄무늬 — 2·4·6번째 행
TOP0, PITCH, BAND_H, TEXT_H = 2200000, 620000, 475488, 329184

prs = Presentation(SRC)
s = prs.slides[5]
sh = {x.name: x for x in s.shapes}
for name, (l, w) in zip(HEAD, COLS):
    sh[name].left, sh[name].width = l, w
for i, row in enumerate(ROWS):
    band_top = TOP0 + i * PITCH
    for name, (l, w) in zip(row, COLS):
        b = sh[name]
        b.left, b.width = l, w
        b.top, b.height = band_top + (BAND_H - TEXT_H) // 2, TEXT_H
    if i in BANDS:
        r = sh[BANDS[i]]
        r.top, r.height = band_top, BAND_H
prs.save(OUT)
print('saved', OUT.name)
