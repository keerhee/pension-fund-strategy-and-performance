# -*- coding: utf-8 -*-
"""강의본·실습·프라이머 덱 in-place 보수 도구 (python-pptx).

W04·W05 보수(2026-09-09)의 lecfix.py 를 W06·W07용으로 다시 만든 것.
  clone_slide(src_prs, i, dst_prs)   다른 덱의 슬라이드를 통째로 복제(그림 없는 슬라이드용)
  move_slide(prs, old, new)          슬라이드 순서 이동(0-based)
  delete_slide(prs, i)
  set_text(shape, text)              첫 run 서식 유지, "\n" 은 문단 분리
  set_table(shape, rows)             표 재작성 — 행 수가 다르면 본문 행을 복제/삭제
  renumber(prs)                      우하단 쪽번호 재계산
  replace_all(prs, old, new)         run 단위 치환(모든 슬라이드·표 포함)
  find(slide, prefix)                텍스트로 도형 찾기
  px(shape)                          (x,y,w,h) 픽셀
"""
import copy, re
from pptx import Presentation
from pptx.util import Emu

def px(sh):
    f = lambda e: int(round(Emu(e).inches * 96))
    return f(sh.left), f(sh.top), f(sh.width), f(sh.height)

def _tf_iter(shape):
    if shape.has_text_frame:
        yield shape.text_frame
    if getattr(shape, "has_table", False) and shape.has_table:
        for r in shape.table.rows:
            for c in r.cells:
                yield c.text_frame
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            yield from _tf_iter(s)

def set_tf(tf, text):
    """text_frame 의 내용을 text 로 바꾼다. 첫 문단 첫 run 의 서식을 유지하고, 줄은 "\n" 으로 나눈다."""
    lines = text.split("\n")
    paras = tf.paragraphs
    p0 = paras[0]
    # 첫 문단: 첫 run 만 남기고 텍스트 교체
    runs = p0.runs
    if runs:
        runs[0].text = lines[0]
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        p0.text = lines[0]
    # 나머지 문단 제거
    for p in paras[1:]:
        p._p.getparent().remove(p._p)
    # 추가 줄: 첫 문단 복제
    for ln in lines[1:]:
        np_ = copy.deepcopy(p0._p)
        p0._p.getparent().append(np_)
        # 복제된 문단의 run 텍스트 교체
        from pptx.text.text import _Paragraph
        para = _Paragraph(np_, p0._parent)
        rr = para.runs
        if rr:
            rr[0].text = ln
            for r in rr[1:]:
                r._r.getparent().remove(r._r)
        else:
            para.text = ln

def set_text(shape, text):
    set_tf(shape.text_frame, text)

def set_table(shape, rows):
    """rows: list of list(str). 첫 행은 머리행. 본문 행은 기존 1행(index 1)을 서식 원형으로 복제한다."""
    tbl = shape.table
    tr_list = tbl._tbl.tr_lst
    n_have, n_want = len(tr_list), len(rows)
    proto = copy.deepcopy(tr_list[1] if n_have > 1 else tr_list[0])
    # 행 수 맞추기
    while len(tbl._tbl.tr_lst) < n_want:
        tbl._tbl.append(copy.deepcopy(proto))
    while len(tbl._tbl.tr_lst) > n_want:
        last = tbl._tbl.tr_lst[-1]
        tbl._tbl.remove(last)
    # 높이 재분배: 기존 총 높이를 유지
    total_h = shape.height
    rh = int(total_h / n_want)
    for tr in tbl._tbl.tr_lst:
        tr.h = rh
    for ri, row in enumerate(rows):
        cells = tbl.rows[ri].cells
        assert len(row) == len(cells), f"row {ri}: {len(row)} vs {len(cells)} cols"
        for ci, val in enumerate(row):
            set_tf(cells[ci].text_frame, val)

def find(slide, prefix, contains=False):
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if (contains and prefix in t) or (not contains and t.startswith(prefix)):
                return sh
    return None

def find_table(slide):
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            return sh
    return None

def replace_all(prs, old, new, count_only=False):
    """run 안에서 먼저 찾고, run 경계에 걸쳐 있으면 그 문단의 run 을 첫 run 서식으로 합쳐서 치환한다."""
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            for tf in _tf_iter(sh):
                for p in tf.paragraphs:
                    hit = False
                    for r in p.runs:
                        if old in r.text:
                            n += 1; hit = True
                            if not count_only:
                                r.text = r.text.replace(old, new)
                    if not hit and p.runs and old in p.text:
                        n += 1
                        if not count_only:
                            txt = p.text.replace(old, new)
                            p.runs[0].text = txt
                            for r in p.runs[1:]:
                                r._r.getparent().remove(r._r)
    return n

def grep(prs, pat):
    out = []
    rx = re.compile(pat)
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            for tf in _tf_iter(sh):
                t = tf.text
                if rx.search(t):
                    out.append((i, " ".join(t.split())[:120]))
    return out

def clone_slide(src_prs, src_idx, dst_prs, layout=None):
    src = src_prs.slides[src_idx]
    layout = layout or dst_prs.slide_layouts[0]
    dst = dst_prs.slides.add_slide(layout)
    # 레이아웃 자리표시자 제거
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        if sh.shape_type == 13:
            raise RuntimeError("그림이 있는 슬라이드는 clone_slide 로 복제하지 않는다")
        dst.shapes._spTree.append(copy.deepcopy(sh._element))
    return dst

def dup_slide(prs, idx):
    """같은 덱 안에서 슬라이드 복제(그림 없는 슬라이드). 맨 뒤에 붙는다."""
    return clone_slide(prs, idx, prs)

def move_slide(prs, old, new):
    sldIdLst = prs.slides._sldIdLst
    el = list(sldIdLst)[old]
    sldIdLst.remove(el)
    sldIdLst.insert(new, el)

def delete_slide(prs, idx):
    sldIdLst = prs.slides._sldIdLst
    el = list(sldIdLst)[idx]
    rId = el.rId
    prs.part.drop_rel(rId)
    sldIdLst.remove(el)

def renumber(prs, x_min=1000, y_min=640):
    n = 0
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            x, y, w, h = px(sh)
            t = sh.text_frame.text.strip()
            if x >= x_min and y >= y_min and re.fullmatch(r"\d{1,3}", t):
                set_text(sh, str(i)); n += 1
    return n

def scan_codes(prs):
    """문자 약호 K1·V1·N1·P1·Q1 스캔."""
    return grep(prs, r"(?<![A-Za-z])[KVNPQ]\d")
