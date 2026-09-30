# -*- coding: utf-8 -*-
"""W06 강의본 3교시 보강 (56 → 63장) — 보충교재 14(강화학습 포트폴리오) · 16(Gârleanu–Pedersen) 반영 (2026-10-01).

새 장표 7장
  30 뒤  상태 분해(외생 MDP)
  31 뒤  기준선 ① 무거래 구간 · ② GP · ③ 네 전략 비교 · 기준선 사다리
  34 뒤  회계 — 시점 · 비용 · 학습 보상 대 실제 순수익
  35 뒤  검증 — 합성시장 · walk-forward · DSR · 시드
기존 장표 수정: 3쪽 용어 ②(기준선 행) · 4쪽 강의 지도 · 29쪽 3교시 표지 · 체크포인트 · 한 장 요약 · 목차 칩 쪽번호.
그림은 _build/gp_art.py 를 GP_WHITE=1 로 돌린 _art/gp_white/ 를 쓴다.
실행: .venv/bin/python _build/lecfix/fix_w06_rl.py   (재실행 가능 — _구판/W06_RL보강_구판_2026-10-01/ 에서 원본 복원)
"""
import os, shutil, subprocess, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from lecfix import *
from slidekit import Kit, place, E
from pptx.util import Pt

R = os.path.dirname(os.path.dirname(HERE))
W = f"{R}/W06_동적포트폴리오와장기투자"
OLD = f"{R}/_구판/W06_RL보강_구판_2026-10-01"; os.makedirs(OLD, exist_ok=True)
ART = f"{R}/_build/_art/gp_white"
name = "W06_동적포트폴리오와장기투자_강의본"

for ext in (".pptx", ".pdf"):
    src, dst = f"{W}/{name}{ext}", f"{OLD}/{name}{ext}"
    if os.path.exists(dst): shutil.copy2(dst, src)
    elif os.path.exists(src): shutil.copy2(src, dst)

prs = Presentation(f"{W}/{name}.pptx")
S = lambda k: prs.slides[k - 1]          # 원판 1-based
assert len(prs.slides) == 56, len(prs.slides)
kit = Kit(prs, panel_idx=[30], table_idx=34, bottom_idx=34)
T = S(35)                                 # 원형: 3교시 · 장벽(그림 없는 장표)


def pic(s, path, x=None, y=196, w=None, h=None, cx=640):
    from PIL import Image
    iw, ih = Image.open(path).size
    if w is None: w = h * iw / ih
    if h is None: h = w * ih / iw
    if x is None: x = cx - w / 2
    s.shapes.add_picture(path, E(x), E(y), E(w), E(h))
    return y + h


# ── N1 · 상태 분해 ────────────────────────────────────────────────
n1 = kit.new(T, "3교시 · 상태", "가격수용자의 상태는 시장과 계좌로 나뉜다",
             "외생 MDP(exogenous MDP) — 시장은 내 행동과 무관하고, 계좌는 알려진 회계식으로 움직인다")
kit.panel(n1, 60, 200, 568, 270, "b", "시장 상태 — 모르는 부분", [
    "금리 · 변동성 · 밸류 · 모멘텀 같은 신호",
    "내 주문이 가격을 바꾸지 않는다 (가격수용자)",
    "과거 경로는 하나뿐 — 학습이 필요한 곳은 여기"], pt=17)
kit.panel(n1, 652, 200, 568, 270, "g", "계좌 상태 — 아는 부분", [
    "현재 비중 · 자산가치 · 낙폭",
    "거래 후 비중 · 비용은 회계식으로 정확히 계산",
    "현재 비중이 빠지면 거래비용을 알 수 없다"], pt=17)
kit.banner(n1, 492, "과거 경로 하나로 모든 행동의 결과를 다시 계산한다 — 탐색은 시뮬레이터 안에서")
kit.bottom(n1, "RL은 예측기가 아니라 제어기 — 신호가 없으면 초과수익도 없고, RL의 몫은 ‘언제 · 얼마나 거래하나’다", y=566)
kit.bottom(n1, "기초(에이전트 · 환경 · 상태 · 행동 · 보상)는 W06 보강 1 ‘처음 배우는 강화학습’", pt=15, y=610)

# ── N2 · 기준선 ① 무거래 구간 ─────────────────────────────────────
n2 = kit.new(T, "3교시 · 기준선 ①", "RL이 넘어야 할 벽 ① — 비례 비용이면 무거래 구간",
             "CRRA 투자자 · 위험자산 1개 — Merton(1969)의 목표 비중과 Davis–Norman(1990)의 무거래 구간")
kit.formula(n2, r"\pi^{*}=\frac{\mu-r}{\gamma\,\sigma^{2}}", cx=330, y=198, pt=28)
kit.formula(n2, r"\mathrm{half\ width}\ \propto\ \left(\frac{c}{\gamma}\right)^{1/3}\left[\pi^{*}(1-\pi^{*})\right]^{2/3}", cx=900, y=198, pt=24)
kit.label(n2, 120, 282, 420, 26, "비용이 없을 때의 목표 비중 (Merton)", pt=14, align="c")
kit.label(n2, 650, 282, 500, 26, "구간 반폭 — 비용 c의 세제곱근에 비례 (점근 근사)", pt=14, align="c")
kit.table(n2, 60, 318, 1161, [
    ["합성시장 (μ 7% · σ 18% · γ 3 · 편도 25bp)", "연율 확실성등가", "비용 손실", "연간 거래량"],
    ["비용 없는 Merton — 상한", "329.8bp", "—", "—"],
    ["매일 Merton 비중으로 리밸런싱", "315.1bp", "14.7bp", "57%"],
    ["*무거래 구간 (반폭 3%p)", "*328.7bp", "*1.1bp", "*3%"]],
    [4.6, 1.6, 1.3, 1.3], rowh=36, pt=15, align="lccc")
kit.banner(n2, 480, "무거래 구간은 비용 손실을 연 14.7bp에서 1.1bp로 줄인다 — RL은 먼저 이 구간을 재현해야 한다")
kit.bottom(n2, "구간 안에서는 가만히, 밖이면 경계까지만 — 2교시 밴드 리밸런싱 · SS2 재균형 밴드의 이론적 근거", y=552)
kit.bottom(n2, "10년 × 20,000경로 · 시뮬레이션 최적 반폭 3~5%p는 점근 공식(4.3%p)과 가깝다 · 보충교재 14 본편 II", pt=14, y=596)

# ── N3 · 기준선 ② GP ─────────────────────────────────────────────
n3 = kit.new(T, "3교시 · 기준선 ②", "RL이 넘어야 할 벽 ② — 이차 비용이면 GP의 조준과 부분 이동",
             "Gârleanu–Pedersen(JF 2013) — 신호가 있고 주문이 클 때의 닫힌 해 · 자세한 내용은 보충교재 16")
kit.formula(n3, r"x_t=\left(1-\frac{a}{\lambda}\right)x_{t-1}+\frac{a}{\lambda}\,\mathrm{aim}_t,\qquad \mathrm{aim}_t=\sum_{\tau\geq 0} z(1-z)^{\tau}\,\mathrm{E}_t\left[x^{*}_{t+\tau}\right]",
            cx=640, y=192, pt=22)
pic(n3, f"{ART}/two_principles.png", y=262, w=1060)
kit.bottom(n3, "γ = 1 · λ = 2 → a/λ = 0.5 : 매번 간격의 절반 · 조준점은 빠른 신호(φ 0.5)를 0.667, 느린 신호(φ 0.05)를 0.952만 반영", pt=15, y=600)
kit.bottom(n3, "비례 비용(수수료) → 구간 · 이차 비용(시장충격) → 부분 이동 — 대형 기금의 큰 주문은 뒤쪽이다", pt=15, y=636)

# ── N4 · 기준선 ③ 네 전략 비교 ─────────────────────────────────────
n4 = kit.new(T, "3교시 · 기준선 ③", "조준 없는 부분 이동은 GP를 이기지 못한다",
             "20만 기간 시뮬레이션 — 목표 추종 · 부분 이동만 · 조준만 · GP의 기간당 순효용 (γ = 1 · λ = 2, 가상 실험)")
pic(n4, f"{ART}/sim_compare.png", y=194, h=392)
kit.banner(n4, 596, "목표 추종 −7.2 · 부분 이동만(최적 속도 0.40) 67.4 · 조준만 42.7 · GP 69.8 — RL이 이 69.8을 재현하는지가 첫 시험", pt=16)

# ── N5 · 기준선 사다리 ───────────────────────────────────────────
n5 = kit.new(T, "3교시 · 기준선 사다리", "RL 결과는 여섯 단계 기준선과 같은 조건에서 비교한다",
             "기준선(baseline) = RL이 같은 비용 · 체결 조건에서 이겨야 할 비교 전략 — 각 단계가 효과 하나씩을 통제한다")
kit.table(n5, 60, 196, 1161, [
    ["단계", "기준선 전략", "이 단계가 통제하는 효과"],
    ["0", "현금 · 매수 후 보유", "시장 노출 그 자체"],
    ["1", "동일가중 · 정기 리밸런싱", "리밸런싱 효과 (SS2 재균형 프리미엄)"],
    ["2", "변동성 목표", "위험 수준 조절 효과"],
    ["3", "수축 MVO (1기간)", "1기간 최적화 효과 — 비용과 내일을 안 본다"],
    ["*4", "*무거래 구간 · GP · MPO", "*비용 인지형 동적 거래 — 여기까지가 공식·수치해"],
    ["5", "강화학습", "위 넷이 설명하지 못한 나머지만 RL의 몫"]],
    [0.6, 3.0, 5.6], rowh=38, pt=16, align="lll")
kit.banner(n5, 480, "4단계를 못 이긴 RL은 단순 규칙을 복잡하게 다시 배운 것이다")
kit.bottom(n5, "원문 Samsudin(2021)의 ‘RL 20% 대 MVO −1%’는 30시점 MVO라는 약한 상대와 비교했다 — 사다리가 없는 비교", y=552)
kit.bottom(n5, "※ 정책경사 알고리즘 안의 ‘baseline’(보상에서 빼는 상태가치 V)과는 다른 말 — 여기서는 비교 전략", pt=14, y=596)

# ── N6 · 회계 ─────────────────────────────────────────────────
n6 = kit.new(T, "3교시 · 회계", "보상보다 회계가 먼저 — 거래 시점과 비용을 고정한다",
             "정보 마감 → 체결 → 보유 구간을 나누고, 학습 보상과 실제 순수익을 따로 저장한다")
kit.panel(n6, 60, 196, 568, 214, "b", "거래 시점 — 미래정보 차단", [
    "전일 종가까지의 정보로 결정 → 오늘 시가 체결",
    "시가에서 다음 시가까지의 수익률을 적용",
    "종가 보고 그 종가에 거래 = 존재할 수 없는 이익"], pt=16)
kit.panel(n6, 652, 196, 568, 214, "g", "비용 — 표류한 비중과 목표의 차이로", [
    "예: 표류 후 40/30/30 → 목표 50/20/30",
    "회전율 20% × 편도 10bp = 비용 2bp",
    "비용 차감 순수익 0.7828%"], pt=16)
kit.table(n6, 60, 424, 1161, [
    ["같은 하루", "실제 순수익", "로그 순수익", "위험 벌점", "회전율 벌점", "학습 보상"],
    ["값", "*0.7828%", "0.7798%", "−0.05%", "−0.01%", "*0.7198%"]],
    [1.2, 1.4, 1.4, 1.2, 1.3, 1.4], rowh=36, pt=16, align="cccccc")
kit.banner(n6, 516, "벌점은 학습 목표에만 들어가는 선호이지 자산에서 빠지는 비용이 아니다 — 두 값을 따로 저장한다")
kit.bottom(n6, "샤프는 미분 샤프(Moody · Saffell 2001)로 학습하고, 성과는 확실성등가로 평가한다 — 현금 편향 정책은 샤프가 불안정", pt=15, y=588)

# ── N7 · 검증 ─────────────────────────────────────────────────
n7 = kit.new(T, "3교시 · 검증", "검증 — 정답이 알려진 시장에서 먼저, 시도 횟수의 대가를 치르고",
             "합성시장 재현 → 시간순 walk-forward → 다중 시도 보정(DSR) → 시드 분포 → 기준선 대비")
kit.table(n7, 60, 196, 1161, [
    ["단계", "무엇을 하나", "합격 기준"],
    ["① 합성시장", "해가 알려진 시장(GBM · GP)에서 먼저 학습", "무거래 구간 · a/λ 재현, 확실성등가 90% 회복"],
    ["② walk-forward", "과거 → 미래 순서로만 학습 · 평가, 테스트 구간을 이어 붙임", "학습 · 평가 구간이 겹치지 않는다"],
    ["③ DSR", "시도한 설정 수만큼 샤프 기준을 올린다 (Deflated Sharpe)", "보정 뒤에도 유의"],
    ["④ 시드 분포", "난수 시드를 여러 개로 다시 학습", "한 번의 운이 아니라 분포로 보고"],
    ["⑤ 기준선 대비", "사다리 4단계와 같은 비용 · 체결 조건", "순효용 · 회전율 모두 우위"]],
    [1.6, 4.4, 3.6], rowh=40, pt=15, align="lll")
kit.banner(n7, 460, "좋은 정책과 조금 나쁜 정책은 데이터로 구별하기 어렵다 — 정답이 있는 시험부터")
kit.bottom(n7, "Lu(2023): 가장 단순한 합성시장에서도 PPO가 약 200만 스텝 필요 — 에피소드를 반복해도 독립된 시장 표본은 늘지 않는다", pt=15, y=532)
kit.bottom(n7, "자세한 실험 · 손계산은 보충교재 14 (본편 II 4부 · 워크북)", pt=15, y=576)

# ── 순서 배치 ────────────────────────────────────────────────────
place(prs, [(n1, 30), (n2, 31), (n3, 31), (n4, 31), (n5, 31), (n6, 34), (n7, 35)])
assert len(prs.slides) == 63
renumber(prs)
idx = {id(s): i for i, s in enumerate(prs.slides, 1)}

# ── 기존 장표 수정 ────────────────────────────────────────────────
t3 = find_table(prs.slides[2])
rows = [[c.text for c in r.cells] for r in t3.table.rows]
k = next(i for i, r in enumerate(rows) if r[0].startswith("MDP"))
rows[k][2] = "3교시 · 기초는 보강 1"
rows.insert(k + 1, ["기준선 (baseline)", "RL이 같은 비용 조건에서 이겨야 할 비교 전략 — 무거래 구간 · GP(조준점으로 부분 이동)", "3교시 · 보충교재 16"])
set_table(t3, rows)
for r in t3.table.rows:
    for c in r.cells:
        for p in c.text_frame.paragraphs:
            for rr in p.runs:
                rr.font.size = Pt(12)

n = replace_all(prs, "MDP → DDPG · 보상 설계 · 4대 장벽 · 하이브리드",
                "MDP · 기준선(무거래 구간 · GP) · DDPG · 보상과 회계 · 검증 · 하이브리드"); assert n == 1, n
n = replace_all(prs, "MDP · DDPG · 4대 장벽 · 하이브리드", "MDP · 기준선(무거래 구간 · GP) · DDPG · 검증 · 하이브리드"); assert n == 1, n
n = replace_all(prs, "MDP의 다섯 요소를 자산배분 맥락으로 구체화하면 ?",
                "RL이 먼저 이겨야 할 기준선 — 무거래 구간과 GP는 각각 어떤 비용 구조의 최적해인가 ?"); assert n == 1, n
n = replace_all(prs, "보상 설계 네 가지 선택지 — 실무가 복합 보상을 쓰는 이유는 ?",
                "학습 보상과 실제 순수익을 왜 따로 저장하는가 — 복합 보상과 회계의 구분 ?"); assert n == 1, n
chk = next(s for s in prs.slides if find(s, "3교시 점검 질문"))
for sh in chk.shapes:
    if sh.has_text_frame and sh.text_frame.text.strip() == "개념" and px(sh)[1] < 300:
        set_text(sh, "기준선")
n = replace_all(prs, "누적 보상 극대화 — 보상 설계가 곧 운용 철학",
                "누적 보상 극대화 — 무거래 구간 · GP 기준선을 같은 조건에서 먼저 이겨야"); assert n == 1, n
n = replace_all(prs, "RL 가치함수", "RL과 기준선"); assert n >= 1, n

# 목차 칩(◀ NN쪽) — 링크 대상 장표의 새 번호로
import re
for s in list(prs.slides)[1:3]:
    for sh in s.shapes:
        if sh.has_text_frame and re.fullmatch(r"◀ \d+쪽", sh.text_frame.text.strip()):
            tgt = sh.click_action.target_slide
            set_text(sh, f"◀ {idx[id(tgt)]}쪽")

prs.save(f"{W}/{name}.pptx")
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", W, f"{W}/{name}.pptx"],
               check=True, capture_output=True)
print("저장:", name, len(prs.slides), "장")
print("새 장표 번호:", [idx[id(s)] for s in (n1, n2, n3, n4, n5, n6, n7)])
