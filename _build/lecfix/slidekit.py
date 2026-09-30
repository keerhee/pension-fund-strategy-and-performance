# -*- coding: utf-8 -*-
"""새 장표를 기존 덱 서식 그대로 만드는 도구 (lecfix 위에 얹는다).

  kit = Kit(prs, panel_idx=[..], table_idx=..)   같은 덱의 장표에서 원형(패널·배너·하단 줄·표)을 모은다
  s = kit.new(tmpl_idx, tag, title, sub)          원형 장표를 복제해 머리(태그·제목·부제)만 남긴다
  kit.panel(s, x, y, w, h, "b"|"g", title, lines, pt=16)
  kit.banner(s, y, text, h=56)
  kit.bottom(s, text)
  kit.table(s, x, y, w, rows, colw, rowh=34, pt=15, align=None)   셀 앞 "*" = 굵게
  kit.formula(s, tex, cx=None, x=None, y=0, pt=26)               mathtext(cm) PNG, 남색 글자 · #F3F4F8 바탕
  kit.label(s, x, y, w, h, text, pt=14, color="6b7280", bold=False, align="l")
좌표는 모두 1280×720 픽셀.
"""
import copy, hashlib, os
from lecfix import *
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PX = 9525
E = lambda v: int(round(v * PX))
ART = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_eq")
os.makedirs(ART, exist_ok=True)


def fillhex(sh):
    try:
        if sh.fill.type == 1:
            return str(sh.fill.fore_color.rgb).upper()
    except Exception:
        pass
    return ""


def render_eq(tex, fs=30, dpi=200):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    matplotlib.rcParams["mathtext.fontset"] = "cm"
    key = hashlib.md5(f"{tex}|{fs}|{dpi}".encode()).hexdigest()[:12]
    out = os.path.join(ART, f"eq_{key}.png")
    if not os.path.exists(out):
        fig = plt.figure(figsize=(0.01, 0.01))
        fig.text(0, 0, f"${tex}$", fontsize=fs, color="#1B2C5E")
        fig.savefig(out, dpi=dpi, bbox_inches="tight", pad_inches=0.12, facecolor="#F3F4F8")
        plt.close(fig)
    return out


class Kit:
    def __init__(self, prs, panel_idx, table_idx, bottom_idx=None):
        self.prs, self.P = prs, {}
        for k in ([panel_idx] if isinstance(panel_idx, int) else panel_idx):
            self._collect(prs.slides[k])
        if bottom_idx is not None:
            self._collect(prs.slides[bottom_idx])
        self.P["table"] = find_table(prs.slides[table_idx])._element
        need = ["rect_b", "rect_g", "title_b", "title_g", "body", "banner_r", "banner_t", "bottom"]
        miss = [n for n in need if n not in self.P]
        assert not miss, f"원형 없음: {miss}"

    def _collect(self, s):
        shapes = list(s.shapes)
        for sh in shapes:
            x, y, w, h = px(sh)
            f = fillhex(sh)
            t = sh.text_frame.text.strip() if sh.has_text_frame else ""
            if f == "EAF1FB" and w > 300 and "rect_b" not in self.P:
                self.P["rect_b"] = sh._element; self._rb = (x, y, w, h)
            elif f == "E6F5F0" and w > 300 and "rect_g" not in self.P:
                self.P["rect_g"] = sh._element; self._rg = (x, y, w, h)
            elif f == "1B2C5E" and w > 900 and h > 40 and "banner_r" not in self.P:
                self.P["banner_r"] = sh._element; self._bn = (x, y, w, h)
        for sh in shapes:
            if not (sh.has_text_frame and sh.text_frame.text.strip()):
                continue
            x, y, w, h = px(sh)
            r = sh.text_frame.paragraphs[0].runs
            sz = r[0].font.size.pt if r and r[0].font.size else 0
            inside = lambda b: b and b[0] <= x <= b[0] + b[2] and b[1] <= y <= b[1] + b[3]
            if sz == 20 and inside(getattr(self, "_rb", None)) and "title_b" not in self.P:
                self.P["title_b"] = sh._element
            elif sz == 20 and inside(getattr(self, "_rg", None)) and "title_g" not in self.P:
                self.P["title_g"] = sh._element
            elif sz in (16, 18) and (inside(getattr(self, "_rb", None)) or inside(getattr(self, "_rg", None))) \
                    and "body" not in self.P and sh.text_frame.paragraphs[0]._p.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}pPr") is not None:
                self.P["body"] = sh._element
            elif inside(getattr(self, "_bn", None)) and "banner_t" not in self.P:
                self.P["banner_t"] = sh._element
            elif 600 <= y < 676 and w > 900 and "bottom" not in self.P and sz == 18:
                self.P["bottom"] = sh._element

    # ---------- basics ----------
    def _put(self, s, key, x, y, w, h):
        el = copy.deepcopy(self.P[key])
        s.shapes._spTree.append(el)
        sh = s.shapes[-1]
        sh.left, sh.top, sh.width, sh.height = E(x), E(y), E(w), E(h)
        return sh

    @staticmethod
    def _size(sh, pt):
        for tf in ([sh.text_frame] if sh.has_text_frame else []):
            for p in tf.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(pt)

    def new(self, tmpl, tag, title, sub):
        """tmpl: 원형 장표(객체 또는 0-based 번호 — 순서를 옮긴 뒤에는 객체를 넘긴다)"""
        if not isinstance(tmpl, int):
            tmpl = list(self.prs.slides).index(tmpl)
        s = dup_slide(self.prs, tmpl)
        for sh in list(s.shapes):
            x, y, w, h = px(sh)
            if 186 <= y < 676:
                sh._element.getparent().remove(sh._element)
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            x, y, w, h = px(sh)
            t = sh.text_frame.text.strip()
            if y < 60 and x < 300 and t:
                set_text(sh, tag)
            elif 70 <= y < 130 and t:
                set_text(sh, title)
            elif 130 <= y < 186 and t:
                set_text(sh, sub)
        # 태그 칩 폭을 글자 수에 맞춘다(12pt 굵게 ≈ 한글 16px · 그 밖 8.5px)
        need = int(sum(16 if ord(ch) > 0x2E80 else 8.5 for ch in tag) + 36)
        for sh in s.shapes:
            x, y, w, h = px(sh)
            if y < 60 and x < 300 and w < need:
                sh.width = E(need)
        return s

    def panel(self, s, x, y, w, h, color, title, lines, pt=16, title_pt=None):
        self._put(s, "rect_" + color, x, y, w, h)
        t = self._put(s, "title_" + color, x + 26, y + 12, w - 50, 38)
        set_text(t, title)
        if title_pt:
            self._size(t, title_pt)
        b = self._put(s, "body", x + 30, y + 54, w - 56, h - 62)
        set_text(b, "\n".join(lines))
        self._size(b, pt)
        return b

    def banner(self, s, y, text, h=56, pt=None):
        self._put(s, "banner_r", 60, y, 1161, h)
        t = self._put(s, "banner_t", 79, y, 1122, h)
        set_text(t, text)
        if pt:
            self._size(t, pt)
        return t

    def bottom(self, s, text, pt=None, y=630):
        t = self._put(s, "bottom", 60, y, 1161, 38)
        set_text(t, text)
        if pt:
            self._size(t, pt)
        return t

    def label(self, s, x, y, w, h, text, pt=14, color="6b7280", bold=False, align="l"):
        tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, ln in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            r = p.add_run()
            r.text = ln
            f = r.font
            f.size, f.bold, f.name = Pt(pt), bold, "Noto Sans CJK KR"
            f.color.rgb = RGBColor.from_string(color.upper())
            rPr = r._r.get_or_add_rPr()
            ea = rPr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}ea", {"typeface": "Noto Sans CJK KR"})
            rPr.append(ea)
        return tb

    def formula(self, s, tex, cx=None, x=None, y=0, pt=26, fs=30, dpi=200):
        from PIL import Image
        path = render_eq(tex, fs, dpi)
        iw, ih = Image.open(path).size
        sc = (96 / dpi) * (pt / fs)
        w, h = iw * sc, ih * sc
        if x is None:
            x = (cx if cx is not None else 640) - w / 2
        s.shapes.add_picture(path, E(x), E(y), E(w), E(h))
        return w, h

    def table(self, s, x, y, w, rows, colw, rowh=34, pt=15, align=None, hdr_pt=None):
        gf = copy.deepcopy(self.P["table"])
        s.shapes._spTree.append(gf)
        sh = s.shapes[-1]
        tbl = sh.table._tbl
        grid = tbl.tblGrid
        gcs = list(grid)
        for g in gcs:
            grid.remove(g)
        tot = sum(colw)
        for cw in colw:
            g = copy.deepcopy(gcs[0]); g.set("w", str(E(w * cw / tot))); grid.append(g)
        trs = list(tbl.tr_lst)
        hdr, b1 = trs[0], trs[1]
        b2 = trs[2] if len(trs) > 2 else trs[1]
        for tr in trs:
            tbl.remove(tr)
        for ri, row in enumerate(rows):
            proto = hdr if ri == 0 else (b1 if ri % 2 == 1 else b2)
            tr = copy.deepcopy(proto)
            tcs = list(tr.tc_lst)
            for tc in tcs:
                tr.remove(tc)
            for _ in row:
                tc = copy.deepcopy(tcs[0])
                for a in ("gridSpan", "hMerge", "vMerge", "rowSpan"):
                    tc.attrib.pop(a, None)
                tr.append(tc)
            tr.h = E(rowh)
            tbl.append(tr)
        sh.left, sh.top, sh.width, sh.height = E(x), E(y), E(w), E(rowh * len(rows))
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row):
                cell = sh.table.cell(ri, ci)
                bold = val.startswith("*")
                set_tf(cell.text_frame, val.lstrip("*"))
                for p in cell.text_frame.paragraphs:
                    if ri > 0 and align:
                        p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align[ci]]
                    for r in p.runs:
                        r.font.size = Pt(hdr_pt if (ri == 0 and hdr_pt) else pt)
                        if ri > 0:
                            r.font.bold = bold
                            if bold:
                                r.font.color.rgb = RGBColor(0x1B, 0x2C, 0x5E)
        return sh


def place(prs, new_slides_after):
    """[(new_slide_obj, after_1based)] — 새 장표(맨 뒤에 붙어 있음)를 지정 위치 뒤로 옮긴다. 뒤쪽부터 처리."""
    ids = list(prs.slides._sldIdLst)
    order = []
    for s, after in new_slides_after:
        sid = next(el for el in ids if prs.slides.get(int(el.get("id"))) is s)
        order.append((after, sid))
    lst = prs.slides._sldIdLst
    # 같은 after 값끼리는 넣은 순서를 지킨다
    for after, sid in sorted(order, key=lambda t: t[0], reverse=True):
        pass
    groups = {}
    for after, sid in order:
        groups.setdefault(after, []).append(sid)
    for after in sorted(groups, reverse=True):
        for sid in groups[after]:
            lst.remove(sid)
        anchor = list(lst)[after - 1]
        pos = list(lst).index(anchor) + 1
        for k, sid in enumerate(groups[after]):
            lst.insert(pos + k, sid)
