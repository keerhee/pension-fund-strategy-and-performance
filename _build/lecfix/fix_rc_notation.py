"""위험기여도(RC) 표기를 공분산 형태로 통일한다 (2026-09-27).

(Σw)_i 는 비전공자에게 '비중의 합'으로 읽히고, 강의 흐름에서도 정의 없이 등장했다.
그래서 RC 를 모두 Cov(r_i, r_p) 로 쓴다.
    Cov(r_i, r_p) = Σ_j w_j σ_ij          (자산 i와 포트폴리오 수익률의 공분산)
    MCTR_i = ∂σ_p/∂w_i = Cov(r_i, r_p)/σ_p
    RC_i   = w_i · MCTR_i = w_i Cov(r_i, r_p)/σ_p,   Σ_i RC_i = σ_p
    RC 비율 = RC_i/σ_p = w_i Cov(r_i, r_p)/σ_p²,   합 1
행렬 표기 (Σw)_i 는 정의 페이지(W03 36쪽)에서 '논문 표기'로 한 번만 소개한다.

  W03 강의본 36쪽 — 정의(MCTR·RC)와 합 = σ_p 의 한 줄 증명
  W05 프라이머 7쪽 — 읽기 쉬운 식 하나, 카드·결론 문장
  W05 강의본 10쪽 — ERC 조건을 Cov 로, 정의 한 줄
  W05 케이스 12쪽 — '합이 1'인 쪽은 RC 비율(RC_i/σ_p)로 기호를 구분

수식 PNG 는 matplotlib mathtext(Computer Modern, 이 맥에는 LaTeX 가 없다)로 굽고,
기존 그림과 같은 배율(픽셀당 4,315 EMU)로 원래 자리의 가운데에 얹는다.
원본은 _구판/RC표기_구판_2026-09-27/ 에 있다. 거기서 다시 돌리면 같은 결과가 나온다.

Run:  <python with python-pptx, matplotlib> fix_rc_notation.py
"""
import copy
import io
import shutil
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
BAK = ROOT / '_구판' / 'RC표기_구판_2026-09-27'
sys.path.insert(0, str(ROOT / '_build'))

EMU_PER_PX = 4315                 # 기존 수식 그림의 배율 (≈ 212 dpi)
NAVY, BLUE, LECT_BG = '#1B2C5E', '#2E5BAA', '#F4F5F9'

COV_I = r'\mathrm{Cov}(r_i,\,r_p)'
COV_J = r'\mathrm{Cov}(r_j,\,r_p)'
EQ = r'\;\;=\;\;'

DECKS = {
    'W03': ('W03_연기금모델과CMA/W03_연기금모델과CMA_강의본.pptx', 36),
    'W05L': ('W05_리스크패리티와HRP/W05_리스크패리티와HRP_강의본.pptx', 10),
    'W05C': ('W05_리스크패리티와HRP/W05_케이스_기대수익률포기_AllWeather와HRP_IC.pptx', 12),
    'W05P': ('W05_리스크패리티와HRP/W05_0_프라이머_리스크패리티는왜주식을줄이나.pptx', 7),
}

# deck -> {picture name: (tex, colour)}
PICS = {
    'W03': {
        'Image 0': (rf'$\mathrm{{MCTR}}_i{EQ}\dfrac{{\partial\sigma_p}}{{\partial w_i}}{EQ}\dfrac{{{COV_I}}}{{\sigma_p}}'
                    rf'\qquad\quad \mathrm{{RC}}_i{EQ}w_i\cdot\mathrm{{MCTR}}_i$', NAVY),
        'Image 1': (rf'$\sum_{{i=1}}^{{N}}\mathrm{{RC}}_i\;=\;\dfrac{{\sigma_p^{{2}}}}{{\sigma_p}}\;=\;\sigma_p$', BLUE),
    },
    'W05L': {
        'Image 0': (rf'$w_i\,{COV_I}{EQ}w_j\,{COV_J},\qquad \forall\, i,\,j$', NAVY),
        'Image 1': (rf'$\mathrm{{RC}}_i{EQ}\dfrac{{w_i\,{COV_I}}}{{\sigma_p}}{EQ}\dfrac{{\sigma_p}}{{N}},\qquad\quad \forall\, i$', BLUE),
    },
    'W05C': {                     # 케이스 덱은 분수를 글자 크기(\frac)로, 조금 작게 그렸다
        'Image 0': (rf'$\frac{{\mathrm{{RC}}_i}}{{\sigma_p}}\;=\;\frac{{w_i\,{COV_I}}}{{\sigma_p^{{2}}}},'
                    rf'\qquad \sum_{{i=1}}^{{N}}\frac{{\mathrm{{RC}}_i}}{{\sigma_p}}\;=\;1$', NAVY),
        'Image 1': (rf'$\mathrm{{ERC}}:\quad \frac{{\mathrm{{RC}}_i}}{{\sigma_p}}\;=\;\frac{{1}}{{N}}\quad \forall\, i'
                    rf'\qquad \left(\rho_{{ij}}=0\;\Rightarrow\; w_i\propto \frac{{1}}{{\sigma_i}}\right)$', NAVY),
    },
}
FS = {'W03': 38, 'W05L': 38, 'W05C': 34}  # 덱마다 원본 그림과 글자 높이가 같아지는 크기

# deck -> list of (shape name, [new paragraph texts])  — 문단 수가 바뀌는 곳은 첫 문단 서식을 복제
TEXTS = {
    'W03': [('Text 12', [
        'Cov(r_{i}, r_{p}) = 자산 i와 포트폴리오 수익률의 공분산 = Σ_{j} w_{j}σ_{ij} (논문의 행렬 표기로는 (Σw)_{i})',
        'Σ_{i} w_{i}Cov(r_{i}, r_{p}) = Cov(r_{p}, r_{p}) = σ_{p}^{2} — 그래서 RC를 모두 더하면 정확히 σ_{p}',
        '비중이 아니라 RC로 보면 진짜 구조가 드러난다'])],
    'W05L': [('Text 12', [
        '모든 자산의 RC가 같아야 한다 — 각 자산이 1/N씩 위험을 담당 : “위험의 1/N” (비중의 1/N과 다르다 — 그 차이가 σ와 ρ의 정보)',
        'Cov(r_{i}, r_{p}) = 자산 i와 포트폴리오 수익률의 공분산 = Σ_{j} w_{j}σ_{ij} (정의는 W3 MCTR · RC)'])],
    'W05C': [('Text 12', [
        'RC 비율(RC_{i}/σ_{p})은 분산의 몫 — 합이 1',
        '현행 SAA: 주식 자본 56.5%, 위험 84.5%'])],
    'W05P': [('Text 11', ['Cov(r_{i}, r_{p})']),
             ('Text 12', ['함께 움직이는 정도']),
             ('Text 20', ['비중이 작아도 많이 출렁이고 포트폴리오와 함께 움직이면 큰 몫을 집니다.'])],
}


def render(tex, colour, bg, fontsize, path):
    f = plt.figure(figsize=(16, 2.5), facecolor=bg)
    f.text(.5, .5, tex, ha='center', va='center', fontsize=fontsize, color=colour)
    f.savefig(path, bbox_inches='tight', facecolor=bg, pad_inches=0.26, dpi=200)
    plt.close(f)
    return path


def swap_picture(slide, pic, png, panel_pad=91440):
    """그림을 새 PNG로 바꾸고, 같은 배율로 원래 자리 가운데에 얹는다. 담긴 패널을 넘으면 줄인다."""
    w_px, h_px = Image.open(png).size
    w, h = w_px * EMU_PER_PX, h_px * EMU_PER_PX
    cx, cy = pic.left + pic.width // 2, pic.top + pic.height // 2
    panels = [sh for sh in slide.shapes if sh.shape_type != 13 and not (sh.has_text_frame and sh.text_frame.text.strip())
              and sh.left <= cx <= sh.left + sh.width and sh.top <= cy <= sh.top + sh.height and sh.height > 0]
    if panels:
        box = min(panels, key=lambda s: s.width * s.height)
        k = min(1.0, (box.width - 2 * panel_pad) / w, (box.height - 2 * panel_pad) / h)
        w, h = int(w * k), int(h * k)
    part, rid = slide.part.get_or_add_image_part(png)
    pic._element.blipFill.find(qn('a:blip')).set(qn('r:embed'), rid)
    pic.left, pic.top, pic.width, pic.height = cx - w // 2, cy - h // 2, w, h


def set_paragraphs(shape, texts):
    tf = shape.text_frame
    ps = tf.paragraphs
    proto = copy.deepcopy(ps[0]._p)
    for p in ps[1:]:
        p._p.getparent().remove(p._p)
    first = tf.paragraphs[0]
    for t in texts[1:]:
        el = copy.deepcopy(proto)
        first._p.getparent().append(el)
    for p, t in zip(tf.paragraphs, texts):
        runs = p.runs
        proto_r = copy.deepcopy(runs[0]._r)
        for r in runs:
            r._r.getparent().remove(r._r)
        anchor = p._p.find(qn('a:endParaRPr'))
        for seg, kind in segments(t):
            r = copy.deepcopy(proto_r)
            r.find(qn('a:t')).text = seg
            rpr = r.find(qn('a:rPr'))
            if kind:
                rpr.set('baseline', '-25000' if kind == 'sub' else '30000')
            elif rpr is not None and 'baseline' in rpr.attrib:
                del rpr.attrib['baseline']
            if anchor is not None:
                anchor.addprevious(r)
            else:
                p._p.append(r)


def segments(t):
    """'Cov(r_{i}, r_{p})' -> [('Cov(r', None), ('i', 'sub'), (', r', None), ('p', 'sub'), (')', None)]"""
    import re
    out, pos = [], 0
    for m in re.finditer(r'([_^])\{([^}]*)\}', t):
        if m.start() > pos:
            out.append((t[pos:m.start()], None))
        out.append((m.group(2), 'sub' if m.group(1) == '_' else 'sup'))
        pos = m.end()
    if pos < len(t):
        out.append((t[pos:], None))
    return out


def main():
    import primer_lib                                   # 프라이머 그림은 원래 빌더 함수로
    tmp = Path(__file__).with_name('_rc_png')
    tmp.mkdir(exist_ok=True)
    primer_lib._OUT = str(tmp)
    BAK.mkdir(parents=True, exist_ok=True)

    for key, (rel, page) in DECKS.items():
        src = ROOT / rel
        bak = BAK / src.name
        if not bak.exists():
            shutil.copy2(src, bak)
            pdf = src.with_suffix('.pdf')
            if pdf.exists():
                shutil.copy2(pdf, BAK / pdf.name)
        prs = Presentation(bak)                          # 언제나 원본에서 출발
        slide = prs.slides[page - 1]
        pics = {sh.name: sh for sh in slide.shapes if sh.shape_type == 13}

        if key == 'W05P':
            png = primer_lib.equation(rf'$RC_i\;=\;w_i\;\times\;\frac{{\mathrm{{Cov}}\left(r_i,\,r_p\right)}}{{\sigma_p}}$',
                                      name='w05p_eq.png', fontsize=44, width=8.0)
            swap_picture(slide, pics['Image 0'], png)
        else:
            for name, (tex, colour) in PICS[key].items():
                fs = FS[key]
                png = render(tex, colour, LECT_BG, fs, str(tmp / f'{key}_{name.replace(" ", "")}.png'))
                swap_picture(slide, pics[name], png)

        shapes = {sh.name: sh for sh in slide.shapes}
        for name, texts in TEXTS[key]:
            set_paragraphs(shapes[name], texts)
        if key == 'W05P':                               # 카드 제목: '이 자산이 진 위험'과 같은 16pt, 글상자를 카드 폭까지
            shapes['Text 12'].text_frame.paragraphs[0].runs[0].font.size = shapes['Text 18'].text_frame.paragraphs[0].runs[0].font.size
            card = shapes['Shape 10']
            for nm in ('Text 11', 'Text 12'):
                shapes[nm].width = card.left + card.width - shapes[nm].left - 91440
        if key == 'W05L':                               # 정의 한 줄이 늘어난 만큼 '읽는 법' 패널을 꼬리말 위까지 늘린다
            shapes['Shape 10'].height = 6400800 - shapes['Shape 10'].top
            shapes['Text 11'].top -= 91440
            shapes['Text 12'].top -= 137160
        if key == 'W03':                                # 세 줄이 된 '핵심 직관' — 패널을 아래 캡션 바로 위까지
            shapes['Shape 10'].height = 5760720 - shapes['Shape 10'].top
            shapes['Text 11'].top -= 91440
            shapes['Text 12'].top -= 91440

        prs.save(src)
        print('saved', src.name, 'p', page)


if __name__ == '__main__':
    main()
