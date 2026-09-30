# -*- coding: utf-8 -*-
"""용어 장이 없던 강의본(W02 · W03 · W07 M9 · W09)에 '부록 · 용어집' 한 장을 덧붙인다.

· 맨 뒤에 붙인다 — 앞에 끼우면 기존 쪽 번호가 밀려 영상 · 대본 · 부록 링크가 어긋난다.
· 모양은 W10~W15의 '부록 A · 용어집'과 같다. 같은 덱의 표 장(W02 · W03 한 장 요약)을 복제하거나,
  그림이 섞인 덱(M9 · W09)은 W10 용어집 장을 복제한다.
· 내용은 glossary_data.json — 이미 붙어 있으면 표만 새로 채운다(다시 돌려도 같은 결과).
· 이것을 돌린 뒤 add_nav.py로 링크와 PDF를 만든다.
"""
import json, os, sys
from pptx import Presentation
from lecfix import dup_slide, clone_slide, set_text, set_table, find, find_table

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
DATA = json.load(open(os.path.join(HERE, "glossary_data.json"), encoding="utf-8"))
REF = "W10_주식퀀트모델/W10_주식퀀트모델_강의본"          # 그림 없는 용어집 원형 (49쪽)

DECKS = {   # 덱 → (용어집 쪽 이름, 복제 원형: ("self", 쪽) 또는 ("ref", 쪽))
    "W02": ("W02_자산배분과CAPM/W02_자산배분과CAPM_강의본", "부록 B · 용어집", ("self", 75)),
    "W03": ("W03_연기금모델과CMA/W03_연기금모델과CMA_강의본", "부록 B · 용어집", ("self", 61)),
    "M9":  ("W07_LDI와GBI/W07_M9_GBI_목표기반투자_강의본", "부록 C · 용어집", ("ref", 49)),
    "W09": ("W09_위험관리와성과평가/W09_위험관리와성과평가_강의본", "부록 B · 용어집", ("ref", 49)),
}


def build(key):
    deck, badge, (kind, n) = DECKS[key]
    path = os.path.join(ROOT, deck + ".pptx")
    prs = Presentation(path)
    last = prs.slides[len(prs.slides) - 1]
    if find(last, badge):
        s = last                                                  # 이미 붙어 있다 — 표만 새로 채운다
    else:
        if kind == "self":
            s = dup_slide(prs, n - 1)
        else:
            s = clone_slide(Presentation(os.path.join(ROOT, REF + ".pptx")), n - 1, prs)
        for sh in [x for x in s.shapes if x.name.startswith("NAV_")]:
            sh._element.getparent().remove(sh._element)
        head = [sh for sh in s.shapes if sh.has_text_frame and sh.top < 914400 * 0.5 and sh.text_frame.text.strip()]
        set_text(head[0], badge)                                   # 왼쪽 위 배지
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            t = sh.text_frame.text.strip()
            if 0.7 < sh.top / 914400 < 1.0:
                set_text(sh, "이번 주의 용어")
            elif 1.4 < sh.top / 914400 < 1.7:
                set_text(sh, "시험 · 과제에서 그대로 쓰는 표준 표기")
            elif sh.top / 914400 > 7.0 and t.isdigit():
                set_text(sh, str(len(prs.slides)))                 # 쪽번호
    set_table(find_table(s), DATA[key])
    prs.save(path)
    print(key, "→", len(prs.slides), "쪽 ·", len(DATA[key]) - 1, "개 용어")


if __name__ == "__main__":
    for k in (sys.argv[1:] or DECKS):
        build(k)
