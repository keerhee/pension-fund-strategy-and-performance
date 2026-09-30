# -*- coding: utf-8 -*-
"""‘FLAM’ 용어 정정 + W09 30장 배너 식 정리 · 첨자 표기 — W09 강의본 · W09 케이스1 · W12 강의본 · W12 케이스 (2026-09-30).
덱은 팩터 노출 통합(β_P = Σ w_i β_i)을 ‘FLAM’이라 불렀지만, FLAM은 Grinold의 Fundamental Law of Active Management
(IR = IC × √BR)를 가리킨다. 팩터 노출 통합은 ‘팩터 렌즈’로 바꾸고, W09 Grinold 장에 FLAM의 올바른 뜻을 적는다.
W15는 fix_qa_w15.py 가 같은 규칙(slidekit.flam_fix)으로 처리한다.
실행: .venv/bin/python _build/lecfix/fix_flam.py   (재실행 가능 — _구판/FLAM정정_구판_2026-09-30/ 에서 원본 복원)
"""
import os, shutil, subprocess, sys
HERE = os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0, HERE)
from lecfix import *
from slidekit import flam_fix, sub_runs, Kit

R = os.path.dirname(os.path.dirname(HERE))
OLD = f"{R}/_구판/FLAM정정_구판_2026-09-30"; os.makedirs(OLD, exist_ok=True)
DECKS = ["W09_위험관리와성과평가/W09_위험관리와성과평가_강의본", "W09_위험관리와성과평가/W09_케이스1_국부펀드2022성과평가_IC",
         "W12_팩터투자/W12_팩터투자_강의본", "W12_팩터투자/W12_케이스_NPS국내주식팩터도입_IC"]


def backup(rel):
    for ext in (".pptx", ".pdf"):
        src, dst = f"{R}/{rel}{ext}", f"{OLD}/{os.path.basename(rel)}{ext}"
        if os.path.exists(dst): shutil.copy2(dst, src)
        elif os.path.exists(src): shutil.copy2(src, dst)


for rel in DECKS:
    backup(rel)
    prs = Presentation(f"{R}/{rel}.pptx")
    n = flam_fix(prs)
    if rel.endswith("W09_위험관리와성과평가_강의본"):
        assert replace_all(prs, "Grinold의 근본법칙 (1989)", "Grinold의 근본법칙 (1989) — FLAM") == 1
        assert replace_all(prs, "Grinold의 Fundamental Law (1989)", "Grinold의 Fundamental Law (1989) — 줄여서 FLAM") == 1
        assert replace_all(prs, "IR = IC√BR, 그리고 팩터 렌즈의 좌표", "IR = IC√BR(FLAM), 그리고 팩터 렌즈의 좌표") == 1
        # 30장 배너 — 위 LaTeX 식과 겹치는 일반 글자 식을 빼고 식을 가리킨다
        assert replace_all(prs, "각 자산을 팩터 노출로 분해 · 포트폴리오 노출 β_p = Σ w_i β_i — 숨은 집중을 드러낸다",
                           "각 자산을 팩터 노출로 분해해 비중대로 더한다 — 자산군 분산 뒤의 숨은 집중을 드러낸다") == 1
        from pptx.util import Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        s30 = prs.slides[29]
        for (x, w, t) in ((100, 580, "각 자산의 수익 = 고유 몫 α + Σ(팩터 민감도 β × 팩터 수익 F) + 잡음 ε"),
                          (690, 500, "기금의 팩터 노출 = 자산별 노출 β를 비중 w로 가중해 더한 값")):
            tb = s30.shapes.add_textbox(int(x * 9525), int(356 * 9525), int(w * 9525), int(60 * 9525))
            tf = tb.text_frame; tf.word_wrap = True
            p0 = tf.paragraphs[0]; p0.alignment = PP_ALIGN.CENTER
            r = p0.add_run(); r.text = t; r.font.size = Pt(14); r.font.name = "Noto Sans CJK KR"
            r.font.color.rgb = RGBColor(0x3B, 0x42, 0x52)
        # 22장 — 두 번째 식 그림이 패널과 겹치던 것을 푼다(패널을 16px 아래로)
        for sh in prs.slides[21].shapes:
            if 390 <= px(sh)[1] < 676:
                sh.top = sh.top + 16 * 9525
        # 23장 — Transfer Coefficient 는 ‘전이계수’
        assert replace_all(prs, "전이비용 (Transfer Coefficient)", "전이계수 (Transfer Coefficient)") == 1
        # 22장 뒤 — Grinold의 두 번째 공식: 예측 규칙 α = σ · IC · z (보충교재 15의 2절)
        kit = Kit(prs, panel_idx=[3, 4, 6, 8, 11], table_idx=10)
        T = prs.slides[3]
        g = kit.new(T, "2교시 · Grinold ②", "두 번째 공식 — 알파 = 변동성 × IC × 점수",
                    "Grinold의 예측 규칙 — 표준화된 신호(점수)를 기대 초과수익(알파)으로 바꾸는 규칙")
        kit.formula(g, r"\alpha_i\ =\ \sigma_i\cdot IC\cdot z_i", cx=640, y=192, pt=32)
        kit.panel(g, 60, 276, 568, 206, "b", "네 가지 가정", [
            "사전 기대 잔차수익 E[r] = 0 (컨센서스) · 변동성 σ",
            "점수 z는 평균 0 · 분산 1로 표준화",
            "상관 corr(r, z) = IC — 종목에 관계없이 같다",
            "(r, z) 결합정규 — 아니면 ‘최적 선형 예측’"], pt=15)
        kit.panel(g, 652, 276, 568, 206, "g", "유도 — 증명은 다음 장", [" "], pt=10)
        kit.formula(g, r"E[r\mid z]=E[r]+\frac{\mathrm{Cov}(r,z)}{\mathrm{Var}(z)}\,(z-E[z])", x=672, y=336, pt=19)
        kit.formula(g, r"\mathrm{Cov}(r,z)=IC\cdot\sigma\ \Rightarrow\ \alpha=0+IC\cdot\sigma\cdot z", x=672, y=408, pt=19)
        kit.table(g, 60, 496, 1161, [
            ["예 (IC = 0.05)", "변동성 σ", "점수 z", "알파 α = σ × 0.05 × z"],
            ["종목 A", "30%", "+2.0", "*+3.0%"],
            ["종목 B", "20%", "+1.0", "*+1.0%"],
            ["종목 C", "25%", "−1.5", "*−1.9%"]],
            [1.6, 1.0, 1.0, 2.0], rowh=30, pt=14, align="lccc")
        kit.bottom(g, "IC가 0.05로 작으면 점수 +2도 알파 3%로 ‘수축’ — 과신을 막는 장치 · 이 알파를 비중으로 옮기면 IR = TC · IC · √BR (보충교재 15)", pt=14, y=626)
        # 증명 — 기본 예측 공식 = 오차 제곱 평균을 최소로 하는 직선(가장 간단한 형태)
        g2 = kit.new(T, "2교시 · Grinold ② 증명", "증명 — 기본 예측 공식은 오차를 가장 작게 만드는 직선",
                     "분포 가정 없이 네 단계 — 정규분포라면 이 직선이 곧 조건부 기대값 E[r | z]")
        steps = [("① 직선으로 예측하고, 오차 제곱의 평균을 최소화한다", r"\hat r=a+b\,z,\qquad \min_{a,b}\ E\left[(r-a-bz)^{2}\right]"),
                 ("② a로 미분해 0 — 예측선은 평균점(μ_z, μ_r)을 지난다", r"a=\mu_r-b\,\mu_z"),
                 ("③ 대입해 b로 미분해 0 — 기울기는 공분산 ÷ 분산", r"E\left[\left((r-\mu_r)-b(z-\mu_z)\right)^{2}\right]=\mathrm{Var}(r)-2b\,\mathrm{Cov}(r,z)+b^{2}\,\mathrm{Var}(z)\ \Rightarrow\ b=\frac{\mathrm{Cov}(r,z)}{\mathrm{Var}(z)}"),
                 ("④ 합치면 기본 예측 공식 → 가정(μ = 0 · Var(z) = 1 · Cov = IC·σ)을 넣으면 예측 규칙", r"\hat r=\mu_r+\frac{\mathrm{Cov}(r,z)}{\mathrm{Var}(z)}(z-\mu_z)\ \ \Rightarrow\ \ \alpha=IC\cdot\sigma\cdot z")]
        for k, (lab, tex) in enumerate(steps):
            y = 190 + k * 104
            kit.label(g2, 60, y, 1161, 26, lab, pt=15, color="1B2C5E", bold=True)
            kit.formula(g2, tex, x=80, y=y + 30, pt=21)
        kit.bottom(g2, "기울기 b는 ‘z가 1 움직일 때 r이 평균적으로 움직이는 양’ — IC가 작으면 기울기도 작아 점수를 줄여 반영한다", pt=15, y=620)
        lst = list(prs.slides)
        move_slide(prs, lst.index(g), 22)
        lst = list(prs.slides)
        move_slide(prs, lst.index(g2), 23)
        renumber(prs)
    n += sub_runs(prs)   # σ_down 등 남은 첨자 표기
    left = grep(prs, "FLAM")
    prs.save(f"{R}/{rel}.pptx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(f"{R}/{rel}.pptx"), f"{R}/{rel}.pptx"],
                   check=True, capture_output=True)
    print(os.path.basename(rel), "치환", n, "· 남은 FLAM", left)
