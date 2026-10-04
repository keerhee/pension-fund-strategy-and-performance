"""W09 강의본 보수(2026-10-05) — 영상 나레이션 집필 중 발견한 것.
 8장 VaR 식의 부호(−inf → inf) · 34장 그림 제목(FLAM → 팩터 렌즈, 용어 규칙) · 46장 메타프롬프트 예시(W06 LDI → W09 위험·성과 엔진)
 1장 시간 표기(× 2 → 2건, 2장 240분과 맞춤) · 31장 '합계' 열이 주식+채권임을 머리글에 밝힘.
실행: <PPTX 가능 파이썬> _build/w09_fix/fix_w09.py   (원본은 _구판/W09_덱보수_구판_2026-10-05/에 한 번만 백업, 재실행해도 결과가 같다)
그림 편집에는 Pillow · matplotlib가 필요하다."""
import io, os, shutil, sys
sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "lecfix"))
from pptx import Presentation
from pptx.util import Emu
from PIL import Image, ImageDraw, ImageFont
import matplotlib; matplotlib.use("Agg"); import matplotlib.pyplot as plt
import lecfix
R = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SRC = os.path.join(R, "W09_위험관리와성과평가", "W09_위험관리와성과평가_강의본.pptx")
BK = os.path.join(R, "_구판", "W09_덱보수_구판_2026-10-05"); os.makedirs(BK, exist_ok=True)
bk = os.path.join(BK, os.path.basename(SRC))
if not os.path.exists(bk): shutil.copy2(SRC, bk)
TTC = os.path.expanduser("~/Library/Fonts/NotoSansCJK.ttc"); REG, BOLD = 26, 36
P = Presentation(bk)
def pic(n, sid): return next(sh for sh in P.slides[n - 1].shapes if sh.shape_id == sid)
def swap_blob(sh, png_bytes):
    part = sh.part.related_part(sh._element.blipFill.blip.rEmbed); part._blob = png_bytes
def to_png(im): b = io.BytesIO(); im.save(b, "PNG"); return b.getvalue()

def redraw(im, box, lines, size, color, bg, bold=False, cy=None, gap=None):
    d = ImageDraw.Draw(im); x0, y0, x1, y1 = box
    d.rectangle(box, fill=bg)
    f = ImageFont.truetype(TTC, size, index=BOLD if bold else REG)
    cys = cy if isinstance(cy, list) else [((y0 + y1) / 2) + (i - (len(lines) - 1) / 2) * (gap or size * 1.6) for i in range(len(lines))]
    for t, y in zip(lines, cys):
        d.text(((x0 + x1) / 2, y), t, font=f, fill=color, anchor="mm")

# 8장 — VaR 식
sh = pic(8, 117)
fig = plt.figure(figsize=(8, 0.6), dpi=200)
with plt.rc_context({"mathtext.fontset": "cm"}):
    fig.text(0.01, 0.5, r"$\mathrm{VaR}_\alpha = \inf\{\,\ell : P(L>\ell) \leq 1-\alpha\,\}$", fontsize=30, color="#1B2A4A", va="center")
buf = io.BytesIO(); fig.savefig(buf, format="png", bbox_inches="tight", pad_inches=0.04, facecolor="white"); plt.close(fig)
im = Image.open(io.BytesIO(buf.getvalue())).convert("RGB")
h = sh.height; sh.width = Emu(int(h * im.width / im.height)); swap_blob(sh, to_png(im))

# 34장 — 그림 제목
sh = pic(34, 477); im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
redraw(im, (300, 16, 1300, 64), ["팩터 렌즈 — 자산군 분산 뒤에 숨은 팩터 집중"], 27, "#222C5E", "#FFFFFF", bold=True, cy=[39])
swap_blob(sh, to_png(im))

# 46장 — 메타프롬프트 예시를 W09 실습(위험·성과·요인분석 엔진)으로
sh = pic(46, 654); im = Image.open(io.BytesIO(sh.image.blob)).convert("RGB")
ys = [277, 312, 347]; C = "#3B4252"
redraw(im, (52, 258, 358, 366), ['"위험·성과 엔진을', '만들어줘 — 4척도·6지표,', 'Brinson·팩터 렌즈"'], 22, C, "#FCF2E0", cy=ys)
redraw(im, (472, 258, 778, 366), ["risk_perf.py (6 Step)", "위험 4척도 · 성과 6지표", "Brinson · 스트레스"], 22, C, "#F4F5F9", cy=ys)
redraw(im, (1312, 258, 1558, 366), ["지표 순위 역전 표", "Brinson 합 검산", "숨은 신용 베타 0.29"], 22, C, "#F4F5F9", cy=ys)
redraw(im, (150, 536, 1450, 578), ['개선 루프 — "케이스①로 Brinson을 다시 돌려줘" (말로 반복 개선)'], 25, "#3FA36F", "#FFFFFF", bold=True, cy=[556])
swap_blob(sh, to_png(im))

# 텍스트
for old, new in [("180분 강의 + 60분 모의 투자위원회(IC) × 2", "180분 강의 + 60분 모의 투자위원회(IC) 2건")]:
    assert lecfix.replace_all(P, old, new) == 1, old
t = next(sh for sh in P.slides[30].shapes if getattr(sh, "has_table", False) and sh.has_table)
c = t.table.cell(0, 2); r = c.text_frame.paragraphs[0].runs[0]; r.text = "합계 (주식 + 채권)"
for x in c.text_frame.paragraphs[0].runs[1:]: x.text = ""
P.save(SRC); print("saved", SRC)
