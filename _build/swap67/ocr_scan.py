# -*- coding: utf-8 -*-
"""모든 PPTX 그림을 뽑아(중복 제거) tesseract kor+eng 로 주차·모듈 표기를 찾는다."""
import glob, os, hashlib, subprocess, re, json, sys
from concurrent.futures import ThreadPoolExecutor
from pptx import Presentation
ROOT=os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))); os.chdir(ROOT)
OUT="/private/tmp/claude-501/-Users-keerhee-Project/4e0f873d-89dd-4ceb-902b-b5b877eda07e/scratchpad/ocr67"; os.makedirs(OUT,exist_ok=True)
SKIP=("Risk_Based_Allocation","Cerberus","Direct_Lending","Residential_Credit","04_RL_Workbook")
files=sorted(f for f in glob.glob("**/*.pptx",recursive=True) if not f.startswith(("_구판","_build","site")) and not any(k in f for k in SKIP))
seen={}
def pics(shapes):
    for sh in shapes:
        if sh.shape_type==13: yield sh
        elif sh.shape_type==6: yield from pics(sh.shapes)
for f in files:
    prs=Presentation(f)
    for i,s in enumerate(prs.slides,1):
        for sh in pics(s.shapes):
            try: blob=sh.image.blob; ext=sh.image.ext
            except Exception: continue
            if len(blob)<8000: continue
            h=hashlib.md5(blob).hexdigest()[:12]
            p=f"{OUT}/{h}.{ext}"
            if h not in seen:
                open(p,"wb").write(blob); seen[h]=[]
            seen[h].append((f,i))
print(len(files),"files",len(seen),"unique images",file=sys.stderr)
pat=re.compile(r"(?<![A-Za-z0-9])(W\s?0?[67]|Week\s?[67]|WEEK\s?[67]|M\s?[789])(?![0-9])|(?<![0-9])[67]\s?주")
def ocr(h):
    p=glob.glob(f"{OUT}/{h}.*")[0]
    try: t=subprocess.run(["tesseract",p,"-","-l","kor+eng","--psm","11"],capture_output=True,text=True,timeout=120).stdout
    except Exception: t=""
    return h,[m.group(0) for m in pat.finditer(t)]
hits={}
with ThreadPoolExecutor(6) as ex:
    for h,m in ex.map(ocr,list(seen)):
        if m: hits[h]={"tokens":m,"where":seen[h]}
json.dump(hits,open(f"{OUT}/hits.json","w"),ensure_ascii=False,indent=1)
for h,v in hits.items(): print(h, v["tokens"][:8], v["where"][:3])
