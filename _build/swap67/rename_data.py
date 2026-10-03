# -*- coding: utf-8 -*-
"""W06↔W07 교환 후속 — 실습·케이스 데이터 파일 이름을 새 주차·모듈 번호로 (2026-10-03).
실습 데이터는 모듈 번호(fml_w{M}_), 케이스 데이터는 주차+모듈(fml_w{W}m{M}_ · fml_w{W}ss2_) 규칙.
  fml_w8_liability → fml_w7_liability · fml_w9_* → fml_w8_* · w9_sim → w8_sim          (M7 LDI · M8 GBI)
  fml_w7m8_ → fml_w6m7_ · fml_w7m9_ → fml_w6m8_ · w7m8_/w7m9_ 스크립트도               (W06 케이스)
  fml_w6_ → fml_w7m9_ · w6_build/compute/results → w7m9_* · fml_w6ss2_ → fml_w7ss2_      (W07 케이스 · SS2)
텍스트(PPTX 본문·표·노트, md·py·html·json·txt·js·csv·tsv)와 실제 파일 이름을 함께 바꾼다. --apply 없으면 dry run."""
import glob, os, re, sys, shutil
from pptx import Presentation
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))); os.chdir(ROOT)
APPLY = "--apply" in sys.argv
BK = "_구판/W06W07교환_구판_2026-10-03/_data_rename"
RX = re.compile(r"(?<![A-Za-z0-9])(fml_)?(w7m8|w7m9|w6ss2|w6|w8|w9)(?=_|\b)(_[A-Za-z0-9_]*)?")
def conv(m):
    fml, core, rest = m.group(1) or "", m.group(2), m.group(3) or ""
    if core == "w7m8": new = "w6m7"
    elif core == "w7m9": new = "w6m8"
    elif core == "w6ss2": new = "w7ss2"
    elif core == "w6":
        if fml or re.match(r"_(build|compute|results)", rest): new = "w7m9"
        else: return m.group(0)
    elif core == "w8":
        if fml and rest.startswith("_liability"): new = "w7"
        else: return m.group(0)
    elif core == "w9":
        if fml or rest.startswith("_sim"): new = "w8"
        else: return m.group(0)
    return fml + new + rest
def sub(t): return RX.sub(conv, t)

def tfs(sh):
    if sh.has_text_frame: yield sh.text_frame
    if getattr(sh, "has_table", False) and sh.has_table:
        for r in sh.table.rows:
            for c in r.cells: yield c.text_frame
    if sh.shape_type == 6:
        for s in sh.shapes: yield from tfs(s)
nt = 0; changed = []
for f in sorted(glob.glob("**/*.pptx", recursive=True)):
    if f.startswith(("_구판", "_build")): continue
    prs = Presentation(f); n = 0
    for s in prs.slides:
        srcs = [tf for sh in s.shapes for tf in tfs(sh)]
        if s.has_notes_slide: srcs.append(s.notes_slide.notes_text_frame)
        for tf in srcs:
            for p in tf.paragraphs:
                t = p.text; new = sub(t)
                if new == t: continue
                n += 1; runs = p.runs
                per = [sub(r.text) for r in runs]
                if "".join(per) == new:
                    for r, x in zip(runs, per): r.text = x
                else:
                    runs[0].text = new
                    for r in runs[1:]: r.text = ""
    if n:
        changed.append(f); nt += n; print(f"{n:3d}  {f}")
        if APPLY:
            b = os.path.join(BK, f); os.makedirs(os.path.dirname(b), exist_ok=True)
            if not os.path.exists(b): shutil.copy2(f, b)
            prs.save(f)
for f in sorted(glob.glob("**/*", recursive=True)):
    if f.startswith(("_구판", "_build/swap67")) or not os.path.isfile(f) or not f.endswith((".md", ".py", ".html", ".json", ".txt", ".js", ".csv", ".tsv")): continue
    try: t = open(f, encoding="utf-8").read()
    except Exception: continue
    new = sub(t)
    if new != t:
        print(f"txt  {f}")
        if APPLY:
            b = os.path.join(BK, f); os.makedirs(os.path.dirname(b), exist_ok=True)
            if not os.path.exists(b): shutil.copy2(f, b)
            open(f, "w", encoding="utf-8").write(new)
# 실제 파일 이름
for f in sorted(glob.glob("W0[67]_*/**/*", recursive=True)):
    if not os.path.isfile(f): continue
    d, b = os.path.split(f); nb = sub(b)
    if nb != b:
        print(f"mv   {f} -> {nb}")
        if APPLY: os.rename(f, os.path.join(d, nb))
print("pptx", len(changed), "paragraphs", nt, "(APPLIED)" if APPLY else "(dry run)")
open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "data_touched.txt"), "w").write("\n".join(changed))
