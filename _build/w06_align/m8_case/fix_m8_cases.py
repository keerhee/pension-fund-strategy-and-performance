# -*- coding: utf-8 -*-
"""W06 M8 케이스 1(디폴트옵션 · 45장) · 케이스 2(H대 · 41→42장)를 개정 강의본(2026-10-05)에 맞춘다. 강의 장 번호는 _build/w06_m8_rework/pages.json 에서 읽는다.

실행: .venv/bin/python _build/w06_align/m8_case/fix_m8_cases.py
  - 항상 _구판/W06_정합전_2026-10-05/ 의 원본에서 읽는다(재실행 가능).
  - 주차 폴더에 같은 이름으로 저장하고, PDF는 Pretendard 임시 사본으로 만든다.
강의본 키: V1(m ≤ 1/x · 기본 답의 역산) · K2(디폴트옵션) · I3(새 상품 (a)) · o34(RB 없음) · FX1 · FX2(Flexicure 다섯 단계) · MC1~MC4(방법 3).
"""
import os, re, shutil, subprocess, sys, tempfile
HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
sys.path.insert(0, f"{R}/_build/lecfix")
from lecfix import *

import json
PG = json.load(open(f"{R}/_build/w06_m8_rework/pages.json", encoding="utf-8"))   # 강의본 장 번호(이름 키 → 실제 번호)
def P(*ks): return " · ".join(str(PG[k]) for k in ks)
MC = f"{PG['MC1']}~{PG['MC4']}"                                                   # 방법 3 몬테카를로
MCL = f"강의 방법 3(몬테카를로, {MC}장)으로 센 값"

W = f"{R}/W06_LDI와GBI"
OLD = f"{R}/_구판/W06_정합전_2026-10-05"
C1 = "W06_M8_케이스1_퇴직연금디폴트옵션도입_IC"
C2 = "W06_M8_케이스2_대학발전기금Yale모델도입_IC"


def rep(prs, a, b, n=None):
    k = replace_all(prs, a, b)
    assert k >= 1, ("not found", a[:60])
    if n is not None:
        assert k == n, (a[:60], k, n)
    return k


def cell_set(slide, r, c, text):
    set_tf(find_table(slide).table.cell(r, c).text_frame, text)


def renum2(prs):
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_text_frame:
                x, y, w, h = px(sh)
                if x >= 1000 and y >= 640 and re.fullmatch(r"\d{1,3}", sh.text_frame.text.strip()):
                    set_text(sh, f"{i:02d}")


def to_pdf(pptx):
    tmp = tempfile.mkdtemp(dir=os.environ.get("TMPDIR"))
    base = os.path.basename(pptx)
    subprocess.run(["/usr/bin/python3", f"{R}/_build/to_pretendard.py", pptx, f"{tmp}/{base}"], check=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{tmp}/lo", "--headless", "--convert-to", "pdf", "--outdir", tmp, f"{tmp}/{base}"],
                   check=True, capture_output=True)
    shutil.copy2(f"{tmp}/{base[:-5]}.pdf", pptx[:-5] + ".pdf")
    shutil.rmtree(tmp)


# ═════════════════════════════ 케이스 1 · 디폴트옵션 ═════════════════════════════
prs = Presentation(f"{OLD}/{C1}.pptx")
S = lambda k: prs.slides[k - 1]

# 7 판정의 말 — 갭 위험을 VaR로 읽는 법(강의 62장), 디폴트옵션 승인 주체(강의 79장)
t = find_table(S(7))
cell_set(S(7), 2, 1, "리밸런싱 사이 PSP 하락이 1/m을 넘으면 Floor가 깨진다 — 거꾸로 m ≤ 1/x (x = 최악 하락률 · VaR)")
cell_set(S(7), 2, 2, "m 2 → 50% · m 3 → 33% · 한 달 −22%")
cell_set(S(7), 5, 1, "사전지정운용제도 — DC·IRP 미지시 적립금의 자동 운용(2022.7 도입, 2023.7 본격 시행) · 상품은 고용노동부 심의위원회가 승인")
# 17 연쇄 — 교시 → 단원
rep(prs, "강의 2교시의 연쇄를 상품 언어로 옮긴다", "강의 단원 ④⑤의 연쇄를 상품 언어로 옮긴다", 1)
# 21 결정 대상 — 승인 주체
rep(prs, "Flexicure형(Floor + Cushion, CPPI)을 어떤 설계로 등록하는가",
    "Flexicure형(Floor + Cushion, CPPI)을 어떤 설계로 등록하는가 — 승인은 고용노동부 심의위원회", 1)
# 23 근거 열 — 교시 → 단원
rep(prs, "강의 1교시 Safety ≥ 95%", "강의 단원 ① Safety ≥ 95%", 1)
rep(prs, "강의 2교시 CPPI의 한계", "강의 단원 ⑤⑥ CPPI의 한계", 1)
# 31 P2 — m 상한을 VaR로
cell_set(S(31), 2, 2, f"m ≤ 1/x — m 2는 한 번에 50% 하락까지 견딘다(강의 {P('V1')}장 ‘기본 답의 역산’) · m 3은 33% · 한 달 −22%에 m 4는 32%가 깨진다")
# 40 기본 답 — 팀 과제 3팀 B안 · 강의 89장 새 상품 (a)
set_text(find(S(40), "교육용 가정(G_S"),
         f"이 답 = 팀 과제 3팀 B안(Floor 90% · m 2 · 0.5%) = 강의 {P('I3')}장 새 상품 (a)의 설계값")
# 42 보너스 — 교시 → 단원
rep(prs, "강의 2교시 — 물가연동국채는 있으나", f"강의 단원 ④({P('o34')}장) — 물가연동국채는 있으나", 1)
# 44 용어집 — 디폴트옵션 도입 시점
rep(prs, "DC · IRP 미지시 적립금의 자동 운용 — 2023.7 본격 시행", "DC · IRP 미지시 적립금의 자동 운용 — 2022.7 도입 · 2023.7 본격 시행", 1)
# 45 마감 — W06↔W07 교환 뒤 주차
rep(prs, "W4 · W5 · W7의 표결과 함께", "W4 · W5 · W6 M7(LDI)의 표결과 함께", 1)

# 방법 3(몬테카를로)으로 센 확률임을 밝힌다 — 숫자는 그대로
rep(prs, "결과는 수수료 0.5%, 4,000 경로 기준(w6m8_compute.py)", f"결과는 수수료 0.5% · 4,000 경로 — {MCL}", 1)
set_text(find(S(23), "임계는 사전에 준다"), f"임계는 사전에 준다 — 확률 조건 ①~④는 {MCL}으로 판정한다")
set_text(find(S(30), "통과 = ①~④ 전부"), f"통과 = ①~④ 전부 · 초기 PSP는 수수료 그로스업 후 Floor 기준 · 확률은 {MCL}")
assert not grep(prs, r"교시"), grep(prs, r"교시")
prs.save(f"{W}/{C1}.pptx"); to_pdf(f"{W}/{C1}.pptx")
print("케이스 1:", len(prs.slides), "장")

# ═════════════════════════════ 케이스 2 · H대 ═════════════════════════════
prs = Presentation(f"{OLD}/{C2}.pptx")
S = lambda k: prs.slides[k - 1]

# 1) 펀드오브펀드 → 재간접 · 세컨더리 (긴 꼴부터)
for a, b in [
    ("간접(펀드오브펀드) 경유", "재간접 · 세컨더리 경유"),
    ("펀드오브펀드 · 세컨더리 펀드로만", "재간접 · 세컨더리 펀드로만"),
    ("직접 · 펀드오브펀드 · 세컨더리", "직접 · 재간접 · 세컨더리"),
    ("접근권 · 펀드오브펀드 · 세컨더리", "접근권 · 재간접 · 세컨더리"),
    ("펀드오브펀드 · 세컨더리", "재간접 · 세컨더리"),
    ("펀드오브펀드 경유", "재간접 · 세컨더리 경유"),
    ("국내 펀드오브펀드 실현치", "국내 재간접 실현치"),
    ("펀드오브펀드는 −2%p로", "재간접은 −2%p로"),
    ("펀드오브펀드로 −2%p", "재간접으로 −2%p"),
    ("펀드오브펀드 −2%p", "재간접 −2%p"),
    ("직접 → 펀드오브펀드", "직접 → 재간접"),
    ("z 4.5 · 펀드오브펀드", "z 4.5 · 재간접"),
    ("지출 4.5 · 펀드오브펀드", "지출 4.5 · 재간접 · 세컨더리"),
]:
    replace_all(prs, a, b)
left = grep(prs, "펀드오브펀드"); assert not left, left

# 2) 한 장 요약 · 대비표 — 과제 4팀 A · B · C와의 관계
rep(prs, "기본 답 — (x, y, z) = (20, 30, 4.5) 조건부 승인 · 분모 효과 시 신규 약정 중단을 의결문에",
    "기본 답 (20, 30, 4.5) = 팀 과제 4팀 B안 · 조건부 승인 · 사모 22%(목표 + 2%p) 넘으면 신규 약정 중단", 1)
rep(prs, "안건 개요 대비표 — 구판의 Yale형 A · GPFG형 B · 하이브리드 C를 좌표로 바꿔 판정한다",
    "안건 개요 대비표 — 구판의 세 이름을 좌표로 · 과제 4팀 A안 = GPFG형 좌표, B안 = 기본 답", 1)
cell_set(S(3), 0, 3, "기본 답 = 과제 4팀 B안")
# 3) 교시 → 단원
rep(prs, "모방의 3대 장벽(강의 3교시 — 접근 · 시간 · 인력)이 조건 ①②③으로 번역된다",
    "모방 3장벽(강의 단원 ⑧ — 접근권 · 시간 지평 · 인력)이 조건 ①②③으로 번역된다", 1)
rep(prs, "기관의 목표 분해 · 지출 규칙 · 시간지평 · 4대 도전(강의 3교시)과 케이스 1의 Floor 언어만이 유효한 논거다.",
    "기관의 목표 분해 · 지출 규칙 · 모방 3장벽 · Flexicure 다섯 단계(강의 단원 ⑧)와 케이스 1의 Floor 언어만이 유효한 논거다.", 1)
rep(prs, "강의 3교시 · 케이스 1", "강의 단원 ⑧ · 케이스 1", 1)
rep(prs, "기관의 목표 분해 · 4대 도전 · Floor 언어 →", "모방 3장벽 · Flexicure 다섯 단계 · Floor 언어 →", 1)
# 4) 조건 ⑤ — Flexicure 규칙(강의 97장)
cell_set(S(20), 5, 1, "충격 뒤에도 (주식 + 사모) ÷ 쿠션 ≤ 2 · 사모 22% 넘으면 신규 약정 중단 · 연 1회 재계산")
set_text(find(S(20), "임계는 사전에 준다"), f"임계는 사전에 준다 — 확률 조건 ②④는 {MCL}으로 판정한다")
rep(prs, "산수 · 조건 ② — 4,000 경로 · 30년 · 80/20 평활 · 패널 연 수익률 재표본(교육용 CMA)",
    f"산수 · 조건 ② — 4,000 경로 · 30년 · 80/20 평활 · 패널 재표본 — {MCL}", 1)
rep(prs, "산수 · 조건 ④ — x 20 · z 4.5 · 1년차 2022형 · 3년 동안 안전 유동자산만으로 지출과 일시금을 낸다",
    f"산수 · 조건 ④ — x 20 · z 4.5 · 1년차 2022형 · 3년 지출 · 일시금 — 확률은 {MCL}", 1)
# 5) 교수용 기본 답 · 의결문
set_text(find(S(35), "① 13년 ≥ 10"),
         "① 13년 ≥ 10 · ② 59% ≥ 50%\n③ 직접 불가 → 재간접 · 세컨더리\n④ y 30에서 100%\n⑤ 2008형 뒤 m 1.53 ≤ 2 · 사모 22% 넘으면 약정 중단 · 연 1회 재계산")
rep(prs, "하이브리드는 이름이 아니라 좌표입니다 — (20, 30, 4.5)라고 써야 심의가 됩니다",
    "(20, 30, 4.5) = 과제 4팀 B안 — 하이브리드는 이름이 아니라 좌표입니다", 1)
set_text(find(S(36), "비유동 대체 상한을 20%로"),
         "비유동 대체 상한을 20%로 하고, 재간접 · 세컨더리 펀드로만 편입한다 (전담 인력 3명 확보 시 재상정).\n"
         "현금 · 국채 등 안전 유동자산을 30% 이상 유지하고, 충격 후 3년간 비유동 자산을 매도하지 않는다.\n"
         "지출률은 4.5%로 하되 80/20 평활을 적용하고, 장학 · 연구 고정 지출 175억(3.5%)을 Floor로 둔다 — 쿠션이 기준 아래면 희망 몫을 동결한다.\n"
         "최악 충격 뒤에도 (주식 + 사모) ÷ 쿠션을 2 이하로 두고, 사모 비중이 22%를 넘으면 신규 약정을 중단하며, 지출률 · 유동성은 연 1회 재계산해 보고한다.")

# 6) 새 장 — Flexicure 다섯 단계(강의 96 · 97장, 과제 4팀 문제 6) : s3 표 장을 복제해 s25 뒤에 둔다
new = dup_slide(prs, 2)
move_slide(prs, len(prs.slides) - 1, 25)           # 0-based 25 → 26번째 장
s = prs.slides[25]
set_text(find(s, "BAF.60080"), "BAF.60080 · W06 M8 IC 제2호 · 안건")
set_text(find(s, "세 가지 이름은"), "Flexicure로 다시 재도 비유동 상한은 20%이고 그것이 과제 B안입니다")
set_text(find(s, "안건 개요 대비표"), f"산수 · 과제 4팀 세 안 = 격자의 세 점 · 단위 억 원 · 2008형: 주식 −40 · 국채 +8 · 현금 +3 · 사모 −10% (강의 {P('FX1', 'FX2')}장)")
tb = find_table(s)
set_table(tb, [
    ["단계", "A안 (0, 35, 4.0) = GPFG형", "B안 (20, 30, 4.5) = 기본 답", "C안 (30, 15, 5.25)"],
    ["① Floor = 3년 지출 536 + 공사비 600 + 사모 × 40%", "1,136", "1,536", "1,736"],
    ["② GHP — 2022형 뒤 현금 · 국채 ≥ 1,136", "1,530 — 통과", "1,318 — 통과", "680 — 미달"],
    ["③ 평시 쿠션 · m (상한 2)", "3,864 · m 0.84", "2,464 · m 1.42", "1,764 · m 2.41 — 초과"],
    ["④ 2008형 뒤 쿠션 · m (사모는 못 판다)", "사모 0 — 주식 매도로 조정", "1,572 · m 1.53 — 통과", "712 · m 4.2 — 탈락"],
    ["④ 2008형 뒤 사모 비중", "0%", "22.5% — 약정 중단", "35.5% — 약정 중단"],
    ["⑤ 지출 z (고정 3.5% = Floor)", "z 4.0", "z 4.5", "z 5.25"],
])
from pptx.util import Inches
for c, wpx in zip(tb.table.columns, (372, 240, 250, 270)):
    c.width = Inches(wpx / 96)
set_text(find(s, "판정 순서"), "비유동 상한 x = 최악 충격 뒤에도 m ≤ 2를 지키는 최대치 — 사모 22% 넘으면 신규 약정 중단 · 쿠션이 기준 아래면 희망 몫 동결")

renum2(prs)
assert not grep(prs, r"[123]교시"), grep(prs, r"[123]교시")
prs.save(f"{W}/{C2}.pptx"); to_pdf(f"{W}/{C2}.pptx")
print("케이스 2:", len(prs.slides), "장")
