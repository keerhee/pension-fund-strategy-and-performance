# -*- coding: utf-8 -*-
"""보충교재 17 · GBI/04 Martellini GBI 수식 도출 III(덱) — 05 노트(2026-10-05 갱신)와 W06 M8 강의본에 맞춘다.
새 장 4개 (05 노트의 새 소단계 · 84장 → 88장)
  31장  3-6 ② (c) VaR로 floor를 움직이는 위험 자본 방식 — 위험 ↑ → floor ↑ · 쿠션 ↓ · 노출 ↓ (σ 15/20/25% 표)
  40장  4-2 단계 1-2 두 자산 손풀이 (w* = 52.2%)
  47장  4-2 한계 — 롱온리 풀이(상속 0 · 8.9 · 91.1%, 연 12bp vs 합산 수준 제약 0)
  48장  4-2 한계 — 엑셀 해 찾기 · scipy SLSQP
고친 장(옛 번호 → 새 번호)
  7 기호 각주의 장 번호 · 30 부제 · 39→41 원문 A·B·C·D와 강의 K = D/C · 40→42 3.9750 위치(p.320, p.317·표는 3.7950)
  42→44 '공매도 허용 시' 합산 효율 · 43→45 각주 장 번호 · 44→46 한계 줄에 47·48장 · 80→84 연 12bp · 83→87 3-6(c) 한 줄
실행: <python-pptx + matplotlib 파이썬> patch_04.py   (늘 _구판/W06_정합전_2026-10-05/보충17_프라이머/GBI/에서 읽는다)
"""
import os
from common import *
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

SRC = os.path.join(BK, "GBI", "04_Martellini_GBI_Derivations_III_zeroone.pptx")
DST = os.path.join(SUP, "GBI", "04_Martellini_GBI_Derivations_III_zeroone.pptx")
FONT = "Noto Sans CJK KR"
INK, TEAL, MUTED, PAPER, HAIR = "102027", "0B6B68", "64757D", "F7F5F0", "DDE3E2"
P = Presentation(SRC)
OLD = list(P.slides)            # 옛 번호(1부터) → OLD[n-1]

def T(slide, sid, text): lecfix.set_text(shape(slide, sid), text)

def eq_in(slide, card_sid, pic_sid, name, tex, fs=28):
    """card 안의 수식 그림을 새 수식으로 바꾼다(카드 배경 F7F5F0)."""
    card = shape(slide, card_sid); remove(shape(slide, pic_sid))
    path, ratio = eq_png(name, tex, fg="#" + INK, bg="#" + PAPER, fs=fs)
    x, y, w, h = lecfix.px(card)
    fit_pic(slide, path, ratio, x + 30, y + 14, w - 60, h - 28)

def table(slide, x, y, w, rows, colw, rowh=38, pt=15):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), E(x), E(y), E(w), E(rowh * len(rows)))
    tb = shp.table
    tblPr = tb._tbl.tblPr
    for k in ("bandRow", "firstRow"): tblPr.set(k, "0")
    tot = sum(colw)
    for j, cw in enumerate(colw): tb.columns[j].width = E(w * cw / tot)
    for i, r in enumerate(rows):
        tb.rows[i].height = E(rowh)
        for j, val in enumerate(r):
            c = tb.cell(i, j); c.fill.solid()
            c.fill.fore_color.rgb = RGBColor.from_string(PAPER if i == 0 else "FFFFFF")
            c.margin_left = c.margin_right = E(8); c.margin_top = c.margin_bottom = E(2)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = c.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
            bold = val.startswith("*"); val = val.lstrip("*")
            run = p.add_run(); run.text = val
            f = run.font; f.name = FONT; f.size = Pt(pt - 2 if i == 0 else pt); f.bold = (i == 0) or bold
            f.color.rgb = RGBColor.from_string(TEAL if i == 0 else INK)
    return shp

# ===================================================================== 새 장
# --- A: 3-6(c) 위험 자본 방식 (원형 38장: 수식 띠 + 흰 카드 + 결론 + 각주)
A = dup_slide(P, 37)
T(A, 3, "위험이 오르면 floor를 올려 손실 예산을 지킵니다")
T(A, 5, "3-6 ② (c) — VaR로 floor를 움직이는 위험 자본 방식 · m과 손실 예산 L은 고정, floor를 매번 다시 정한다")
eq_in(A, 7, 8, "s04_rc", r"E\,x=L,\quad E=m\,C\quad\Longrightarrow\quad \mathrm{floor}=A-\frac{L}{m\,x},\qquad x=1.645\,\sigma-\mu")
remove(shape(A, 10)); remove(shape(A, 11))
table(A, 84, 372, 1101, [
    ["σ (위험)", "VaR 비율 x", "쿠션 C = L ÷ (m x)", "floor = A − C", "위험 노출 E = m C", "VaR 손실 E x"],
    ["15% (하락)", "18.7%", "28.8", "71.2 ↓", "86.4 ↑", "16.1"],
    ["20% (기준)", "26.9%", "20.0", "80.0", "60.0", "16.1"],
    ["*25% (상승)", "*35.1%", "*15.3", "*84.7 ↑", "*46.0 ↓", "*16.1"],
], [1.3, 1.2, 1.6, 1.3, 1.6, 1.2], rowh=44, pt=17)
T(A, 12, "위험 ↑ → floor ↑ · 쿠션 ↓ · 노출 ↓ — 위험이 오를 때 floor를 낮추면 위험 자본과 반대입니다.")
T(A, 13, "※ A = 100 · m = 3 · μ = 6% · 무위험 0 · 1년 · 정규, 기준 σ 20%에서 C = 20이 되도록 L = 3 × 26.9% × 20 = 16.1 · floor는 필수 목표의 가격(GHP) 아래로 내리지 않는다 · (a)는 floor 고정·m 변동, (c)는 m 고정·floor 변동")

# --- B: 4-2 단계 1-2 두 자산 손풀이
B = dup_slide(P, 37)
T(B, 3, "자산을 둘로 줄이면 주식 비중 하나를 손으로 풉니다")
T(B, 5, "4-2 단계 1-2 — 채권형(5%, 5%) + 저위험 주식형(10%, 20%), 상관 0 · 은퇴 계정 H = −10%, α = 5%")
eq_in(B, 7, 8, "s04_two", r"m(w)=5\%+5\%\,w,\qquad \sigma(w)=\sqrt{(1-w)^2\times 0.05^2+w^2\times 0.20^2}")
T(B, 10, "시험해 보고, 제약이 등호로 걸릴 때까지 주식 비중 w를 올린다")
T(B, 11, "w = 0.5: 평균 7.5%, 표준편차 10.31% → 하위 5% = −9.46% (통과, 여유 0.54%p)\n"
         "w = 0.6: 평균 8.0%, 표준편차 12.17% → 하위 5% = −12.01% (위반)\n"
         "등호로 놓고 제곱한 2차식의 양수 근 → w* = 52.2% (평균 7.61%, 표준편차 10.70%, 하위 5% −10.00%)")
T(B, 12, "자산이 셋 이상이면 같은 일을 행렬로 합니다 — 그것이 단계 2–3입니다.")
T(B, 13, "※ 이 장의 m(w)는 포트폴리오의 평균 수익률입니다 — CPPI 승수 m(3절)과 다른 뜻입니다.")

# --- C: 4-2 한계 — 롱온리 풀이
C = dup_slide(P, 37)
T(C, 3, "공매도를 막으면 상속 계정은 0 · 8.9 · 91.1%가 됩니다")
T(C, 5, "4-2 한계 — 롱온리로 풀기 · 상속 계정 H = −15%, α = 20% → z = 0.8416 (강의 재계산)")
eq_in(C, 7, 8, "s04_lo", r"m(w)=10\%+15\%\,w,\qquad \sigma^2(w)=0.04-0.04\,w+0.25\,w^2")
T(C, 10, "① 채권 −78.9%는 위반 → ② 채권을 0에 묶는다(KKT 활성 제약) → ③ 저위험 1 − w, 고위험 w로 다시 푼다")
T(C, 11, "w = 0.8 → 하위 20% −12.50% (통과) · w = 1.0 → −17.08% (위반) · 등호 2차식의 근 0.911\n"
         "해 0 · 8.9 · 91.1% — 기대 23.67%, 표준편차 45.94% (공매도 허용 해는 26.35%)\n"
         "합산 롱온리 13.30% (표준편차 19.38%) vs 같은 표준편차의 효율선 13.43% → 연 12bp")
T(C, 12, "계정별로 막으면 연 12bp를 잃고, 합산 수준에서만 막으면 이 예에서는 0입니다.")
T(C, 13, "※ 은퇴 · 교육 계정은 해가 원래 전부 양수라 그대로다 · 원문 §V(pp. 328–329)는 결과만 · 원문 합산 σ 19.89%는 19.38%의 오기로 보인다")

# --- D: 엑셀 해 찾기 · SLSQP (원형 30장: 두 카드)
D = dup_slide(P, 29)
for sid in (13, 14, 15, 16, 17, 18): remove(shape(D, sid))
T(D, 3, "엑셀 해 찾기나 SLSQP로 같은 롱온리 해를 얻습니다")
T(D, 5, "4-2 한계 — 상속 계정을 수치 최적화로 · 닫힌 해 g + v/γ는 공매도를 허용할 때만 있다")
for sid in (7, 10): shape(D, sid).height = E(330)
for sid in (9, 12): shape(D, sid).height = E(250)
T(D, 8, "엑셀 해 찾기 (한국어판)")
T(D, 9, "B2:B4 비중 · C2:C4 기대수익 μ · E2:G4 공분산 Σ\n"
        "B6 합계 · B7 기대수익(SUMPRODUCT) · B9 표준편차\n"
        "B10 = B7 − NORM.S.INV(0.8) × B9 ≥ B11(−15%)\n"
        "목표 B7 최대 · 제한 B6 = 1 · 해법 GRG 비선형\n"
        "‘음이 아닌 수’ 체크 → 0 · 8.9 · 91.1%")
T(D, 11, "파이썬 scipy.optimize.minimize")
T(D, 12, "목적 −w·μ, method = 'SLSQP'\n"
         "bounds = [(0, 1)] × 3 — 롱온리\n"
         "제약: 비중 합 = 1 · 하위 20% 경계 ≥ −15%\n"
         "결과 0, 0.089, 0.911 (기대 23.67%)\n"
         "bounds를 빼면 −78.9 · 96.2 · 82.7%")
T(D, 19, "‘음이 아닌 수’ 체크와 bounds (0, 1)이 같은 일을 합니다 — 체크 하나로 두 해가 갈립니다.")

# ===================================================================== 고친 장(옛 번호)
s = OLD[6]
T(s, 8, "※ 두 뜻으로 쓰는 글자 — A·C·K: 41~45장은 평균-분산 해의 상수, 그 밖은 자산·쿠션·행사가(60장) · m: 39~42·47장은 평균(mₚ, m(w), m(γ)), 그 밖은 CPPI 승수 · Cₜ: 72~74장은 그해 인출액, 그 밖의 C는 쿠션")
s = OLD[29]
T(s, 5, "3-6 ② — (a) gap risk를 막는 m의 상한, (b) 2절과 같은 해를 주는 위험 예산 · (c) floor를 움직이는 위험 자본 방식은 다음 장")
s = OLD[38]
T(s, 17, "※ A·C·K는 평균-분산 상수(원문 각주 6은 A · B · C · D, D = BC − A² — 강의 K = D/C), m(γ)는 평균 — 41~45장의 이 글자는 자산 A·쿠션 C·행사가 K·CPPI 승수 m과 다른 뜻입니다.")
s = OLD[39]
T(s, 11, "γ = 3.795 → 채권 53.9% · 저위험 26.6% · 고위험 19.5%\n"
         "평균 10.23%, 표준편차 12.30% → 10.23 − 1.645 × 12.30 = −10.00%\n"
         "원문 p.320의 3.9750은 오기다 — p.317과 표 1 · 2는 3.7950")
T(s, 13, "※ A·C·K는 41장의 평균-분산 상수입니다 — 자산 A·쿠션 C·행사가 K와 다른 뜻입니다.")
s = OLD[41]
T(s, 3, "공매도를 허용하면 계정을 합쳐도 같은 프론티어 위에 있습니다")
T(s, 5, "4-2 단계 5 — 공매도 허용: 각 계정이 g + v/γ 꼴이므로 합쳐도 같은 꼴 (원문 p.315 ii)b · pp. 320–321)")
T(s, 13, "공매도를 허용하면 계정을 나눠 관리해도 효율을 잃지 않습니다 — 롱온리는 47장.")
s = OLD[42]
lecfix.set_text(lecfix.find(s, "※ A·C·K는 39장"), "※ A·C·K는 41장의 평균-분산 상수입니다 — 자산 A·쿠션 C·행사가 K와 다른 뜻입니다.")
s = OLD[43]
T(s, 27, "정규분포 가정 — 꼬리가 두꺼우면 미달이 더 잦다\n빈도만 제한, 크기는 제한 안 함\n공매도 금지 시 직접 최적화(47 · 48장)\n한 기간 모형 — 동적 전략은 밖")
s = OLD[79]
T(s, 18, "계정별 롱온리 제약이면 정확히는 성립하지 않습니다(이 예 연 12bp, 합산 수준 제약이면 0).")
s = OLD[82]
note = s.shapes.add_textbox(E(79), E(652), E(1016), E(26))
tf = note.text_frame; tf.word_wrap = True
r = tf.paragraphs[0].add_run()
r.text = "※ 3-6(c) 한 줄 — 위험이 커져도 손실 예산을 일정하게: VaR 연동 floor(위험 자본 방식 CPPI, 31장), 손잡이는 손실 예산 L"
r.font.name = FONT; r.font.size = Pt(12); r.font.color.rgb = RGBColor.from_string(MUTED)

# ===================================================================== 순서 · 쪽번호
n = len(P.slides._sldIdLst)
ids = list(P.slides._sldIdLst)
a, b, c, d = ids[n - 4], ids[n - 3], ids[n - 2], ids[n - 1]
lst = P.slides._sldIdLst
for el in (a, b, c, d): lst.remove(el)
order = list(lst)
def after(old_no):   # 옛 번호 old_no 다음 위치
    return order.index(OLD_ID[old_no]) + 1
OLD_ID = {i + 1: el for i, el in enumerate(order)}
for el, old_no in ((d, 44), (c, 44), (b, 38), (a, 30)):   # 뒤에서부터 끼운다 — c 다음에 d가 오도록 d 먼저
    lst.insert(list(lst).index(OLD_ID[old_no]) + 1, el)
print("renumber", lecfix.renumber(P))
P.save(DST)
print("saved", DST, len(P.slides._sldIdLst))
