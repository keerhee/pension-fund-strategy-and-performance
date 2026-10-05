# -*- coding: utf-8 -*-
"""W06 정합(2026-10-05) 보충교재 17 · 프라이머 패치 공용 — 경로 · 수식 PNG · 그림 있는 슬라이드 복제."""
import copy, os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(R, "_build", "lecfix"))
import lecfix  # noqa
from pptx.util import Emu, Pt
BK = os.path.join(R, "_구판", "W06_정합전_2026-10-05", "보충17_프라이머")
SUP = os.path.join(R, "보충교재", "17_LDI_GBI 수식도출")
ART = os.path.join(HERE, "_art"); os.makedirs(ART, exist_ok=True)
PX = 9525
E = lambda v: int(round(v * PX))
FONT_DIR = os.path.expanduser("~/Library/Fonts")

def mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager as fm
    for f in ("Pretendard-Regular.otf", "Pretendard-SemiBold.otf", "Pretendard-Bold.otf"):
        p = os.path.join(FONT_DIR, f)
        if os.path.exists(p): fm.fontManager.addfont(p)
    matplotlib.rcParams["mathtext.fontset"] = "cm"
    matplotlib.rcParams["font.family"] = "Pretendard"
    return plt

def eq_png(name, tex, fg="#102027", bg="#F7F5F0", fs=30, dpi=220):
    """한 줄 mathtext 수식을 카드 배경색으로 구운 PNG. 반환: (경로, 가로/세로)"""
    plt = mpl()
    fig = plt.figure(figsize=(0.01, 0.01), dpi=dpi)
    fig.patch.set_facecolor(bg)
    fig.text(0, 0, f"${tex}$", fontsize=fs, color=fg)
    out = os.path.join(ART, name + ".png")
    fig.savefig(out, dpi=dpi, facecolor=bg, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)
    from PIL import Image
    w, h = Image.open(out).size
    return out, w / h

def add_pic(slide, path, x, y, w=None, h=None, ratio=None):
    """픽셀 좌표(1280×720)로 그림을 넣는다. w 또는 h 하나만 주면 비율로 맞춘다."""
    if w is not None and h is None: h = w / ratio
    if h is not None and w is None: w = h * ratio
    return slide.shapes.add_picture(path, E(x), E(y), E(w), E(h))

def fit_pic(slide, path, ratio, bx, by, bw, bh):
    """상자 (bx, by, bw, bh) 안에 비율을 지켜 가운데 맞춤으로 그림을 넣는다."""
    w, h = bw, bw / ratio
    if h > bh: h, w = bh, bh * ratio
    return add_pic(slide, path, bx + (bw - w) / 2, by + (bh - h) / 2, w, h)

def dup_slide(prs, idx):
    """같은 덱 안에서 슬라이드 복제(그림 포함 — 이미지 관계를 새로 건다). 맨 뒤에 붙는다."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    src = prs.slides[idx]
    dst = prs.slides.add_slide(src.slide_layout)
    for sh in list(dst.shapes):
        sh._element.getparent().remove(sh._element)
    for sh in src.shapes:
        el = copy.deepcopy(sh._element)
        dst.shapes._spTree.append(el)
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
    rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    for blip in dst.shapes._spTree.iter(ns):
        old = blip.get(rel)
        part = src.part.related_part(old)
        blip.set(rel, dst.part.relate_to(part, RT.IMAGE))
    # 배경(<p:bg>)도 옮긴다
    P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
    sbg = src._element.find(P + "cSld/" + P + "bg")
    if sbg is not None:
        dc = dst._element.find(P + "cSld")
        old = dc.find(P + "bg")
        if old is not None: dc.remove(old)
        dc.insert(0, copy.deepcopy(sbg))
    return dst

def shape(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid: return sh
    raise KeyError(sid)

def remove(sh):
    sh._element.getparent().remove(sh._element)

def move_px(sh, dx=0, dy=0):
    sh.left = sh.left + E(dx); sh.top = sh.top + E(dy)
