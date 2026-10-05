"""W06 M7 실습데이터 덱 — 2026-10-05 개정 강의본과 정합(python-pptx, 백업에서 읽는다).
실행: .venv/bin/python _build/w06_align/m7-data/patch_m7_data.py
"""
import os, sys
HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(R, "_build", "lecfix"))
import lecfix
from pptx import Presentation

BASE = "W06_M7_실습데이터_LDI부채연계투자"
SRC = os.path.join(R, "_구판", "W06_정합전_2026-10-05", BASE + ".pptx")
DST = os.path.join(R, "W06_LDI와GBI", BASE + ".pptx")

FIX = [
    (2, "케이스 재현 — 2,122조 · 69%", "케이스 재현 — 69%(우리 계산)"),
    (11, "69% · 95%", "69% · 95% (우리 계산)"),
    (11, "안 B", "안 B 국고채 교체"),
    (11, "안 C", "안 C IRS"),
    # 강의본에서 "FR −18%p 대 −2%p"는 빠졌다 — 단원 ⑦ · ⑧의 "2022 영국 일부 기금 FR −30%p"로 맞춘다
    (17, "강의본의 “FR −18%p 대 −2%p”는 단일 기금 모형으로는 재현되지 않습니다",
         "2022 영국 “일부 기금 FR −30%p”는 단일 기금 모형으로는 재현되지 않습니다"),
    (17, "−18%p의 붕괴", "−30%p의 붕괴"),
]


def rep(slide, old, new, exact=False):
    n = 0
    for sh in slide.shapes:
        for tf in lecfix._tf_iter(sh):
            for p in tf.paragraphs:
                t = p.text
                if exact and t.strip() != old:
                    continue
                if old in t and p.runs:
                    p.runs[0].text = t.replace(old, new); n += 1
                    for r in p.runs[1:]:
                        r._r.getparent().remove(r._r)
    return n


P = Presentation(SRC)
for no, old, new in FIX:
    n = rep(P.slides[no - 1], old, new, exact=old in ("안 B", "안 C"))
    if n == 0:
        sys.exit(f"[{no}] 찾지 못함: {old}")
    print(f"s{no:02d} ×{n}  {old[:40]}")
for i, t in lecfix.grep(P, r"−18|현물 (듀레이션 )?연장|교시"):
    print("남음", i, t)
P.save(DST)
print("저장", DST)
