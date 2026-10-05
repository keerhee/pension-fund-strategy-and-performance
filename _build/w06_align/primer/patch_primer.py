# -*- coding: utf-8 -*-
"""W06 프라이머(16장) — 2026-10-05 개정 M7 · M8 강의본과 정합. 빌드 스크립트가 없는 덱(_build/w07_primer.js는 옛 판)이라 python-pptx로 고친다.
  · 9장 방법 1 숫자: 9.10억 · 0.90억 · 49/22/29 → 9.13억 · 0.87억 · 49/22/28 (Market은 GHP + 주식 67% 2.24억, 균형형 2.22억은 후보 비교)
        그림 안 '합 9.10억' 글자를 덮어 다시 쓴다 · 텍스트 수식 6 × e^(−0.2) → mathtext 그림
  · 12장 95% · 69%는 우리 계산 표기 · 헤지 3%(2027 목표 비중) · 'M7 4교시 IC' → 단원 ⑨ 모의 IC
  · 13 · 15 · 16장 '교시' → 단원(M7 ①③⑦⑨ · M8 ③⑩) · 팀 번호(3팀 디폴트옵션 · 4팀 H대)
  · '돈' → 자금 · 금액, 밑줄 기호(D_A · W_T · z_p · D_H) 글자 수식 제거, 6장 기호 상자 넘침 정리
실행: <python-pptx + matplotlib 파이썬> patch_primer.py   (늘 _구판/W06_정합전_2026-10-05/보충17_프라이머/에서 읽는다)
"""
import io, os, sys
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "sup17"))
from common import *
from pptx import Presentation
from PIL import Image, ImageDraw, ImageFont

NAME = "W06_0_프라이머_LDI와GBI는왜부채와목표에서시작하나.pptx"
SRC = os.path.join(BK, NAME)
DST = os.path.join(R, "W06_LDI와GBI", NAME)
P = Presentation(SRC)
S = lambda n: P.slides[n - 1]
def T(n, sid, text): lecfix.set_text(shape(S(n), sid), text)
def cell(n, r, c, text):
    tb = lecfix.find_table(S(n)).table
    lecfix.set_tf(tb.cell(r, c).text_frame, text)

# 2 · 3 · 4 · 5 — '돈', 글자 수식
T(2, 19, "지킬 것을 먼저 안전자산으로 복제하고, 남는 자금으로 수익을 추구합니다")
T(3, 3, "연기금은 벌어들인 수익이 아니라 부채를 갚을 수 있는 비율로 평가받습니다")
T(4, 5, "기본 지식 — 멀리 있는 금액일수록 금리에 크게 흔들린다 · 연금은 수십 년 뒤에 나간다")
T(4, 12, "자산 듀레이션 ≈ 3년")
T(4, 14, "부채 듀레이션 ≈ 20년")
T(5, 8, "위험 = T년 뒤 자산이 목표 G에 못 미칠 확률")
# 6 — 핵심 수식
T(6, 11, "② 헤지 비중 (아래첨자 H = 헤지 채권)")
T(6, 14, "③ 목표 G를 T년 뒤 확률 p로 이루는 데 오늘 떼어 둘 자금 K")
T(6, 18, "m 성장률 · σ 변동성\nz 확률의 z값 · 95%면 1.645"); shape(S(6), 18).height = E(72)
# 9 — 방법 1
T(9, 3, "목표마다 공식 비중으로 사면 9.13억이 들고 배분은 49 / 22 / 28입니다")
s9 = S(9); box = shape(s9, 9); x, y, w, h = lecfix.px(box)
path, ratio = eq_png("primer_k_safety", r"6\times e^{-0.2}=4.91", fg="#102027", bg="#FFFFFF", fs=28)
remove(box); fit_pic(s9, path, ratio, x, y + 4, w - 40, h - 8)
T(9, 10, "GHP 성장률 2% · 변동성 0\nMarket은 GHP + 주식 67% → 2.24억\n남는 0.87억은 꿈의 계좌로")
T(9, 11, "49 / 22 / 28은 결과다")
pic = [sh for sh in s9.shapes if sh.shape_type == 13 and sh.width > E(500)][0]
im = Image.open(io.BytesIO(pic.image.blob)).convert("RGBA")
d = ImageDraw.Draw(im)
bg = im.getpixel((1000, 20))
d.rectangle([1030, 25, 1735, 105], fill=bg)
f = ImageFont.truetype(os.path.join(FONT_DIR, "Pretendard-Bold.otf"), 44)
d.text((1690, 64), "라임 = 그 목표에서 가장 싼 것", font=f, fill=(11, 107, 104, 255), anchor="rm")
out = os.path.join(ART, "primer_p9_chart.png"); im.save(out)
lecfix.replace_picture(pic, out, fit=False)
# 10 — 반전
T(10, 16, "반전 ③ 변동성이 자금을 아낀다")
# 12 — 국민연금
T(12, 14, "헤지 비율 (현 보유 · 2027 목표)")
cell(12, 0, 2, "적립비율 (우리 계산)")
T(12, 17, "국민연금은 적립비율을 공시하지 않습니다 — 공시 지표는 적립배율(적립금 ÷ 그해 지출)입니다")
T(12, 18, "M7 단원 ⑨ 모의 IC — 세 안 모두 헤지 3%에서 출발해, 이 갭을 얼마나 무엇으로 메울지 표결합니다")
# 13 — 처방
T(13, 20, "M8 단원 ⑩ 모의 IC — 디폴트옵션 Flexicure(3팀) · H대 발전기금 Yale 모델(4팀)")
# 14 — 용어집
cell(14, 3, 2, "금리 1%p에 가치가 변하는 % · 자산 D − 부채 D")
# 15 — 자기 점검
T(15, 5, "자기 점검 — 답은 가운데 칸, 자세한 설명은 강의본의 해당 단원")
for r, v in ((1, "M7 단원 ①"), (2, "M7 단원 ③"), (3, "M7 단원 ⑦"), (4, "M8 단원 ③")):
    cell(15, r, 2, v)
# 16 — 마감
T(16, 7, "마지막 단원은 모의 투자위원회 — M7 국민연금 LDI · M8 디폴트옵션 Flexicure · H대 발전기금 Yale 모델")

left = lecfix.grep(P, r"돈|교시|9\.10|0\.90억|_[A-Za-z]")
print("left:", left)
P.save(DST); print("saved", DST)
