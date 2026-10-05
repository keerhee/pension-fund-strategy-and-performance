# -*- coding: utf-8 -*-
"""W06 M8 실습데이터 덱 — 강의본(2026-10-05 개정) 정합.
실행: .venv/bin/python _build/w06_align/m8data/fix_m8data.py
원본은 _구판/W06_정합전_2026-10-05/에서 읽는다(덱 · w8_sim.py).
 1) w8_sim.py의 K 블록을 강의본 단원 ③ 방법 1(w* · g · K)로 바꾸고 실행 → w8_sim_results.json
 2) 10장 수식 PNG · 11장 차트를 다시 그린다
 3) 3 · 7 · 10 · 11 · 18장 문구, 17장 '선택 과제'(방법 2 + CPPI · 위험 자본 floor) 신설 → 20장
"""
import os, sys, json, re, shutil, subprocess
from math import sqrt
HERE = os.path.dirname(os.path.abspath(__file__))
R = os.path.abspath(os.path.join(HERE, "..", "..", ".."))
sys.path.insert(0, os.path.join(R, "_build", "lecfix"))
import lecfix
from pptx import Presentation
from pptx.util import Inches, Pt
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager

W = os.path.join(R, "W06_LDI와GBI"); BK = os.path.join(R, "_구판", "W06_정합전_2026-10-05")
DECK = "W06_M8_실습데이터_GBI목표기반투자"
OUT = os.path.join(HERE, "out"); os.makedirs(OUT, exist_ok=True)
# 강의본 장 번호는 하드코딩하지 않고 재번호 결과(이름 키 → 실제 번호)에서 읽는다
PG = json.load(open(os.path.join(R, "_build", "w06_m8_rework", "pages.json"), encoding="utf8"))

# ---------- 1) w8_sim.py K 블록 교체 + 실행 ----------
src = open(os.path.join(BK, "w8_sim.py"), encoding="utf8").read()
OLD_K = src[src.index("# 이순자 Safety 미달"):src.index("json.dump(out")]
NEW_K = '''# 필요 자본 K — 강의본 단원 ③ 방법 1(목표마다 주식 비중 w를 골라 K를 가장 작게):
#   K = G·e^(−T·g), g = r + wλ − ½w²σ² − z·wσ/√T, w* = λ/σ² − z/(σ√T)를 0~1로 자른다.
#   실습 가정에서 r = GHP 실질 1.5%, 주식 로그수익 6% = r + λ − ½σ² → λ = 6.5%.
#   후보 비교(강의본 26장 · 부록 B-2와 같은 방식): 균형형(m 4.5% · σ 9%) · 주식 · GHP 한 가지씩만 쓸 때의 K
from math import exp, sqrt
z = {0.7: 0.524, 0.3: -0.524}
LAM = MU - R + SIG**2/2
def K(G, p, m, s, T=10): return G*exp(-m*T + z[p]*s*sqrt(T))
def K1(G, p, T=10):
    w_raw = LAM/SIG**2 - z[p]/(SIG*sqrt(T)); w = min(max(w_raw, 0.0), 1.0)
    g = R + w*LAM - 0.5*w**2*SIG**2 - z[p]*w*SIG/sqrt(T)
    return G*exp(-T*g), w_raw, w, g
km, wmr, wm, gm = K1(1.0, 0.7); ka, war, wa, ga = K1(0.5, 0.3)
out["K"] = dict(lam=LAM, market_m1=km, w_market_raw=wmr, w_market=wm, g_market=gm,
                asp_m1=ka, w_asp_raw=war, w_asp=wa, g_asp=ga,
                total_m1=F0 + km + ka, spare_m1=A0 - (F0 + km + ka),
                market_bal=K(1.0, 0.7, 0.045, 0.09), market_eq=K(1.0, 0.7, 0.06, 0.2), market_ghp=1.0/(1+R)**10,
                asp_eq=K(0.5, 0.3, 0.06, 0.2), asp_bal=K(0.5, 0.3, 0.045, 0.09))
'''
src = src.replace(OLD_K, NEW_K)
src = src.replace('GHP(물가연동국채) 실질 1.5% = Floor 할인율, 월간 리밸런싱',
                  'GHP(물가연동국채) 실질 1.5% = Floor 할인율, 필요 자본 K는 강의본 방법 1, 월간 리밸런싱')
src = src.replace("(강의본 26장 · 부록 B-2", f"(강의본 {PG['o21']}장 · 부록 B-2")
open(os.path.join(W, "w8_sim.py"), "w", encoding="utf8").write(src)
subprocess.run([sys.executable, "w8_sim.py"], cwd=W, check=True)
J = json.load(open(os.path.join(W, "w8_sim_results.json")))
Kd = J["K"]
F0 = J["F0"]
print("K:", {k: round(v, 4) for k, v in Kd.items()})
f2 = lambda x: f"{x:.2f}"
KM, KA, TOT, SPARE = f2(Kd["market_m1"]), f2(Kd["asp_m1"]), f2(Kd["total_m1"]), f2(Kd["spare_m1"])
WM = f2(Kd["w_market"]); WAR = f2(Kd["w_asp_raw"])
GM, GA = f"{Kd['g_market']:.4f}", f"{Kd['g_asp']:.4f}"
assert (KM, KA, TOT, SPARE) == ("0.76", "0.20", "4.46", "0.54"), (KM, KA, TOT, SPARE)

# ---------- 2) 그림 ----------
for f in ["Pretendard-Regular.otf", "Pretendard-Bold.otf"]:
    font_manager.fontManager.addfont(os.path.join("/Users/keerhee/Library/Fonts", f))
plt.rcParams["font.family"] = "Pretendard"
plt.rcParams["mathtext.fontset"] = "cm"
BG, INK, TEAL, LIME, BLUE, GRY, GRY2, DARK, MUTED = "#F7F5F0", "#102027", "#0B6B68", "#B7F34A", "#4677F5", "#DDE3E2", "#9FB0B2", "#071A1D", "#64757D"

def eq_png(tex, path, fs=22):
    fig = plt.figure(figsize=(0.01, 0.01)); fig.patch.set_facecolor(BG)
    fig.text(0, 0, tex, fontsize=fs, color=INK)
    fig.savefig(path, dpi=300, facecolor=BG, bbox_inches="tight", pad_inches=0.06); plt.close(fig)

E1 = os.path.join(OUT, "eq_general.png"); E2 = os.path.join(OUT, "eq_market.png"); E3 = os.path.join(OUT, "eq_asp.png")
eq_png(r"$w^{*}=\dfrac{\lambda}{\sigma^{2}}-\dfrac{z_p}{\sigma\sqrt{T}},\qquad g_p=r+w\lambda-\dfrac{1}{2}w^{2}\sigma^{2}-\dfrac{z_p\,w\,\sigma}{\sqrt{T}},\qquad K=G\,e^{-T g_p}$", E1)
eq_png(r"$w^{*}=1.625-0.524/0.632=%s,\qquad K_{\mathrm{Mkt}}=1.0\times e^{-10\times %s}=%s$" % (WM, GM[:6], KM), E2)
eq_png(r"$w^{*}=1.625+0.524/0.632=%s\ \rightarrow\ 1,\qquad K_{\mathrm{Asp}}=0.5\times e^{-10\times %s}=%s$" % (WAR, GA[:6], KA), E3)

C11 = os.path.join(OUT, "chart11.png")
fig, (a1, a2) = plt.subplots(1, 2, figsize=(13.0, 4.37), gridspec_kw=dict(width_ratios=[1, 1.05], wspace=0.12))
fig.patch.set_facecolor(BG)
for a in (a1, a2):
    a.set_facecolor(BG)
    for s in ("top", "right", "left"): a.spines[s].set_visible(False)
    a.spines["bottom"].set_color(MUTED); a.set_yticks([]); a.tick_params(axis="x", length=0, labelsize=15)
sf, mk, asp = F0, Kd["market_m1"], Kd["asp_m1"]
a1.bar(0, sf, 0.6, color=TEAL, edgecolor=DARK, lw=1)
a1.bar(0, mk, 0.6, bottom=sf, color=BLUE, edgecolor=DARK, lw=1)
a1.bar(0, asp, 0.6, bottom=sf + mk, color=LIME, edgecolor=DARK, lw=1)
a1.bar(3.0, 5.0, 0.6, color=DARK)
a1.text(0, sf + mk + asp + 0.08, TOT, ha="center", va="bottom", fontsize=19, fontweight="bold", color=TEAL)
a1.text(3.0, 5.08, "5.00", ha="center", va="bottom", fontsize=19, fontweight="bold", color=DARK)
for y, lab, ly in [(sf / 2, f"Safety Floor {sf:.2f}", sf / 2), (sf + mk / 2, f"Market K {KM} (주식 {round(Kd['w_market']*100)}%)", sf + mk / 2 - 0.02),
                   (sf + mk + asp / 2, f"Aspirational K {KA}", sf + mk + asp / 2 + 0.42)]:
    a1.plot([0.31, 0.45], [y, ly], color=MUTED, lw=1)
    a1.text(0.48, ly, lab, va="center", fontsize=14, fontweight="bold", color=INK)
a1.set_xticks([0, 3.0]); a1.set_xticklabels(["필요 자본 합 (방법 1)", "금융자산"]); a1.set_xlim(-0.5, 3.5); a1.set_ylim(0, 5.9)
a1.set_title("3버킷 필요 자본 vs 금융자산 (억 원)", fontsize=15, color=INK, pad=8)
# 오른쪽: 후보 비교 — 한 가지 포트폴리오만 쓸 때의 K
x = [0, 1, 2]; v70 = [Kd["market_ghp"], Kd["market_bal"], Kd["market_eq"]]; v30 = [None, Kd["asp_bal"], Kd["asp_eq"]]
c70 = [GRY, LIME, GRY]; c30 = [None, GRY2, LIME]
for i in x:
    a2.bar(i - 0.2, v70[i], 0.38, color=c70[i], edgecolor=DARK, lw=1)
    a2.text(i - 0.2, v70[i] + 0.012, f2(v70[i]), ha="center", va="bottom", fontsize=15, fontweight="bold", color=INK)
    if v30[i] is not None:
        a2.bar(i + 0.2, v30[i], 0.38, color=c30[i], edgecolor=DARK, lw=1)
        a2.text(i + 0.2, v30[i] + 0.012, f2(v30[i]), ha="center", va="bottom", fontsize=15, fontweight="bold", color=INK)
a2.set_xticks(x); a2.set_xticklabels(["GHP\n(σ 0)", "균형형\n(4.5%, 9%)", "주식\n(6%, 20%)"]); a2.set_ylim(0, 1.12)
a2.set_title("후보 비교 — 한 가지만 쓸 때의 K (라임 = 후보 중 가장 싼 것)", fontsize=15, color=INK, pad=8)
from matplotlib.patches import Patch
a2.legend(handles=[Patch(facecolor=GRY, edgecolor=DARK, label="의료 1억 · 70%"), Patch(facecolor=GRY2, edgecolor=DARK, label="손주 0.5억 · 30%")],
          loc="upper center", ncol=2, frameon=False, fontsize=13, bbox_to_anchor=(0.5, 1.0))
fig.savefig(C11, dpi=200, facecolor=BG, bbox_inches="tight", pad_inches=0.1); plt.close(fig)

# ---------- 3) 덱 ----------
P = Presentation(os.path.join(BK, DECK + ".pptx"))
S = lambda n: P.slides[n - 1]
def put(slide, prefix, text, contains=False):
    sh = lecfix.find(slide, prefix, contains)
    if sh is None: sys.exit(f"못 찾음: {prefix}")
    lecfix.set_text(sh, text); return sh

# 3장 — 옛 교시 표기
put(S(3), "강의 2교시의 CPPI 규칙", "강의본 단원 ⑤의 CPPI 규칙 — 매달 인출(부족분)을 빼면서 돌린다")
# 7장 — Step 3
put(S(7), "의료비 1억(70%)", "의료비 1억(70%) · 손주 교육 5천(30%)의 필요 자본 K — 강의본 방법 1(목표마다 w*)")
# 10장 — 방법 1로 다시
s = S(10)
put(s, "Step 3 —", "Step 3 — 필요 자본은 강의본 방법 1로 목표마다 가장 싼 비중을 찾아 셉니다")
put(s, "강의 1교시 필요 자본 공식", f"강의본 단원 ③ 방법 1({PG['M1']}장 손계산과 같은 순서) — 10년 뒤(75세) 목표 G를 확률 p로 이룰 오늘 자금")
by = {sh.name: sh for sh in s.shapes}
cards = [("Shape 5", "Text 6", "Image 0", 2.20, 1.30,
          f"방법 1 — w*는 0~1로 자른다 · 실습 가정 r 1.5% · σ 20% · T 10 · λ = 6% − 1.5% + ½ × 0.2² = 6.5%", E1),
         ("Shape 7", "Text 8", "Image 1", 3.62, 1.22,
          f"Market — 의료비 1억 · 70%(z 0.524) → GHP {100-round(Kd['w_market']*100)}% + 주식 {round(Kd['w_market']*100)}%, 확률 조정 성장률 g = {Kd['g_market']*100:.2f}%", E2),
         ("Shape 9", "Text 10", "Image 2", 4.96, 1.22,
          f"Aspirational — 손주 교육 0.5억 · 30%(z −0.524) → 전액 주식, g = {Kd['g_asp']*100:.2f}%", E3)]
from PIL import Image
for box, lab, img, top, h, text, png in cards:
    b, l, p = by[box], by[lab], by[img]
    b.left, b.top, b.width, b.height = Inches(0.88), Inches(top), Inches(11.47), Inches(h)
    l.left, l.top, l.width = Inches(1.08), Inches(top + 0.08), Inches(11.07)
    lecfix.set_text(l, text)
    iw, ih = Image.open(png).size
    ah = h - 0.48; aw = 11.0
    sc = min(aw / (iw / 300), ah / (ih / 300))
    nw, nh = iw / 300 * sc, ih / 300 * sc
    nw, nh = min(nw, iw / 300 * 0.62 * 1.0 if False else nw), nh
    p.left = Inches(0.88 + (11.47 - nw) / 2); p.top = Inches(top + 0.40 + (ah - nh) / 2); p.width = Inches(nw); p.height = Inches(nh)
    rId = p._element.blipFill.blip.rEmbed; p.part.related_part(rId)._blob = open(png, "rb").read()
put(s, "Floor 3.50 + 0.74", f"Floor 3.50 + {KM} + {KA} = {TOT}억 ≤ 5억 — 균형형 0.74는 후보 비교용, 합계는 방법 1 공식값")
# 11장 — 차트
s = S(11)
put(s, "세 버킷의 필요 자본 4.44억", f"세 버킷의 필요 자본 {TOT}억이 금융자산 5억 안에 들어옵니다")
put(s, "왼쪽: 버킷별", "왼쪽: 버킷별 필요 자본의 합(방법 1 공식값) · 오른쪽: 후보 비교 — 확률 70%면 GHP(0.86)보다 위험자산 쪽이 싸다")
put(s, "여유 0.56억", f"여유 {SPARE}억 · 균형형 0.74는 후보 비교용 — 변동성 낮은 상품이 더 싸다 · 30%는 주식 0.20 < 균형형 0.27")
pic = lecfix.pictures(s)[0]; pic.left, pic.top, pic.width, pic.height = Inches(1.14), Inches(2.30), Inches(10.95), Inches(3.70)
lecfix.replace_picture(pic, C11)
# 18장 — 합 · 여유
s = S(18)
put(s, "Floor 3.50억 먼저", f"Floor 3.50억 먼저\nK 합 {TOT}억\n여유 {SPARE}억")
# 13 · 14장 — '달성' 열에 기준을 붙인다(강의본 MXG장 방식): 열 이름 % = 지켜야 할 기준, 표 안 % = 센 달성 확률
from pptx.dml.color import RGBColor
CRIT = {"Market 달성": ("Market 달성 확률\n잉여 ≥ 1.0억 · 기준 70%", 70),
        "Aspirational 달성": ("Aspirational 달성 확률\n잉여 ≥ 1.5억 · 기준 30%", 30)}
LEGEND = "열 이름 % = 지켜야 할 기준 · 표 안 % = 1,000 경로 중 그 금액을 넘은 비율(센 달성 확률) · 빨간 칸 = 미달(여기선 없음)"
for k in (13, 14):
    s = S(k); tb = lecfix.find_table(s).table
    for ci, cell in enumerate(tb.rows[0].cells):
        h = cell.text.strip()
        if h == "Floor 위반" and k == 13: lecfix.set_tf(cell.text_frame, "Floor 위반\n(월간 갭 σ 20%)")
        if h not in CRIT: continue
        head, need = CRIT[h]; lecfix.set_tf(cell.text_frame, head)
        for r in list(tb.rows)[1:]:
            v = float(r.cells[ci].text.strip().rstrip("%"))
            if v < need:
                for p in r.cells[ci].text_frame.paragraphs:
                    for run in p.runs: run.font.bold = True; run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    # 폭 · 높이 — 기준 열을 넓히고 행을 낮춘다(아래 문구와 겹치지 않게)
    total = sum(c.width for c in tb.columns); n = len(tb.columns)
    wts = [1.0] * (n - 2) + [1.9, 1.9]
    for c, w in zip(tb.columns, wts): c.width = int(total * w / sum(wts))
    for ri, row in enumerate(tb.rows):
        row.height = Inches(0.62 if ri == 0 else 0.34)
        for cell in row.cells:
            for p in cell.text_frame.paragraphs:
                for run in p.runs: run.font.size = Pt(13 if ri == 0 else 15)
sh = lecfix.find(S(13), "Market = 75세 잉여"); lecfix.set_text(sh, LEGEND)
sh = lecfix.find(S(14), "Flexicure: 매년 초"); lecfix.set_text(sh, "Flexicure: 매년 초 Floor 재설정, m = 1 · 고정 40/60은 TDF형 비교 기준 · 열 이름 % = 기준, 표 안 % = 센 달성 확률")
# 17장 신설 — 선택 과제 (16장 복제)
new = lecfix.dup_slide(P, 15)
import copy
bg = P.slides[15]._element.cSld.bg   # 배경(크림)도 옮긴다
if bg is not None: new._element.cSld.insert(0, copy.deepcopy(bg))
lecfix.move_slide(P, len(P.slides) - 1, 16)
s = S(17)
put(s, "BAF.60080", "BAF.60080 · W06 M8 실습데이터 · 선택 과제")
put(s, "선택 확장 —", "선택 과제 — 방법 2와 위험 자본 floor로 다시 봅니다")
put(s, "기본 실습은 로그정규", "기본 실습은 방법 1 + CPPI · 선택 과제는 강의본 숫자를 먼저 재현한 뒤 이순자 씨에게 옮긴다")
tb = lecfix.find_table(s)
lecfix.set_table(tb, [
    ["선택 과제", "먼저 재현할 강의본 숫자", "이순자 씨에게 옮겨 볼 질문"],
    [f"A · 방법 2 + CPPI\n단원 ③ {PG['DM1']}~{PG['DM2']} · {PG['FD']}장",
     "Das–Markowitz 세 계정 γ 3.795 · 2.706 · 0.877 → 합산(60 · 20 · 20) γ 2.174 · 비중 24.2 · 42.2 · 33.7%(공매도 허용)\nCPPI 결합: A 100 · F 80 · m 3 → E 60",
     "Floor 3.50억은 CPPI로 두고 의료 · 손주 목표를 (H, α) 계정으로 적으면, 위험 블록 구성이 방법 1과 어떻게 다른가"],
    [f"B · 위험 자본 floor\n단원 ⑤ {PG['V1']}~{PG['RC']}장",
     "L 고정 · F = A − L ÷ (m × x), x = 1.645σ − μ\nA 100 · m 3 · L 16.1: σ 15 / 20 / 25% → F 71.2 / 80 / 84.7",
     "위험이 오르면 floor를 올려 노출을 줄인다 — 필수 Floor 3.50억을 하한으로 두고 m 3 결과와 비교"]])
tbl = tb.table
for c, wd in zip(tbl.columns, [2.25, 5.70, 3.84]): c.width = Inches(wd)
put(s, "[입력] 추가 지시", "[입력] 추가 지시 · 선택 과제 A")
put(s, "주식 수익률을 W06_M8", "강의본 Das–Markowitz 예(자산 5·5 / 10·20 / 25·50%, 상관 0.2)를 먼저 재현해 γ 3.795 · 2.706 · 0.877을 확인하고,\n같은 순서로 이순자 씨의 쿠션(위험 블록)을 의료 · 손주 계정의 (H, α)로 나눠 구성해줘.")
put(s, "※ 숫자는 기본 실습과", "※ 기대 결과는 강의본 숫자까지만 — 이순자 씨에게 옮긴 결과는 정답표가 없다 · 고른 (H, α)와 가정을 함께 밝힌다")
n = lecfix.renumber(P)
# 남은 옛 표현 점검
left = lecfix.grep(P, r"교시|4\.44|0\.56억|균형형\(m 4\.5%, σ 9%\)\)?$")
print("쪽번호", n, "· 남은 옛 표현", left)
dst = os.path.join(W, DECK + ".pptx"); P.save(dst)

# ---------- 4) PDF (Pretendard) ----------
TMP = os.path.join(OUT, "pdf"); shutil.rmtree(TMP, ignore_errors=True); os.makedirs(TMP)
tp = os.path.join(TMP, DECK + ".pptx")
subprocess.run(["/usr/bin/python3", os.path.join(R, "_build", "to_pretendard.py"), dst, tp], check=True)
import tempfile; LO = tempfile.mkdtemp()   # 공백 · 대괄호 없는 경로의 전용 LibreOffice 프로필(동시 실행 충돌 방지)
subprocess.run(["soffice", f"-env:UserInstallation=file://{LO}", "--headless", "--convert-to", "pdf", "--outdir", TMP, tp], check=True, capture_output=True)
shutil.copy2(os.path.join(TMP, DECK + ".pdf"), os.path.join(W, DECK + ".pdf"))
print("saved", dst, "slides", len(P.slides))

# ---------- 5) 대본의 강의본 장 참조 (pages.json 기준, 재실행해도 같은 결과) ----------
NARR = "/Users/keerhee/Project/lectures/pension-w06n-m8-data/script/narr.py"
t = open(NARR, encoding="utf8").read()
t, n1 = re.subn(r"강의본 단원 3의 \d+장에서 \d+장(?:, 그리고 \d+장)?입니다\.",
                f"강의본 단원 3의 {PG['DM1']}장에서 {PG['DM2']}장, 그리고 {PG['FD']}장입니다.", t)
t, n2 = re.subn(r"강의본 단원 5의 \d+장과 \d+장입니다\.", f"강의본 단원 5의 {PG['V1']}장과 {PG['RC']}장입니다.", t)
assert n1 == 1 and n2 == 1, (n1, n2)
open(NARR, "w", encoding="utf8").write(t); print("narr 장 참조", n1 + n2)

# 13 · 14장 대본 — 기준과 센 확률의 구분(한 문장, 재실행해도 한 번만)
t = open(NARR, encoding="utf8").read()
ADD = {"\"5단계 승수 민감도의 전체 표입니다.\",": "열 이름의 퍼센트는 지켜야 할 기준이고, 표 안의 퍼센트는 1,000개 경로 가운데 그 금액을 넘은 비율, 곧 시뮬레이션이 센 달성 확률입니다.",
       "\"6단계입니다.\",": "이 표도 열 이름의 퍼센트는 기준, 표 안의 퍼센트는 시뮬레이션이 센 달성 확률입니다."}
for a, sent in ADD.items():
    if sent not in t:
        assert t.count(a) == 1, a
        t = t.replace(a, a + f' "{sent}",')
open(NARR, "w", encoding="utf8").write(t)
